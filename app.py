# -*- coding: utf-8 -*-
"""
ALLUCO — Planning Laquage Agentic IA V5
=======================================
Un seul fichier Python, prêt pour GitHub / Streamlit Cloud.

Entrée:
    Bd-Client-S36.xlsx (versionné avec le dépôt, aucun upload utilisateur)

Sortie:
    Planning automatique Lundi -> Samedi, affiché dans Streamlit et exportable Excel.

Politique couleur:
    - 1 couleur/jour est fortement privilégiée;
    - 2 couleurs/jour maximum (contrainte dure).

Architecture agentique déterministe:
    Données -> Référentiel source -> Quantités -> Priorités -> Scénarios ->
    OR-Tools / fallback -> Couleurs -> Réparation -> Critique -> Validation -> Export.

La valeur "Confiance règles = 100%" signifie que toutes les règles du moteur ont été
validées (capacité, <=2 couleurs/jour, pas de doublon, jours actifs, etc.). Elle ne
constitue pas une garantie de réalité terrain si les données sources sont incorrectes.
"""
from __future__ import annotations

import hashlib
import io
import json
import math
import os
import re
import sys
import time
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, replace as dc_replace
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:
    import tomllib  # Python 3.11+
except Exception:  # pragma: no cover
    tomllib = None

try:
    import streamlit as st
except Exception:  # permet les tests CLI sans Streamlit installé
    st = None

try:
    from ortools.sat.python import cp_model
    ORTOOLS_AVAILABLE = True
except Exception:
    cp_model = None
    ORTOOLS_AVAILABLE = False


# =============================================================================
# 1) CONFIGURATION
# =============================================================================
VERSION = "5.0.0"
APP_NAME = "ALLUCO — Planning Laquage IA"
APP_SUBTITLE = "Agentic AI · Planification automatique"
ROOT_DIR = Path(__file__).resolve().parent
CONFIG_PATH = ROOT_DIR / "config.toml"

DAYS = ["LUNDI", "MARDI", "MERCREDI", "JEUDI", "VENDREDI", "SAMEDI"]
SHEET_NAMES = [f"Planning {d.capitalize()}" for d in DAYS[:-1]] + ["Planning SAMEDI"]

OUTPUT_COLUMNS = [
    "NumCommande", "DateCréation", "NomClient", "Article", "Article/int", "Couleur", "Nuance",
    "QteCommandé", "ResteALivrer", "Prelevé", "reservation brut", "NumOF", "ProdStatut",
    "QteCommencé", "QteRestante", "QteRèçu", "ReserverBR", "StockPhysique", "Reserver",
    "Lancement", "Re-laquage", "PoidsUn", "PoidsT", "Poudre", "Barre/bal", "Nbre Bal", "tps", "Stock brut",
]

DEFAULT_NUANCE = {
    "BLC": 1.0, "R9016": 2.0, "R1013": 5.0, "R1019": 7.0,
    "SAND": 9.0, "FRENE": 11.0, "TECK": 13.0, "ACAJOU": 15.0,
    "NOYER": 16.0, "NOCE": 17.0, "R8019": 19.0, "TRESOR": 20.0,
    "GRIS": 22.0, "GREY": 23.0, "GRISG": 24.0, "R7016": 27.0,
    "CHPG": 29.0, "COOL": 31.0, "N02": 33.0, "N07": 34.0,
    "N22": 35.0, "NOIR": 37.0, "DARK": 38.0,
    "ANOD": 40.0, "ANODN": 41.0, "ABRONZE": 42.0,
}

# Référentiel atelier embarqué, utilisé uniquement quand Barre/bal n'est pas disponible
# dans le fichier commandes. Il évite de dépendre d'un ancien Planning S36 dans GitHub.
FAMILY_BARS_PER_BAL = {
    "EC": 13, "FR": 13, "CSQ": 13, "FSQ": 13, "CO": 13,
    "LM": 17, "GL": 20, "P": 14, "PL": 20, "C": 10,
    "LMDP": 14, "LMDPF": 20, "AL": 10, "LMMO": 19,
    "MR": 9, "PR": 6, "PCN": 800, "T": 800,
}


def _load_toml() -> Dict[str, Any]:
    if tomllib is None or not CONFIG_PATH.exists():
        return {}
    try:
        with open(CONFIG_PATH, "rb") as f:
            return tomllib.load(f)
    except Exception:
        return {}


APP_CONFIG = _load_toml()
DATA_CFG = APP_CONFIG.get("data", {})
PLAN_CFG = APP_CONFIG.get("planning", {})
UI_CFG = APP_CONFIG.get("ui", {})

SOURCE_FILENAME = str(DATA_CFG.get("source_file", "Bd-Client-S36.xlsx"))
SOURCE_PATH = ROOT_DIR / SOURCE_FILENAME

DEFAULT_YEAR = int(PLAN_CFG.get("year", 2026))
DEFAULT_WEEK = int(PLAN_CFG.get("week", 36))
DEFAULT_CAPACITY_H = float(PLAN_CFG.get("capacity_weekday_h", 15.0))
DEFAULT_SATURDAY_ENABLED = bool(PLAN_CFG.get("saturday_enabled", True))
DEFAULT_SATURDAY_CAPACITY_H = float(PLAN_CFG.get("saturday_capacity_h", 15.0))
DEFAULT_MIN_PER_BAL = float(PLAN_CFG.get("minutes_per_bal", 4.0))
DEFAULT_POWDER_COEFF = float(PLAN_CFG.get("powder_coeff", 0.052))
DEFAULT_CLEANING_MIN = int(PLAN_CFG.get("cleaning_min", 15))
DEFAULT_TARGET_UTIL = float(PLAN_CFG.get("target_utilization", 0.94))
DEFAULT_SOLVER_SECONDS = float(PLAN_CFG.get("solver_seconds", 18.0))
DEFAULT_AUTO_GENERATE = bool(PLAN_CFG.get("auto_generate", True))
DEFAULT_ALLOW_RELAQUAGE = bool(PLAN_CFG.get("allow_relaquage", False))
DEFAULT_MAX_JOBS = int(PLAN_CFG.get("max_jobs", 1600))
DEFAULT_POOL_FACTOR = float(PLAN_CFG.get("pool_factor", 2.7))
PREFERRED_COLORS_PER_DAY = 1
HARD_MAX_COLORS_PER_DAY = 2


# =============================================================================
# 2) HELPERS
# =============================================================================
def norm_text(v: Any) -> str:
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return ""
    return str(v).strip()


def norm_key(v: Any) -> str:
    s = unicodedata.normalize("NFKD", norm_text(v)).encode("ascii", "ignore").decode("ascii")
    s = s.lower().strip()
    return re.sub(r"[^a-z0-9]+", "_", s).strip("_")


def to_float(v: Any, default: float = 0.0) -> float:
    try:
        if v is None or (isinstance(v, float) and np.isnan(v)):
            return default
        if isinstance(v, str):
            s = v.strip().replace(" ", "").replace(",", ".")
            if not s or s.upper() in {"#N/A", "N/A", "NA", "NONE", "NAN"}:
                return default
            return float(s)
        return float(v)
    except Exception:
        return default


def to_int(v: Any, default: int = 0) -> int:
    return int(round(to_float(v, float(default))))


def parse_date(v: Any) -> Optional[pd.Timestamp]:
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return None
    try:
        if isinstance(v, (int, float, np.integer, np.floating)):
            n = float(v)
            if 20000 <= n <= 80000:
                ts = pd.Timestamp("1899-12-30") + pd.to_timedelta(n, unit="D")
            else:
                return None
        else:
            ts = pd.to_datetime(v, errors="coerce")
        if pd.isna(ts) or ts.year <= 1900:
            return None
        return pd.Timestamp(ts)
    except Exception:
        return None


def iso_week_dates(year: int, week: int) -> Dict[int, date]:
    monday = date.fromisocalendar(int(year), int(week), 1)
    return {i: monday + timedelta(days=i) for i in range(6)}


def _looks_like_color_token(token: str) -> bool:
    t = norm_text(token).upper()
    if not t or t == "BRUT":
        return False
    # Écarter les suffixes qui ressemblent clairement à des dimensions/références produit.
    if re.search(r"\d+(?:[.,]\d+)?X\d+", t) or "/" in t or "." in t:
        return False
    if t in DEFAULT_NUANCE:
        return True
    if re.fullmatch(r"R\d{4}", t) or re.fullmatch(r"N\d{2}", t):
        return True
    # Noms de teintes libres (OTARIE, SWEET, GALET, WENGE, etc.).
    return bool(re.fullmatch(r"[A-Z][A-Z0-9]{1,14}", t))


def split_article(article: Any) -> Tuple[str, str]:
    s = norm_text(article)
    if "-" not in s:
        return s, ""
    left, right = s.rsplit("-", 1)
    color = right.strip().upper()
    if not _looks_like_color_token(color):
        return s, ""
    return left.strip(), color


def article_family(article_internal: str) -> str:
    s = norm_text(article_internal).upper()
    m = re.match(r"([A-Z]+)", s)
    return m.group(1) if m else ""


def stable_unknown_nuance(color: str) -> float:
    # déterministe pour que deux exécutions donnent le même ordre.
    h = hashlib.sha256(norm_text(color).upper().encode("utf-8")).hexdigest()
    return 50.0 + (int(h[:6], 16) % 4000) / 100.0


def bytes_from_path(path: Path) -> bytes:
    return path.read_bytes()


def format_num(v: float, digits: int = 0) -> str:
    if digits == 0:
        return f"{v:,.0f}".replace(",", " ")
    return f"{v:,.{digits}f}".replace(",", " ")


# =============================================================================
# 3) AGENT DONNÉES — lecture du fichier GitHub
# =============================================================================
def _compact_sheet_rows(ws, header_row: int = 1, key_col: int = 2, blank_stop: int = 180) -> pd.DataFrame:
    header = [c.value for c in next(ws.iter_rows(min_row=header_row, max_row=header_row))]
    while header and header[-1] is None:
        header.pop()
    if not header:
        return pd.DataFrame()

    rows: List[Tuple[Any, ...]] = []
    blanks = 0
    seen = False
    for row in ws.iter_rows(min_row=header_row + 1, max_col=len(header), values_only=True):
        key = row[key_col - 1] if len(row) >= key_col else None
        if key is None and all(v is None for v in row):
            if seen:
                blanks += 1
                if blanks >= blank_stop:
                    break
            continue
        blanks = 0
        seen = True
        rows.append(tuple(row[: len(header)]))
    return pd.DataFrame(rows, columns=header)


def load_source_workbook(data: bytes) -> pd.DataFrame:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    candidate = None
    if "Feuil1" in wb.sheetnames:
        candidate = wb["Feuil1"]
    else:
        for ws in wb.worksheets:
            first = [norm_key(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
            if "numcommande" in first and "article" in first:
                candidate = ws
                break
    if candidate is None:
        raise ValueError("Feuille commandes introuvable (NumCommande / Article).")
    df = _compact_sheet_rows(candidate)
    if df.empty:
        raise ValueError("Le fichier commandes est vide.")
    required = ["NumCommande", "DateCréation", "NomClient", "Article", "QteCommandé", "ResteALivrer"]
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise ValueError("Colonnes obligatoires manquantes: " + ", ".join(missing))
    return df


# =============================================================================
# 4) AGENT RÉFÉRENTIEL — apprentissage uniquement depuis l'input
# =============================================================================
@dataclass
class TechnicalMaster:
    weight_by_article: Dict[str, float]
    weight_by_family: Dict[str, float]
    powder_coeff: float
    minutes_per_bal: float
    source_weight_samples: int


def learn_source_master(source: pd.DataFrame, powder_coeff: float, minutes_per_bal: float) -> TechnicalMaster:
    rows = []
    for _, r in source.iterrows():
        article_internal, _ = split_article(r.get("Article"))
        w = to_float(r.get("PoidArticle"), 0.0)
        if article_internal and w > 0:
            rows.append((article_internal.upper(), article_family(article_internal), w))
    if not rows:
        return TechnicalMaster({}, {}, powder_coeff, minutes_per_bal, 0)

    d = pd.DataFrame(rows, columns=["article", "family", "weight"])
    exact = d.groupby("article")["weight"].median().to_dict()
    fam_stats = d[d["family"] != ""].groupby("family")["weight"].agg(["median", "count"])
    fam = {k: float(v["median"]) for k, v in fam_stats.iterrows() if int(v["count"]) >= 3}
    return TechnicalMaster(exact, fam, powder_coeff, minutes_per_bal, len(d))


def infer_unit_weight(article_internal: str, direct: float, master: TechnicalMaster) -> Tuple[float, str]:
    if direct > 0:
        return direct, "source"
    art = norm_text(article_internal).upper()
    if art in master.weight_by_article:
        return float(master.weight_by_article[art]), "article_source"
    fam = article_family(art)
    if fam in master.weight_by_family:
        return float(master.weight_by_family[fam]), "famille_source"
    return 0.0, "inconnu"


def infer_bars_per_bal(article_internal: str, unit_weight: float) -> Tuple[int, str]:
    fam = article_family(article_internal)
    if fam in FAMILY_BARS_PER_BAL:
        return int(FAMILY_BARS_PER_BAL[fam]), "référentiel"
    if unit_weight <= 0:
        return 13, "fallback"
    # fallback physique prudent, borné pour éviter des valeurs absurdes.
    if unit_weight <= 0.08:
        return 200, "poids"
    if unit_weight < 1.2:
        return 25, "poids"
    if unit_weight < 2.0:
        return 20, "poids"
    if unit_weight < 3.2:
        return 17, "poids"
    if unit_weight < 8.0:
        return 13, "poids"
    if unit_weight < 10.0:
        return 9, "poids"
    return 6, "poids"


def color_nuance(color: str) -> Tuple[float, bool]:
    c = norm_text(color).upper()
    if c in DEFAULT_NUANCE:
        return DEFAULT_NUANCE[c], True
    return stable_unknown_nuance(c), False


# =============================================================================
# 5) AGENT QUANTITÉS + PRIORITÉS
# =============================================================================
@dataclass(frozen=True)
class PlannerConfig:
    year: int
    week: int
    capacity_h: float = DEFAULT_CAPACITY_H
    saturday_enabled: bool = DEFAULT_SATURDAY_ENABLED
    saturday_capacity_h: float = DEFAULT_SATURDAY_CAPACITY_H
    cleaning_min: int = DEFAULT_CLEANING_MIN
    minutes_per_bal: float = DEFAULT_MIN_PER_BAL
    powder_coeff: float = DEFAULT_POWDER_COEFF
    target_utilization: float = DEFAULT_TARGET_UTIL
    solver_seconds: float = DEFAULT_SOLVER_SECONDS
    max_jobs: int = DEFAULT_MAX_JOBS
    pool_factor: float = DEFAULT_POOL_FACTOR
    allow_relaquage: bool = DEFAULT_ALLOW_RELAQUAGE
    strategy: str = "Auto — meilleur compromis"
    force_commands: Tuple[str, ...] = ()
    exclude_commands: Tuple[str, ...] = ()


def day_capacity_h(cfg: PlannerConfig, d: int) -> float:
    if d == 5:
        return cfg.saturday_capacity_h if cfg.saturday_enabled else 0.0
    return cfg.capacity_h


def day_capacity_min(cfg: PlannerConfig, d: int) -> int:
    return max(0, int(round(day_capacity_h(cfg, d) * 60)))


def due_date_from_row(row: pd.Series) -> Optional[pd.Timestamp]:
    for c in ["DateLivraisonConfirmé", "DateExpeditionConfirmé", "DateExpeditionDemandé"]:
        if c in row.index:
            d = parse_date(row.get(c))
            if d is not None:
                return d
    return None


def _ready_score(prod_status: str, reservation_flag: str, reserver_br: float, stock: float, qte_recue: float, remaining: float) -> float:
    s = 0.0
    ps = norm_key(prod_status)
    if "commenc" in ps:
        s += 280
    elif "cree" in ps:
        s += 220
    elif ps in {"", "_", "-"}:
        s += 80
    if norm_key(reservation_flag) in {"oui", "yes", "1", "true"}:
        s += 260
    if remaining > 0:
        s += min(260, (reserver_br / remaining) * 260) if reserver_br > 0 else 0
        s += min(100, (stock / remaining) * 100) if stock > 0 else 0
        s += min(120, (qte_recue / remaining) * 120) if qte_recue > 0 else 0
    return s


def _priority_score(created: Optional[pd.Timestamp], due: Optional[pd.Timestamp], week_start: pd.Timestamp,
                    prod_status: str, reservation_flag: str, reserver_br: float, stock: float,
                    qte_recue: float, remaining: float) -> Tuple[float, int, str]:
    score = 100.0
    reasons: List[str] = []
    overdue_days = 0
    if due is not None:
        overdue_days = max(0, (week_start.date() - due.date()).days)
        days_to_due = (due.date() - week_start.date()).days
        if overdue_days > 0:
            score += 2200 + overdue_days * 260
            reasons.append(f"retard {overdue_days}j")
        elif days_to_due <= 1:
            score += 1500
            reasons.append("échéance immédiate")
        elif days_to_due <= 5:
            score += 950 - max(0, days_to_due) * 80
            reasons.append("échéance semaine")
        elif days_to_due <= 12:
            score += 300
            reasons.append("échéance proche")
    else:
        score += 40
        reasons.append("date non confirmée")

    if created is not None:
        age = max(0, (week_start.date() - created.date()).days)
        score += min(650, age * 8)
        if age > 30:
            reasons.append("commande ancienne")

    ready = _ready_score(prod_status, reservation_flag, reserver_br, stock, qte_recue, remaining)
    score += ready
    if ready >= 350:
        reasons.append("matière/OF prêt")
    elif ready < 100:
        reasons.append("préparation faible")

    return float(score), int(overdue_days), " · ".join(reasons[:4])


def build_candidate_lines(source: pd.DataFrame, master: TechnicalMaster, cfg: PlannerConfig) -> Tuple[pd.DataFrame, Dict[str, Any]]:
    week_dates = iso_week_dates(cfg.year, cfg.week)
    week_start = pd.Timestamp(week_dates[0])
    rows: List[Dict[str, Any]] = []
    excluded = Counter()
    force_set = {norm_text(x).upper() for x in cfg.force_commands if norm_text(x)}
    exclude_set = {norm_text(x).upper() for x in cfg.exclude_commands if norm_text(x)}
    quality = Counter()

    for src_idx, src in source.iterrows():
        cmd = norm_text(src.get("NumCommande")).upper()
        if not cmd:
            excluded["commande vide"] += 1
            continue
        if cmd in exclude_set:
            excluded["exclusion manuelle"] += 1
            continue

        etat = norm_key(src.get("EtatCommande"))
        ligne_etat = norm_key(src.get("EtatLigneCommande"))
        remaining = max(0.0, to_float(src.get("ResteALivrer")))
        if etat and "encours" not in etat:
            excluded["commande non encours"] += 1
            continue
        if ligne_etat and "encours" not in ligne_etat:
            excluded["ligne non encours"] += 1
            continue
        if remaining <= 0:
            excluded["reste nul"] += 1
            continue

        article = norm_text(src.get("Article"))
        article_internal, color = split_article(article)
        if not article_internal or not color or color == "BRUT":
            excluded["article/couleur non planifiable"] += 1
            continue

        prod_status = norm_text(src.get("ProdStatut"))
        if "declare_termine" in norm_key(prod_status) or "declare termine" in norm_key(prod_status).replace("_", " "):
            excluded["OF terminé"] += 1
            continue

        created = parse_date(src.get("DateCréation"))
        due = due_date_from_row(src)
        qte_commanded = max(0.0, to_float(src.get("QteCommandé")))
        qte_recue = max(0.0, to_float(src.get("QteRèçu")))
        qte_commence = max(0.0, to_float(src.get("QteCommencé")))
        qte_restante = max(0.0, to_float(src.get("QteRestante")))
        reservation_flag = norm_text(src.get(" 2")) or "non"
        reserver_br = max(0.0, to_float(src.get("ReserverBR")))
        stock_phys = max(0.0, to_float(src.get("StockPhysique")))
        reserver = max(0.0, to_float(src.get("Reserver")))

        direct_weight = max(0.0, to_float(src.get("PoidArticle")))
        unit_weight, weight_source = infer_unit_weight(article_internal, direct_weight, master)
        if weight_source == "inconnu":
            quality["poids_inconnu"] += 1
        elif weight_source != "source":
            quality["poids_infere_source"] += 1

        bars, bars_source = infer_bars_per_bal(article_internal, unit_weight)
        if bars_source != "référentiel":
            quality["barres_estimees"] += 1

        relaquage = 0.0
        if cfg.allow_relaquage and stock_phys > 0:
            # conservateur: maximum 25% du besoin client et jamais plus que le stock physique.
            relaquage = min(stock_phys, remaining * 0.25, remaining)

        launch = max(0.0, remaining - relaquage)
        launch_i = int(math.ceil(launch - 1e-9))
        relaq_i = int(math.floor(relaquage + 1e-9))
        nbal = int(math.ceil(launch_i / bars)) if launch_i > 0 and bars > 0 else 0
        duration_h = nbal * master.minutes_per_bal / 60.0
        weight_total = (launch_i + relaq_i) * unit_weight if unit_weight > 0 else 0.0
        powder = weight_total * master.powder_coeff
        nuance, nuance_known = color_nuance(color)
        if not nuance_known:
            quality["nuance_inconnue"] += 1

        priority, overdue_days, reason = _priority_score(
            created, due, week_start, prod_status, reservation_flag, reserver_br,
            stock_phys, qte_recue, remaining
        )
        if cmd in force_set:
            priority += 100000
            reason = "FORCÉ · " + reason

        line_id = hashlib.sha1(
            f"{src_idx}|{cmd}|{article}|{norm_text(src.get('NumOF'))}".encode("utf-8")
        ).hexdigest()[:16]

        rows.append({
            "NumCommande": cmd,
            "DateCréation": created.to_pydatetime() if created is not None else None,
            "NomClient": norm_text(src.get("NomClient")),
            "Article": article,
            "Article/int": article_internal,
            "Couleur": color,
            "Nuance": round(float(nuance), 2),
            "QteCommandé": int(round(qte_commanded)),
            "ResteALivrer": int(round(remaining)),
            "Prelevé": to_int(src.get("Prelevé")),
            "reservation brut": reservation_flag,
            "NumOF": norm_text(src.get("NumOF")),
            "ProdStatut": prod_status,
            "QteCommencé": int(round(qte_commence)),
            "QteRestante": int(round(qte_restante)),
            "QteRèçu": int(round(qte_recue)),
            "ReserverBR": int(round(reserver_br)),
            "StockPhysique": int(round(stock_phys)),
            "Reserver": int(round(reserver)),
            "Lancement": launch_i,
            "Re-laquage": relaq_i,
            "PoidsUn": round(unit_weight, 4),
            "PoidsT": round(weight_total, 3),
            "Poudre": round(powder, 3),
            "Barre/bal": int(bars),
            "Nbre Bal": int(nbal),
            "tps": round(duration_h, 4),
            "Stock brut": int(round(reserver_br)),
            "_line_id": line_id,
            "_source_index": int(src_idx),
            "_due": due.to_pydatetime() if due is not None else None,
            "_score": round(priority, 3),
            "_overdue_days": overdue_days,
            "_reason": reason,
            "_forced": cmd in force_set,
            "_weight_source": weight_source,
            "_bars_source": bars_source,
            "_nuance_known": nuance_known,
        })

    df = pd.DataFrame(rows)
    if not df.empty:
        df = df.sort_values(["_score", "_due", "DateCréation"], ascending=[False, True, True], na_position="last").reset_index(drop=True)

    info = {
        "source_rows": int(len(source)),
        "eligible_lines": int(len(df)),
        "excluded": dict(excluded),
        "quality": dict(quality),
        "source_weight_samples": master.source_weight_samples,
    }
    return df, info


# =============================================================================
# 6) AGENT JOBS — cohérence commande + couleur
# =============================================================================
@dataclass
class Job:
    job_id: str
    line_indices: List[int]
    command: str
    color: str
    duration_min: int
    score: float
    due: Optional[datetime]
    created: Optional[datetime]
    forced: bool = False


def _make_job(cmd: str, color: str, idxs: List[int], sub: pd.DataFrame, chunk_no: int) -> Job:
    dues = [x for x in sub["_due"].tolist() if x is not None and not pd.isna(x)]
    created = [x for x in sub["DateCréation"].tolist() if x is not None and not pd.isna(x)]
    return Job(
        job_id=f"{cmd}|{color}|{chunk_no}",
        line_indices=list(idxs),
        command=str(cmd),
        color=str(color),
        duration_min=max(1, int(round(pd.to_numeric(sub["tps"], errors="coerce").fillna(0).sum() * 60))),
        score=float(sub["_score"].max() + min(400.0, sub["_score"].mean() * 0.04) + len(sub) * 3),
        due=min(dues) if dues else None,
        created=min(created) if created else None,
        forced=bool(sub["_forced"].any()),
    )


def build_jobs(lines: pd.DataFrame, cfg: PlannerConfig) -> Tuple[List[Job], List[int]]:
    if lines.empty:
        return [], []
    max_day_cap = max(day_capacity_min(cfg, d) for d in range(6))
    jobs: List[Job] = []
    oversized: List[int] = []

    for (cmd, color), g in lines.groupby(["NumCommande", "Couleur"], sort=False):
        g = g.sort_values(["_score", "_due"], ascending=[False, True], na_position="last")
        chunk: List[int] = []
        chunk_min = 0
        chunk_no = 1
        for i, r in g.iterrows():
            mins = max(1, int(round(to_float(r["tps"]) * 60)))
            if mins > max_day_cap:
                oversized.append(int(i))
                continue
            if chunk and chunk_min + mins > max_day_cap:
                sub = lines.loc[chunk]
                jobs.append(_make_job(cmd, color, chunk, sub, chunk_no))
                chunk_no += 1
                chunk = []
                chunk_min = 0
            chunk.append(int(i))
            chunk_min += mins
        if chunk:
            jobs.append(_make_job(cmd, color, chunk, lines.loc[chunk], chunk_no))
    return jobs, oversized


def select_candidate_pool(jobs: List[Job], cfg: PlannerConfig) -> Tuple[List[Job], List[Job]]:
    week_capacity = sum(day_capacity_min(cfg, d) for d in range(6))
    ordered = sorted(jobs, key=lambda j: (not j.forced, -j.score, j.due or datetime.max, j.created or datetime.max, j.duration_min))
    selected: List[Job] = []
    minutes = 0
    target = week_capacity * max(1.6, cfg.pool_factor)
    for job in ordered:
        if len(selected) >= cfg.max_jobs:
            break
        selected.append(job)
        minutes += job.duration_min
        if minutes >= target and len(selected) >= 300:
            break
    selected_ids = {j.job_id for j in selected}
    outside = [j for j in jobs if j.job_id not in selected_ids]
    return selected, outside


# =============================================================================
# 7) AGENT PLANIFICATEUR — OR-Tools + mono-couleur prioritaire
# =============================================================================
SCENARIO_WEIGHTS = {
    "Délais clients": {"unscheduled": 180, "late": 180, "two_color": 80, "color": 15, "balance": 1, "split": 20, "early": 0},
    "Mono-couleur": {"unscheduled": 125, "late": 90, "two_color": 250000, "color": 90, "balance": 1, "split": 35, "early": 1},
    "Équilibre": {"unscheduled": 150, "late": 125, "two_color": 230, "color": 45, "balance": 2, "split": 28, "early": 1},
}


def _job_unscheduled_penalty(job: Job) -> int:
    return max(10000, int(round(25000 + job.score * 150 + min(900, job.duration_min) * 8)))


def assign_jobs_ortools(jobs: List[Job], cfg: PlannerConfig) -> Tuple[Dict[str, int], List[str], str]:
    if not ORTOOLS_AVAILABLE or not jobs:
        return {}, [], ""

    w = SCENARIO_WEIGHTS.get(cfg.strategy, SCENARIO_WEIGHTS["Équilibre"])
    week_dates = iso_week_dates(cfg.year, cfg.week)
    colors = sorted({j.color for j in jobs})
    commands = sorted({j.command for j in jobs})
    by_color: Dict[str, List[int]] = defaultdict(list)
    by_command: Dict[str, List[int]] = defaultdict(list)
    for ji, job in enumerate(jobs):
        by_color[job.color].append(ji)
        by_command[job.command].append(ji)

    model = cp_model.CpModel()
    x: Dict[Tuple[int, int], Any] = {}
    u: Dict[int, Any] = {}
    y: Dict[Tuple[int, str], Any] = {}
    objective: List[Any] = []

    for ji, job in enumerate(jobs):
        u[ji] = model.NewBoolVar(f"u_{ji}")
        for d in range(6):
            x[(ji, d)] = model.NewBoolVar(f"x_{ji}_{d}")
            if day_capacity_min(cfg, d) <= 0:
                model.Add(x[(ji, d)] == 0)
        model.Add(sum(x[(ji, d)] for d in range(6)) + u[ji] == 1)
        if job.forced:
            model.Add(u[ji] == 0)
        objective.append(u[ji] * _job_unscheduled_penalty(job) * w["unscheduled"])

        if job.due is not None:
            due_d = job.due.date()
            for d in range(6):
                late = max(0, (week_dates[d] - due_d).days)
                early = max(0, (due_d - week_dates[d]).days - 2)
                if late:
                    objective.append(x[(ji, d)] * late * w["late"] * 140)
                if early and w["early"]:
                    objective.append(x[(ji, d)] * early * w["early"] * 10)

    for d in range(6):
        cap = day_capacity_min(cfg, d)
        if cap <= 0:
            continue
        active = []
        for color in colors:
            y[(d, color)] = model.NewBoolVar(f"y_{d}_{norm_key(color)}")
            idxs = by_color[color]
            for ji in idxs:
                model.Add(x[(ji, d)] <= y[(d, color)])
            model.Add(y[(d, color)] <= sum(x[(ji, d)] for ji in idxs))
            active.append(y[(d, color)])
            objective.append(y[(d, color)] * w["color"] * 100)

        n_colors = sum(active)
        color_limit = PREFERRED_COLORS_PER_DAY if cfg.strategy == "Mono-couleur" else HARD_MAX_COLORS_PER_DAY
        model.Add(n_colors <= color_limit)
        second_color = model.NewBoolVar(f"second_color_{d}")
        model.Add(n_colors <= 1 + second_color)
        objective.append(second_color * w["two_color"] * 1000)

        prod = sum(jobs[ji].duration_min * x[(ji, d)] for ji in range(len(jobs)))
        model.Add(prod + cfg.cleaning_min * second_color <= cap)

        target = int(round(cap * cfg.target_utilization))
        load = model.NewIntVar(0, cap, f"load_{d}")
        model.Add(load == prod + cfg.cleaning_min * second_color)
        dev = model.NewIntVar(0, cap, f"dev_{d}")
        model.Add(dev >= target - load)
        model.Add(dev >= load - target)
        objective.append(dev * w["balance"])

    # même commande: limiter la dispersion sur plusieurs jours.
    for ci, cmd in enumerate(commands):
        idxs = by_command[cmd]
        if len(idxs) <= 1:
            continue
        presents = []
        for d in range(6):
            p = model.NewBoolVar(f"cmd_{ci}_{d}")
            for ji in idxs:
                model.Add(x[(ji, d)] <= p)
            model.Add(p <= sum(x[(ji, d)] for ji in idxs))
            presents.append(p)
        extra = model.NewIntVar(0, 5, f"split_{ci}")
        model.Add(extra >= sum(presents) - 1)
        objective.append(extra * w["split"] * 500)

    model.Minimize(sum(objective))
    solver = cp_model.CpSolver()
    solver.parameters.max_time_in_seconds = max(3.0, float(cfg.solver_seconds))
    solver.parameters.num_search_workers = max(1, min(8, os.cpu_count() or 2))
    solver.parameters.random_seed = 36
    status = solver.Solve(model)
    if status not in (cp_model.OPTIMAL, cp_model.FEASIBLE):
        return {}, [], ""

    assignments: Dict[str, int] = {}
    unscheduled: List[str] = []
    for ji, job in enumerate(jobs):
        if solver.Value(u[ji]):
            unscheduled.append(job.job_id)
            continue
        for d in range(6):
            if solver.Value(x[(ji, d)]):
                assignments[job.job_id] = d
                break
        if job.job_id not in assignments:
            unscheduled.append(job.job_id)
    return assignments, unscheduled, "OR-Tools CP-SAT Agentic"


def _fallback_day_cost(job: Job, d: int, loads: List[int], colors_by_day: List[set], commands_by_day: List[set], cfg: PlannerConfig) -> float:
    cap = day_capacity_min(cfg, d)
    if cap <= 0:
        return float("inf")
    new_color = job.color not in colors_by_day[d]
    ncolors = len(colors_by_day[d]) + (1 if new_color else 0)
    if ncolors > HARD_MAX_COLORS_PER_DAY:
        return float("inf")
    cleaning = cfg.cleaning_min if ncolors == 2 else 0
    current_cleaning = cfg.cleaning_min if len(colors_by_day[d]) == 2 else 0
    projected = loads[d] - current_cleaning + job.duration_min + cleaning
    if projected > cap:
        return float("inf")

    week_day = iso_week_dates(cfg.year, cfg.week)[d]
    late = max(0, (week_day - job.due.date()).days) if job.due else 0
    mono_pen = 0
    if not colors_by_day[d]:
        mono_pen = 80
    elif new_color:
        # En mode mono-couleur, une deuxième couleur n’est utilisée qu’en dernier recours.
        mono_pen = 1_000_000_000 if cfg.strategy == "Mono-couleur" else 120_000
    else:
        mono_pen = -25_000
    due_pen = late * (45000 if cfg.strategy == "Délais clients" else 28000)
    group_bonus = -2500 if job.command in commands_by_day[d] else 0
    target = cap * cfg.target_utilization
    balance_pen = abs(target - projected) * (2 if cfg.strategy == "Équilibre" else 1)
    return due_pen + mono_pen + balance_pen + group_bonus - job.score * 5


def assign_jobs_fallback(jobs: List[Job], cfg: PlannerConfig) -> Tuple[Dict[str, int], List[str], str]:
    loads = [0] * 6
    colors_by_day: List[set] = [set() for _ in range(6)]
    commands_by_day: List[set] = [set() for _ in range(6)]
    assignments: Dict[str, int] = {}
    unscheduled: List[str] = []
    ordered = sorted(jobs, key=lambda j: (not j.forced, -j.score, j.due or datetime.max, -j.duration_min))

    for job in ordered:
        options = []
        for d in range(6):
            if cfg.strategy == "Mono-couleur" and colors_by_day[d] and job.color not in colors_by_day[d]:
                continue
            cost = _fallback_day_cost(job, d, loads, colors_by_day, commands_by_day, cfg)
            if math.isfinite(cost):
                options.append((cost, d))
        if not options:
            unscheduled.append(job.job_id)
            continue
        _, d = min(options, key=lambda x: (x[0], x[1]))
        before_n = len(colors_by_day[d])
        colors_by_day[d].add(job.color)
        after_n = len(colors_by_day[d])
        if before_n < 2 <= after_n:
            loads[d] += cfg.cleaning_min
        loads[d] += job.duration_min
        commands_by_day[d].add(job.command)
        assignments[job.job_id] = d

    return assignments, unscheduled, "Heuristique Agentic robuste"


def assign_jobs(jobs: List[Job], cfg: PlannerConfig) -> Tuple[Dict[str, int], List[str], str]:
    if ORTOOLS_AVAILABLE:
        a, u, engine = assign_jobs_ortools(jobs, cfg)
        if engine:
            return a, u, engine
    return assign_jobs_fallback(jobs, cfg)


# =============================================================================
# 8) AGENT RÉPARATION — mono-couleur + backlog urgent
# =============================================================================
def _state_from_assignments(jobs: List[Job], assignments: Dict[str, int], cfg: PlannerConfig):
    by_id = {j.job_id: j for j in jobs}
    loads = [0] * 6
    colors: List[set] = [set() for _ in range(6)]
    day_jobs: List[List[str]] = [[] for _ in range(6)]
    for jid, d in assignments.items():
        j = by_id[jid]
        day_jobs[d].append(jid)
        loads[d] += j.duration_min
        colors[d].add(j.color)
    for d in range(6):
        if len(colors[d]) == 2:
            loads[d] += cfg.cleaning_min
    return by_id, loads, colors, day_jobs


def _can_place(job: Job, d: int, loads: List[int], colors: List[set], cfg: PlannerConfig, remove_job: Optional[Job] = None) -> bool:
    cap = day_capacity_min(cfg, d)
    if cap <= 0:
        return False
    current_colors = set(colors[d])
    current_load = loads[d]
    if remove_job is not None:
        current_load -= remove_job.duration_min
        # exact color recalculation requires day job list; caller only uses remove_job from same color swap conservatively.
    new_colors = set(current_colors)
    new_colors.add(job.color)
    if len(new_colors) > HARD_MAX_COLORS_PER_DAY:
        return False
    # cleaning state estimated conservatively: if 2 colors after placement, reserve cleaning.
    clean_after = cfg.cleaning_min if len(new_colors) == 2 else 0
    clean_before = cfg.cleaning_min if len(current_colors) == 2 else 0
    prod_before = current_load - clean_before
    return prod_before + job.duration_min + clean_after <= cap + 1e-9


def repair_assignments(jobs: List[Job], assignments: Dict[str, int], unscheduled: List[str], cfg: PlannerConfig) -> Tuple[Dict[str, int], List[str], List[str]]:
    assignments = dict(assignments)
    unscheduled_set = set(unscheduled)
    log: List[str] = []
    by_id = {j.job_id: j for j in jobs}

    # Boucle 1: essayer de transformer les jours à 2 couleurs en jours mono-couleur.
    for _ in range(4):
        changed = False
        by_id, loads, colors, day_jobs = _state_from_assignments(jobs, assignments, cfg)
        for d in range(6):
            if len(colors[d]) != 2:
                continue
            color_minutes = Counter()
            for jid in day_jobs[d]:
                j = by_id[jid]
                color_minutes[j.color] += j.duration_min
            minority = min(color_minutes, key=color_minutes.get)
            move_ids = [jid for jid in day_jobs[d] if by_id[jid].color == minority and not by_id[jid].forced]
            # on ne déplace la couleur minoritaire que si toutes ses lignes peuvent aller ensemble ailleurs.
            moved_to = None
            for target in range(6):
                if target == d or day_capacity_min(cfg, target) <= 0:
                    continue
                target_colors = colors[target]
                if target_colors and minority not in target_colors:
                    continue  # mono-couleur privilégiée: pas créer un nouveau 2e coloris ici.
                total = sum(by_id[jid].duration_min for jid in move_ids)
                clean_before = cfg.cleaning_min if len(target_colors) == 2 else 0
                prod_target = loads[target] - clean_before
                projected_colors = set(target_colors) | ({minority} if move_ids else set())
                clean_after = cfg.cleaning_min if len(projected_colors) == 2 else 0
                if prod_target + total + clean_after <= day_capacity_min(cfg, target):
                    moved_to = target
                    break
            if moved_to is not None and move_ids:
                for jid in move_ids:
                    assignments[jid] = moved_to
                log.append(f"Réduction couleur: {minority} déplacée de {DAYS[d]} vers {DAYS[moved_to]} pour obtenir un jour mono-couleur.")
                changed = True
                break
        if not changed:
            break

    # Boucle 2: récupérer les jobs en retard / forte priorité si une place compatible existe.
    for _ in range(8):
        if not unscheduled_set:
            break
        by_id, loads, colors, day_jobs = _state_from_assignments(jobs, assignments, cfg)
        backlog = sorted((by_id[jid] for jid in unscheduled_set if jid in by_id), key=lambda j: (not j.forced, -j.score, j.due or datetime.max))
        changed = False
        for job in backlog[:120]:
            candidates = []
            for d in range(6):
                if day_capacity_min(cfg, d) <= 0:
                    continue
                # préférence absolue même couleur, puis jour vide. En scénario Mono-couleur,
                # l'agent de réparation n'a pas le droit de réintroduire une 2e couleur.
                if cfg.strategy == "Mono-couleur" and colors[d] and job.color not in colors[d]:
                    continue
                color_class = 0 if job.color in colors[d] else 1 if not colors[d] else 2
                if _can_place(job, d, loads, colors, cfg):
                    planned_date = iso_week_dates(cfg.year, cfg.week)[d]
                    late = max(0, (planned_date - job.due.date()).days) if job.due else 0
                    candidates.append((late, color_class, loads[d], d))
            if candidates:
                _, _, _, d = min(candidates)
                assignments[job.job_id] = d
                unscheduled_set.remove(job.job_id)
                log.append(f"Backlog récupéré: {job.command} / {job.color} placé {DAYS[d]}.")
                changed = True
                break
        if not changed:
            break

    return assignments, sorted(unscheduled_set), log


# =============================================================================
# 9) AGENT COULEURS / SÉQUENÇAGE
# =============================================================================
def order_colors(colors: Sequence[str]) -> List[str]:
    unique = list(dict.fromkeys(norm_text(c).upper() for c in colors if norm_text(c)))
    return sorted(unique, key=lambda c: (color_nuance(c)[0], c))


def sequence_days(lines: pd.DataFrame, jobs: List[Job], assignments: Dict[str, int]) -> Dict[int, pd.DataFrame]:
    job_map = {j.job_id: j for j in jobs}
    day_indices: Dict[int, List[int]] = {d: [] for d in range(6)}
    for jid, d in assignments.items():
        if jid in job_map:
            day_indices[d].extend(job_map[jid].line_indices)

    result: Dict[int, pd.DataFrame] = {}
    for d in range(6):
        if not day_indices[d]:
            result[d] = lines.iloc[0:0].copy()
            continue
        df = lines.loc[day_indices[d]].copy()
        route = order_colors(df["Couleur"].tolist())
        rank = {c: i for i, c in enumerate(route)}
        df["_color_rank"] = df["Couleur"].astype(str).str.upper().map(rank).fillna(999)
        df["_due_sort"] = pd.to_datetime(df["_due"], errors="coerce")
        df = df.sort_values(
            ["_color_rank", "_due_sort", "_score", "NumCommande", "Article"],
            ascending=[True, True, False, True, True], na_position="last"
        ).drop(columns=["_color_rank", "_due_sort"]).reset_index(drop=True)
        df["_planned_day"] = DAYS[d]
        result[d] = df
    return result


# =============================================================================
# 10) AGENT CRITIQUE / VALIDATION
# =============================================================================
def business_day_df(df: pd.DataFrame) -> pd.DataFrame:
    if df is None or df.empty:
        return pd.DataFrame(columns=OUTPUT_COLUMNS)
    out = df.copy()
    for c in OUTPUT_COLUMNS:
        if c not in out.columns:
            out[c] = None
    return out[OUTPUT_COLUMNS].reset_index(drop=True)


def planning_metrics(days: Dict[int, pd.DataFrame], cfg: PlannerConfig, unscheduled: Optional[pd.DataFrame] = None) -> Dict[str, Any]:
    day_metrics = []
    total_load = 0.0
    total_cleaning = 0.0
    late_planned = 0
    late_days_sum = 0
    two_color_days = 0
    week_dates = iso_week_dates(cfg.year, cfg.week)

    for d in range(6):
        df = days[d]
        production = float(pd.to_numeric(df.get("tps", pd.Series(dtype=float)), errors="coerce").fillna(0).sum())
        colors = list(dict.fromkeys(df.get("Couleur", pd.Series(dtype=str)).dropna().astype(str).str.upper().tolist()))
        cleaning_h = cfg.cleaning_min / 60.0 if len(colors) == 2 else 0.0
        if len(colors) == 2:
            two_color_days += 1
        total = production + cleaning_h
        cap = day_capacity_h(cfg, d)
        total_load += total
        total_cleaning += cleaning_h

        if not df.empty:
            due = pd.to_datetime(df.get("_due"), errors="coerce")
            for dt in due.dropna():
                l = max(0, (week_dates[d] - dt.date()).days)
                if l > 0:
                    late_planned += 1
                    late_days_sum += l

        day_metrics.append({
            "Jour": DAYS[d],
            "Date": week_dates[d].strftime("%d/%m/%Y"),
            "Charge production h": round(production, 2),
            "Nettoyage h": round(cleaning_h, 2),
            "Charge totale h": round(total, 2),
            "Capacité h": round(cap, 2),
            "Charge %": round(total / cap * 100, 1) if cap > 0 else 0.0,
            "Couleurs": " → ".join(colors),
            "Nb couleurs": len(colors),
            "Lignes": int(len(df)),
        })

    total_capacity = sum(day_capacity_h(cfg, d) for d in range(6))
    overdue_uns = 0
    if unscheduled is not None and not unscheduled.empty and "_overdue_days" in unscheduled.columns:
        overdue_uns = int((pd.to_numeric(unscheduled["_overdue_days"], errors="coerce").fillna(0) > 0).sum())

    return {
        "days": day_metrics,
        "total_load_h": round(total_load, 2),
        "capacity_h": round(total_capacity, 2),
        "utilization_pct": round(total_load / total_capacity * 100, 1) if total_capacity else 0.0,
        "cleaning_h": round(total_cleaning, 2),
        "two_color_days": two_color_days,
        "mono_color_days": sum(1 for x in day_metrics if x["Nb couleurs"] == 1),
        "late_planned_lines": late_planned,
        "late_days_sum": late_days_sum,
        "overdue_unscheduled": overdue_uns,
    }


def validate_plan(days: Dict[int, pd.DataFrame], lines: pd.DataFrame, cfg: PlannerConfig) -> Tuple[List[str], List[str], int]:
    hard: List[str] = []
    soft: List[str] = []
    metrics = planning_metrics(days, cfg)

    # Contrôles durs.
    for dm in metrics["days"]:
        if dm["Capacité h"] <= 0 and dm["Lignes"] > 0:
            hard.append(f"{dm['Jour']}: journée désactivée utilisée.")
        if dm["Charge totale h"] > dm["Capacité h"] + 1e-6:
            hard.append(f"{dm['Jour']}: surcharge {dm['Charge totale h']:.2f}h > {dm['Capacité h']:.2f}h.")
        if dm["Nb couleurs"] > HARD_MAX_COLORS_PER_DAY:
            hard.append(f"{dm['Jour']}: {dm['Nb couleurs']} couleurs > maximum 2.")
        if dm["Nb couleurs"] == 2:
            soft.append(f"{dm['Jour']}: 2 couleurs utilisées; 1 couleur reste la cible préférée.")

    planned_frames = [d for d in days.values() if d is not None and not d.empty]
    planned = pd.concat(planned_frames, ignore_index=True) if planned_frames else pd.DataFrame()
    if not planned.empty:
        if planned["_line_id"].duplicated().any():
            hard.append("Doublon de ligne détecté dans le planning.")
        source_ids = set(lines["_line_id"].astype(str))
        if not set(planned["_line_id"].astype(str)).issubset(source_ids):
            hard.append("Une ligne planifiée ne provient pas du pool éligible.")
        for c in ["Lancement", "Nbre Bal", "tps"]:
            vals = pd.to_numeric(planned[c], errors="coerce")
            if vals.isna().any() or (vals < 0).any():
                hard.append(f"Valeurs invalides dans {c}.")
        if any(list(business_day_df(days[d]).columns) != OUTPUT_COLUMNS for d in range(6)):
            hard.append("Format métier des 28 colonnes non conforme.")

    # Score = couverture stricte des règles, pas une probabilité de réussite terrain.
    confidence = 100 if not hard else max(0, 100 - min(100, len(hard) * 25))
    return hard, soft, confidence


def data_quality_notes(lines: pd.DataFrame, quality: Dict[str, Any]) -> List[str]:
    notes: List[str] = []
    q = quality.get("quality", {})
    if q.get("poids_inconnu", 0):
        notes.append(f"{q['poids_inconnu']} ligne(s) sans PoidsUn fiable: PoidsT/Poudre à contrôler.")
    if q.get("poids_infere_source", 0):
        notes.append(f"{q['poids_infere_source']} PoidsUn inféré(s) depuis d'autres lignes du même fichier source.")
    if q.get("barres_estimees", 0):
        notes.append(f"{q['barres_estimees']} Barre/bal estimé(s) par référentiel atelier / poids.")
    if q.get("nuance_inconnue", 0):
        notes.append(f"{q['nuance_inconnue']} nuance(s) non référencée(s): ordre couleur déterministe utilisé.")
    return notes


# =============================================================================
# 11) AGENT SCÉNARIOS — orchestration complète
# =============================================================================
def _assemble(lines: pd.DataFrame, quality: Dict[str, Any], selected_jobs: List[Job], outside_jobs: List[Job], oversized: List[int], cfg: PlannerConfig) -> Dict[str, Any]:
    assignments, unsched_ids, engine = assign_jobs(selected_jobs, cfg)
    assignments, unsched_ids, repair_log = repair_assignments(selected_jobs, assignments, unsched_ids, cfg)
    days = sequence_days(lines, selected_jobs, assignments)

    job_map = {j.job_id: j for j in selected_jobs}
    unscheduled_idx: List[int] = []
    for jid in unsched_ids:
        if jid in job_map:
            unscheduled_idx.extend(job_map[jid].line_indices)
    for job in outside_jobs:
        unscheduled_idx.extend(job.line_indices)
    unscheduled_idx.extend(oversized)
    unscheduled_idx = list(dict.fromkeys(unscheduled_idx))
    unscheduled = lines.loc[unscheduled_idx].copy() if unscheduled_idx else lines.iloc[0:0].copy()

    metrics = planning_metrics(days, cfg, unscheduled)
    hard, soft, confidence = validate_plan(days, lines, cfg)
    notes = data_quality_notes(lines, quality)

    weighted_backlog = float(pd.to_numeric(unscheduled.get("_score", pd.Series(dtype=float)), errors="coerce").fillna(0).sum()) if not unscheduled.empty else 0.0
    scenario_score = (
        metrics["overdue_unscheduled"] * 1_000_000
        + metrics["late_days_sum"] * 80_000
        + metrics["two_color_days"] * 35_000
        + len(unscheduled) * 1_000
        + weighted_backlog * 0.2
        - metrics["utilization_pct"] * 250
    )

    return {
        "config": cfg,
        "days": days,
        "unscheduled": unscheduled,
        "metrics": metrics,
        "hard_errors": hard,
        "soft_warnings": soft,
        "data_notes": notes,
        "confidence": confidence,
        "engine": engine,
        "repair_log": repair_log,
        "scenario_score": scenario_score,
        "quality": quality,
    }


def generate_agentic_plan(source: pd.DataFrame, cfg: PlannerConfig) -> Dict[str, Any]:
    t_start = time.perf_counter()
    master = learn_source_master(source, cfg.powder_coeff, cfg.minutes_per_bal)
    lines, quality = build_candidate_lines(source, master, cfg)
    if lines.empty:
        raise ValueError("Aucune ligne éligible à planifier dans la source.")

    jobs, oversized = build_jobs(lines, cfg)
    selected_jobs, outside_jobs = select_candidate_pool(jobs, cfg)
    strategies = ["Délais clients", "Mono-couleur", "Équilibre"] if cfg.strategy == "Auto — meilleur compromis" else [cfg.strategy]
    per_seconds = max(3.0, cfg.solver_seconds / max(1, len(strategies)))

    results = []
    for strategy in strategies:
        scfg = dc_replace(cfg, strategy=strategy, solver_seconds=per_seconds)
        t0 = time.perf_counter()
        r = _assemble(lines, quality, selected_jobs, outside_jobs, oversized, scfg)
        r["elapsed_s"] = round(time.perf_counter() - t0, 2)
        results.append(r)

    # Priorité absolue à la validité puis à la politique atelier mono-couleur.
    # Les retards départagent ensuite les propositions de même qualité couleur.
    best = min(results, key=lambda r: (
        len(r["hard_errors"]),
        r["metrics"]["two_color_days"],
        r["metrics"]["overdue_unscheduled"],
        r["metrics"]["late_days_sum"],
        r["scenario_score"],
    ))

    selected_strategy = best["config"].strategy
    scenario_table = pd.DataFrame([
        {
            "Scénario": r["config"].strategy,
            "Moteur": r["engine"],
            "Confiance règles %": r["confidence"],
            "Charge h": r["metrics"]["total_load_h"],
            "Utilisation %": r["metrics"]["utilization_pct"],
            "Jours mono-couleur": r["metrics"]["mono_color_days"],
            "Jours 2 couleurs": r["metrics"]["two_color_days"],
            "Retards backlog": r["metrics"]["overdue_unscheduled"],
            "Retard planifié (jours)": r["metrics"]["late_days_sum"],
            "Backlog": len(r["unscheduled"]),
            "Temps s": r["elapsed_s"],
        }
        for r in results
    ])
    best["selected_strategy"] = selected_strategy
    best["scenario_table"] = scenario_table
    best["config"] = cfg
    best["master"] = master
    best["lines"] = lines
    best["total_elapsed_s"] = round(time.perf_counter() - t_start, 2)

    best["steps"] = [
        ("Agent Données", "OK", f"{len(source):,} lignes lues depuis GitHub; {len(lines):,} lignes éligibles."),
        ("Agent Référentiel", "OK", f"{master.source_weight_samples:,} poids source appris; cadence {master.minutes_per_bal:.1f} min/bal; poudre {master.powder_coeff:.3f}."),
        ("Agent Quantités", "OK", "Lancement, poids, poudre, balancelles et temps recalculés automatiquement."),
        ("Agent Priorités", "OK", "Délais, retards, ancienneté, OF et disponibilité matière scorés."),
        ("Agent Scénarios", "OK", f"{len(results)} scénario(s) comparé(s); {best['selected_strategy']} retenu."),
        ("Agent Planificateur", "OK", f"{best['engine']} · contrainte dure: maximum 2 couleurs/jour."),
        ("Agent Couleurs", "OK", f"{best['metrics']['mono_color_days']} jour(s) mono-couleur; {best['metrics']['two_color_days']} jour(s) à 2 couleurs."),
        ("Agent Réparation", "OK", f"{len(best['repair_log'])} correction(s) autonome(s)."),
        ("Agent Validation", "OK" if best["confidence"] == 100 else "ERREUR", f"Confiance règles {best['confidence']}% · {len(best['hard_errors'])} erreur(s) dure(s)."),
    ]
    return best


# =============================================================================
# 12) EXPLICATIONS
# =============================================================================
def explanation_table(result: Dict[str, Any]) -> pd.DataFrame:
    rows = []
    week_dates = iso_week_dates(result["config"].year, result["config"].week)
    for d in range(6):
        df = result["days"][d]
        if df.empty:
            continue
        for _, r in df.iterrows():
            due = parse_date(r.get("_due"))
            rows.append({
                "Jour": DAYS[d].title(),
                "Date": week_dates[d].strftime("%d/%m/%Y"),
                "Commande": r["NumCommande"],
                "Client": r["NomClient"],
                "Couleur": r["Couleur"],
                "Article": r["Article/int"],
                "Échéance": due.strftime("%d/%m/%Y") if due is not None else "—",
                "Score IA": round(to_float(r.get("_score")), 1),
                "Raison": norm_text(r.get("_reason")),
            })
    return pd.DataFrame(rows)


# =============================================================================
# 13) EXPORT EXCEL
# =============================================================================
def export_planning_excel(result: Dict[str, Any]) -> bytes:
    cfg: PlannerConfig = result["config"]
    out = io.BytesIO()
    wb = Workbook()
    wb.remove(wb.active)

    navy = "163A5F"
    blue = "155EEF"
    light = "EEF4FF"
    green = "ECFDF3"
    amber = "FFFAEB"
    red = "FEF3F2"
    white = "FFFFFF"
    line = Side(style="thin", color="D0D5DD")

    # Résumé IA
    ws = wb.create_sheet("Résumé IA")
    ws["A1"] = "ALLUCO — Planning Laquage Agentic IA"
    ws["A1"].font = Font(size=16, bold=True, color=white)
    ws["A1"].fill = PatternFill("solid", fgColor=navy)
    ws.merge_cells("A1:F1")
    summary = [
        ("Semaine", f"S{cfg.week} / {cfg.year}"),
        ("Scénario retenu", result.get("selected_strategy", cfg.strategy)),
        ("Moteur", result["engine"]),
        ("Confiance règles", f"{result['confidence']}%"),
        ("Politique couleur", "1 couleur/jour privilégiée · 2 maximum"),
        ("Charge semaine", f"{result['metrics']['total_load_h']:.2f} h"),
        ("Capacité", f"{result['metrics']['capacity_h']:.2f} h"),
        ("Utilisation", f"{result['metrics']['utilization_pct']:.1f}%"),
        ("Jours mono-couleur", result['metrics']['mono_color_days']),
        ("Jours à 2 couleurs", result['metrics']['two_color_days']),
        ("Backlog", len(result["unscheduled"])),
        ("Retards backlog", result['metrics']['overdue_unscheduled']),
    ]
    for i, (k, v) in enumerate(summary, start=3):
        ws.cell(i, 1, k).font = Font(bold=True, color=navy)
        ws.cell(i, 2, v)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 34

    # Contrôles
    row = 17
    ws.cell(row, 1, "Contrôles").font = Font(bold=True, color=navy, size=12)
    row += 1
    if result["hard_errors"]:
        for msg in result["hard_errors"]:
            ws.cell(row, 1, "ERREUR")
            ws.cell(row, 2, msg)
            ws.cell(row, 1).fill = PatternFill("solid", fgColor=red)
            row += 1
    else:
        ws.cell(row, 1, "OK")
        ws.cell(row, 2, "Toutes les règles du moteur sont validées.")
        ws.cell(row, 1).fill = PatternFill("solid", fgColor=green)
        row += 1
    for msg in result["soft_warnings"] + result["data_notes"]:
        ws.cell(row, 1, "INFO")
        ws.cell(row, 2, msg)
        ws.cell(row, 1).fill = PatternFill("solid", fgColor=amber)
        row += 1

    numeric_int_cols = {"QteCommandé", "ResteALivrer", "Prelevé", "QteCommencé", "QteRestante", "QteRèçu", "ReserverBR", "StockPhysique", "Reserver", "Lancement", "Re-laquage", "Barre/bal", "Nbre Bal", "Stock brut"}
    numeric_dec_cols = {"Nuance", "PoidsUn", "PoidsT", "Poudre", "tps"}

    for d, sheet_name in enumerate(SHEET_NAMES):
        ws = wb.create_sheet(sheet_name)
        df = business_day_df(result["days"][d])
        for ci, col in enumerate(OUTPUT_COLUMNS, 1):
            cell = ws.cell(1, ci, col)
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(color=white, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=line)
        previous_color = None
        for ri, (_, r) in enumerate(df.iterrows(), start=2):
            current_color = norm_text(r.get("Couleur"))
            change = previous_color is not None and current_color != previous_color
            for ci, col in enumerate(OUTPUT_COLUMNS, 1):
                v = r.get(col)
                if isinstance(v, pd.Timestamp):
                    v = v.to_pydatetime()
                cell = ws.cell(ri, ci, v)
                cell.border = Border(bottom=Side(style="hair", color="EAECF0"))
                if change:
                    cell.fill = PatternFill("solid", fgColor=light)
                if col in numeric_int_cols and v is not None:
                    cell.number_format = "0"
                elif col in numeric_dec_cols and v is not None:
                    cell.number_format = "0.000"
                elif col == "DateCréation" and isinstance(v, datetime):
                    cell.number_format = "dd/mm/yyyy"
            previous_color = current_color
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(OUTPUT_COLUMNS))}{max(1, len(df)+1)}"
        widths = {1: 16, 2: 13, 3: 23, 4: 25, 5: 18, 6: 12, 7: 10, 12: 14, 13: 16}
        for ci in range(1, len(OUTPUT_COLUMNS) + 1):
            ws.column_dimensions[get_column_letter(ci)].width = widths.get(ci, 13)

    # Backlog
    ws = wb.create_sheet("Backlog")
    backlog_cols = ["NumCommande", "NomClient", "Article", "Couleur", "ResteALivrer", "NumOF", "ProdStatut", "tps", "_overdue_days", "_score", "_reason"]
    backlog = result["unscheduled"].copy()
    cols = [c for c in backlog_cols if c in backlog.columns]
    for ci, c in enumerate(cols, 1):
        ws.cell(1, ci, c).fill = PatternFill("solid", fgColor=navy)
        ws.cell(1, ci).font = Font(color=white, bold=True)
    if not backlog.empty:
        backlog = backlog.sort_values("_score", ascending=False)
        for ri, (_, r) in enumerate(backlog[cols].iterrows(), start=2):
            for ci, c in enumerate(cols, 1):
                ws.cell(ri, ci, r.get(c))
    ws.freeze_panes = "A2"

    # Explications IA
    expl = explanation_table(result)
    ws = wb.create_sheet("Décisions IA")
    if not expl.empty:
        for ci, c in enumerate(expl.columns, 1):
            ws.cell(1, ci, c).fill = PatternFill("solid", fgColor=blue)
            ws.cell(1, ci).font = Font(color=white, bold=True)
        for ri, (_, r) in enumerate(expl.iterrows(), start=2):
            for ci, c in enumerate(expl.columns, 1):
                ws.cell(ri, ci, r.get(c))
        ws.freeze_panes = "A2"
        for ci in range(1, len(expl.columns) + 1):
            ws.column_dimensions[get_column_letter(ci)].width = 18 if ci != len(expl.columns) else 50

    wb.save(out)
    return out.getvalue()


# =============================================================================
# 14) UI / DESIGN
# =============================================================================
CSS = r"""
<style>
:root{--bg:#F5F7FB;--card:#FFFFFF;--ink:#172033;--muted:#667085;--line:#E4E9F0;--blue:#155EEF;--navy:#14365A;--green:#067647;--amber:#B54708;--red:#B42318}
html,body,.stApp,[data-testid="stAppViewContainer"]{background:var(--bg)!important;color:var(--ink)!important;color-scheme:light!important}
header[data-testid="stHeader"]{background:rgba(245,247,251,.94)!important;box-shadow:none!important}
/* Masquer uniquement les actions Streamlit à droite du header.
   IMPORTANT : on ne masque PAS tout le header / stToolbar afin de garder
   le bouton natif qui permet de fermer et rouvrir la sidebar. */
#MainMenu,
footer,
[data-testid="stAppDeployButton"],
[data-testid="stHeaderActionElements"],
[data-testid="stToolbarActions"],
.stToolbarActions,
[data-testid="stToolbarActionButton"],
.stToolbarActionButton,
[data-testid="stMainMenu"],
[data-testid="stStatusWidget"] {
    display: none !important;
    visibility: hidden !important;
}

/* Share */
header button[title="Share"],
header button[aria-label="Share"],
header button[title*="share" i],
header button[aria-label*="share" i] {
    display: none !important;
}

/* Etoile / favori */
header button[title*="favorite" i],
header button[aria-label*="favorite" i],
header button[title*="favourite" i],
header button[aria-label*="favourite" i],
header button[title*="star" i],
header button[aria-label*="star" i] {
    display: none !important;
}

/* Edit */
header button[title="Edit"],
header button[aria-label="Edit"],
header button[title*="edit" i],
header button[aria-label*="edit" i] {
    display: none !important;
}

/* GitHub */
header a[href*="github.com"],
header a[title*="github" i],
header a[aria-label*="github" i],
header button[title*="github" i],
header button[aria-label*="github" i] {
    display: none !important;
}

/* Menu trois points */
header [data-testid="stMainMenu"] {
    display: none !important;
}
.block-container{max-width:1520px;padding-top:1rem;padding-bottom:2rem}
section[data-testid="stSidebar"],section[data-testid="stSidebar"]>div{background:#FFF!important;border-right:1px solid var(--line)}
h1,h2,h3,h4,h5,h6,p,label,span,div{color:var(--ink)}
[data-testid="stCaptionContainer"],.stCaption{color:var(--muted)!important}
.brand{display:flex;align-items:center;gap:.75rem;margin:.2rem 0 1rem}.brandmark{width:44px;height:44px;border-radius:13px;background:linear-gradient(145deg,#123B67,#155EEF);color:#fff!important;display:flex;align-items:center;justify-content:center;font-weight:950;font-size:1.4rem;box-shadow:0 8px 20px rgba(21,94,239,.2)}.brandname{font-weight:950;font-size:1.05rem;letter-spacing:.06em}.brandsub{font-size:.72rem;color:var(--muted)!important}
.topbar{display:flex;justify-content:space-between;align-items:center;margin:.2rem 0 1rem}.title{font-size:1.55rem;font-weight:950;letter-spacing:-.025em}.subtitle{font-size:.86rem;color:var(--muted)!important;margin-top:.15rem}.weekbadge{background:#EEF4FF;color:#155EEF!important;border:1px solid #C7D7FE;padding:.42rem .72rem;border-radius:999px;font-weight:850;font-size:.78rem}
.hero{background:linear-gradient(135deg,#123B67 0%,#155EEF 100%);border-radius:20px;padding:1.3rem 1.45rem;margin-bottom:1rem;box-shadow:0 12px 30px rgba(21,94,239,.12)}.hero *{color:#fff!important}.hero-title{font-weight:950;font-size:1.28rem}.hero-sub{opacity:.9;font-size:.89rem;margin-top:.35rem;max-width:1050px}
.card{background:#FFF;border:1px solid var(--line);border-radius:16px;padding:1rem 1.1rem;box-shadow:0 1px 3px rgba(16,24,40,.035)}
.kpi{background:#FFF;border:1px solid var(--line);border-radius:15px;padding:.86rem 1rem;min-height:98px}.kpi-l{font-size:.70rem;font-weight:850;color:var(--muted)!important;text-transform:uppercase;letter-spacing:.05em}.kpi-v{font-size:1.42rem;font-weight:950;margin-top:.25rem}.kpi-s{font-size:.72rem;color:var(--muted)!important;margin-top:.15rem}
.status-ok{display:inline-flex;align-items:center;gap:.35rem;color:#067647!important;background:#ECFDF3;border:1px solid #ABEFC6;padding:.34rem .58rem;border-radius:999px;font-size:.75rem;font-weight:850}.status-warn{display:inline-flex;align-items:center;gap:.35rem;color:#B54708!important;background:#FFFAEB;border:1px solid #FEDF89;padding:.34rem .58rem;border-radius:999px;font-size:.75rem;font-weight:850}.status-err{display:inline-flex;align-items:center;gap:.35rem;color:#B42318!important;background:#FEF3F2;border:1px solid #FECDCA;padding:.34rem .58rem;border-radius:999px;font-size:.75rem;font-weight:850}
.agent{background:#FFF;border:1px solid var(--line);border-radius:12px;padding:.68rem .8rem;margin:.35rem 0}.agent-ok{border-left:4px solid #12B76A}.agent-warn{border-left:4px solid #F79009}.agent b{font-size:.83rem}.agent small{color:var(--muted)!important}
[data-testid="stMetric"]{background:#FFF;border:1px solid var(--line);border-radius:14px;padding:.72rem .9rem}.stButton>button{border-radius:10px;font-weight:800;border:1px solid #D5DCE6;min-height:42px}.stButton>button[kind="primary"]{background:#155EEF;border-color:#155EEF;color:#fff!important}.stButton>button[kind="primary"] *{color:#fff!important}
[data-testid="stDataFrame"],[data-testid="stDataEditor"]{background:#FFF;border:1px solid var(--line);border-radius:14px;overflow:hidden}button[data-baseweb="tab"]{font-weight:800}button[data-baseweb="tab"][aria-selected="true"]{color:#155EEF!important}.stDownloadButton button{border-radius:10px;font-weight:850}
@media(max-width:900px){.block-container{padding-left:.75rem;padding-right:.75rem}.topbar{align-items:flex-start;flex-direction:column;gap:.55rem}.hero{padding:1rem}}
</style>
"""


def _brand() -> None:
    st.markdown("<div class='brand'><div class='brandmark'>A</div><div><div class='brandname'>ALLUCO</div><div class='brandsub'>Planning Laquage Agentic IA</div></div></div>", unsafe_allow_html=True)


def _top(title: str, subtitle: str, cfg: PlannerConfig) -> None:
    st.markdown(f"<div class='topbar'><div><div class='title'>{title}</div><div class='subtitle'>{subtitle}</div></div><div class='weekbadge'>S{cfg.week} · {cfg.year}</div></div>", unsafe_allow_html=True)


def _kpi(label: str, value: str, sub: str = "") -> None:
    st.markdown(f"<div class='kpi'><div class='kpi-l'>{label}</div><div class='kpi-v'>{value}</div><div class='kpi-s'>{sub}</div></div>", unsafe_allow_html=True)


def _parse_commands(text: str) -> Tuple[str, ...]:
    return tuple(dict.fromkeys(x.strip().upper() for x in re.split(r"[,;\n]+", text or "") if x.strip()))


def _source_signature(path: Path, cfg: PlannerConfig) -> str:
    payload = [str(path.resolve()), str(path.stat().st_mtime_ns), str(path.stat().st_size), json.dumps(cfg.__dict__, sort_keys=True, default=str)]
    return hashlib.sha256("|".join(payload).encode("utf-8")).hexdigest()


if st is not None:
    @st.cache_data(show_spinner=False)
    def _cached_source(path: str, mtime_ns: int, size: int) -> pd.DataFrame:
        return load_source_workbook(Path(path).read_bytes())


def render_planning(result: Dict[str, Any], cfg: PlannerConfig) -> None:
    m = result["metrics"]
    hero_msg = "Le moteur choisit automatiquement les commandes, privilégie une seule couleur par jour et n'autorise jamais plus de deux couleurs. Il compare plusieurs stratégies, répare la proposition puis la valide avant affichage."
    st.markdown(f"<div class='hero'><div class='hero-title'>Proposition IA recommandée · {result['selected_strategy']}</div><div class='hero-sub'>{hero_msg}</div></div>", unsafe_allow_html=True)

    cols = st.columns(6)
    values = [
        ("Confiance règles", f"{result['confidence']}%", "validation déterministe"),
        ("Charge", f"{m['total_load_h']:.1f} h", f"sur {m['capacity_h']:.1f} h"),
        ("Utilisation", f"{m['utilization_pct']:.0f}%", "semaine"),
        ("Mono-couleur", str(m["mono_color_days"]), "jour(s)"),
        ("2 couleurs", str(m["two_color_days"]), "maximum autorisé"),
        ("Backlog", str(len(result["unscheduled"])), "ligne(s) hors semaine"),
    ]
    for col, (a, b, s) in zip(cols, values):
        with col:
            _kpi(a, b, s)

    st.caption("Confiance 100% = toutes les règles codées sont validées. La qualité du résultat dépend aussi de la qualité des données du fichier source.")
    st.write("")

    left, right = st.columns([3, 2])
    with left:
        chart = pd.DataFrame({
            "Jour": [x["Jour"].title() for x in m["days"]],
            "Charge h": [x["Charge totale h"] for x in m["days"]],
            "Capacité h": [x["Capacité h"] for x in m["days"]],
        }).set_index("Jour")
        st.markdown("#### Charge par jour")
        st.bar_chart(chart)
    with right:
        st.markdown("#### Validation IA")
        if result["confidence"] == 100:
            st.markdown("<span class='status-ok'>● 100% des règles validées</span>", unsafe_allow_html=True)
        else:
            st.markdown("<span class='status-err'>● Contrôle bloquant détecté</span>", unsafe_allow_html=True)
        st.write("")
        st.caption(f"Moteur: {result['engine']}")
        st.caption("Politique couleur: 1 couleur privilégiée · 2 maximum")
        st.caption(f"Temps calcul: {result['total_elapsed_s']:.2f} s")
        st.caption(f"Retards backlog: {m['overdue_unscheduled']}")

    if result["hard_errors"]:
        for msg in result["hard_errors"]:
            st.error(msg)
    if result["soft_warnings"]:
        with st.expander("Informations planning"):
            for msg in result["soft_warnings"]:
                st.write("• " + msg)
    if result["data_notes"]:
        with st.expander("Qualité / estimations de données"):
            for msg in result["data_notes"]:
                st.write("• " + msg)

    st.markdown("#### Planning détaillé")
    tabs = st.tabs([f"{DAYS[d].title()} · {m['days'][d]['Date'][:5]}" for d in range(6)])
    for d, tab in enumerate(tabs):
        with tab:
            dm = m["days"][d]
            a, b, c, e = st.columns(4)
            a.metric("Charge", f"{dm['Charge totale h']:.2f} h", f"cap. {dm['Capacité h']:.1f} h")
            b.metric("Utilisation", f"{dm['Charge %']:.0f}%")
            c.metric("Couleurs", dm["Nb couleurs"])
            e.metric("Lignes", dm["Lignes"])
            if dm["Couleurs"]:
                st.caption("Séquence: " + dm["Couleurs"])
            st.dataframe(business_day_df(result["days"][d]), hide_index=True, use_container_width=True, height=490)

    st.write("")
    excel = export_planning_excel(result)
    st.download_button(
        "⬇ Télécharger le planning Excel",
        data=excel,
        file_name=f"Planning_IA_S{cfg.week}_{cfg.year}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
        use_container_width=True,
    )


def render_ui() -> None:
    if st is None:
        raise RuntimeError("Streamlit n'est pas installé. Lancez: pip install -r requirements.txt")

    st.set_page_config(page_title=APP_NAME, page_icon="🤖", layout="wide", initial_sidebar_state="expanded")
    st.markdown(CSS, unsafe_allow_html=True)

    if not SOURCE_PATH.exists():
        cfg = PlannerConfig(DEFAULT_YEAR, DEFAULT_WEEK)
        _top("Fichier source absent", "Le planning ne peut pas être calculé.", cfg)
        st.error(f"Ajoutez `{SOURCE_FILENAME}` à côté de `app.py` dans GitHub.")
        return

    today = date.today().isocalendar()
    with st.sidebar:
        _brand()
        nav = st.radio("Navigation", ["🤖 Planning IA", "🏠 Tableau de bord", "🧠 Analyse IA", "⚙️ Paramètres"], label_visibility="collapsed")
        st.divider()
        year = int(st.number_input("Année", 2024, 2035, DEFAULT_YEAR or int(today.year), 1))
        week = int(st.number_input("Semaine", 1, 53, DEFAULT_WEEK or int(today.week), 1))
        objective = st.selectbox("Objectif", ["Auto — meilleur compromis", "Délais clients", "Mono-couleur", "Équilibre"])
        with st.expander("Réglages production"):
            cap = float(st.number_input("Capacité Lun–Ven (h)", 1.0, 24.0, DEFAULT_CAPACITY_H, 0.5))
            sat = st.checkbox("Production samedi", value=DEFAULT_SATURDAY_ENABLED)
            sat_cap = float(st.number_input("Capacité samedi (h)", 0.0, 24.0, DEFAULT_SATURDAY_CAPACITY_H, 0.5, disabled=not sat))
            cleaning = int(st.number_input("Nettoyage si 2e couleur (min)", 0, 120, DEFAULT_CLEANING_MIN, 5))
            solver = float(st.slider("Budget optimisation IA (s)", 6, 60, int(DEFAULT_SOLVER_SECONDS), 3))
            st.info("Règle fixe: 1 couleur/jour privilégiée, 2 maximum.")
        with st.expander("Contraintes manuelles"):
            force_text = st.text_area("Forcer commandes cette semaine", placeholder="VTE2601234, VTE2605678")
            exclude_text = st.text_area("Exclure commandes", placeholder="VTE2609999")
        st.divider()
        st.markdown("<span class='status-ok'>● Base GitHub connectée</span>", unsafe_allow_html=True)
        st.caption(SOURCE_FILENAME)
        st.caption("OR-Tools: " + ("actif" if ORTOOLS_AVAILABLE else "fallback local"))
        st.caption(f"Version {VERSION}")

    cfg = PlannerConfig(
        year=year, week=week, capacity_h=cap,
        saturday_enabled=sat, saturday_capacity_h=sat_cap,
        cleaning_min=cleaning, minutes_per_bal=DEFAULT_MIN_PER_BAL,
        powder_coeff=DEFAULT_POWDER_COEFF, target_utilization=DEFAULT_TARGET_UTIL,
        solver_seconds=solver, max_jobs=DEFAULT_MAX_JOBS, pool_factor=DEFAULT_POOL_FACTOR,
        allow_relaquage=DEFAULT_ALLOW_RELAQUAGE, strategy=objective,
        force_commands=_parse_commands(force_text), exclude_commands=_parse_commands(exclude_text),
    )

    try:
        source = _cached_source(str(SOURCE_PATH), SOURCE_PATH.stat().st_mtime_ns, SOURCE_PATH.stat().st_size)
    except Exception as exc:
        _top("Erreur de lecture", "Le moteur a arrêté le calcul pour protéger le planning.", cfg)
        st.exception(exc)
        return

    signature = _source_signature(SOURCE_PATH, cfg)

    def ensure_plan(force: bool = False):
        if force or st.session_state.get("plan_signature") != signature or "plan_result" not in st.session_state:
            with st.spinner("Agents IA: données → priorités → scénarios → optimisation → mono-couleur → réparation → validation..."):
                st.session_state["plan_result"] = generate_agentic_plan(source, cfg)
                st.session_state["plan_signature"] = signature
        return st.session_state.get("plan_result")

    if nav == "🤖 Planning IA":
        _top("Planning IA", "Le fichier est lu depuis GitHub. Aucun upload manuel.", cfg)
        c1, c2, c3 = st.columns([2, 1, 1])
        with c1:
            st.markdown(f"<div class='card'><b>Source production</b><br><span class='subtitle'>{SOURCE_FILENAME}</span><br><span class='subtitle'>{format_num(len(source))} lignes détectées</span></div>", unsafe_allow_html=True)
        with c2:
            if st.button("↻ Régénérer", use_container_width=True):
                ensure_plan(force=True)
        with c3:
            st.markdown("<div class='card'><b>Couleurs / jour</b><br><span class='subtitle'>1 privilégiée</span><br><span class='subtitle'>2 maximum</span></div>", unsafe_allow_html=True)

        result = ensure_plan(False) if DEFAULT_AUTO_GENERATE else st.session_state.get("plan_result")
        if result:
            render_planning(result, cfg)

    elif nav == "🏠 Tableau de bord":
        _top("Tableau de bord", "Vue rapide des commandes et de la capacité de la semaine.", cfg)
        master = learn_source_master(source, cfg.powder_coeff, cfg.minutes_per_bal)
        lines, quality = build_candidate_lines(source, master, cfg)
        overdue = int((pd.to_numeric(lines.get("_overdue_days"), errors="coerce").fillna(0) > 0).sum()) if not lines.empty else 0
        unique_colors = int(lines["Couleur"].nunique()) if not lines.empty else 0
        cols = st.columns(5)
        vals = [
            ("Commandes", format_num(source["NumCommande"].nunique()), "source GitHub"),
            ("Lignes éligibles", format_num(len(lines)), "à planifier"),
            ("Retards", str(overdue), "au début de semaine"),
            ("Couleurs", str(unique_colors), "backlog éligible"),
            ("Capacité", f"{sum(day_capacity_h(cfg,d) for d in range(6)):.0f} h", "semaine"),
        ]
        for col, val in zip(cols, vals):
            with col:
                _kpi(*val)
        result = ensure_plan(False)
        if result:
            st.write("")
            render_planning(result, cfg)

    elif nav == "🧠 Analyse IA":
        _top("Analyse IA", "Comparaison des stratégies, décisions des agents et backlog.", cfg)
        result = ensure_plan(False)
        st.markdown("#### Comparaison des scénarios")
        st.dataframe(result["scenario_table"], hide_index=True, use_container_width=True)
        left, right = st.columns(2)
        with left:
            st.markdown("#### Rapport des agents")
            for agent, status, msg in result["steps"]:
                cls = "agent agent-ok" if status == "OK" else "agent agent-warn"
                st.markdown(f"<div class='{cls}'><b>{agent} · {status}</b><br><small>{msg}</small></div>", unsafe_allow_html=True)
        with right:
            st.markdown("#### Réparations autonomes")
            if result["repair_log"]:
                for msg in result["repair_log"]:
                    st.write("• " + msg)
            else:
                st.success("Aucune réparation supplémentaire nécessaire.")
            if result["data_notes"]:
                with st.expander("Qualité des données"):
                    for msg in result["data_notes"]:
                        st.write("• " + msg)

        st.markdown(f"#### Backlog hors semaine — {len(result['unscheduled'])} ligne(s)")
        if result["unscheduled"].empty:
            st.success("Tout le pool prioritaire tient dans la semaine.")
        else:
            cols = [c for c in ["NumCommande", "NomClient", "Article", "Couleur", "ResteALivrer", "NumOF", "ProdStatut", "tps", "_overdue_days", "_score", "_reason"] if c in result["unscheduled"].columns]
            show = result["unscheduled"][cols].sort_values("_score", ascending=False).head(1000).rename(columns={"_overdue_days": "Retard jours", "_score": "Score IA", "_reason": "Raison IA"})
            st.dataframe(show, hide_index=True, use_container_width=True, height=520)

        with st.expander("Pourquoi les commandes sont placées ainsi ?"):
            expl = explanation_table(result)
            st.dataframe(expl.head(1500), hide_index=True, use_container_width=True, height=520)

    else:
        _top("Paramètres", "Règles transparentes utilisées par le moteur.", cfg)
        rules = pd.DataFrame([
            ["Fichier source", SOURCE_FILENAME],
            ["Historique planning", "Aucun fichier historique requis"],
            ["Couleurs / jour", "1 privilégiée · 2 maximum (dur)"],
            ["Capacité Lun–Ven", f"{cfg.capacity_h:.1f} h/j"],
            ["Samedi", f"{'Actif' if cfg.saturday_enabled else 'Inactif'} · {cfg.saturday_capacity_h:.1f} h"],
            ["Cadence", f"{cfg.minutes_per_bal:.1f} min/bal"],
            ["Poudre", f"coefficient {cfg.powder_coeff:.3f}"],
            ["Nettoyage", f"{cfg.cleaning_min} min uniquement lorsqu'un jour contient 2 couleurs"],
            ["Optimisation", "3 scénarios + CP-SAT OR-Tools + réparation agentique"],
            ["Confiance 100%", "Toutes les règles du moteur validées; ce n'est pas une garantie de données terrain"],
        ], columns=["Paramètre", "Valeur"])
        st.dataframe(rules, hide_index=True, use_container_width=True)
        st.info("Pour mettre à jour les commandes, remplacez Bd-Client-S36.xlsx dans GitHub puis redeployez. L'utilisateur n'importe aucun fichier dans l'application.")


# =============================================================================
# 15) CLI / TESTS AUTOMATIQUES
# =============================================================================
def cli_generate(source_path: str, output_path: str, year: int, week: int) -> Dict[str, Any]:
    source = load_source_workbook(Path(source_path).read_bytes())
    cfg = PlannerConfig(year=year, week=week, strategy="Auto — meilleur compromis", solver_seconds=8)
    result = generate_agentic_plan(source, cfg)
    Path(output_path).write_bytes(export_planning_excel(result))
    return result


def self_test(source_path: Optional[str] = None) -> None:
    checks: List[str] = []

    def check(name: str, condition: bool, detail: Any = None):
        if not condition:
            raise AssertionError(f"{name}: {detail}")
        checks.append(name)
        print(f"[OK] {name}")

    check("Article split", split_article("LMMO-S758-BLC") == ("LMMO-S758", "BLC"))
    check("ISO semaine", iso_week_dates(2026, 36)[0] == date(2026, 8, 31))
    check("Max couleur constant", HARD_MAX_COLORS_PER_DAY == 2)
    check("Préférence mono-couleur", PREFERRED_COLORS_PER_DAY == 1)

    if source_path:
        src = load_source_workbook(Path(source_path).read_bytes())
        cfg = PlannerConfig(2026, 36, strategy="Auto — meilleur compromis", solver_seconds=6, max_jobs=900, pool_factor=2.0)
        master = learn_source_master(src, cfg.powder_coeff, cfg.minutes_per_bal)
        lines, quality = build_candidate_lines(src, master, cfg)
        check("Source réelle lue", len(src) > 100, len(src))
        check("Lignes éligibles", not lines.empty)
        check("Apprentissage poids source", master.source_weight_samples > 100, master.source_weight_samples)
        result = generate_agentic_plan(src, cfg)
        check("Planning non vide", sum(len(x) for x in result["days"].values()) > 0)
        check("Confiance règles 100%", result["confidence"] == 100, result["hard_errors"])
        check("Capacité respectée", all(d["Charge totale h"] <= d["Capacité h"] + 1e-6 for d in result["metrics"]["days"] if d["Capacité h"] > 0), result["metrics"]["days"])
        check("Deux couleurs maximum", all(d["Nb couleurs"] <= 2 for d in result["metrics"]["days"]), result["metrics"]["days"])
        check("Format 28 colonnes", all(list(business_day_df(result["days"][d]).columns) == OUTPUT_COLUMNS for d in range(6)))
        planned_frames = [d for d in result["days"].values() if not d.empty]
        planned = pd.concat(planned_frames, ignore_index=True) if planned_frames else pd.DataFrame()
        check("Aucun doublon", planned.empty or not planned["_line_id"].duplicated().any())
        out = export_planning_excel(result)
        check("Export Excel", len(out) > 7000, len(out))
        wb = load_workbook(io.BytesIO(out), read_only=True, data_only=True)
        check("6 feuilles planning", all(s in wb.sheetnames for s in SHEET_NAMES), wb.sheetnames)
        check("Aucun historique requis", not (ROOT_DIR / "Planning S36.xlsx").exists())

    print(f"\n{len(checks)} test(s) OK")


if __name__ == "__main__":
    if "--self-test" in sys.argv:
        src = sys.argv[sys.argv.index("--input") + 1] if "--input" in sys.argv else None
        self_test(src)
    elif "--generate" in sys.argv:
        src = sys.argv[sys.argv.index("--input") + 1]
        out = sys.argv[sys.argv.index("--output") + 1]
        year = int(sys.argv[sys.argv.index("--year") + 1])
        week = int(sys.argv[sys.argv.index("--week") + 1])
        r = cli_generate(src, out, year, week)
        print("Planning généré:", out)
        print("Moteur:", r["engine"], "| confiance règles:", r["confidence"], "%")
    else:
        if st is None:
            print("Installez les dépendances: pip install -r requirements.txt")
        else:
            render_ui()
