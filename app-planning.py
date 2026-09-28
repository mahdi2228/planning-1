# -*- coding: utf-8 -*-
"""
ALLUCO - Planning Laquage Agentic IA (Final Core)
=================================================

Objectif
--------
Transformer une base de preparation finale en planning de laquage deterministe,
avec validation stricte et audit de reference optionnel.

Principes de fiabilite
----------------------
1) Toujours lire la feuille finale quand elle existe:
   - "preparation pour planning VF" (Base 1)
   - "version final" (Base 2)
2) Conserver les donnees metier deja preparees (Lancement, Barre/bal, Nbre Bal).
3) Cadence atelier par defaut: 4 min / balancelle.
4) Le coeur de planification est deterministe (OR-Tools CP-SAT si disponible,
   heuristique deterministe sinon). Un LLM n'est jamais autorise a contourner
   les contraintes du moteur.
5) Separer clairement:
   - validation des regles du moteur;
   - conformite a un planning de reference.
   "Regles OK" ne veut jamais dire "realite terrain garantie".
6) Si un planning de reference contient des lignes absentes de l'input, la
   conformite 100% est impossible: l'application le signale explicitement.

Execution
---------
Streamlit:
    streamlit run alluco_planner_final.py

CLI:
    python alluco_planner_final.py --generate --input "Base 2.xlsx" \
        --output planning.xlsx --year 2026 --week 40 --spillover 1

Audit historique:
    python alluco_planner_final.py --generate --input "Base 1.xlsx" \
        --output planning_s38.xlsx --year 2026 --week 38 \
        --reference "Planning S38.xlsx" --spillover 0
"""
from __future__ import annotations

import argparse
import hashlib
import io
import math
import re
import unicodedata
from collections import Counter, OrderedDict, defaultdict, deque
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:
    from ortools.sat.python import cp_model
    ORTOOLS_AVAILABLE = True
except Exception:
    cp_model = None
    ORTOOLS_AVAILABLE = False

try:
    import streamlit as st
except Exception:
    st = None

VERSION = "7.0.0-final-safe"
APP_NAME = "ALLUCO - Planning Laquage Agentic IA"

OUTPUT_COLUMNS = [
    "NumCommande", "DateCreation", "NomClient", "Article", "Article/int", "Couleur", "Nuance",
    "QteCommandee", "ResteALivrer", "Preleve", "reservation brut", "NumOF", "ProdStatut",
    "QteCommence", "QteRestante", "QteRecu", "ReserverBR", "StockPhysique", "Reserver",
    "Lancement", "Re-laquage", "PoidsUn", "PoidsT", "Poudre", "Barre/bal", "Nbre Bal", "tps",
    "Stock brut", "moyenne vente", "% laque",
]

# Libelles d'export proches des fichiers atelier.
EXPORT_HEADERS = {
    "NumCommande": "Num Commande", "DateCreation": "DateCreation", "NomClient": "Nom Client",
    "Article": "Article", "Article/int": "Article/int", "Couleur": "Couleur", "Nuance": "Nuance",
    "QteCommandee": "Qte Commandee", "ResteALivrer": "Reste A Livrer", "Preleve": "Preleve",
    "reservation brut": "reservation brut", "NumOF": "Num OF", "ProdStatut": "Prod Statut",
    "QteCommence": "Qte Commencee", "QteRestante": "QteRestante", "QteRecu": "QteRecu",
    "ReserverBR": "ReserverBR", "StockPhysique": "StockPhysique", "Reserver": "Reserver",
    "Lancement": "Lancement", "Re-laquage": "Re-laquage", "PoidsUn": "PoidsUn", "PoidsT": "PoidsT",
    "Poudre": "Poudre", "Barre/bal": "Barre/bal", "Nbre Bal": "Nbre Bal", "tps": "tps",
    "Stock brut": "Stock brut", "moyenne vente": "moyenne vente", "% laque": "% laque",
}

# Normalisation des en-tetes rencontres dans les fichiers fournis.
HEADER_ALIASES = {
    "numcommande": "NumCommande", "num_commande": "NumCommande",
    "datecreation": "DateCreation", "date_creation": "DateCreation",
    "nomclient": "NomClient", "nom_client": "NomClient",
    "article": "Article", "article_int": "Article/int", "articleint": "Article/int",
    "couleur": "Couleur", "nuance": "Nuance",
    "qtecommandee": "QteCommandee", "qte_commandee": "QteCommandee",
    "restealivrer": "ResteALivrer", "reste_a_livrer": "ResteALivrer",
    "preleve": "Preleve", "reservation_brut": "reservation brut",
    "numof": "NumOF", "num_of": "NumOF", "prodstatut": "ProdStatut", "prod_statut": "ProdStatut",
    "qtecommencee": "QteCommence", "qte_commencee": "QteCommence",
    "qterestante": "QteRestante", "qte_restante": "QteRestante",
    "qterecu": "QteRecu", "qte_recu": "QteRecu",
    "reserverbr": "ReserverBR", "stockphysique": "StockPhysique", "reserver": "Reserver",
    "lancement": "Lancement", "re_laquage": "Re-laquage", "relaquage": "Re-laquage",
    "poidsun": "PoidsUn", "poidst": "PoidsT", "poudre": "Poudre",
    "barre_bal": "Barre/bal", "barrebal": "Barre/bal",
    "nbre_bal": "Nbre Bal", "nbrebal": "Nbre Bal", "tps": "tps",
    "stock_brut": "Stock brut", "moyenne_vente": "moyenne vente", "laque": "% laque",
}

FINAL_SHEET_PRIORITY = [
    "preparation_pour_planning_vf",
    "version_final",
    "planning_vf",
    "preparation_planning_vf",
    "feuil1",
]

WHITE_ALIASES = {"BLC", "BLANC", "WHITE", "R9016"}
BLACK_ALIASES = {"NOIR", "DARK", "BLACK", "R9005"}


def norm_text(v: Any) -> str:
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except Exception:
        pass
    return str(v).strip()


def norm_key(v: Any) -> str:
    s = unicodedata.normalize("NFKD", norm_text(v)).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")


def to_float(v: Any, default: float = 0.0) -> float:
    try:
        if v is None:
            return default
        if isinstance(v, str):
            s = v.strip().replace(" ", "").replace(",", ".")
            if not s or s.upper() in {"#N/A", "N/A", "NA", "NONE", "NAN", "#VALUE!", "#REF!"}:
                return default
            return float(s)
        if pd.isna(v):
            return default
        return float(v)
    except Exception:
        return default


def to_int(v: Any, default: int = 0) -> int:
    return int(round(to_float(v, float(default))))


def parse_date(v: Any) -> Optional[pd.Timestamp]:
    if v is None:
        return None
    try:
        if isinstance(v, (int, float, np.integer, np.floating)):
            n = float(v)
            if 20000 <= n <= 80000:
                return pd.Timestamp("1899-12-30") + pd.to_timedelta(n, unit="D")
            return None
        ts = pd.to_datetime(v, errors="coerce", dayfirst=False)
        if pd.isna(ts) or ts.year <= 1900:
            return None
        return pd.Timestamp(ts)
    except Exception:
        return None


def iso_week_monday(year: int, week: int) -> date:
    return date.fromisocalendar(int(year), int(week), 1)


def split_article(article: Any) -> Tuple[str, str]:
    s = norm_text(article)
    if "-" not in s:
        return s, ""
    left, right = s.rsplit("-", 1)
    token = right.strip().upper()
    # Suffixes de couleur observes: noms courts / Rxxxx / Nxx.
    if re.fullmatch(r"R\d{4}", token) or re.fullmatch(r"N\d{2}", token) or re.fullmatch(r"[A-Z][A-Z0-9]{1,14}", token):
        return left.strip(), token
    return s, ""


def _header_name(v: Any) -> str:
    k = norm_key(v)
    return HEADER_ALIASES.get(k, norm_text(v))


def _sheet_headers(ws, max_cols: int = 31) -> List[Any]:
    return [c.value for c in next(ws.iter_rows(min_row=1, max_row=1, max_col=max_cols))]


def _looks_like_prepared(ws) -> bool:
    names = {_header_name(x) for x in _sheet_headers(ws)}
    return {"Article", "Lancement", "Nbre Bal"}.issubset(names)


def choose_source_sheet(wb) -> Tuple[Any, str]:
    by_key = {norm_key(ws.title): ws for ws in wb.worksheets}
    for k in FINAL_SHEET_PRIORITY:
        if k in by_key and _looks_like_prepared(by_key[k]):
            return by_key[k], "final_prepared"
    # Ensuite chercher une vraie feuille finale/prepared, sans se fier a max_column Excel.
    for ws in wb.worksheets:
        if _looks_like_prepared(ws):
            return ws, "prepared_detected"
    # Dernier recours: commandes brutes, mais ce mode est explicitement marque comme moins fiable.
    for ws in wb.worksheets:
        names = {_header_name(x) for x in _sheet_headers(ws)}
        if {"NumCommande", "Article"}.issubset(names):
            return ws, "raw_orders"
    raise ValueError("Aucune feuille exploitable trouvee (Article / NumCommande).")


def _read_compact_rows(ws, max_cols: int = 31, blank_stop: int = 120) -> pd.DataFrame:
    raw_headers = _sheet_headers(ws, max_cols=max_cols)
    headers = [_header_name(x) for x in raw_headers]
    # Eviter les colonnes totalement vides en fin de plage.
    last = 0
    for i, h in enumerate(headers, 1):
        if norm_text(h):
            last = i
    last = max(1, last)
    headers = headers[:last]

    rows: List[List[Any]] = []
    excel_rows: List[int] = []
    seen = False
    blanks = 0
    article_idx = headers.index("Article") if "Article" in headers else None
    cmd_idx = headers.index("NumCommande") if "NumCommande" in headers else None

    for erow, values in enumerate(ws.iter_rows(min_row=2, max_col=last, values_only=True), start=2):
        vals = list(values[:last])
        art = vals[article_idx] if article_idx is not None and article_idx < len(vals) else None
        cmd = vals[cmd_idx] if cmd_idx is not None and cmd_idx < len(vals) else None
        meaningful = bool(norm_text(art) or norm_text(cmd))
        if not meaningful:
            if seen:
                blanks += 1
                if blanks >= blank_stop:
                    break
            continue
        seen = True
        blanks = 0
        rows.append(vals)
        excel_rows.append(erow)

    if not rows:
        raise ValueError(f"Feuille '{ws.title}' vide ou non exploitable.")
    df = pd.DataFrame(rows, columns=headers)
    # Si des en-tetes dupliques existent, garder la premiere occurrence.
    df = df.loc[:, ~df.columns.duplicated()].copy()
    df["_excel_row"] = excel_rows
    return df


def _learn_article_defaults(df: pd.DataFrame) -> Tuple[Dict[str, int], Dict[str, float]]:
    bars: Dict[str, List[int]] = defaultdict(list)
    weights: Dict[str, List[float]] = defaultdict(list)
    for _, r in df.iterrows():
        art = norm_text(r.get("Article/int")).upper()
        if not art:
            art, _ = split_article(r.get("Article"))
            art = art.upper()
        b = to_int(r.get("Barre/bal"), 0)
        w = to_float(r.get("PoidsUn"), 0.0)
        if art and b > 0:
            bars[art].append(b)
        if art and w > 0:
            weights[art].append(w)
    return (
        {k: int(round(float(np.median(v)))) for k, v in bars.items() if v},
        {k: float(np.median(v)) for k, v in weights.items() if v},
    )


def prepare_source(data: bytes, minutes_per_bal: float = 4.0, additions: Optional[bytes] = None) -> Tuple[pd.DataFrame, Dict[str, Any]]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws, mode = choose_source_sheet(wb)
    df = _read_compact_rows(ws)
    df["_source_sheet"] = ws.title
    df["_source_mode"] = mode
    df["_source_order"] = np.arange(len(df), dtype=int)
    df["_origin"] = "base"

    if additions:
        wb2 = load_workbook(io.BytesIO(additions), read_only=True, data_only=True)
        ws2, _ = choose_source_sheet(wb2)
        ad = _read_compact_rows(ws2)
        ad["_source_sheet"] = ws2.title
        ad["_source_mode"] = "additions"
        ad["_source_order"] = np.arange(len(df), len(df) + len(ad), dtype=int)
        ad["_origin"] = "addition"
        df = pd.concat([df, ad], ignore_index=True, sort=False)

    for col in OUTPUT_COLUMNS:
        if col not in df.columns:
            df[col] = None

    # Deriver Article/int et Couleur uniquement s'ils manquent.
    for idx in df.index:
        article = norm_text(df.at[idx, "Article"])
        art_int = norm_text(df.at[idx, "Article/int"])
        color = norm_text(df.at[idx, "Couleur"]).upper()
        if not art_int or not color:
            ai, co = split_article(article)
            if not art_int:
                art_int = ai
            if not color:
                color = co
        df.at[idx, "Article"] = article
        df.at[idx, "Article/int"] = art_int
        df.at[idx, "Couleur"] = color
        df.at[idx, "NumCommande"] = norm_text(df.at[idx, "NumCommande"]).upper()
        df.at[idx, "NumOF"] = norm_text(df.at[idx, "NumOF"]).upper()

    bars_master, weight_master = _learn_article_defaults(df)
    warnings: List[str] = []
    unknown_bars = 0
    recalculated_bales = 0
    tps_inconsistent = 0

    for idx in df.index:
        art = norm_text(df.at[idx, "Article/int"]).upper()
        launch = to_float(df.at[idx, "Lancement"], -1)
        if launch < 0:
            # Base brute: fallback prudent sur reste a livrer.
            launch = max(0.0, to_float(df.at[idx, "ResteALivrer"], to_float(df.at[idx, "QteRestante"], 0.0)))
        launch_i = max(0, int(math.ceil(launch - 1e-9)))
        df.at[idx, "Lancement"] = launch_i
        df.at[idx, "Re-laquage"] = max(0, to_int(df.at[idx, "Re-laquage"], 0))

        bars = to_int(df.at[idx, "Barre/bal"], 0)
        if bars <= 0:
            bars = bars_master.get(art, 0)
        if bars <= 0:
            bars = 13
            unknown_bars += 1
        df.at[idx, "Barre/bal"] = bars

        nbal_raw = df.at[idx, "Nbre Bal"]
        nbal = to_int(nbal_raw, -1)
        if nbal < 0 or (norm_text(nbal_raw).upper() in {"#N/A", "N/A", "NA"}):
            nbal = int(math.ceil(launch_i / bars - 1e-12)) if launch_i > 0 else 0
            recalculated_bales += 1
        nbal = max(0, nbal)
        # Une ligne sans lancement ne doit jamais consommer de capacite automatiquement.
        if launch_i <= 0:
            nbal = 0
        df.at[idx, "Nbre Bal"] = nbal

        expected_tps = nbal * float(minutes_per_bal) / 60.0
        raw_tps = to_float(df.at[idx, "tps"], expected_tps)
        if abs(raw_tps - expected_tps) > 1e-6:
            tps_inconsistent += 1
        # La cadence du moteur est la source de verite; l'incoherence est auditee.
        df.at[idx, "tps"] = round(expected_tps, 12)

        weight = to_float(df.at[idx, "PoidsUn"], 0.0)
        if weight <= 0 and art in weight_master:
            weight = weight_master[art]
            df.at[idx, "PoidsUn"] = weight
        if weight > 0:
            qty = launch_i + max(0, to_int(df.at[idx, "Re-laquage"], 0))
            if to_float(df.at[idx, "PoidsT"], 0.0) <= 0:
                df.at[idx, "PoidsT"] = round(qty * weight, 3)

        raw_id = f"{int(df.at[idx, '_source_order'])}|{df.at[idx,'NumCommande']}|{article}|{df.at[idx,'NumOF']}"
        df.at[idx, "_line_id"] = hashlib.sha1(raw_id.encode("utf-8")).hexdigest()[:16]

    if unknown_bars:
        warnings.append(f"{unknown_bars} ligne(s) sans Barre/bal fiable: fallback 13 utilise.")
    if recalculated_bales:
        warnings.append(f"Nbre Bal recalcule pour {recalculated_bales} ligne(s) manquante(s)/N-A.")
    if tps_inconsistent:
        warnings.append(f"{tps_inconsistent} valeur(s) tps source differaient de la cadence {minutes_per_bal:g} min/bal; tps normalise.")

    # Colonnes internes necessaires au moteur.
    df["_bales"] = pd.to_numeric(df["Nbre Bal"], errors="coerce").fillna(0).clip(lower=0).astype(int)
    df["_launch"] = pd.to_numeric(df["Lancement"], errors="coerce").fillna(0).clip(lower=0).astype(int)
    # Une ligne a lancement 0 peut rester necessaire dans le planning quand elle
    # appartient a une commande/couleur qui contient d'autres lignes actives.
    # Exemple observe: accessoires d'une meme commande DARK. En revanche une
    # ligne 0 isolee ne doit pas consommer ni polluer le planning.
    has_article = df["Article"].astype(str).str.strip() != ""
    direct_active = (df["_launch"] > 0) & has_article
    group_has_active: Dict[Tuple[str, str], bool] = {}
    for (cmd, color), g in df.groupby(["NumCommande", "Couleur"], dropna=False, sort=False):
        key = (norm_text(cmd).upper(), norm_text(color).upper())
        group_has_active[key] = bool((pd.to_numeric(g["_launch"], errors="coerce").fillna(0) > 0).any())
    contextual = []
    for _, r in df.iterrows():
        cmd = norm_text(r.get("NumCommande")).upper()
        color = norm_text(r.get("Couleur")).upper()
        contextual.append(bool(cmd and group_has_active.get((cmd, color), False)))
    df["_active"] = has_article & (direct_active | pd.Series(contextual, index=df.index))
    info = {
        "source_sheet": ws.title,
        "source_mode": mode,
        "rows": int(len(df)),
        "active_rows": int(df["_active"].sum()),
        "warnings": warnings,
    }
    return df.reset_index(drop=True), info


@dataclass(frozen=True)
class PlannerConfig:
    year: int
    week: int
    minutes_per_bal: float = 4.0
    target_bales_per_day: int = 250
    max_bales_per_day: int = 270
    max_colors_per_day: int = 8
    base_workdays: int = 5
    max_spillover_days: int = 1
    auto_spillover: bool = True
    forbid_white_black_same_day: bool = True
    solver_seconds: float = 15.0
    preserve_source_order_weight: int = 2
    color_activation_weight: int = 80
    balance_weight: int = 4
    unscheduled_weight: int = 100000


@dataclass
class Job:
    job_id: str
    rows: List[int]
    color: str
    article: str
    bales: int
    source_order: int
    priority: int


def _row_priority(r: pd.Series) -> int:
    score = 1000
    if norm_key(r.get("reservation brut")) in {"oui", "yes", "1", "true"}:
        score += 500
    status = norm_key(r.get("ProdStatut"))
    if "commenc" in status:
        score += 250
    elif "cree" in status:
        score += 150
    # Les additions manuelles sont prioritaires.
    if norm_text(r.get("_origin")) == "addition":
        score += 5000
    return score


def build_jobs(df: pd.DataFrame, cfg: PlannerConfig) -> Tuple[List[Job], List[int]]:
    active = df[df["_active"]].copy().sort_values("_source_order", kind="stable")
    jobs: List[Job] = []
    zero_capacity_rows: List[int] = []

    # Groupes contigus par couleur + article interne: on preserve l'intention de la base finale.
    current_key: Optional[Tuple[str, str]] = None
    current_rows: List[int] = []

    def flush(rows: List[int]) -> None:
        if not rows:
            return
        # Couper seulement si un groupe depasse a lui seul la capacite dure.
        chunk: List[int] = []
        load = 0
        chunk_no = 1
        for ridx in rows:
            b = int(df.at[ridx, "_bales"])
            if b > cfg.max_bales_per_day:
                zero_capacity_rows.append(ridx)
                continue
            if chunk and load + b > cfg.max_bales_per_day:
                sub = df.loc[chunk]
                jobs.append(Job(
                    job_id=f"{norm_text(sub.iloc[0]['Couleur'])}|{norm_text(sub.iloc[0]['Article/int'])}|{int(sub.iloc[0]['_source_order'])}|{chunk_no}",
                    rows=list(chunk), color=norm_text(sub.iloc[0]["Couleur"]).upper(),
                    article=norm_text(sub.iloc[0]["Article/int"]).upper(), bales=int(load),
                    source_order=int(sub["_source_order"].min()), priority=max(_row_priority(r) for _, r in sub.iterrows()),
                ))
                chunk_no += 1
                chunk = []
                load = 0
            chunk.append(int(ridx)); load += b
        if chunk:
            sub = df.loc[chunk]
            jobs.append(Job(
                job_id=f"{norm_text(sub.iloc[0]['Couleur'])}|{norm_text(sub.iloc[0]['Article/int'])}|{int(sub.iloc[0]['_source_order'])}|{chunk_no}",
                rows=list(chunk), color=norm_text(sub.iloc[0]["Couleur"]).upper(),
                article=norm_text(sub.iloc[0]["Article/int"]).upper(), bales=int(load),
                source_order=int(sub["_source_order"].min()), priority=max(_row_priority(r) for _, r in sub.iterrows()),
            ))

    for ridx, r in active.iterrows():
        key = (norm_text(r.get("Couleur")).upper(), norm_text(r.get("Article/int")).upper())
        if current_key is None:
            current_key = key
        if key != current_key:
            flush(current_rows)
            current_rows = []
            current_key = key
        current_rows.append(int(ridx))
    flush(current_rows)
    return jobs, zero_capacity_rows


def _color_class(c: str) -> str:
    u = norm_text(c).upper()
    if u in WHITE_ALIASES:
        return "WHITE"
    if u in BLACK_ALIASES:
        return "BLACK"
    return "OTHER"


def _day_labels(cfg: PlannerConfig, n_days: int) -> List[Tuple[str, date]]:
    mon = iso_week_monday(cfg.year, cfg.week)
    labels: List[Tuple[str, date]] = []
    french_days = ["LUNDI", "MARDI", "MERCREDI", "JEUDI", "VENDREDI"]
    for d in range(min(cfg.base_workdays, n_days)):
        dt = mon + timedelta(days=d)
        label = french_days[d] if d < len(french_days) else f"JOUR {d+1}"
        labels.append((label, dt))
    # Spillover = lundi(s) suivant(s), pas samedi.
    extra = n_days - len(labels)
    for k in range(extra):
        dt = mon + timedelta(days=7 * (k + 1))
        labels.append(("LUNDI", dt))
    return labels


def _required_days(jobs: Sequence[Job], cfg: PlannerConfig) -> int:
    base = max(1, cfg.base_workdays)
    if not cfg.auto_spillover or cfg.max_spillover_days <= 0:
        return base
    total = sum(j.bales for j in jobs)
    needed = int(math.ceil(total / max(1, cfg.max_bales_per_day)))
    return min(base + cfg.max_spillover_days, max(base, needed))


def assign_jobs_ortools(jobs: List[Job], cfg: PlannerConfig, n_days: int) -> Tuple[Dict[str, int], List[str], str]:
    if not ORTOOLS_AVAILABLE or not jobs:
        return {}, [], ""
    model = cp_model.CpModel()
    x: Dict[Tuple[int, int], Any] = {}
    u: Dict[int, Any] = {}
    colors = sorted({j.color for j in jobs if j.color})
    y: Dict[Tuple[int, str], Any] = {}
    objective: List[Any] = []

    # Ranks normalises pour favoriser l'ordre de la base sans le rendre bloquant.
    ordered_ids = {j.job_id: rank for rank, j in enumerate(sorted(jobs, key=lambda z: z.source_order))}

    for ji, job in enumerate(jobs):
        u[ji] = model.NewBoolVar(f"u_{ji}")
        for d in range(n_days):
            x[(ji, d)] = model.NewBoolVar(f"x_{ji}_{d}")
        model.Add(sum(x[(ji, d)] for d in range(n_days)) + u[ji] == 1)
        # Backlog tres couteux, surtout pour les lignes prioritaires.
        objective.append(u[ji] * (cfg.unscheduled_weight + job.priority * 50 + job.bales * 100))
        rank = ordered_ids[job.job_id]
        # Plus une ligne est tot dans la base, plus elle prefere les premiers jours.
        for d in range(n_days):
            objective.append(x[(ji, d)] * d * cfg.preserve_source_order_weight * max(1, len(jobs) - rank))

    by_color: Dict[str, List[int]] = defaultdict(list)
    for ji, j in enumerate(jobs):
        if j.color:
            by_color[j.color].append(ji)

    for d in range(n_days):
        model.Add(sum(jobs[ji].bales * x[(ji, d)] for ji in range(len(jobs))) <= cfg.max_bales_per_day)
        active_colors = []
        for color in colors:
            var = model.NewBoolVar(f"y_{d}_{norm_key(color)}")
            y[(d, color)] = var
            idxs = by_color[color]
            for ji in idxs:
                model.Add(x[(ji, d)] <= var)
            model.Add(var <= sum(x[(ji, d)] for ji in idxs))
            active_colors.append(var)
            objective.append(var * cfg.color_activation_weight)
        if active_colors:
            model.Add(sum(active_colors) <= cfg.max_colors_per_day)

        if cfg.forbid_white_black_same_day:
            whites = [y[(d, c)] for c in colors if _color_class(c) == "WHITE"]
            blacks = [y[(d, c)] for c in colors if _color_class(c) == "BLACK"]
            if whites and blacks:
                has_white = model.NewBoolVar(f"has_white_{d}")
                has_black = model.NewBoolVar(f"has_black_{d}")
                model.AddMaxEquality(has_white, whites)
                model.AddMaxEquality(has_black, blacks)
                model.Add(has_white + has_black <= 1)

        load = model.NewIntVar(0, cfg.max_bales_per_day, f"load_{d}")
        model.Add(load == sum(jobs[ji].bales * x[(ji, d)] for ji in range(len(jobs))))
        dev = model.NewIntVar(0, cfg.max_bales_per_day, f"dev_{d}")
        model.Add(dev >= cfg.target_bales_per_day - load)
        model.Add(dev >= load - cfg.target_bales_per_day)
        objective.append(dev * cfg.balance_weight)

    # Eviter de disperser une meme couleur sur trop de jours, sans l'interdire.
    for color in colors:
        use_days = [y[(d, color)] for d in range(n_days)]
        extra = model.NewIntVar(0, max(0, n_days - 1), f"color_spread_{norm_key(color)}")
        model.Add(extra >= sum(use_days) - 1)
        objective.append(extra * 25)

    model.Minimize(sum(objective))
    solver = cp_model.CpSolver()
    solver.parameters.max_time_in_seconds = max(2.0, float(cfg.solver_seconds))
    solver.parameters.num_search_workers = 1  # determinisme reproductible
    solver.parameters.random_seed = 0
    status = solver.Solve(model)
    if status not in (cp_model.OPTIMAL, cp_model.FEASIBLE):
        return {}, [j.job_id for j in jobs], ""

    assignments: Dict[str, int] = {}
    backlog: List[str] = []
    for ji, job in enumerate(jobs):
        if solver.Value(u[ji]):
            backlog.append(job.job_id)
            continue
        placed = False
        for d in range(n_days):
            if solver.Value(x[(ji, d)]):
                assignments[job.job_id] = d
                placed = True
                break
        if not placed:
            backlog.append(job.job_id)
    return assignments, backlog, "OR-Tools CP-SAT deterministe"


def assign_jobs_fallback(jobs: List[Job], cfg: PlannerConfig, n_days: int) -> Tuple[Dict[str, int], List[str], str]:
    loads = [0] * n_days
    colors: List[set] = [set() for _ in range(n_days)]
    assignments: Dict[str, int] = {}
    backlog: List[str] = []
    ordered = sorted(jobs, key=lambda j: (-j.priority, j.source_order, -j.bales, j.job_id))
    for job in ordered:
        options = []
        for d in range(n_days):
            if loads[d] + job.bales > cfg.max_bales_per_day:
                continue
            new_colors = set(colors[d]) | ({job.color} if job.color else set())
            if len(new_colors) > cfg.max_colors_per_day:
                continue
            if cfg.forbid_white_black_same_day:
                classes = {_color_class(c) for c in new_colors}
                if "WHITE" in classes and "BLACK" in classes:
                    continue
            color_pen = 0 if job.color in colors[d] else cfg.color_activation_weight
            dev = abs(cfg.target_bales_per_day - (loads[d] + job.bales))
            order_pen = d * cfg.preserve_source_order_weight * max(1, len(jobs) - job.source_order)
            options.append((dev * cfg.balance_weight + color_pen + order_pen, d))
        if not options:
            backlog.append(job.job_id)
            continue
        _, d = min(options, key=lambda z: (z[0], z[1]))
        assignments[job.job_id] = d
        loads[d] += job.bales
        if job.color:
            colors[d].add(job.color)
    return assignments, backlog, "Heuristique deterministe"


def build_plan(df: pd.DataFrame, cfg: PlannerConfig) -> Dict[str, Any]:
    jobs, oversized_rows = build_jobs(df, cfg)
    n_days = _required_days(jobs, cfg)
    assignments, backlog_job_ids, engine = assign_jobs_ortools(jobs, cfg, n_days)
    if not engine:
        assignments, backlog_job_ids, engine = assign_jobs_fallback(jobs, cfg, n_days)

    job_map = {j.job_id: j for j in jobs}
    day_rows: Dict[int, List[int]] = defaultdict(list)
    for jid, d in assignments.items():
        day_rows[d].extend(job_map[jid].rows)
    for d in day_rows:
        # Ordre stable de la base finale a l'interieur de la journee.
        day_rows[d].sort(key=lambda i: int(df.at[i, "_source_order"]))

    days: Dict[int, pd.DataFrame] = {}
    for d in range(n_days):
        rows = day_rows.get(d, [])
        out = df.loc[rows].copy().reset_index(drop=True) if rows else df.iloc[0:0].copy()
        out["_planned_day"] = d
        days[d] = out

    backlog_rows: List[int] = list(oversized_rows)
    for jid in backlog_job_ids:
        backlog_rows.extend(job_map[jid].rows)
    # Lignes actives non capturees par les jobs = anomalie defensive.
    planned_ids = {i for rows in day_rows.values() for i in rows}
    active_ids = set(df.index[df["_active"]].tolist())
    backlog_rows.extend(sorted(active_ids - planned_ids - set(backlog_rows)))
    backlog_rows = sorted(set(backlog_rows), key=lambda i: int(df.at[i, "_source_order"]))
    backlog = df.loc[backlog_rows].copy().reset_index(drop=True) if backlog_rows else df.iloc[0:0].copy()

    labels = _day_labels(cfg, n_days)
    hard_errors: List[str] = []
    metrics: List[Dict[str, Any]] = []
    seen_line_ids: set = set()
    for d in range(n_days):
        day = days[d]
        bales = int(day["_bales"].sum()) if not day.empty else 0
        colors = list(OrderedDict.fromkeys(norm_text(x).upper() for x in day["Couleur"].tolist() if norm_text(x)))
        classes = {_color_class(c) for c in colors}
        if bales > cfg.max_bales_per_day:
            hard_errors.append(f"Jour {d+1}: {bales} bal > capacite {cfg.max_bales_per_day}.")
        if len(colors) > cfg.max_colors_per_day:
            hard_errors.append(f"Jour {d+1}: {len(colors)} couleurs > max {cfg.max_colors_per_day}.")
        if cfg.forbid_white_black_same_day and "WHITE" in classes and "BLACK" in classes:
            hard_errors.append(f"Jour {d+1}: BLANC et NOIR/DARK le meme jour.")
        for lid in day["_line_id"].tolist():
            if lid in seen_line_ids:
                hard_errors.append(f"Doublon planifie: {lid}.")
            seen_line_ids.add(lid)
        metrics.append({
            "Jour": labels[d][0], "Date": labels[d][1].strftime("%d/%m/%Y"),
            "Lignes": int(len(day)), "Balancelles": bales,
            "Temps h": round(bales * cfg.minutes_per_bal / 60.0, 2),
            "Charge %": round(100.0 * bales / cfg.max_bales_per_day, 1) if cfg.max_bales_per_day else 0.0,
            "Nb couleurs": len(colors), "Couleurs": " -> ".join(colors),
        })

    # Integrite: aucune ligne active ne doit disparaitre silencieusement.
    planned_active = sum(len(x) for x in days.values())
    if planned_active + len(backlog) != int(df["_active"].sum()):
        hard_errors.append("Integrite lignes: planifie + backlog != lignes actives source.")

    return {
        "config": cfg,
        "engine": engine,
        "days": days,
        "day_labels": labels,
        "backlog": backlog,
        "metrics": metrics,
        "hard_errors": list(dict.fromkeys(hard_errors)),
        "rules_ok": not hard_errors,
        "jobs": jobs,
        "source": df,
    }


def _planning_sheet_rows(ws, max_cols: int = 31) -> pd.DataFrame:
    raw = _read_compact_rows(ws, max_cols=max_cols)
    for col in OUTPUT_COLUMNS:
        if col not in raw.columns:
            raw[col] = None
    return raw


def load_reference_planning(data: bytes) -> List[pd.DataFrame]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    days: List[pd.DataFrame] = []
    for ws in wb.worksheets:
        title = norm_key(ws.title)
        if "planning" not in title and not title.startswith("lundi"):
            continue
        if "samedi" in title:
            continue
        try:
            d = _planning_sheet_rows(ws)
        except Exception:
            continue
        if not d.empty:
            d["_reference_sheet"] = ws.title
            days.append(d)
    if not days:
        raise ValueError("Aucune feuille de planning de reference trouvee.")
    return days


def _row_match_key(r: pd.Series) -> Tuple[str, str, str]:
    return (
        norm_text(r.get("NumCommande")).upper(),
        norm_text(r.get("Article")).upper(),
        norm_text(r.get("NumOF")).upper(),
    )


def _tag_occurrences(frames: Sequence[pd.DataFrame]) -> List[Tuple[Tuple[str, str, str], int, pd.Series]]:
    counts: Counter = Counter()
    out: List[Tuple[Tuple[str, str, str], int, pd.Series]] = []
    for day, df in enumerate(frames):
        for _, r in df.iterrows():
            k = _row_match_key(r)
            counts[k] += 1
            out.append((k, counts[k], r))
    return out


def audit_against_reference(plan: Dict[str, Any], reference_data: bytes) -> Dict[str, Any]:
    ref_days = load_reference_planning(reference_data)
    gen_days = [plan["days"][d] for d in sorted(plan["days"])]
    source = plan["source"]

    # Multisets, indispensable pour les lignes stock sans commande/OF.
    source_counter = Counter(_row_match_key(r) for _, r in source.iterrows() if norm_text(r.get("Article")))
    ref_counter = Counter(_row_match_key(r) for df in ref_days for _, r in df.iterrows())
    gen_counter = Counter(_row_match_key(r) for df in gen_days for _, r in df.iterrows())

    missing_from_input = sum(max(0, n - source_counter[k]) for k, n in ref_counter.items())
    missing_from_plan = sum(max(0, n - gen_counter[k]) for k, n in ref_counter.items())
    extra_in_plan = sum(max(0, n - ref_counter[k]) for k, n in gen_counter.items())
    ref_total = sum(ref_counter.values())

    # Matching occurrence par occurrence pour jour / Nbre Bal / tps.
    ref_queues: Dict[Tuple[str, str, str], deque] = defaultdict(deque)
    for d, df in enumerate(ref_days):
        for _, r in df.iterrows():
            ref_queues[_row_match_key(r)].append((d, r))

    same_day = 0
    matched = 0
    nbal_same = 0
    tps_same = 0
    for gd, df in enumerate(gen_days):
        for _, r in df.iterrows():
            k = _row_match_key(r)
            if not ref_queues[k]:
                continue
            rd, rr = ref_queues[k].popleft()
            matched += 1
            if gd == rd:
                same_day += 1
            if to_int(r.get("Nbre Bal"), -9999) == to_int(rr.get("Nbre Bal"), -9998):
                nbal_same += 1
            if abs(to_float(r.get("tps"), -9999.0) - to_float(rr.get("tps"), -9998.0)) <= 1e-8:
                tps_same += 1

    coverage = 100.0 * (ref_total - missing_from_plan) / ref_total if ref_total else 100.0
    input_coverage = 100.0 * (ref_total - missing_from_input) / ref_total if ref_total else 100.0
    day_match = 100.0 * same_day / matched if matched else 0.0
    nbal_match = 100.0 * nbal_same / matched if matched else 0.0
    tps_match = 100.0 * tps_same / matched if matched else 0.0

    exact = (
        missing_from_input == 0 and missing_from_plan == 0 and extra_in_plan == 0
        and matched == ref_total and same_day == ref_total and nbal_same == ref_total and tps_same == ref_total
    )
    return {
        "reference_rows": ref_total,
        "input_coverage_pct": round(input_coverage, 2),
        "planned_coverage_pct": round(coverage, 2),
        "matched_rows": matched,
        "same_day_pct": round(day_match, 2),
        "nbal_match_pct": round(nbal_match, 2),
        "tps_match_pct": round(tps_match, 2),
        "missing_from_input": missing_from_input,
        "missing_from_plan": missing_from_plan,
        "extra_in_plan": extra_in_plan,
        "exact_reference_match": exact,
    }


def _write_sheet(wb: Workbook, title: str, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(title[:31])
    navy, white, light = "163A5F", "FFFFFF", "EEF4FF"
    thin = Side(style="hair", color="D9E1F2")
    for ci, col in enumerate(OUTPUT_COLUMNS, 1):
        c = ws.cell(1, ci, EXPORT_HEADERS.get(col, col))
        c.fill = PatternFill("solid", fgColor=navy)
        c.font = Font(color=white, bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    prev_color = None
    for ri, (_, row) in enumerate(df.iterrows(), start=2):
        color = norm_text(row.get("Couleur"))
        changed = prev_color is not None and color != prev_color
        for ci, col in enumerate(OUTPUT_COLUMNS, 1):
            v = row.get(col)
            if isinstance(v, pd.Timestamp):
                v = v.to_pydatetime()
            if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
                v = None
            cell = ws.cell(ri, ci, v)
            cell.border = Border(bottom=thin)
            if changed:
                cell.fill = PatternFill("solid", fgColor=light)
            if col == "DateCreation" and isinstance(v, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm:ss"
            elif col in {"Nuance", "PoidsUn", "PoidsT", "Poudre", "tps", "moyenne vente", "% laque"} and isinstance(v, (int, float)):
                cell.number_format = "0.000"
            elif col in {"QteCommandee", "ResteALivrer", "Preleve", "QteCommence", "QteRestante", "QteRecu", "ReserverBR", "StockPhysique", "Reserver", "Lancement", "Re-laquage", "Barre/bal", "Nbre Bal", "Stock brut"} and isinstance(v, (int, float)):
                cell.number_format = "0"
        prev_color = color
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(OUTPUT_COLUMNS))}{max(1, len(df)+1)}"
    widths = {1: 16, 2: 20, 3: 25, 4: 25, 5: 18, 6: 12, 7: 10, 12: 14, 13: 16}
    for ci in range(1, len(OUTPUT_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(ci)].width = widths.get(ci, 13)


def export_excel(plan: Dict[str, Any], audit: Optional[Dict[str, Any]] = None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for d in sorted(plan["days"]):
        label, dt = plan["day_labels"][d]
        title = f"Planning {label.capitalize()} {dt.strftime('%d %m %Y')}"
        _write_sheet(wb, title, plan["days"][d][OUTPUT_COLUMNS])
    if not plan["backlog"].empty:
        _write_sheet(wb, "Backlog", plan["backlog"][OUTPUT_COLUMNS])

    ws = wb.create_sheet("Validation")
    ws.append(["Controle", "Valeur"])
    ws.append(["Version", VERSION])
    ws.append(["Moteur", plan["engine"]])
    ws.append(["Regles moteur", "OK" if plan["rules_ok"] else "NON OK"])
    ws.append(["Erreurs dures", " | ".join(plan["hard_errors"])])
    ws.append(["Backlog lignes", len(plan["backlog"])])
    if audit:
        for k, v in audit.items():
            ws.append([f"Reference - {k}", v])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="163A5F")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 70

    out = io.BytesIO(); wb.save(out); return out.getvalue()


def generate_from_bytes(base_data: bytes, cfg: PlannerConfig, additions: Optional[bytes] = None,
                        reference: Optional[bytes] = None) -> Dict[str, Any]:
    source, source_info = prepare_source(base_data, cfg.minutes_per_bal, additions=additions)
    plan = build_plan(source, cfg)
    plan["source_info"] = source_info
    audit = audit_against_reference(plan, reference) if reference else None
    plan["reference_audit"] = audit
    # Statut global sans promesse trompeuse.
    if not plan["rules_ok"]:
        status = "BLOCKED_RULES"
    elif audit is not None and not audit["exact_reference_match"]:
        status = "BLOCKED_REFERENCE_MISMATCH"
    elif audit is not None and audit["exact_reference_match"]:
        status = "REFERENCE_EXACT"
    else:
        status = "RULES_VALIDATED"
    plan["status"] = status
    return plan


def _summary_text(plan: Dict[str, Any]) -> str:
    audit = plan.get("reference_audit")
    parts = [
        f"status={plan['status']}", f"engine={plan['engine']}",
        f"source={plan['source_info']['source_sheet']}",
        f"active={plan['source_info']['active_rows']}", f"backlog={len(plan['backlog'])}",
    ]
    if audit:
        parts += [
            f"input_ref={audit['input_coverage_pct']}%",
            f"planned_ref={audit['planned_coverage_pct']}%",
            f"same_day={audit['same_day_pct']}%",
            f"nbal={audit['nbal_match_pct']}%",
        ]
    return " | ".join(parts)


def render_streamlit() -> None:
    if st is None:
        raise RuntimeError("Streamlit n'est pas installe.")
    st.set_page_config(page_title=APP_NAME, layout="wide")
    st.title(APP_NAME)
    st.caption(f"Version {VERSION} - moteur deterministe + agents de validation")

    with st.sidebar:
        st.subheader("Entrees")
        base = st.file_uploader("Base Excel", type=["xlsx"], key="base")
        additions = st.file_uploader("Ajouts urgents (optionnel)", type=["xlsx"], key="add")
        reference = st.file_uploader("Planning de reference (audit optionnel)", type=["xlsx"], key="ref")
        st.subheader("Parametres atelier")
        today = date.today().isocalendar()
        year = st.number_input("Annee ISO", value=int(today.year), step=1)
        week = st.number_input("Semaine ISO", min_value=1, max_value=53, value=int(today.week), step=1)
        target = st.number_input("Cible balancelles/jour", min_value=50, max_value=500, value=250, step=5)
        max_bales = st.number_input("Maximum balancelles/jour", min_value=int(target), max_value=600, value=max(270, int(target)), step=5)
        max_colors = st.number_input("Maximum couleurs/jour", min_value=1, max_value=20, value=8, step=1)
        spill = st.number_input("Lundis de debordement max", min_value=0, max_value=3, value=1, step=1)
        run = st.button("Generer et valider", type="primary", use_container_width=True)

    if not base:
        st.info("Charge une Base 1 / Base 2. La feuille finale est detectee automatiquement.")
        return
    if run or "last_plan" not in st.session_state:
        cfg = PlannerConfig(
            year=int(year), week=int(week), target_bales_per_day=int(target), max_bales_per_day=int(max_bales),
            max_colors_per_day=int(max_colors), max_spillover_days=int(spill), auto_spillover=int(spill) > 0,
        )
        try:
            plan = generate_from_bytes(
                base.getvalue(), cfg,
                additions=additions.getvalue() if additions else None,
                reference=reference.getvalue() if reference else None,
            )
            st.session_state["last_plan"] = plan
        except Exception as exc:
            st.exception(exc)
            return
    plan = st.session_state["last_plan"]

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Statut", plan["status"])
    c2.metric("Lignes actives", plan["source_info"]["active_rows"])
    c3.metric("Backlog", len(plan["backlog"]))
    c4.metric("Moteur", plan["engine"])
    st.caption(f"Feuille source: {plan['source_info']['source_sheet']} ({plan['source_info']['source_mode']})")

    if plan["hard_errors"]:
        st.error("\n".join(plan["hard_errors"]))
    for w in plan["source_info"]["warnings"]:
        st.warning(w)

    audit = plan.get("reference_audit")
    if audit:
        st.subheader("Audit reference")
        st.dataframe(pd.DataFrame([audit]), use_container_width=True, hide_index=True)
        if audit["missing_from_input"]:
            st.error(
                f"Le planning de reference contient {audit['missing_from_input']} ligne(s) absente(s) de la base. "
                "Une generation exacte est impossible sans fichier d'ajouts/urgence."
            )

    st.subheader("Charge journaliere")
    st.dataframe(pd.DataFrame(plan["metrics"]), use_container_width=True, hide_index=True)
    tabs = st.tabs([f"{lab} {dt.strftime('%d/%m')}" for lab, dt in plan["day_labels"]])
    for d, tab in enumerate(tabs):
        with tab:
            st.dataframe(plan["days"][d][OUTPUT_COLUMNS], use_container_width=True, hide_index=True)
    if not plan["backlog"].empty:
        with st.expander(f"Backlog ({len(plan['backlog'])} lignes)"):
            st.dataframe(plan["backlog"][OUTPUT_COLUMNS], use_container_width=True, hide_index=True)

    excel = export_excel(plan, audit)
    st.download_button("Telecharger Excel", excel, file_name=f"Planning_S{plan['config'].week}_{plan['config'].year}.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def cli() -> None:
    parser = argparse.ArgumentParser(description=APP_NAME)
    parser.add_argument("--generate", action="store_true")
    parser.add_argument("--input")
    parser.add_argument("--output")
    parser.add_argument("--reference")
    parser.add_argument("--additions")
    parser.add_argument("--year", type=int)
    parser.add_argument("--week", type=int)
    parser.add_argument("--spillover", type=int, default=1)
    parser.add_argument("--target-bales", type=int, default=250)
    parser.add_argument("--max-bales", type=int, default=270)
    parser.add_argument("--max-colors", type=int, default=8)
    parser.add_argument("--self-test", action="store_true")
    args = parser.parse_args()

    if args.self_test:
        assert split_article("LMMO-S758-BLC") == ("LMMO-S758", "BLC")
        assert iso_week_monday(2026, 38) == date(2026, 9, 14)
        assert PlannerConfig(2026, 40).minutes_per_bal == 4.0
        print("Self-test OK")
        return

    if args.generate:
        if not all([args.input, args.output, args.year, args.week]):
            parser.error("--input --output --year --week requis avec --generate")
        cfg = PlannerConfig(
            year=args.year, week=args.week, target_bales_per_day=args.target_bales,
            max_bales_per_day=args.max_bales, max_colors_per_day=args.max_colors,
            max_spillover_days=args.spillover, auto_spillover=args.spillover > 0,
        )
        base_data = Path(args.input).read_bytes()
        ref_data = Path(args.reference).read_bytes() if args.reference else None
        add_data = Path(args.additions).read_bytes() if args.additions else None
        plan = generate_from_bytes(base_data, cfg, additions=add_data, reference=ref_data)
        Path(args.output).write_bytes(export_excel(plan, plan.get("reference_audit")))
        print(_summary_text(plan))
        if plan["status"].startswith("BLOCKED"):
            raise SystemExit(2)
        return

    if st is not None:
        render_streamlit()
    else:
        parser.print_help()


if __name__ == "__main__":
    cli()
