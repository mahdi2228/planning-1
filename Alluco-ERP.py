# -*- coding: utf-8 -*-
"""
ALLUCO Industrial ERP IA - V8 monofichier
==========================================

Flux cible:
AX / Excel -> Commandes -> Planning IA -> Magasin J-2 -> Laquage -> Qualite
-> Re-laquage -> Logistique -> Livraison.

Le fichier reste volontairement monolithique, mais il est structure en sections.
La base SQLite ``alluco.db`` est creee automatiquement au premier lancement.

Lancement:
    pip install -r requirements.txt
    streamlit run alluco_erp.py

Tests:
    python alluco_erp.py --self-test
"""
from __future__ import annotations

import base64
import hashlib
import hmac
import html
import io
import json
import math
import os
import re
import secrets
import sqlite3
import sys
import time
import unicodedata
from collections import Counter, defaultdict
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

try:
    from zoneinfo import ZoneInfo
except Exception:
    ZoneInfo = None

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:
    import streamlit as st
except Exception:
    st = None

try:
    from ortools.sat.python import cp_model
    ORTOOLS_AVAILABLE = True
except Exception:
    cp_model = None
    ORTOOLS_AVAILABLE = False

try:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfgen import canvas as pdf_canvas
    REPORTLAB_AVAILABLE = True
except Exception:
    A4 = landscape = pdf_canvas = None
    REPORTLAB_AVAILABLE = False

# =============================================================================
# 01. CONFIGURATION
# =============================================================================
VERSION = "8.1.0"
APP_NAME = "ALLUCO - Industrial ERP IA"
APP_SUBTITLE = "Planning · Magasin · Laquage · Qualite · Logistique"
ROOT_DIR = Path(__file__).resolve().parent
DB_PATH = Path(os.environ.get("ALLUCO_DB_PATH", str(ROOT_DIR / "alluco.db")))
APP_TIMEZONE = os.environ.get("ALLUCO_TIMEZONE", "Africa/Tunis")
CLIENT_ACCESS_CODE = os.environ.get("ALLUCO_CLIENT_ACCESS_CODE", "")

# Identite visuelle. Priorite aux fichiers locaux si presents; sinon URLs GitHub RAW.
LOGO_DARK_URL = "https://raw.githubusercontent.com/mahdi2228/planning-1/main/logo_dark.png"
LOGO_WHITE_URL = "https://raw.githubusercontent.com/mahdi2228/planning-1/main/logo_white.png"
LOGO_DARK_PATH = ROOT_DIR / "logo_dark.png"
LOGO_WHITE_PATH = ROOT_DIR / "logo_white.png"

DAYS = ["LUNDI", "MARDI", "MERCREDI", "JEUDI", "VENDREDI"]
DEFAULT_CAPACITY_H = 16.0
DEFAULT_MINUTES_PER_BAL = 4.0
DEFAULT_CLEANING_MIN = 15
DEFAULT_POWDER_COEFF = 0.052
DEFAULT_TARGET_UTIL = 0.94
DEFAULT_SOLVER_SECONDS = 12.0
HARD_MAX_COLORS_PER_DAY = 4
PBKDF2_ITERATIONS = 310_000

ROLES = [
    "ADMIN", "DIRECTION", "PLANNING", "MAGASIN", "LAQUAGE",
    "QUALITE", "LOGISTIQUE", "COMMERCIAL", "CLIENT",
]

ORDER_STATUSES = [
    "A_PLANIFIER", "PLANIFIE", "MATIERE_A_PREPARER", "MATIERE_PRETE",
    "EN_COURS_LAQUAGE", "CONTROLE_QUALITE", "CONFORME", "BLOQUE",
    "RE_LAQUAGE", "PRET_EXPEDITION", "PLANIFIE_LOGISTIQUE", "CHARGE",
    "EXPEDIE", "LIVRE",
]

PLANNING_ENTRY_STATUSES = [
    "PLANIFIE", "MATIERE_A_PREPARER", "MATIERE_PRETE", "EN_COURS_LAQUAGE",
    "CONTROLE_QUALITE", "CONFORME", "RE_LAQUAGE", "PRET_EXPEDITION",
    "PLANIFIE_LOGISTIQUE", "CHARGE", "EXPEDIE", "LIVRE", "BLOQUE",
]

MAGASIN_STATUSES = ["A_PREPARER", "EN_PREPARATION", "PRET", "BLOQUE"]
RELAQUAGE_STATUSES = ["A_PLANIFIER", "PLANIFIE", "EN_COURS", "CONTROLE", "CONFORME", "ANNULE"]
TRIP_STATUSES = ["BROUILLON", "PLANIFIE", "CHARGE", "EXPEDIE", "LIVRE", "ANNULE"]

DEFAULT_NUANCE = {
    "BLC": 1.0, "R9016": 2.0, "R1013": 5.0, "R1019": 7.0,
    "SAND": 9.0, "FRENE": 11.0, "TECK": 13.0, "ACAJOU": 15.0,
    "NOYER": 16.0, "NOCE": 17.0, "R8019": 19.0, "TRESOR": 20.0,
    "GRIS": 22.0, "GREY": 23.0, "GRISG": 24.0, "R7016": 27.0,
    "CHPG": 29.0, "COOL": 31.0, "N02": 33.0, "N07": 34.0,
    "N22": 35.0, "NOIR": 37.0, "DARK": 38.0,
    "ANOD": 40.0, "ANODN": 41.0, "ABRONZE": 42.0,
}

ARTICLE_BARS_PER_BAL = {"EC40100": 14}
FAMILY_BARS_PER_BAL = {
    "EC": 13, "FR": 13, "CSQ": 13, "FSQ": 13, "CO": 13,
    "LM": 17, "GL": 20, "P": 14, "PL": 20, "C": 10,
    "LMDP": 14, "LMDPF": 20, "AL": 10, "LMMO": 19,
    "MR": 9, "PR": 6, "PCN": 800, "T": 800,
}

WHITE_COLOR_ALIASES = {"BLC", "BLANC", "WHITE", "R9016"}
BLACK_COLOR_ALIASES = {"NOIR", "DARK", "BLACK", "R9005"}

PLANNING_COLUMNS = [
    "NumCommande", "DateCreation", "NomClient", "Article", "ArticleInt",
    "Couleur", "Nuance", "QteCommandee", "ResteALivrer", "Preleve",
    "ReservationBrut", "NumOF", "ProdStatut", "QteCommencee", "QteRestante",
    "QteRecue", "ReserverBR", "StockPhysique", "Reserver", "Lancement",
    "ReLaquage", "PoidsUn", "PoidsT", "Poudre", "BarreBal", "NbreBal",
    "TempsH", "StockBrut", "DateLivraison", "Destination",
]

# =============================================================================
# 02. HELPERS
# =============================================================================
def app_now() -> datetime:
    if ZoneInfo is not None:
        try:
            return datetime.now(ZoneInfo(APP_TIMEZONE))
        except Exception:
            pass
    return datetime.now()


def app_today() -> date:
    return app_now().date()


def norm_text(v: Any) -> str:
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except Exception:
        pass
    return str(v).strip()


def norm_key(v: Any) -> str:
    s = unicodedata.normalize("NFKD", norm_text(v)).encode("ascii", "ignore").decode("ascii")
    s = s.lower().strip()
    return re.sub(r"[^a-z0-9]+", "_", s).strip("_")


def to_float(v: Any, default: float = 0.0) -> float:
    try:
        if v is None or (isinstance(v, float) and np.isnan(v)):
            return default
        if isinstance(v, str):
            s = v.strip().replace(" ", "").replace(",", ".")
            if not s or s.upper() in {"#N/A", "N/A", "NA", "NONE", "NAN", "-"}:
                return default
            return float(s)
        return float(v)
    except Exception:
        return default


def to_int(v: Any, default: int = 0) -> int:
    return int(round(to_float(v, float(default))))


def parse_date(v: Any) -> Optional[pd.Timestamp]:
    if v is None:
        return None
    try:
        if isinstance(v, (int, float, np.integer, np.floating)):
            n = float(v)
            if 20000 <= n <= 80000:
                ts = pd.Timestamp("1899-12-30") + pd.to_timedelta(n, unit="D")
            else:
                return None
        else:
            raw = str(v).strip()
            if re.fullmatch(r"\d{4}-\d{2}-\d{2}(?:[ T].*)?", raw):
                ts = pd.to_datetime(v, errors="coerce")
            else:
                ts = pd.to_datetime(v, errors="coerce", dayfirst=True)
        if pd.isna(ts) or ts.year <= 1900:
            return None
        return pd.Timestamp(ts)
    except Exception:
        return None


def db_date(v: Any) -> Optional[str]:
    d = parse_date(v)
    return d.date().isoformat() if d is not None else None


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def json_dumps(v: Any) -> str:
    return json.dumps(v, ensure_ascii=False, default=str, separators=(",", ":"))


def esc(v: Any) -> str:
    return html.escape(norm_text(v), quote=True)


def file_data_uri(path: Path) -> str:
    """Retourne un data URI pour une image locale; chaîne vide en cas d'échec."""
    try:
        if path.is_file():
            data = path.read_bytes()
            if data:
                mime = "image/png" if path.suffix.lower() == ".png" else "image/jpeg"
                return f"data:{mime};base64,{base64.b64encode(data).decode('ascii')}"
    except Exception:
        pass
    return ""


def brand_logo_src(dark_mode: bool) -> str:
    """Logo blanc sur fond sombre, logo sombre sur fond clair.

    Les fichiers locaux sont utilisés en priorité pour permettre un fonctionnement
    hors-ligne. S'ils sont absents, l'URL RAW GitHub fournie est utilisée.
    """
    if dark_mode:
        return file_data_uri(LOGO_WHITE_PATH) or LOGO_WHITE_URL
    return file_data_uri(LOGO_DARK_PATH) or LOGO_DARK_URL


def brand_html(dark_mode: bool, compact: bool = False) -> str:
    src = esc(brand_logo_src(dark_mode))
    compact_cls = " brand-compact" if compact else ""
    return (
        f"<div class='brand-shell{compact_cls}'>"
        f"<img class='brand-logo' src='{src}' alt='ALLUCO'>"
        f"<div class='brand-copy'><div class='brand-title'>Industrial ERP IA</div>"
        f"<div class='brand-sub'>Planning · Magasin · Laquage · Qualité · Logistique</div></div></div>"
    )


def article_family(article_internal: str) -> str:
    m = re.match(r"([A-Z]+)", norm_text(article_internal).upper())
    return m.group(1) if m else ""


def looks_like_color_token(token: str) -> bool:
    t = norm_text(token).upper()
    if not t or t == "BRUT":
        return False
    if re.search(r"\d+(?:[.,]\d+)?X\d+", t) or "/" in t or "." in t:
        return False
    if t in DEFAULT_NUANCE:
        return True
    if re.fullmatch(r"R\d{4}", t) or re.fullmatch(r"N\d{2}", t):
        return True
    return bool(re.fullmatch(r"[A-Z][A-Z0-9]{1,14}", t))


def split_article(article: Any) -> Tuple[str, str]:
    s = norm_text(article)
    if "-" not in s:
        return s, ""
    left, right = s.rsplit("-", 1)
    c = right.strip().upper()
    if not looks_like_color_token(c):
        return s, ""
    return left.strip(), c


def color_nuance(color: str) -> float:
    c = norm_text(color).upper()
    if c in DEFAULT_NUANCE:
        return DEFAULT_NUANCE[c]
    h = hashlib.sha256(c.encode("utf-8")).hexdigest()
    return 50.0 + (int(h[:6], 16) % 4000) / 100.0


def infer_bars_per_bal(article_internal: str, unit_weight: float = 0.0) -> int:
    art = norm_text(article_internal).upper()
    if art in ARTICLE_BARS_PER_BAL:
        return ARTICLE_BARS_PER_BAL[art]
    fam = article_family(art)
    if fam in FAMILY_BARS_PER_BAL:
        return FAMILY_BARS_PER_BAL[fam]
    if unit_weight <= 0:
        return 13
    if unit_weight <= 0.08:
        return 200
    if unit_weight < 1.2:
        return 25
    if unit_weight < 2.0:
        return 20
    if unit_weight < 3.2:
        return 17
    if unit_weight < 8.0:
        return 13
    if unit_weight < 10.0:
        return 9
    return 6


def color_class(color: Any) -> str:
    c = norm_text(color).upper()
    if c in WHITE_COLOR_ALIASES:
        return "WHITE"
    if c in BLACK_COLOR_ALIASES:
        return "BLACK"
    return "OTHER"


def white_black_conflict(colors_a: Iterable[Any], colors_b: Optional[Iterable[Any]] = None) -> bool:
    a = {color_class(c) for c in colors_a if norm_text(c)}
    if colors_b is None:
        return "WHITE" in a and "BLACK" in a
    b = {color_class(c) for c in colors_b if norm_text(c)}
    return ("WHITE" in a and "BLACK" in b) or ("BLACK" in a and "WHITE" in b)


def subtract_business_days(d: date, n: int) -> date:
    cur = d
    left = int(n)
    while left > 0:
        cur -= timedelta(days=1)
        if cur.weekday() < 5:
            left -= 1
    return cur


def next_planning_period(reference: Optional[date] = None) -> Tuple[int, int, date, date]:
    ref = reference or app_today()
    days_until_monday = (7 - ref.weekday()) % 7
    if days_until_monday == 0:
        days_until_monday = 7
    start = ref + timedelta(days=days_until_monday)
    end = start + timedelta(days=4)
    iso = start.isocalendar()
    return int(iso.year), int(iso.week), start, end


def iso_week_dates(year: int, week: int) -> Dict[int, date]:
    monday = date.fromisocalendar(int(year), int(week), 1)
    return {i: monday + timedelta(days=i) for i in range(5)}

# =============================================================================
# 03. DATABASE
# =============================================================================
def get_db() -> sqlite3.Connection:
    conn = sqlite3.connect(str(DB_PATH), timeout=30, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("PRAGMA journal_mode = WAL")
    conn.execute("PRAGMA synchronous = NORMAL")
    return conn


@contextmanager
def db_conn():
    conn = get_db()
    try:
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def q_all(sql: str, params: Sequence[Any] = ()) -> List[Dict[str, Any]]:
    with db_conn() as conn:
        return [dict(r) for r in conn.execute(sql, params).fetchall()]


def q_one(sql: str, params: Sequence[Any] = ()) -> Optional[Dict[str, Any]]:
    with db_conn() as conn:
        r = conn.execute(sql, params).fetchone()
        return dict(r) if r else None


def init_db() -> None:
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    ddl = """
    CREATE TABLE IF NOT EXISTS users (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT NOT NULL UNIQUE,
        password_hash TEXT NOT NULL,
        role TEXT NOT NULL,
        active INTEGER NOT NULL DEFAULT 1,
        created_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS audit_log (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        ts TEXT NOT NULL,
        username TEXT,
        action TEXT NOT NULL,
        entity TEXT NOT NULL,
        entity_id TEXT,
        before_json TEXT,
        after_json TEXT,
        reason TEXT
    );
    CREATE TABLE IF NOT EXISTS notifications (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        user_role TEXT NOT NULL,
        event_type TEXT NOT NULL,
        title TEXT NOT NULL,
        message TEXT NOT NULL,
        entity TEXT,
        entity_id TEXT,
        created_at TEXT NOT NULL,
        read_at TEXT
    );
    CREATE TABLE IF NOT EXISTS imports (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        filename TEXT NOT NULL,
        sha256 TEXT NOT NULL,
        source_sheet TEXT,
        imported_at TEXT NOT NULL,
        imported_by TEXT,
        row_count INTEGER NOT NULL DEFAULT 0,
        valid_rows INTEGER NOT NULL DEFAULT 0,
        invalid_rows INTEGER NOT NULL DEFAULT 0,
        active INTEGER NOT NULL DEFAULT 1
    );
    CREATE INDEX IF NOT EXISTS idx_import_sha ON imports(sha256);
    CREATE TABLE IF NOT EXISTS order_lines (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        import_id INTEGER NOT NULL,
        source_row INTEGER,
        line_uid TEXT NOT NULL,
        num_commande TEXT,
        date_creation TEXT,
        nom_client TEXT,
        article TEXT,
        article_int TEXT,
        couleur TEXT,
        nuance REAL,
        qte_commandee REAL DEFAULT 0,
        reste_a_livrer REAL DEFAULT 0,
        preleve REAL DEFAULT 0,
        reservation_brut TEXT,
        num_of TEXT,
        prod_statut TEXT,
        qte_commencee REAL DEFAULT 0,
        qte_restante REAL DEFAULT 0,
        qte_recue REAL DEFAULT 0,
        reserver_br REAL DEFAULT 0,
        stock_physique REAL DEFAULT 0,
        reserver REAL DEFAULT 0,
        lancement REAL DEFAULT 0,
        relaquage REAL DEFAULT 0,
        poids_un REAL DEFAULT 0,
        poids_t REAL DEFAULT 0,
        poudre REAL DEFAULT 0,
        barre_bal REAL DEFAULT 0,
        nbre_bal REAL DEFAULT 0,
        temps_h REAL DEFAULT 0,
        stock_brut REAL DEFAULT 0,
        date_livraison TEXT,
        destination TEXT,
        status TEXT NOT NULL DEFAULT 'A_PLANIFIER',
        active INTEGER NOT NULL DEFAULT 1,
        updated_at TEXT NOT NULL,
        FOREIGN KEY(import_id) REFERENCES imports(id)
    );
    CREATE INDEX IF NOT EXISTS idx_order_active ON order_lines(active);
    CREATE INDEX IF NOT EXISTS idx_order_cmd ON order_lines(num_commande);
    CREATE TABLE IF NOT EXISTS planning_versions (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        year INTEGER NOT NULL,
        week INTEGER NOT NULL,
        version_no INTEGER NOT NULL,
        status TEXT NOT NULL,
        source_import_id INTEGER,
        created_at TEXT NOT NULL,
        created_by TEXT,
        published_at TEXT,
        notes TEXT,
        parent_version_id INTEGER,
        FOREIGN KEY(source_import_id) REFERENCES imports(id),
        FOREIGN KEY(parent_version_id) REFERENCES planning_versions(id)
    );
    CREATE INDEX IF NOT EXISTS idx_planning_period ON planning_versions(year, week, status);
    CREATE TABLE IF NOT EXISTS planning_entries (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        planning_version_id INTEGER NOT NULL,
        order_line_id INTEGER,
        relaquage_id INTEGER,
        planned_date TEXT NOT NULL,
        sequence_no INTEGER NOT NULL DEFAULT 0,
        num_commande TEXT,
        nom_client TEXT,
        article TEXT,
        article_int TEXT,
        couleur TEXT,
        num_of TEXT,
        planned_qty REAL DEFAULT 0,
        poids_t REAL DEFAULT 0,
        poudre REAL DEFAULT 0,
        nbre_bal REAL DEFAULT 0,
        duration_h REAL DEFAULT 0,
        priority_score REAL DEFAULT 0,
        decision_reason TEXT,
        status TEXT NOT NULL DEFAULT 'PLANIFIE',
        FOREIGN KEY(planning_version_id) REFERENCES planning_versions(id),
        FOREIGN KEY(order_line_id) REFERENCES order_lines(id)
    );
    CREATE INDEX IF NOT EXISTS idx_entry_version ON planning_entries(planning_version_id);
    CREATE INDEX IF NOT EXISTS idx_entry_date ON planning_entries(planned_date);
    CREATE TABLE IF NOT EXISTS magasin_preparations (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        planning_entry_id INTEGER NOT NULL UNIQUE,
        prep_date TEXT NOT NULL,
        status TEXT NOT NULL DEFAULT 'A_PREPARER',
        required_qty REAL DEFAULT 0,
        reserved_qty REAL DEFAULT 0,
        missing_qty REAL DEFAULT 0,
        location TEXT,
        updated_at TEXT NOT NULL,
        updated_by TEXT,
        FOREIGN KEY(planning_entry_id) REFERENCES planning_entries(id)
    );
    CREATE TABLE IF NOT EXISTS quality_incidents (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        planning_entry_id INTEGER,
        order_line_id INTEGER,
        inspected_at TEXT NOT NULL,
        inspector TEXT,
        result TEXT NOT NULL,
        defect_reason TEXT,
        qty_affected REAL DEFAULT 0,
        comment TEXT,
        status TEXT NOT NULL DEFAULT 'OUVERT',
        FOREIGN KEY(planning_entry_id) REFERENCES planning_entries(id),
        FOREIGN KEY(order_line_id) REFERENCES order_lines(id)
    );
    CREATE TABLE IF NOT EXISTS relaquage_orders (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        quality_incident_id INTEGER,
        order_line_id INTEGER,
        original_command TEXT,
        article TEXT,
        article_int TEXT,
        couleur TEXT,
        num_of TEXT,
        qty REAL NOT NULL DEFAULT 0,
        reason TEXT,
        priority TEXT NOT NULL DEFAULT 'HAUTE',
        status TEXT NOT NULL DEFAULT 'A_PLANIFIER',
        requested_date TEXT,
        planned_date TEXT,
        created_at TEXT NOT NULL,
        updated_at TEXT NOT NULL,
        FOREIGN KEY(quality_incident_id) REFERENCES quality_incidents(id),
        FOREIGN KEY(order_line_id) REFERENCES order_lines(id)
    );
    CREATE TABLE IF NOT EXISTS trucks (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        code TEXT NOT NULL UNIQUE,
        plate TEXT,
        capacity_kg REAL NOT NULL,
        active INTEGER NOT NULL DEFAULT 1
    );
    CREATE TABLE IF NOT EXISTS drivers (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT NOT NULL UNIQUE,
        phone TEXT,
        active INTEGER NOT NULL DEFAULT 1
    );
    CREATE TABLE IF NOT EXISTS transport_trips (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        trip_date TEXT NOT NULL,
        truck_id INTEGER NOT NULL,
        driver_id INTEGER,
        destination TEXT NOT NULL,
        status TEXT NOT NULL DEFAULT 'PLANIFIE',
        created_by TEXT,
        created_at TEXT NOT NULL,
        FOREIGN KEY(truck_id) REFERENCES trucks(id),
        FOREIGN KEY(driver_id) REFERENCES drivers(id)
    );
    CREATE TABLE IF NOT EXISTS shipments (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        trip_id INTEGER NOT NULL,
        planning_entry_id INTEGER NOT NULL,
        num_commande TEXT,
        nom_client TEXT,
        destination TEXT,
        weight_kg REAL NOT NULL DEFAULT 0,
        status TEXT NOT NULL DEFAULT 'PLANIFIE',
        created_at TEXT NOT NULL,
        UNIQUE(trip_id, planning_entry_id),
        FOREIGN KEY(trip_id) REFERENCES transport_trips(id),
        FOREIGN KEY(planning_entry_id) REFERENCES planning_entries(id)
    );
    CREATE TABLE IF NOT EXISTS settings (
        key TEXT PRIMARY KEY,
        value TEXT
    );
    """
    with db_conn() as conn:
        conn.executescript(ddl)

# =============================================================================
# 04. AUTHENTIFICATION / RBAC
# =============================================================================
def make_password_hash(password: str, iterations: int = PBKDF2_ITERATIONS) -> str:
    if len(password) < 10:
        raise ValueError("Le mot de passe doit contenir au moins 10 caracteres.")
    salt = secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, iterations)
    return f"pbkdf2_sha256${iterations}${salt.hex()}${digest.hex()}"


def verify_password_hash(password: str, encoded: str) -> bool:
    try:
        scheme, iterations, salt_hex, digest_hex = encoded.split("$", 3)
        if scheme != "pbkdf2_sha256":
            return False
        expected = bytes.fromhex(digest_hex)
        actual = hashlib.pbkdf2_hmac(
            "sha256", password.encode("utf-8"), bytes.fromhex(salt_hex), int(iterations)
        )
        return hmac.compare_digest(actual, expected)
    except Exception:
        return False


def users_exist() -> bool:
    row = q_one("SELECT COUNT(*) AS n FROM users")
    return bool(row and int(row["n"]) > 0)


def create_user(username: str, password: str, role: str, actor: str = "SYSTEM") -> int:
    username = norm_text(username)
    role = norm_text(role).upper()
    if not username:
        raise ValueError("Utilisateur vide.")
    if role not in ROLES:
        raise ValueError("Role invalide.")
    encoded = make_password_hash(password)
    with db_conn() as conn:
        cur = conn.execute(
            "INSERT INTO users(username,password_hash,role,active,created_at) VALUES(?,?,?,?,?)",
            (username, encoded, role, 1, app_now().isoformat(timespec="seconds")),
        )
        user_id = int(cur.lastrowid)
    audit(actor, "CREATE_USER", "user", user_id, None, {"username": username, "role": role})
    return user_id


def authenticate(username: str, password: str) -> Optional[Dict[str, Any]]:
    row = q_one("SELECT * FROM users WHERE username=? AND active=1", (norm_text(username),))
    if not row or not verify_password_hash(password, row["password_hash"]):
        return None
    return row


def role_can(role: str, area: str) -> bool:
    role = norm_text(role).upper()
    if role == "ADMIN":
        return True
    matrix = {
        "DIRECTION": {"dashboard", "balancelles", "analysis", "notifications"},
        "PLANNING": {"dashboard", "planning", "balancelles", "analysis", "notifications", "orders"},
        "MAGASIN": {"dashboard", "magasin", "notifications"},
        "LAQUAGE": {"dashboard", "balancelles", "laquage", "notifications"},
        "QUALITE": {"dashboard", "quality", "relaquage", "notifications"},
        "LOGISTIQUE": {"dashboard", "logistics", "notifications"},
        "COMMERCIAL": {"dashboard", "orders", "notifications"},
        "CLIENT": set(),
    }
    return area in matrix.get(role, set())

# =============================================================================
# 05. AUDIT / NOTIFICATIONS
# =============================================================================
def audit(username: str, action: str, entity: str, entity_id: Any,
          before: Any = None, after: Any = None, reason: str = "") -> None:
    with db_conn() as conn:
        conn.execute(
            "INSERT INTO audit_log(ts,username,action,entity,entity_id,before_json,after_json,reason) VALUES(?,?,?,?,?,?,?,?)",
            (
                app_now().isoformat(timespec="seconds"), norm_text(username), norm_text(action),
                norm_text(entity), norm_text(entity_id),
                json_dumps(before) if before is not None else None,
                json_dumps(after) if after is not None else None,
                norm_text(reason),
            ),
        )


def notify(role: str, event_type: str, title: str, message: str,
           entity: str = "", entity_id: Any = "") -> None:
    role = norm_text(role).upper()
    if role not in ROLES:
        return
    with db_conn() as conn:
        conn.execute(
            "INSERT INTO notifications(user_role,event_type,title,message,entity,entity_id,created_at) VALUES(?,?,?,?,?,?,?)",
            (role, event_type, title, message, entity, norm_text(entity_id), app_now().isoformat(timespec="seconds")),
        )

# =============================================================================
# 06. IMPORT AX / EXCEL
# =============================================================================
HEADER_ALIASES: Dict[str, str] = {
    "numcommande": "NumCommande", "num_commande": "NumCommande", "commande": "NumCommande",
    "datecreation": "DateCreation", "date_creation": "DateCreation",
    "nomclient": "NomClient", "nom_client": "NomClient", "client": "NomClient",
    "article": "Article", "article_int": "ArticleInt", "articleint": "ArticleInt",
    "couleur": "Couleur", "nuance": "Nuance",
    "qtecommande": "QteCommandee", "qte_commandee": "QteCommandee", "qtecommande": "QteCommandee",
    "restealivrer": "ResteALivrer", "reste_a_livrer": "ResteALivrer",
    "preleve": "Preleve", "prelevee": "Preleve",
    "reservation_brut": "ReservationBrut", "reservationbrut": "ReservationBrut", "2": "ReservationBrut",
    "numof": "NumOF", "num_of": "NumOF",
    "prodstatut": "ProdStatut", "prod_statut": "ProdStatut",
    "qtecommence": "QteCommencee", "qte_commencee": "QteCommencee",
    "qterestante": "QteRestante", "qte_restante": "QteRestante",
    "qterecu": "QteRecue", "qte_recu": "QteRecue", "qterecue": "QteRecue",
    "reserverbr": "ReserverBR", "stockphysique": "StockPhysique", "stock_physique": "StockPhysique",
    "reserver": "Reserver", "lancement": "Lancement", "re_laquage": "ReLaquage", "relaquage": "ReLaquage",
    "poidsun": "PoidsUn", "poidarticle": "PoidsUn", "poids_un": "PoidsUn",
    "poidst": "PoidsT", "poids_t": "PoidsT", "poudre": "Poudre",
    "barre_bal": "BarreBal", "barrebal": "BarreBal", "nbre_bal": "NbreBal", "nbrebal": "NbreBal",
    "tps": "TempsH", "temps": "TempsH", "temps_h": "TempsH",
    "stock_brut": "StockBrut", "stockbrut": "StockBrut",
    "datelivraisonconfirme": "DateLivraison", "date_livraison_confirme": "DateLivraison",
    "dateexpeditionconfirme": "DateLivraison", "date_expedition_confirme": "DateLivraison",
    "dateexpeditiondemande": "DateLivraison", "date_expedition_demande": "DateLivraison",
    "destination": "Destination", "ville": "Destination", "gouvernorat": "Destination",
    "adresselivraison": "Destination", "adresse_livraison": "Destination",
}


def canonical_header(v: Any) -> str:
    k = norm_key(v)
    return HEADER_ALIASES.get(k, norm_text(v))


def compact_sheet(ws, blank_stop: int = 250) -> pd.DataFrame:
    raw_header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    while raw_header and raw_header[-1] is None:
        raw_header.pop()
    if not raw_header:
        return pd.DataFrame()
    header: List[str] = []
    seen: Counter = Counter()
    for h in raw_header:
        name = canonical_header(h) or f"COL_{len(header)+1}"
        seen[name] += 1
        if seen[name] > 1:
            name = f"{name}_{seen[name]}"
        header.append(name)
    rows: List[Tuple[Any, ...]] = []
    blanks = 0
    started = False
    for row in ws.iter_rows(min_row=2, max_col=len(header), values_only=True):
        vals = tuple(row[:len(header)])
        if all(v is None for v in vals):
            if started:
                blanks += 1
                if blanks >= blank_stop:
                    break
            continue
        started = True
        blanks = 0
        rows.append(vals)
    return pd.DataFrame(rows, columns=header)


def choose_source_sheet(wb) -> Any:
    preferred_keys = [
        "preparation_pour_planning_vf", "version_0", "extraction_ax", "feuil1",
    ]
    by_key = {norm_key(s): s for s in wb.sheetnames}
    for k in preferred_keys:
        if k in by_key:
            return wb[by_key[k]]
    best = None
    best_score = -1
    for ws in wb.worksheets:
        try:
            hdr = [canonical_header(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
            keys = {norm_key(x) for x in hdr}
            score = (3 if "numcommande" in keys else 0) + (3 if "article" in keys else 0) + (2 if "restealivrer" in keys else 0)
            if score > best_score:
                best_score = score
                best = ws
        except Exception:
            continue
    if best is None or best_score < 3:
        raise ValueError("Aucune feuille AX exploitable detectee.")
    return best


def read_ax_workbook(data: bytes) -> Tuple[pd.DataFrame, str]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = choose_source_sheet(wb)
    df = compact_sheet(ws)
    if df.empty:
        raise ValueError("La feuille source est vide.")
    # Re-normaliser quelques variantes qui ont garde leur libelle original.
    ren = {c: canonical_header(c) for c in df.columns}
    df = df.rename(columns=ren)
    if "Article" not in df.columns:
        raise ValueError("Colonne Article introuvable.")
    if "NumCommande" not in df.columns:
        df["NumCommande"] = ""
    if "ResteALivrer" not in df.columns:
        if "QteCommandee" in df.columns:
            df["ResteALivrer"] = df["QteCommandee"]
        else:
            raise ValueError("Colonne ResteALivrer / QteCommandee introuvable.")
    return df, ws.title


def import_ax_bytes(data: bytes, filename: str, username: str) -> Dict[str, Any]:
    digest = sha256_bytes(data)
    latest = q_one("SELECT * FROM imports WHERE active=1 ORDER BY id DESC LIMIT 1")
    if latest and latest.get("sha256") == digest:
        return {"import_id": latest["id"], "rows": latest["row_count"], "sheet": latest["source_sheet"], "reused": True}

    df, sheet_name = read_ax_workbook(data)
    now = app_now().isoformat(timespec="seconds")
    valid = 0
    invalid = 0
    with db_conn() as conn:
        conn.execute("UPDATE imports SET active=0 WHERE active=1")
        conn.execute("UPDATE order_lines SET active=0 WHERE active=1")
        cur = conn.execute(
            "INSERT INTO imports(filename,sha256,source_sheet,imported_at,imported_by,row_count,valid_rows,invalid_rows,active) VALUES(?,?,?,?,?,?,?,?,1)",
            (filename, digest, sheet_name, now, username, int(len(df)), 0, 0),
        )
        import_id = int(cur.lastrowid)

        for idx, r in df.iterrows():
            article = norm_text(r.get("Article"))
            article_int = norm_text(r.get("ArticleInt"))
            color = norm_text(r.get("Couleur")).upper()
            parsed_int, parsed_color = split_article(article)
            if not article_int:
                article_int = parsed_int
            if not color:
                color = parsed_color
            remaining = max(0.0, to_float(r.get("ResteALivrer")))
            if not article or remaining <= 0:
                invalid += 1
                continue
            command = norm_text(r.get("NumCommande")).upper()
            if not command:
                command = f"STOCK-{import_id}-{int(idx)+1:05d}"
            unit_weight = max(0.0, to_float(r.get("PoidsUn")))
            bars = max(0, to_int(r.get("BarreBal"))) or infer_bars_per_bal(article_int, unit_weight)
            launch_src = max(0.0, to_float(r.get("Lancement")))
            relaq = max(0.0, to_float(r.get("ReLaquage")))
            launch = launch_src if launch_src > 0 else max(0.0, remaining - relaq)
            nbal_src = max(0.0, to_float(r.get("NbreBal")))
            nbal = nbal_src if nbal_src > 0 else (math.ceil(launch / bars) if launch > 0 and bars > 0 else 0)
            tps_src = max(0.0, to_float(r.get("TempsH")))
            tps = tps_src if tps_src > 0 else nbal * DEFAULT_MINUTES_PER_BAL / 60.0
            poids_t_src = max(0.0, to_float(r.get("PoidsT")))
            poids_t = poids_t_src if poids_t_src > 0 else launch * unit_weight
            poudre_src = max(0.0, to_float(r.get("Poudre")))
            poudre = poudre_src if poudre_src > 0 else poids_t * DEFAULT_POWDER_COEFF
            line_uid = hashlib.sha1(
                f"{digest}|{idx}|{command}|{article}|{norm_text(r.get('NumOF'))}".encode("utf-8")
            ).hexdigest()
            vals = (
                import_id, int(idx) + 2, line_uid, command, db_date(r.get("DateCreation")),
                norm_text(r.get("NomClient")), article, article_int, color,
                to_float(r.get("Nuance"), color_nuance(color) if color else 0.0),
                max(0.0, to_float(r.get("QteCommandee"))), remaining,
                max(0.0, to_float(r.get("Preleve"))), norm_text(r.get("ReservationBrut")),
                norm_text(r.get("NumOF")), norm_text(r.get("ProdStatut")),
                max(0.0, to_float(r.get("QteCommencee"))), max(0.0, to_float(r.get("QteRestante"))),
                max(0.0, to_float(r.get("QteRecue"))), max(0.0, to_float(r.get("ReserverBR"))),
                max(0.0, to_float(r.get("StockPhysique"))), max(0.0, to_float(r.get("Reserver"))),
                launch, relaq, unit_weight, poids_t, poudre, bars, nbal, tps,
                max(0.0, to_float(r.get("StockBrut"), to_float(r.get("ReserverBR")))),
                db_date(r.get("DateLivraison")), norm_text(r.get("Destination")),
                "A_PLANIFIER", 1, now,
            )
            conn.execute(
                """INSERT INTO order_lines(
                    import_id,source_row,line_uid,num_commande,date_creation,nom_client,article,article_int,couleur,nuance,
                    qte_commandee,reste_a_livrer,preleve,reservation_brut,num_of,prod_statut,qte_commencee,qte_restante,
                    qte_recue,reserver_br,stock_physique,reserver,lancement,relaquage,poids_un,poids_t,poudre,barre_bal,
                    nbre_bal,temps_h,stock_brut,date_livraison,destination,status,active,updated_at
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                vals,
            )
            valid += 1
        conn.execute("UPDATE imports SET valid_rows=?, invalid_rows=? WHERE id=?", (valid, invalid, import_id))

    audit(username, "IMPORT_AX", "import", import_id, None, {"filename": filename, "rows": len(df), "valid": valid, "invalid": invalid})
    notify("PLANNING", "AX_IMPORTED", "Nouvelle base AX importee", f"{filename}: {valid} lignes valides.", "import", import_id)
    return {"import_id": import_id, "rows": len(df), "valid": valid, "invalid": invalid, "sheet": sheet_name, "reused": False}


def active_import() -> Optional[Dict[str, Any]]:
    return q_one("SELECT * FROM imports WHERE active=1 ORDER BY id DESC LIMIT 1")


def load_active_orders_df() -> pd.DataFrame:
    rows = q_all("SELECT * FROM order_lines WHERE active=1 ORDER BY id")
    return pd.DataFrame(rows)

# =============================================================================
# 07. PLANNING ENGINE
# =============================================================================
@dataclass(frozen=True)
class PlannerConfig:
    year: int
    week: int
    capacity_h: float = DEFAULT_CAPACITY_H
    cleaning_min: int = DEFAULT_CLEANING_MIN
    minutes_per_bal: float = DEFAULT_MINUTES_PER_BAL
    powder_coeff: float = DEFAULT_POWDER_COEFF
    target_utilization: float = DEFAULT_TARGET_UTIL
    solver_seconds: float = DEFAULT_SOLVER_SECONDS
    max_colors_per_day: int = HARD_MAX_COLORS_PER_DAY
    strategy: str = "EQUILIBRE"
    force_commands: Tuple[str, ...] = ()
    exclude_commands: Tuple[str, ...] = ()


@dataclass
class Job:
    job_id: str
    line_ids: List[int]
    command: str
    color: str
    duration_min: int
    score: float
    due: Optional[date]
    relaquage_ids: List[int]


def planning_priority(row: pd.Series, week_start: date) -> Tuple[float, str, int]:
    score = 100.0
    reasons: List[str] = []
    due = parse_date(row.get("date_livraison"))
    overdue = 0
    if due is not None:
        delta = (due.date() - week_start).days
        if delta < 0:
            overdue = -delta
            score += 2200 + overdue * 260
            reasons.append(f"retard {overdue}j")
        elif delta <= 1:
            score += 1500
            reasons.append("echeance immediate")
        elif delta <= 5:
            score += 950 - delta * 80
            reasons.append("echeance semaine")
        elif delta <= 12:
            score += 300
            reasons.append("echeance proche")
    else:
        reasons.append("date non confirmee")

    created = parse_date(row.get("date_creation"))
    if created is not None:
        age = max(0, (week_start - created.date()).days)
        score += min(650, age * 8)
        if age > 30:
            reasons.append("commande ancienne")

    remaining = max(1.0, to_float(row.get("reste_a_livrer"), 1.0))
    ready = 0.0
    ps = norm_key(row.get("prod_statut"))
    if "commenc" in ps:
        ready += 280
    elif "cree" in ps:
        ready += 220
    if norm_key(row.get("reservation_brut")) in {"oui", "yes", "1", "true"}:
        ready += 260
    ready += min(260, max(0.0, to_float(row.get("reserver_br"))) / remaining * 260)
    ready += min(100, max(0.0, to_float(row.get("stock_physique"))) / remaining * 100)
    ready += min(120, max(0.0, to_float(row.get("qte_recue"))) / remaining * 120)
    score += ready
    if ready >= 350:
        reasons.append("matiere/OF pret")
    elif ready < 100:
        reasons.append("preparation faible")
    return score, " · ".join(reasons[:4]), overdue


def prepare_planning_lines(cfg: PlannerConfig) -> pd.DataFrame:
    source = load_active_orders_df()
    if source.empty:
        return pd.DataFrame()
    start = date.fromisocalendar(cfg.year, cfg.week, 1)
    force = {norm_text(x).upper() for x in cfg.force_commands if norm_text(x)}
    exclude = {norm_text(x).upper() for x in cfg.exclude_commands if norm_text(x)}
    rows: List[Dict[str, Any]] = []

    # Mediane par article pour les poids absents.
    known = source[pd.to_numeric(source["poids_un"], errors="coerce").fillna(0) > 0].copy()
    article_weights = known.groupby("article_int")["poids_un"].median().to_dict() if not known.empty else {}

    for _, r in source.iterrows():
        cmd = norm_text(r.get("num_commande")).upper()
        if cmd in exclude:
            continue
        remaining = max(0.0, to_float(r.get("reste_a_livrer")))
        color = norm_text(r.get("couleur")).upper()
        if remaining <= 0 or not color or color == "BRUT":
            continue
        art_int = norm_text(r.get("article_int"))
        unit_weight = max(0.0, to_float(r.get("poids_un")))
        if unit_weight <= 0:
            unit_weight = max(0.0, to_float(article_weights.get(art_int), 0.0))
        bars = max(1, to_int(r.get("barre_bal"), 0) or infer_bars_per_bal(art_int, unit_weight))
        relaq = max(0.0, to_float(r.get("relaquage")))
        launch = max(0.0, to_float(r.get("lancement"))) or max(0.0, remaining - relaq)
        nbal = max(0, to_int(r.get("nbre_bal"), 0)) or int(math.ceil(launch / bars))
        duration_h = max(0.0, to_float(r.get("temps_h"))) or nbal * cfg.minutes_per_bal / 60.0
        poids_t = max(0.0, to_float(r.get("poids_t"))) or launch * unit_weight
        poudre = max(0.0, to_float(r.get("poudre"))) or poids_t * cfg.powder_coeff
        score, reason, overdue = planning_priority(r, start)
        if cmd in force:
            score += 100000
            reason = "FORCE · " + reason
        rows.append({
            "_kind": "ORDER", "_order_line_id": int(r["id"]), "_relaquage_id": None,
            "NumCommande": cmd, "DateCreation": r.get("date_creation"), "NomClient": norm_text(r.get("nom_client")),
            "Article": norm_text(r.get("article")), "ArticleInt": art_int, "Couleur": color,
            "Nuance": to_float(r.get("nuance"), color_nuance(color)), "QteCommandee": to_float(r.get("qte_commandee")),
            "ResteALivrer": remaining, "Preleve": to_float(r.get("preleve")), "ReservationBrut": norm_text(r.get("reservation_brut")),
            "NumOF": norm_text(r.get("num_of")), "ProdStatut": norm_text(r.get("prod_statut")),
            "QteCommencee": to_float(r.get("qte_commencee")), "QteRestante": to_float(r.get("qte_restante")),
            "QteRecue": to_float(r.get("qte_recue")), "ReserverBR": to_float(r.get("reserver_br")),
            "StockPhysique": to_float(r.get("stock_physique")), "Reserver": to_float(r.get("reserver")),
            "Lancement": launch, "ReLaquage": relaq, "PoidsUn": unit_weight, "PoidsT": poids_t,
            "Poudre": poudre, "BarreBal": bars, "NbreBal": nbal, "TempsH": duration_h,
            "StockBrut": to_float(r.get("stock_brut")), "DateLivraison": r.get("date_livraison"),
            "Destination": norm_text(r.get("destination")), "_score": score, "_reason": reason,
            "_overdue_days": overdue, "_due": db_date(r.get("date_livraison")),
        })

    # Re-laquages ouverts reinjectes avec priorite forte.
    relaqs = q_all("""
        SELECT r.*, o.nom_client, o.destination, o.poids_un, o.barre_bal
        FROM relaquage_orders r
        LEFT JOIN order_lines o ON o.id=r.order_line_id
        WHERE r.status='A_PLANIFIER'
        ORDER BY r.id
    """)
    for r in relaqs:
        qty = max(0.0, to_float(r.get("qty")))
        if qty <= 0:
            continue
        weight = max(0.0, to_float(r.get("poids_un")))
        bars = max(1, to_int(r.get("barre_bal"), 0) or infer_bars_per_bal(norm_text(r.get("article_int")), weight))
        nbal = int(math.ceil(qty / bars))
        duration_h = nbal * cfg.minutes_per_bal / 60.0
        poids_t = qty * weight
        color = norm_text(r.get("couleur")).upper()
        rows.append({
            "_kind": "RELAQUAGE", "_order_line_id": r.get("order_line_id"), "_relaquage_id": int(r["id"]),
            "NumCommande": norm_text(r.get("original_command")), "DateCreation": r.get("created_at"),
            "NomClient": norm_text(r.get("nom_client")), "Article": norm_text(r.get("article")),
            "ArticleInt": norm_text(r.get("article_int")), "Couleur": color, "Nuance": color_nuance(color),
            "QteCommandee": qty, "ResteALivrer": qty, "Preleve": 0, "ReservationBrut": "RELAQUAGE",
            "NumOF": norm_text(r.get("num_of")), "ProdStatut": "RELAQUAGE", "QteCommencee": 0,
            "QteRestante": qty, "QteRecue": qty, "ReserverBR": qty, "StockPhysique": qty,
            "Reserver": qty, "Lancement": 0, "ReLaquage": qty, "PoidsUn": weight,
            "PoidsT": poids_t, "Poudre": poids_t * cfg.powder_coeff, "BarreBal": bars,
            "NbreBal": nbal, "TempsH": duration_h, "StockBrut": qty,
            "DateLivraison": r.get("requested_date"), "Destination": norm_text(r.get("destination")),
            "_score": 5000.0 + (1000 if norm_text(r.get("priority")).upper() == "HAUTE" else 0),
            "_reason": "re-laquage qualite prioritaire", "_overdue_days": 0,
            "_due": r.get("requested_date"),
        })

    out = pd.DataFrame(rows)
    if not out.empty:
        out = out.sort_values(["_score", "DateLivraison", "NumCommande"], ascending=[False, True, True], na_position="last").reset_index(drop=True)
    return out


def build_jobs(lines: pd.DataFrame, cfg: PlannerConfig) -> Tuple[List[Job], List[int]]:
    if lines.empty:
        return [], []
    max_min = int(round(cfg.capacity_h * 60))
    jobs: List[Job] = []
    oversized: List[int] = []
    for (cmd, color), g in lines.groupby(["NumCommande", "Couleur"], sort=False):
        current: List[int] = []
        current_min = 0
        chunk = 1
        for idx, r in g.iterrows():
            mins = max(1, int(round(to_float(r.get("TempsH")) * 60)))
            if mins > max_min:
                oversized.append(int(idx))
                continue
            if current and current_min + mins > max_min:
                sub = lines.loc[current]
                dues = [parse_date(x) for x in sub["_due"].tolist() if norm_text(x)]
                jobs.append(Job(
                    f"{cmd}|{color}|{chunk}", current.copy(), str(cmd), str(color), current_min,
                    float(sub["_score"].max()), min([x.date() for x in dues if x is not None], default=None),
                    [int(x) for x in sub["_relaquage_id"].dropna().tolist()],
                ))
                chunk += 1
                current, current_min = [], 0
            current.append(int(idx))
            current_min += mins
        if current:
            sub = lines.loc[current]
            dues = [parse_date(x) for x in sub["_due"].tolist() if norm_text(x)]
            jobs.append(Job(
                f"{cmd}|{color}|{chunk}", current.copy(), str(cmd), str(color), current_min,
                float(sub["_score"].max()), min([x.date() for x in dues if x is not None], default=None),
                [int(x) for x in sub["_relaquage_id"].dropna().tolist()],
            ))
    return jobs, oversized


def assign_jobs_ortools(jobs: List[Job], cfg: PlannerConfig) -> Tuple[Dict[str, int], List[str], str]:
    if not ORTOOLS_AVAILABLE or not jobs:
        return {}, [j.job_id for j in jobs], ""
    dates = iso_week_dates(cfg.year, cfg.week)
    colors = sorted({j.color for j in jobs})
    model = cp_model.CpModel()
    x: Dict[Tuple[int, int], Any] = {}
    u: Dict[int, Any] = {}
    y: Dict[Tuple[int, str], Any] = {}
    objective: List[Any] = []

    for ji, job in enumerate(jobs):
        u[ji] = model.NewBoolVar(f"u_{ji}")
        for d in range(5):
            x[(ji, d)] = model.NewBoolVar(f"x_{ji}_{d}")
        model.Add(sum(x[(ji, d)] for d in range(5)) + u[ji] == 1)
        uns_pen = max(10000, int(30000 + job.score * 120 + job.duration_min * 6))
        objective.append(u[ji] * uns_pen)
        if job.due:
            for d in range(5):
                late = max(0, (dates[d] - job.due).days)
                if late:
                    objective.append(x[(ji, d)] * late * 90000)

    for d in range(5):
        active_colors = []
        for color in colors:
            y[(d, color)] = model.NewBoolVar(f"y_{d}_{norm_key(color)}")
            idxs = [i for i, j in enumerate(jobs) if j.color == color]
            for ji in idxs:
                model.Add(x[(ji, d)] <= y[(d, color)])
            model.Add(y[(d, color)] <= sum(x[(ji, d)] for ji in idxs))
            active_colors.append(y[(d, color)])
            objective.append(y[(d, color)] * 2500)
        n_colors = sum(active_colors)
        model.Add(n_colors <= cfg.max_colors_per_day)
        has_color = model.NewBoolVar(f"has_color_{d}")
        model.Add(n_colors >= has_color)
        model.Add(n_colors <= cfg.max_colors_per_day * has_color)
        extra = model.NewIntVar(0, max(0, cfg.max_colors_per_day - 1), f"extra_{d}")
        model.Add(extra == n_colors - has_color)
        prod = sum(jobs[ji].duration_min * x[(ji, d)] for ji in range(len(jobs)))
        cap = int(round(cfg.capacity_h * 60))
        model.Add(prod + cfg.cleaning_min * extra <= cap)
        target = int(round(cap * cfg.target_utilization))
        load = model.NewIntVar(0, cap, f"load_{d}")
        model.Add(load == prod + cfg.cleaning_min * extra)
        dev = model.NewIntVar(0, cap, f"dev_{d}")
        model.Add(dev >= target - load)
        model.Add(dev >= load - target)
        objective.append(dev * 3)

    white = [c for c in colors if color_class(c) == "WHITE"]
    black = [c for c in colors if color_class(c) == "BLACK"]
    for d in range(5):
        for wc in white:
            for bc in black:
                model.Add(y[(d, wc)] + y[(d, bc)] <= 1)
    for d in range(4):
        for wc in white:
            for bc in black:
                model.Add(y[(d, wc)] + y[(d + 1, bc)] <= 1)
                model.Add(y[(d, bc)] + y[(d + 1, wc)] <= 1)

    model.Minimize(sum(objective))
    solver = cp_model.CpSolver()
    solver.parameters.max_time_in_seconds = max(3.0, cfg.solver_seconds)
    solver.parameters.num_search_workers = max(1, min(8, os.cpu_count() or 2))
    solver.parameters.random_seed = 8
    status = solver.Solve(model)
    if status not in (cp_model.OPTIMAL, cp_model.FEASIBLE):
        return {}, [j.job_id for j in jobs], ""
    assignments: Dict[str, int] = {}
    uns: List[str] = []
    for ji, job in enumerate(jobs):
        if solver.Value(u[ji]):
            uns.append(job.job_id)
            continue
        for d in range(5):
            if solver.Value(x[(ji, d)]):
                assignments[job.job_id] = d
                break
        if job.job_id not in assignments:
            uns.append(job.job_id)
    return assignments, uns, "OR-Tools CP-SAT"


def assign_jobs_fallback(jobs: List[Job], cfg: PlannerConfig) -> Tuple[Dict[str, int], List[str], str]:
    dates = iso_week_dates(cfg.year, cfg.week)
    loads = [0] * 5
    colors: List[set] = [set() for _ in range(5)]
    assignments: Dict[str, int] = {}
    uns: List[str] = []
    cap = int(round(cfg.capacity_h * 60))
    ordered = sorted(jobs, key=lambda j: (-j.score, j.due or date.max, -j.duration_min))
    for job in ordered:
        options = []
        for d in range(5):
            new_colors = set(colors[d]) | {job.color}
            if len(new_colors) > cfg.max_colors_per_day or white_black_conflict(new_colors):
                continue
            if d > 0 and white_black_conflict(new_colors, colors[d - 1]):
                continue
            if d < 4 and white_black_conflict(new_colors, colors[d + 1]):
                continue
            clean_before = cfg.cleaning_min * max(0, len(colors[d]) - 1)
            clean_after = cfg.cleaning_min * max(0, len(new_colors) - 1)
            prod_before = loads[d] - clean_before
            projected = prod_before + job.duration_min + clean_after
            if projected > cap:
                continue
            late = max(0, (dates[d] - job.due).days) if job.due else 0
            new_color_pen = 15000 if job.color not in colors[d] and colors[d] else 0
            options.append((late * 100000 + new_color_pen + abs(cap * cfg.target_utilization - projected) - job.score * 5, d, projected))
        if not options:
            uns.append(job.job_id)
            continue
        _, d, projected = min(options, key=lambda x: (x[0], x[1]))
        assignments[job.job_id] = d
        colors[d].add(job.color)
        loads[d] = int(projected)
    return assignments, uns, "Heuristique robuste"


def generate_plan(cfg: PlannerConfig) -> Dict[str, Any]:
    t0 = time.perf_counter()
    lines = prepare_planning_lines(cfg)
    if lines.empty:
        raise ValueError("Aucune ligne eligible a planifier.")
    jobs, oversized = build_jobs(lines, cfg)
    assignments, uns_ids, engine = assign_jobs_ortools(jobs, cfg)
    if not engine:
        assignments, uns_ids, engine = assign_jobs_fallback(jobs, cfg)
    job_map = {j.job_id: j for j in jobs}
    day_indices: Dict[int, List[int]] = {d: [] for d in range(5)}
    for jid, d in assignments.items():
        day_indices[d].extend(job_map[jid].line_ids)
    uns_idx: List[int] = []
    for jid in uns_ids:
        if jid in job_map:
            uns_idx.extend(job_map[jid].line_ids)
    uns_idx.extend(oversized)
    uns_idx = list(dict.fromkeys(uns_idx))
    unscheduled = lines.loc[uns_idx].copy() if uns_idx else lines.iloc[0:0].copy()

    days: Dict[int, pd.DataFrame] = {}
    metrics: List[Dict[str, Any]] = []
    hard: List[str] = []
    dates = iso_week_dates(cfg.year, cfg.week)
    for d in range(5):
        df = lines.loc[day_indices[d]].copy() if day_indices[d] else lines.iloc[0:0].copy()
        if not df.empty:
            df["_color_rank"] = df["Couleur"].map(lambda c: color_nuance(c))
            df = df.sort_values(["_color_rank", "DateLivraison", "_score", "NumCommande"], ascending=[True, True, False, True], na_position="last").drop(columns=["_color_rank"]).reset_index(drop=True)
            df["_planned_date"] = dates[d].isoformat()
        days[d] = df
        colors = list(dict.fromkeys(df["Couleur"].tolist())) if not df.empty else []
        prod_h = float(pd.to_numeric(df.get("TempsH", pd.Series(dtype=float)), errors="coerce").fillna(0).sum())
        cleaning_h = cfg.cleaning_min * max(0, len(colors) - 1) / 60.0
        total_h = prod_h + cleaning_h
        if len(colors) > cfg.max_colors_per_day:
            hard.append(f"{DAYS[d]}: trop de couleurs.")
        if total_h > cfg.capacity_h + 1e-6:
            hard.append(f"{DAYS[d]}: surcharge {total_h:.2f} h > {cfg.capacity_h:.2f} h.")
        if white_black_conflict(colors):
            hard.append(f"{DAYS[d]}: conflit BLANC / NOIR-DARK.")
        metrics.append({
            "Jour": DAYS[d], "Date": dates[d].strftime("%d/%m/%Y"), "Charge production h": round(prod_h, 2),
            "Nettoyage h": round(cleaning_h, 2), "Charge totale h": round(total_h, 2),
            "Capacite h": cfg.capacity_h, "Charge %": round(total_h / cfg.capacity_h * 100, 1) if cfg.capacity_h else 0,
            "Couleurs": " -> ".join(colors), "Nb couleurs": len(colors), "Lignes": len(df),
        })
    for d in range(4):
        c1 = set(days[d]["Couleur"].tolist()) if not days[d].empty else set()
        c2 = set(days[d + 1]["Couleur"].tolist()) if not days[d + 1].empty else set()
        if white_black_conflict(c1, c2):
            hard.append(f"Transition interdite {DAYS[d]} -> {DAYS[d+1]}: BLANC / NOIR-DARK.")
    total_load = sum(x["Charge totale h"] for x in metrics)
    capacity = cfg.capacity_h * 5
    return {
        "config": cfg, "lines": lines, "days": days, "unscheduled": unscheduled,
        "engine": engine, "hard_errors": list(dict.fromkeys(hard)),
        "confidence": 100 if not hard else 0,
        "metrics": {
            "days": metrics, "total_load_h": round(total_load, 2), "capacity_h": round(capacity, 2),
            "utilization_pct": round(total_load / capacity * 100, 1) if capacity else 0,
            "backlog": int(len(unscheduled)),
        },
        "elapsed_s": round(time.perf_counter() - t0, 3),
    }

# =============================================================================
# 08. PLANNING VERSIONING / PUBLICATION
# =============================================================================
def next_planning_version_no(year: int, week: int) -> int:
    row = q_one("SELECT COALESCE(MAX(version_no),0)+1 AS n FROM planning_versions WHERE year=? AND week=?", (year, week))
    return int(row["n"] if row else 1)


def current_published_version(year: Optional[int] = None, week: Optional[int] = None) -> Optional[Dict[str, Any]]:
    if year is not None and week is not None:
        return q_one(
            "SELECT * FROM planning_versions WHERE year=? AND week=? AND status='PUBLIE' ORDER BY version_no DESC LIMIT 1",
            (year, week),
        )
    return q_one("SELECT * FROM planning_versions WHERE status='PUBLIE' ORDER BY published_at DESC,id DESC LIMIT 1")


def published_entries(version_id: int) -> pd.DataFrame:
    rows = q_all("SELECT * FROM planning_entries WHERE planning_version_id=? ORDER BY planned_date,sequence_no,id", (version_id,))
    return pd.DataFrame(rows)


def publish_plan(result: Dict[str, Any], username: str, notes: str = "") -> int:
    cfg: PlannerConfig = result["config"]
    if result.get("hard_errors"):
        raise ValueError("Publication impossible: le planning contient une erreur bloquante.")
    imp = active_import()
    if not imp:
        raise ValueError("Aucune base AX active.")
    version_no = next_planning_version_no(cfg.year, cfg.week)
    now = app_now().isoformat(timespec="seconds")
    previous = current_published_version(cfg.year, cfg.week)
    with db_conn() as conn:
        if previous:
            conn.execute("UPDATE planning_versions SET status='ARCHIVE' WHERE id=?", (previous["id"],))
        cur = conn.execute(
            """INSERT INTO planning_versions(year,week,version_no,status,source_import_id,created_at,created_by,published_at,notes,parent_version_id)
               VALUES(?,?,?,?,?,?,?,?,?,?)""",
            (cfg.year, cfg.week, version_no, "PUBLIE", imp["id"], now, username, now, notes, previous["id"] if previous else None),
        )
        version_id = int(cur.lastrowid)
        for d in range(5):
            df = result["days"].get(d, pd.DataFrame())
            if df is None or df.empty:
                continue
            planned_date = date.fromisocalendar(cfg.year, cfg.week, d + 1).isoformat()
            for seq, (_, r) in enumerate(df.iterrows(), start=1):
                cur_e = conn.execute(
                    """INSERT INTO planning_entries(
                        planning_version_id,order_line_id,relaquage_id,planned_date,sequence_no,num_commande,nom_client,
                        article,article_int,couleur,num_of,planned_qty,poids_t,poudre,nbre_bal,duration_h,
                        priority_score,decision_reason,status
                    ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (
                        version_id, int(r["_order_line_id"]) if pd.notna(r.get("_order_line_id")) else None,
                        int(r["_relaquage_id"]) if pd.notna(r.get("_relaquage_id")) else None,
                        planned_date, seq, norm_text(r.get("NumCommande")), norm_text(r.get("NomClient")),
                        norm_text(r.get("Article")), norm_text(r.get("ArticleInt")), norm_text(r.get("Couleur")),
                        norm_text(r.get("NumOF")), to_float(r.get("Lancement")) + to_float(r.get("ReLaquage")),
                        to_float(r.get("PoidsT")), to_float(r.get("Poudre")), to_float(r.get("NbreBal")),
                        to_float(r.get("TempsH")), to_float(r.get("_score")), norm_text(r.get("_reason")), "PLANIFIE",
                    ),
                )
                entry_id = int(cur_e.lastrowid)
                if pd.notna(r.get("_relaquage_id")):
                    conn.execute(
                        "UPDATE relaquage_orders SET status='PLANIFIE',planned_date=?,updated_at=? WHERE id=?",
                        (planned_date, now, int(r["_relaquage_id"])),
                    )
                else:
                    prep_date = subtract_business_days(date.fromisoformat(planned_date), 2).isoformat()
                    required = to_float(r.get("Lancement"))
                    reserved = min(required, max(to_float(r.get("ReserverBR")), to_float(r.get("QteRecue"))))
                    missing = max(0.0, required - reserved)
                    status = "A_PREPARER" if missing <= 0 else "BLOQUE"
                    conn.execute(
                        """INSERT INTO magasin_preparations(planning_entry_id,prep_date,status,required_qty,reserved_qty,missing_qty,location,updated_at,updated_by)
                           VALUES(?,?,?,?,?,?,?,?,?)""",
                        (entry_id, prep_date, status, required, reserved, missing, "", now, username),
                    )
                    if r.get("_order_line_id"):
                        conn.execute(
                            "UPDATE order_lines SET status='PLANIFIE',updated_at=? WHERE id=?",
                            (now, int(r["_order_line_id"])),
                        )
    audit(username, "PUBLISH_PLAN", "planning_version", version_id, previous, {"year": cfg.year, "week": cfg.week, "version": version_no}, notes)
    notify("MAGASIN", "PLANNING_PUBLISHED", f"Planning S{cfg.week} publie", "Les preparations J-2 ont ete generees.", "planning_version", version_id)
    notify("LAQUAGE", "PLANNING_PUBLISHED", f"Planning S{cfg.week} publie", f"Version {version_no} disponible.", "planning_version", version_id)
    return version_id


def revise_planning_entry(version_id: int, entry_id: int, new_date: date, username: str, reason: str) -> int:
    old_version = q_one("SELECT * FROM planning_versions WHERE id=? AND status='PUBLIE'", (version_id,))
    old_entry = q_one("SELECT * FROM planning_entries WHERE id=? AND planning_version_id=?", (entry_id, version_id))
    if not old_version or not old_entry:
        raise ValueError("Version ou ligne de planning introuvable.")
    if not reason.strip():
        raise ValueError("Le motif de modification est obligatoire.")
    version_no = next_planning_version_no(int(old_version["year"]), int(old_version["week"]))
    now = app_now().isoformat(timespec="seconds")
    entries = q_all("SELECT * FROM planning_entries WHERE planning_version_id=? ORDER BY planned_date,sequence_no,id", (version_id,))
    with db_conn() as conn:
        conn.execute("UPDATE planning_versions SET status='ARCHIVE' WHERE id=?", (version_id,))
        cur = conn.execute(
            """INSERT INTO planning_versions(year,week,version_no,status,source_import_id,created_at,created_by,published_at,notes,parent_version_id)
               VALUES(?,?,?,?,?,?,?,?,?,?)""",
            (old_version["year"], old_version["week"], version_no, "PUBLIE", old_version["source_import_id"], now, username, now, reason, version_id),
        )
        new_vid = int(cur.lastrowid)
        new_entry_id_for_changed = None
        for e in entries:
            planned_date = new_date.isoformat() if int(e["id"]) == entry_id else e["planned_date"]
            cur_e = conn.execute(
                """INSERT INTO planning_entries(planning_version_id,order_line_id,relaquage_id,planned_date,sequence_no,num_commande,nom_client,
                    article,article_int,couleur,num_of,planned_qty,poids_t,poudre,nbre_bal,duration_h,priority_score,decision_reason,status)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    new_vid, e["order_line_id"], e["relaquage_id"], planned_date, e["sequence_no"], e["num_commande"], e["nom_client"],
                    e["article"], e["article_int"], e["couleur"], e["num_of"], e["planned_qty"], e["poids_t"], e["poudre"],
                    e["nbre_bal"], e["duration_h"], e["priority_score"], e["decision_reason"], e["status"],
                ),
            )
            new_eid = int(cur_e.lastrowid)
            if int(e["id"]) == entry_id:
                new_entry_id_for_changed = new_eid
            if e["relaquage_id"]:
                conn.execute("UPDATE relaquage_orders SET planned_date=?,updated_at=? WHERE id=?", (planned_date, now, e["relaquage_id"]))
            else:
                old_prep = q_one("SELECT * FROM magasin_preparations WHERE planning_entry_id=?", (e["id"],))
                prep_date = subtract_business_days(date.fromisoformat(planned_date), 2).isoformat()
                required = to_float(e["planned_qty"])
                reserved = min(required, to_float(old_prep["reserved_qty"]) if old_prep else 0.0)
                missing = max(0.0, required - reserved)
                status = (old_prep["status"] if old_prep else "A_PREPARER")
                if missing > 0:
                    status = "BLOQUE"
                conn.execute(
                    """INSERT INTO magasin_preparations(planning_entry_id,prep_date,status,required_qty,reserved_qty,missing_qty,location,updated_at,updated_by)
                       VALUES(?,?,?,?,?,?,?,?,?)""",
                    (new_eid, prep_date, status, required, reserved, missing, (old_prep or {}).get("location", ""), now, username),
                )
    audit(username, "REVISE_PLAN", "planning_entry", entry_id, old_entry, {"new_version": new_vid, "new_date": new_date.isoformat()}, reason)
    notify("MAGASIN", "PLANNING_CHANGED", "Planning modifie", f"{old_entry['num_commande']} passe du {old_entry['planned_date']} au {new_date.isoformat()}.", "planning_version", new_vid)
    notify("LAQUAGE", "PLANNING_CHANGED", "Planning modifie", f"Nouvelle version {version_no} publiee.", "planning_version", new_vid)
    return new_vid

# =============================================================================
# 09. MAGASIN J-2
# =============================================================================
def magasin_rows() -> pd.DataFrame:
    v = current_published_version()
    if not v:
        return pd.DataFrame()
    rows = q_all("""
        SELECT m.*, e.num_commande,e.nom_client,e.article,e.article_int,e.couleur,e.num_of,e.planned_date,e.status AS planning_status
        FROM magasin_preparations m
        JOIN planning_entries e ON e.id=m.planning_entry_id
        WHERE e.planning_version_id=?
        ORDER BY m.prep_date,e.planned_date,e.sequence_no
    """, (v["id"],))
    return pd.DataFrame(rows)


def update_magasin(prep_id: int, status: str, reserved_qty: float, location: str, username: str) -> None:
    old = q_one("SELECT * FROM magasin_preparations WHERE id=?", (prep_id,))
    if not old:
        raise ValueError("Preparation introuvable.")
    required = to_float(old["required_qty"])
    reserved = max(0.0, min(required, float(reserved_qty)))
    missing = max(0.0, required - reserved)
    if status == "PRET" and missing > 0:
        raise ValueError("Impossible de declarer PRET avec une quantite manquante.")
    if missing > 0:
        status = "BLOQUE"
    now = app_now().isoformat(timespec="seconds")
    with db_conn() as conn:
        conn.execute(
            "UPDATE magasin_preparations SET status=?,reserved_qty=?,missing_qty=?,location=?,updated_at=?,updated_by=? WHERE id=?",
            (status, reserved, missing, norm_text(location), now, username, prep_id),
        )
        entry = conn.execute("SELECT * FROM planning_entries WHERE id=?", (old["planning_entry_id"],)).fetchone()
        if entry:
            entry_status = "MATIERE_PRETE" if status == "PRET" else ("BLOQUE" if status == "BLOQUE" else "MATIERE_A_PREPARER")
            conn.execute("UPDATE planning_entries SET status=? WHERE id=?", (entry_status, entry["id"]))
            if entry["order_line_id"]:
                conn.execute("UPDATE order_lines SET status=?,updated_at=? WHERE id=?", (entry_status, now, entry["order_line_id"]))
    audit(username, "UPDATE_MAGASIN", "magasin_preparation", prep_id, old, {"status": status, "reserved_qty": reserved, "missing_qty": missing, "location": location})
    if missing > 0:
        notify("PLANNING", "MATERIAL_SHORTAGE", "Manque matiere", f"Preparation #{prep_id}: manque {missing:.0f} unite(s).", "magasin_preparation", prep_id)

# =============================================================================
# 10. LAQUAGE / QUALITE / RE-LAQUAGE
# =============================================================================
def laquage_rows() -> pd.DataFrame:
    v = current_published_version()
    if not v:
        return pd.DataFrame()
    rows = q_all("""
        SELECT e.*, m.status AS magasin_status,m.missing_qty,m.location
        FROM planning_entries e
        LEFT JOIN magasin_preparations m ON m.planning_entry_id=e.id
        WHERE e.planning_version_id=?
        ORDER BY e.planned_date,e.sequence_no
    """, (v["id"],))
    return pd.DataFrame(rows)


def update_laquage_status(entry_id: int, new_status: str, username: str) -> None:
    if new_status not in {"MATIERE_PRETE", "EN_COURS_LAQUAGE", "CONTROLE_QUALITE", "BLOQUE"}:
        raise ValueError("Statut laquage invalide.")
    old = q_one("SELECT * FROM planning_entries WHERE id=?", (entry_id,))
    if not old:
        raise ValueError("Ligne planning introuvable.")
    if new_status == "EN_COURS_LAQUAGE" and old["status"] not in {"MATIERE_PRETE", "PLANIFIE"}:
        raise ValueError("La matiere doit etre prete avant demarrage.")
    now = app_now().isoformat(timespec="seconds")
    with db_conn() as conn:
        conn.execute("UPDATE planning_entries SET status=? WHERE id=?", (new_status, entry_id))
        if old["order_line_id"]:
            conn.execute("UPDATE order_lines SET status=?,updated_at=? WHERE id=?", (new_status, now, old["order_line_id"]))
        if old["relaquage_id"]:
            relaq_status = "EN_COURS" if new_status == "EN_COURS_LAQUAGE" else ("CONTROLE" if new_status == "CONTROLE_QUALITE" else "PLANIFIE")
            conn.execute("UPDATE relaquage_orders SET status=?,updated_at=? WHERE id=?", (relaq_status, now, old["relaquage_id"]))
    audit(username, "UPDATE_LAQUAGE", "planning_entry", entry_id, old, {"status": new_status})
    if new_status == "CONTROLE_QUALITE":
        notify("QUALITE", "QUALITY_REQUIRED", "Controle qualite requis", f"{old['num_commande']} / {old['article_int']} attend le controle.", "planning_entry", entry_id)


def quality_pending_rows() -> pd.DataFrame:
    v = current_published_version()
    if not v:
        return pd.DataFrame()
    rows = q_all("""
        SELECT * FROM planning_entries
        WHERE planning_version_id=? AND status IN ('CONTROLE_QUALITE','RE_LAQUAGE')
        ORDER BY planned_date,sequence_no
    """, (v["id"],))
    return pd.DataFrame(rows)


def record_quality(entry_id: int, result: str, defect_reason: str, qty_affected: float, comment: str, username: str) -> int:
    result = norm_text(result).upper()
    if result not in {"CONFORME", "NON_CONFORME"}:
        raise ValueError("Resultat qualite invalide.")
    entry = q_one("SELECT * FROM planning_entries WHERE id=?", (entry_id,))
    if not entry:
        raise ValueError("Ligne planning introuvable.")
    now = app_now().isoformat(timespec="seconds")
    affected = max(0.0, min(to_float(entry["planned_qty"]), float(qty_affected)))
    if result == "NON_CONFORME" and affected <= 0:
        raise ValueError("Quantite affectee obligatoire pour une non-conformite.")
    with db_conn() as conn:
        cur = conn.execute(
            """INSERT INTO quality_incidents(planning_entry_id,order_line_id,inspected_at,inspector,result,defect_reason,qty_affected,comment,status)
               VALUES(?,?,?,?,?,?,?,?,?)""",
            (entry_id, entry["order_line_id"], now, username, result, norm_text(defect_reason), affected, norm_text(comment), "CLOTURE" if result == "CONFORME" else "OUVERT"),
        )
        incident_id = int(cur.lastrowid)
        if result == "CONFORME":
            new_status = "PRET_EXPEDITION"
            conn.execute("UPDATE planning_entries SET status=? WHERE id=?", (new_status, entry_id))
            if entry["order_line_id"]:
                conn.execute("UPDATE order_lines SET status=?,updated_at=? WHERE id=?", (new_status, now, entry["order_line_id"]))
            if entry["relaquage_id"]:
                conn.execute("UPDATE relaquage_orders SET status='CONFORME',updated_at=? WHERE id=?", (now, entry["relaquage_id"]))
        else:
            new_status = "RE_LAQUAGE"
            conn.execute("UPDATE planning_entries SET status=? WHERE id=?", (new_status, entry_id))
            if entry["order_line_id"]:
                conn.execute("UPDATE order_lines SET status=?,updated_at=? WHERE id=?", (new_status, now, entry["order_line_id"]))
            if entry["relaquage_id"]:
                # Un re-laquage peut lui-meme etre refuse: on cree un nouvel ordre lie a la meme ligne.
                pass
            conn.execute(
                """INSERT INTO relaquage_orders(quality_incident_id,order_line_id,original_command,article,article_int,couleur,num_of,qty,reason,priority,status,requested_date,planned_date,created_at,updated_at)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    incident_id, entry["order_line_id"], entry["num_commande"], entry["article"], entry["article_int"],
                    entry["couleur"], entry["num_of"], affected, norm_text(defect_reason), "HAUTE", "A_PLANIFIER",
                    (app_today() + timedelta(days=1)).isoformat(), None, now, now,
                ),
            )
    audit(username, "QUALITY_INSPECTION", "planning_entry", entry_id, entry, {"result": result, "defect_reason": defect_reason, "qty": affected}, comment)
    if result == "CONFORME":
        notify("LOGISTIQUE", "ORDER_READY_TO_SHIP", "Produit pret a expedier", f"{entry['num_commande']} est conforme.", "planning_entry", entry_id)
    else:
        notify("PLANNING", "RELAQUAGE_CREATED", "Re-laquage cree", f"{entry['num_commande']}: {affected:.0f} unite(s) a re-laquer.", "quality_incident", incident_id)
        notify("LAQUAGE", "RELAQUAGE_CREATED", "Re-laquage cree", f"Defaut: {defect_reason}.", "quality_incident", incident_id)
    return incident_id


def relaquage_rows() -> pd.DataFrame:
    return pd.DataFrame(q_all("SELECT * FROM relaquage_orders ORDER BY CASE status WHEN 'A_PLANIFIER' THEN 0 WHEN 'PLANIFIE' THEN 1 ELSE 2 END,id DESC"))

# =============================================================================
# 11. LOGISTIQUE
# =============================================================================
def add_truck(code: str, plate: str, capacity_kg: float, username: str) -> int:
    if capacity_kg <= 0:
        raise ValueError("Capacite camion invalide.")
    with db_conn() as conn:
        cur = conn.execute("INSERT INTO trucks(code,plate,capacity_kg,active) VALUES(?,?,?,1)", (norm_text(code).upper(), norm_text(plate).upper(), float(capacity_kg)))
        tid = int(cur.lastrowid)
    audit(username, "CREATE_TRUCK", "truck", tid, None, {"code": code, "plate": plate, "capacity_kg": capacity_kg})
    return tid


def add_driver(name: str, phone: str, username: str) -> int:
    with db_conn() as conn:
        cur = conn.execute("INSERT INTO drivers(name,phone,active) VALUES(?,?,1)", (norm_text(name), norm_text(phone)))
        did = int(cur.lastrowid)
    audit(username, "CREATE_DRIVER", "driver", did, None, {"name": name, "phone": phone})
    return did


def ready_for_shipping() -> pd.DataFrame:
    v = current_published_version()
    if not v:
        return pd.DataFrame()
    rows = q_all("""
        SELECT e.*, o.destination
        FROM planning_entries e
        LEFT JOIN order_lines o ON o.id=e.order_line_id
        WHERE e.planning_version_id=? AND e.status='PRET_EXPEDITION'
          AND NOT EXISTS (SELECT 1 FROM shipments s WHERE s.planning_entry_id=e.id AND s.status NOT IN ('ANNULE','LIVRE'))
        ORDER BY COALESCE(o.destination,''),e.num_commande,e.id
    """, (v["id"],))
    return pd.DataFrame(rows)


def create_trip(trip_date: date, truck_id: int, driver_id: Optional[int], destination: str,
                entry_ids: Sequence[int], username: str) -> int:
    if not entry_ids:
        raise ValueError("Selectionnez au moins une commande.")
    truck = q_one("SELECT * FROM trucks WHERE id=? AND active=1", (truck_id,))
    if not truck:
        raise ValueError("Camion introuvable.")
    placeholders = ",".join("?" for _ in entry_ids)
    rows = q_all(f"SELECT e.*,o.destination FROM planning_entries e LEFT JOIN order_lines o ON o.id=e.order_line_id WHERE e.id IN ({placeholders})", tuple(entry_ids))
    if len(rows) != len(set(entry_ids)):
        raise ValueError("Une ligne selectionnee est introuvable.")
    total = sum(to_float(r.get("poids_t")) for r in rows)
    if total > to_float(truck["capacity_kg"]) + 1e-6:
        raise ValueError(f"Capacite depassee: {total:.1f} kg > {to_float(truck['capacity_kg']):.1f} kg.")
    now = app_now().isoformat(timespec="seconds")
    with db_conn() as conn:
        cur = conn.execute(
            "INSERT INTO transport_trips(trip_date,truck_id,driver_id,destination,status,created_by,created_at) VALUES(?,?,?,?,?,?,?)",
            (trip_date.isoformat(), truck_id, driver_id, norm_text(destination), "PLANIFIE", username, now),
        )
        trip_id = int(cur.lastrowid)
        for r in rows:
            conn.execute(
                "INSERT INTO shipments(trip_id,planning_entry_id,num_commande,nom_client,destination,weight_kg,status,created_at) VALUES(?,?,?,?,?,?,?,?)",
                (trip_id, r["id"], r["num_commande"], r["nom_client"], norm_text(destination) or norm_text(r.get("destination")), to_float(r.get("poids_t")), "PLANIFIE", now),
            )
            conn.execute("UPDATE planning_entries SET status='PLANIFIE_LOGISTIQUE' WHERE id=?", (r["id"],))
            if r["order_line_id"]:
                conn.execute("UPDATE order_lines SET status='PLANIFIE_LOGISTIQUE',updated_at=? WHERE id=?", (now, r["order_line_id"]))
    fill = total / to_float(truck["capacity_kg"]) * 100 if to_float(truck["capacity_kg"]) else 0
    audit(username, "CREATE_TRIP", "transport_trip", trip_id, None, {"destination": destination, "truck": truck["code"], "weight": total, "fill_pct": fill})
    notify("LOGISTIQUE", "TRIP_CREATED", "Tournee creee", f"{destination}: {total:.0f} kg, remplissage {fill:.1f}%.", "transport_trip", trip_id)
    return trip_id


def trips_df() -> pd.DataFrame:
    rows = q_all("""
        SELECT t.*,tr.code AS truck_code,tr.plate,tr.capacity_kg,d.name AS driver_name,
               COALESCE(SUM(s.weight_kg),0) AS load_kg,COUNT(s.id) AS shipment_count
        FROM transport_trips t
        JOIN trucks tr ON tr.id=t.truck_id
        LEFT JOIN drivers d ON d.id=t.driver_id
        LEFT JOIN shipments s ON s.trip_id=t.id
        GROUP BY t.id
        ORDER BY t.trip_date DESC,t.id DESC
    """)
    df = pd.DataFrame(rows)
    if not df.empty:
        df["fill_pct"] = np.where(pd.to_numeric(df["capacity_kg"], errors="coerce") > 0,
                                  pd.to_numeric(df["load_kg"], errors="coerce") / pd.to_numeric(df["capacity_kg"], errors="coerce") * 100, 0)
    return df


def update_trip_status(trip_id: int, status: str, username: str) -> None:
    if status not in TRIP_STATUSES:
        raise ValueError("Statut tournee invalide.")
    old = q_one("SELECT * FROM transport_trips WHERE id=?", (trip_id,))
    if not old:
        raise ValueError("Tournee introuvable.")
    entry_status = {"CHARGE": "CHARGE", "EXPEDIE": "EXPEDIE", "LIVRE": "LIVRE"}.get(status)
    now = app_now().isoformat(timespec="seconds")
    with db_conn() as conn:
        conn.execute("UPDATE transport_trips SET status=? WHERE id=?", (status, trip_id))
        conn.execute("UPDATE shipments SET status=? WHERE trip_id=?", (status, trip_id))
        if entry_status:
            entry_ids = [r[0] for r in conn.execute("SELECT planning_entry_id FROM shipments WHERE trip_id=?", (trip_id,)).fetchall()]
            for eid in entry_ids:
                conn.execute("UPDATE planning_entries SET status=? WHERE id=?", (entry_status, eid))
                row = conn.execute("SELECT order_line_id FROM planning_entries WHERE id=?", (eid,)).fetchone()
                if row and row[0]:
                    conn.execute("UPDATE order_lines SET status=?,updated_at=? WHERE id=?", (entry_status, now, row[0]))
    audit(username, "UPDATE_TRIP", "transport_trip", trip_id, old, {"status": status})

# =============================================================================
# 12. KPI / ANALYSE IA DETERMINISTE
# =============================================================================
def dashboard_kpis() -> Dict[str, Any]:
    imp = active_import()
    active_orders = q_one("SELECT COUNT(DISTINCT num_commande) AS n FROM order_lines WHERE active=1")
    overdue = q_one("SELECT COUNT(*) AS n FROM order_lines WHERE active=1 AND date_livraison IS NOT NULL AND date(date_livraison)<date('now') AND status!='LIVRE'")
    blocked = q_one("SELECT COUNT(*) AS n FROM magasin_preparations WHERE status='BLOQUE'")
    relaq = q_one("SELECT COUNT(*) AS n FROM relaquage_orders WHERE status NOT IN ('CONFORME','ANNULE')")
    ready = q_one("SELECT COALESCE(SUM(poids_t),0) AS kg FROM planning_entries WHERE status='PRET_EXPEDITION'")
    trips = q_one("SELECT COUNT(*) AS n FROM transport_trips WHERE status IN ('PLANIFIE','CHARGE')")
    return {
        "active_import": imp,
        "orders": int((active_orders or {}).get("n", 0)),
        "overdue": int((overdue or {}).get("n", 0)),
        "blocked": int((blocked or {}).get("n", 0)),
        "relaquage_open": int((relaq or {}).get("n", 0)),
        "ready_kg": float((ready or {}).get("kg", 0.0)),
        "active_trips": int((trips or {}).get("n", 0)),
    }


def ai_recommendations() -> List[Dict[str, str]]:
    recs: List[Dict[str, str]] = []
    blocked = q_all("""
        SELECT m.id,m.missing_qty,e.num_commande,e.planned_date
        FROM magasin_preparations m JOIN planning_entries e ON e.id=m.planning_entry_id
        WHERE m.status='BLOQUE' AND m.missing_qty>0 ORDER BY e.planned_date LIMIT 10
    """)
    for r in blocked:
        recs.append({
            "niveau": "CRITIQUE", "type": "Matiere",
            "message": f"{r['num_commande']} manque {to_float(r['missing_qty']):.0f} unite(s) avant le {r['planned_date']}.",
            "action": "Verifier stock/reservation ou replanifier la commande.",
        })
    today = app_today().isoformat()
    overdue = q_all("""
        SELECT num_commande,nom_client,date_livraison,reste_a_livrer,status
        FROM order_lines WHERE active=1 AND date_livraison IS NOT NULL AND date(date_livraison)<date(?) AND status!='LIVRE'
        ORDER BY date_livraison LIMIT 10
    """, (today,))
    for r in overdue:
        recs.append({
            "niveau": "HAUTE", "type": "Retard",
            "message": f"{r['num_commande']} ({r['nom_client']}) est en retard depuis {r['date_livraison']}.",
            "action": "Prioriser planning ou confirmer une nouvelle date client.",
        })
    trips = trips_df()
    if not trips.empty:
        for _, r in trips[trips["status"].isin(["PLANIFIE", "CHARGE"])].iterrows():
            fill = to_float(r.get("fill_pct"))
            if fill < 75:
                recs.append({
                    "niveau": "OPPORTUNITE", "type": "Logistique",
                    "message": f"Tournee #{int(r['id'])} vers {r['destination']} chargee a {fill:.1f}%.",
                    "action": "Chercher une commande prete compatible avant depart.",
                })
    if not recs:
        recs.append({"niveau": "OK", "type": "Systeme", "message": "Aucune alerte prioritaire detectee.", "action": "Continuer le suivi operationnel."})
    return recs

# =============================================================================
# 13. EXPORTS
# =============================================================================
def result_day_export_df(df: pd.DataFrame) -> pd.DataFrame:
    if df is None or df.empty:
        return pd.DataFrame(columns=PLANNING_COLUMNS)
    out = df.copy()
    for c in PLANNING_COLUMNS:
        if c not in out.columns:
            out[c] = None
    return out[PLANNING_COLUMNS]


def export_plan_excel(result: Dict[str, Any]) -> bytes:
    out = io.BytesIO()
    wb = Workbook()
    wb.remove(wb.active)
    navy, blue, white, green, amber = "163A5F", "155EEF", "FFFFFF", "ECFDF3", "FFFAEB"
    line = Side(style="thin", color="D0D5DD")
    ws = wb.create_sheet("Resume")
    ws["A1"] = "ALLUCO - Planning Laquage IA"
    ws["A1"].fill = PatternFill("solid", fgColor=navy)
    ws["A1"].font = Font(color=white, bold=True, size=15)
    ws.merge_cells("A1:D1")
    cfg: PlannerConfig = result["config"]
    summary = [
        ("Semaine", f"S{cfg.week}/{cfg.year}"), ("Moteur", result["engine"]),
        ("Confiance regles", f"{result['confidence']}%"), ("Charge", f"{result['metrics']['total_load_h']:.2f} h"),
        ("Capacite", f"{result['metrics']['capacity_h']:.2f} h"), ("Utilisation", f"{result['metrics']['utilization_pct']:.1f}%"),
        ("Backlog", len(result["unscheduled"])), ("Temps calcul", f"{result['elapsed_s']:.2f} s"),
    ]
    for i, (k, v) in enumerate(summary, start=3):
        ws.cell(i, 1, k).font = Font(bold=True, color=navy)
        ws.cell(i, 2, v)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 36

    dates = iso_week_dates(cfg.year, cfg.week)
    for d in range(5):
        ws = wb.create_sheet(f"{DAYS[d].title()} {dates[d].strftime('%d-%m')}")
        df = result_day_export_df(result["days"].get(d, pd.DataFrame()))
        for ci, col in enumerate(df.columns, 1):
            c = ws.cell(1, ci, col)
            c.fill = PatternFill("solid", fgColor=navy)
            c.font = Font(color=white, bold=True)
            c.alignment = Alignment(horizontal="center", wrap_text=True)
            c.border = Border(bottom=line)
        for ri, (_, row) in enumerate(df.iterrows(), start=2):
            for ci, col in enumerate(df.columns, 1):
                v = row.get(col)
                if isinstance(v, pd.Timestamp):
                    v = v.to_pydatetime()
                ws.cell(ri, ci, v)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{max(1,len(df)+1)}"
        for ci in range(1, len(df.columns) + 1):
            ws.column_dimensions[get_column_letter(ci)].width = 15

    ws = wb.create_sheet("Backlog")
    backlog = result_day_export_df(result["unscheduled"])
    for ci, col in enumerate(backlog.columns, 1):
        c = ws.cell(1, ci, col); c.fill = PatternFill("solid", fgColor=blue); c.font = Font(color=white, bold=True)
    for ri, (_, row) in enumerate(backlog.iterrows(), start=2):
        for ci, col in enumerate(backlog.columns, 1):
            ws.cell(ri, ci, row.get(col))
    wb.save(out)
    return out.getvalue()


def export_published_excel(version_id: int) -> bytes:
    v = q_one("SELECT * FROM planning_versions WHERE id=?", (version_id,))
    if not v:
        raise ValueError("Version introuvable.")
    entries = published_entries(version_id)
    out = io.BytesIO(); wb = Workbook(); wb.remove(wb.active)
    for planned_date, g in entries.groupby("planned_date", sort=True):
        ws = wb.create_sheet(date.fromisoformat(planned_date).strftime("%a %d-%m")[:31])
        cols = ["num_commande", "nom_client", "article", "article_int", "couleur", "num_of", "planned_qty", "poids_t", "poudre", "nbre_bal", "duration_h", "status"]
        for ci, c in enumerate(cols, 1):
            ws.cell(1, ci, c).font = Font(bold=True)
        for ri, (_, r) in enumerate(g.iterrows(), start=2):
            for ci, c in enumerate(cols, 1):
                ws.cell(ri, ci, r.get(c))
        ws.freeze_panes = "A2"
    wb.save(out); return out.getvalue()


def export_plan_pdf(result: Dict[str, Any]) -> bytes:
    if not REPORTLAB_AVAILABLE:
        raise RuntimeError("ReportLab non installe.")
    out = io.BytesIO(); c = pdf_canvas.Canvas(out, pagesize=landscape(A4)); width, height = landscape(A4)
    cfg: PlannerConfig = result["config"]
    c.setTitle(f"ALLUCO Planning S{cfg.week}/{cfg.year}")
    c.setFont("Helvetica-Bold", 17); c.drawString(35, height - 45, f"ALLUCO - Planning Laquage IA S{cfg.week}/{cfg.year}")
    c.setFont("Helvetica", 9); c.drawString(35, height - 63, f"Moteur: {result['engine']} - confiance regles: {result['confidence']}%")
    y = height - 95
    for dm in result["metrics"]["days"]:
        c.setFont("Helvetica-Bold", 10); c.drawString(35, y, f"{dm['Jour']} {dm['Date']} - charge {dm['Charge totale h']:.2f}h / {dm['Capacite h']:.2f}h - {dm['Couleurs']}")
        y -= 16
    for d in range(5):
        c.showPage(); y = height - 40
        dm = result["metrics"]["days"][d]
        c.setFont("Helvetica-Bold", 15); c.drawString(35, y, f"{dm['Jour']} - {dm['Date']}"); y -= 22
        df = result["days"][d]
        c.setFont("Helvetica-Bold", 7)
        headers = ["Commande", "Client", "Article", "Couleur", "OF", "Qte", "h"]
        xs = [35, 130, 300, 450, 520, 640, 720]
        for x, h in zip(xs, headers): c.drawString(x, y, h)
        y -= 12; c.setFont("Helvetica", 7)
        for _, r in df.iterrows():
            if y < 35:
                c.showPage(); y = height - 40; c.setFont("Helvetica", 7)
            vals = [norm_text(r.get("NumCommande"))[:15], norm_text(r.get("NomClient"))[:26], norm_text(r.get("ArticleInt"))[:24], norm_text(r.get("Couleur"))[:12], norm_text(r.get("NumOF"))[:18], f"{to_float(r.get('Lancement')) + to_float(r.get('ReLaquage')):.0f}", f"{to_float(r.get('TempsH')):.2f}"]
            for x, v in zip(xs, vals): c.drawString(x, y, v)
            y -= 11
    c.save(); return out.getvalue()

# =============================================================================
# 14. UI / DESIGN
# =============================================================================
def app_css(dark_mode: bool = False) -> str:
    if dark_mode:
        theme = {
            "bg": "#0B1220", "surface": "#111827", "surface2": "#172033", "ink": "#F4F7FB",
            "muted": "#AAB4C3", "line": "#2A3547", "primary": "#5B8CFF", "primary2": "#2F6BFF",
            "soft": "#142342", "success": "#4ADE80", "warning": "#FBBF24", "danger": "#FB7185",
            "hero1": "#10243E", "hero2": "#1D4ED8", "shadow": "rgba(0,0,0,.28)", "input": "#0F172A",
        }
        scheme = "dark"
    else:
        theme = {
            "bg": "#F4F7FB", "surface": "#FFFFFF", "surface2": "#F8FAFC", "ink": "#162033",
            "muted": "#667085", "line": "#E1E7EF", "primary": "#155EEF", "primary2": "#0B5ED7",
            "soft": "#EEF4FF", "success": "#067647", "warning": "#B54708", "danger": "#B42318",
            "hero1": "#123B67", "hero2": "#155EEF", "shadow": "rgba(16,24,40,.08)", "input": "#FFFFFF",
        }
        scheme = "light"
    return f"""
    <style>
      :root{{--bg:{theme['bg']};--surface:{theme['surface']};--surface2:{theme['surface2']};--ink:{theme['ink']};--muted:{theme['muted']};--line:{theme['line']};--primary:{theme['primary']};--primary2:{theme['primary2']};--soft:{theme['soft']};--success:{theme['success']};--warning:{theme['warning']};--danger:{theme['danger']};--input:{theme['input']};--shadow:{theme['shadow']};}}
      html,body,.stApp,[data-testid="stAppViewContainer"]{{background:var(--bg)!important;color:var(--ink)!important;color-scheme:{scheme}!important}}
      .block-container{{max-width:1580px;padding-top:1.05rem;padding-bottom:2.2rem}}
      #MainMenu,footer,[data-testid="stDeployButton"],[data-testid="stToolbar"],[data-testid="stDecoration"]{{display:none!important}}
      header[data-testid="stHeader"]{{background:transparent!important}}
      h1,h2,h3,h4,h5,h6,p,label,span,div{{color:var(--ink)}}
      [data-testid="stCaptionContainer"],.stCaption,.erp-muted{{color:var(--muted)!important}}

      section[data-testid="stSidebar"]{{background:var(--surface)!important;border-right:1px solid var(--line)!important}}
      section[data-testid="stSidebar"]>div{{background:var(--surface)!important}}
      .brand-shell{{display:flex;align-items:center;gap:.8rem;padding:.2rem 0 .85rem}}
      .brand-shell.brand-compact{{padding:.05rem 0 .45rem}}
      .brand-logo{{width:185px;max-width:72%;height:64px;object-fit:contain;object-position:left center;display:block}}
      .brand-copy{{min-width:0}}.brand-title{{font-size:.82rem;font-weight:900;letter-spacing:.03em}}
      .brand-sub{{font-size:.67rem;color:var(--muted)!important;line-height:1.25;margin-top:.12rem}}

      .alluco-hero{{position:relative;overflow:hidden;background:linear-gradient(135deg,{theme['hero1']} 0%,{theme['hero2']} 100%);padding:1.18rem 1.35rem;border-radius:18px;margin-bottom:1rem;box-shadow:0 12px 28px rgba(21,94,239,.13)}}
      .alluco-hero:after{{content:'';position:absolute;width:220px;height:220px;border-radius:50%;right:-80px;top:-130px;background:rgba(255,255,255,.09)}}
      .alluco-hero *{{color:#fff!important}}.alluco-title{{font-size:1.48rem;font-weight:950;letter-spacing:-.015em}}.alluco-sub{{opacity:.91;font-size:.86rem;margin-top:.25rem;max-width:1050px}}

      .erp-card{{border:1px solid var(--line);border-radius:14px;padding:.9rem 1rem;background:var(--surface);margin-bottom:.7rem;box-shadow:0 1px 3px var(--shadow)}}
      .erp-badge{{display:inline-flex;align-items:center;padding:.28rem .55rem;border-radius:999px;background:var(--soft);color:var(--primary)!important;font-weight:850;font-size:.75rem}}
      .erp-ok{{background:color-mix(in srgb,var(--success) 13%,var(--surface));color:var(--success)!important}}
      .erp-warn{{background:color-mix(in srgb,var(--warning) 13%,var(--surface));color:var(--warning)!important}}
      .erp-err{{background:color-mix(in srgb,var(--danger) 13%,var(--surface));color:var(--danger)!important}}

      [data-testid="stMetric"]{{border:1px solid var(--line);border-radius:14px;padding:.68rem .85rem;background:var(--surface);box-shadow:0 1px 3px var(--shadow)}}
      [data-testid="stMetricLabel"]{{color:var(--muted)!important}}
      [data-testid="stDataFrame"],[data-testid="stDataEditor"]{{border:1px solid var(--line);border-radius:13px;overflow:hidden;background:var(--surface)!important}}
      [data-baseweb="input"]>div,[data-baseweb="select"]>div,textarea,input{{background:var(--input)!important;color:var(--ink)!important;border-color:var(--line)!important}}
      [data-baseweb="popover"],[data-baseweb="menu"],[role="listbox"]{{background:var(--surface)!important;color:var(--ink)!important}}
      button[data-baseweb="tab"]{{font-weight:800;color:var(--muted)!important}}
      button[data-baseweb="tab"][aria-selected="true"]{{color:var(--primary)!important}}
      .stButton>button,.stDownloadButton>button{{border-radius:10px;font-weight:800;min-height:40px;border:1px solid var(--line);background:var(--surface);color:var(--ink)!important}}
      .stButton>button[kind="primary"]{{background:var(--primary);border-color:var(--primary);color:#fff!important}}
      .stButton>button[kind="primary"] *{{color:#fff!important}}
      hr{{border-color:var(--line)!important}}

      .bal-summary{{display:flex;gap:.45rem;align-items:center;flex-wrap:wrap;margin:.2rem 0 .7rem}}
      .bal-chip{{display:inline-flex;align-items:center;gap:.28rem;border:1px solid var(--line);background:var(--surface);border-radius:999px;padding:.3rem .58rem;font-size:.75rem;font-weight:800;color:var(--ink)!important}}
      .bal-chip strong{{color:var(--primary)!important}}
      @media(max-width:900px){{.block-container{{padding-left:.7rem;padding-right:.7rem}}.brand-logo{{width:145px;height:54px}}.alluco-hero{{padding:1rem 1.05rem}}}}
    </style>
    """


def hero(title: str, subtitle: str = "") -> None:
    st.markdown(
        f"<div class='alluco-hero'><div class='alluco-title'>{esc(title)}</div><div class='alluco-sub'>{esc(subtitle)}</div></div>",
        unsafe_allow_html=True,
    )


def flash_error(exc: Exception, prefix: str = "Erreur") -> None:
    ref = "ERR-" + hashlib.sha256(f"{type(exc).__name__}|{exc}|{time.time_ns()}".encode()).hexdigest()[:10].upper()
    st.error(f"{prefix}. Reference: {ref}")
    if st.session_state.get("role") == "ADMIN":
        st.caption(f"Detail admin: {type(exc).__name__}: {exc}")


def bootstrap_admin_ui() -> None:
    st.markdown(brand_html(bool(st.session_state.get("ui_dark_mode", False))), unsafe_allow_html=True)
    hero("Initialisation ALLUCO ERP", "Premiere utilisation: creez le compte administrateur. Aucun mot de passe n'est code dans le fichier.")
    with st.form("bootstrap_admin"):
        username = st.text_input("Utilisateur administrateur", value="admin")
        p1 = st.text_input("Mot de passe", type="password")
        p2 = st.text_input("Confirmer le mot de passe", type="password")
        submitted = st.form_submit_button("Creer l'administrateur", type="primary")
    if submitted:
        try:
            if p1 != p2:
                raise ValueError("Les mots de passe ne correspondent pas.")
            create_user(username, p1, "ADMIN", "SYSTEM")
            st.success("Administrateur cree. Connectez-vous.")
            st.rerun()
        except Exception as exc:
            st.error(str(exc))


def login_ui() -> None:
    st.markdown(brand_html(bool(st.session_state.get("ui_dark_mode", False))), unsafe_allow_html=True)
    hero("ALLUCO Industrial ERP IA", "Acces interne securise")
    with st.form("login"):
        username = st.text_input("Utilisateur")
        password = st.text_input("Mot de passe", type="password")
        submit = st.form_submit_button("Se connecter", type="primary", use_container_width=True)
    if submit:
        user = authenticate(username, password)
        if not user:
            st.error("Identifiants incorrects.")
            return
        st.session_state["user_id"] = user["id"]
        st.session_state["username"] = user["username"]
        st.session_state["role"] = user["role"]
        audit(user["username"], "LOGIN", "session", user["id"])
        st.rerun()


def client_portal_ui() -> None:
    st.markdown(brand_html(bool(st.session_state.get("ui_dark_mode", False))), unsafe_allow_html=True)
    hero("Suivi client", "Consultation du dernier planning publie")
    if not CLIENT_ACCESS_CODE:
        st.info("Le portail client est desactive. Configurez ALLUCO_CLIENT_ACCESS_CODE pour l'activer.")
        return
    code = st.text_input("Code d'acces", type="password")
    command = st.text_input("Numero de commande").strip().upper()
    if st.button("Rechercher", type="primary"):
        if not hmac.compare_digest(code, CLIENT_ACCESS_CODE):
            st.error("Code d'acces invalide.")
            return
        v = current_published_version()
        if not v:
            st.info("Aucun planning publie.")
            return
        rows = q_all(
            "SELECT num_commande,nom_client,article_int,couleur,planned_date,status FROM planning_entries WHERE planning_version_id=? AND UPPER(num_commande)=? ORDER BY planned_date,sequence_no",
            (v["id"], command),
        )
        if not rows:
            st.warning("Commande non trouvee dans le planning publie.")
            return
        safe = pd.DataFrame(rows)
        st.success(f"Commande {command} trouvee - planning S{v['week']}/{v['year']} V{v['version_no']}")
        st.dataframe(safe, hide_index=True, use_container_width=True)


def sidebar_navigation() -> str:
    role = st.session_state.get("role", "")
    username = st.session_state.get("username", "")
    with st.sidebar:
        dark_mode = bool(st.session_state.get("ui_dark_mode", False))
        st.markdown(brand_html(dark_mode, compact=True), unsafe_allow_html=True)
        st.markdown(f"<div class='erp-badge'>{esc(username)} · {esc(role)}</div>", unsafe_allow_html=True)
        st.toggle("Mode sombre", key="ui_dark_mode", help="Basculer entre le thème clair et le thème sombre")
        st.divider()
        options: List[Tuple[str, str]] = []
        candidates = [
            ("dashboard", "▣ Tableau de bord"), ("orders", "📋 Commandes / AX"),
            ("planning", "🤖 Planning IA"), ("balancelles", "⚖️ Balancelles"),
            ("magasin", "📦 Magasin J-2"), ("laquage", "🎨 Laquage"),
            ("quality", "✓ Qualité"), ("relaquage", "↻ Re-laquage"),
            ("logistics", "🚚 Logistique"), ("analysis", "🧠 Analyse IA"),
            ("notifications", "🔔 Notifications"), ("admin", "⚙ Administration"),
        ]
        for key, label in candidates:
            if key == "admin" and role == "ADMIN":
                options.append((key, label))
            elif role_can(role, key):
                options.append((key, label))
        labels = [x[1] for x in options]
        selected_label = st.radio("Navigation", labels, label_visibility="collapsed") if labels else ""
        selected = next((k for k, l in options if l == selected_label), "dashboard")
        st.divider()
        st.caption(f"Version {VERSION} · SQLite")
        if st.button("Se déconnecter", use_container_width=True):
            for k in ["user_id", "username", "role", "plan_result", "plan_signature"]:
                st.session_state.pop(k, None)
            st.rerun()
    return selected


# =============================================================================
# 15. UI PAGES
# =============================================================================
def page_dashboard() -> None:
    hero("Tableau de bord", "Vue operationnelle de la chaine planning -> magasin -> laquage -> qualite -> logistique")
    k = dashboard_kpis()
    cols = st.columns(6)
    vals = [
        ("Commandes actives", k["orders"]), ("Retards", k["overdue"]), ("Magasin bloque", k["blocked"]),
        ("Re-laquages", k["relaquage_open"]), ("Pret expedition", f"{k['ready_kg'] / 1000:.2f} t"), ("Tournees actives", k["active_trips"]),
    ]
    for c, (label, value) in zip(cols, vals):
        c.metric(label, value)
    imp = k.get("active_import")
    if imp:
        st.caption(f"Base active: {imp['filename']} · feuille {imp['source_sheet']} · import {imp['imported_at']} · {imp['valid_rows']} lignes valides")
    else:
        st.warning("Aucune base AX active. Importez un fichier dans Commandes / AX.")

    v = current_published_version()
    if v:
        st.markdown("#### Planning officiel")
        entries = published_entries(v["id"])
        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Semaine", f"S{v['week']}/{v['year']}")
        c2.metric("Version", f"V{v['version_no']}")
        c3.metric("Lignes", len(entries))
        c4.metric("Publie", v["published_at"][:16].replace("T", " ") if v.get("published_at") else "-")
        if not entries.empty:
            chart = entries.groupby("planned_date")["duration_h"].sum().to_frame("Charge h")
            st.bar_chart(chart)
    st.markdown("#### Alertes prioritaires")
    st.dataframe(pd.DataFrame(ai_recommendations()), hide_index=True, use_container_width=True)


def page_orders() -> None:
    hero("Commandes / AX", "Importer la base AX, controler les donnees et alimenter le planning")
    role = st.session_state.get("role")
    username = st.session_state.get("username", "")
    if role in {"ADMIN", "PLANNING"}:
        uploaded = st.file_uploader("Base Excel AX", type=["xlsx"], help="Feuilles supportees: preparation pour planning VF, version 0, extraction ax, Feuil1 ou feuille detectee automatiquement.")
        if uploaded is not None and st.button("Importer et activer cette base", type="primary"):
            try:
                res = import_ax_bytes(uploaded.getvalue(), uploaded.name, username)
                if res.get("reused"):
                    st.info("Cette base est deja la base active.")
                else:
                    st.success(f"Import termine: {res['valid']} valides, {res['invalid']} ignorees · feuille {res['sheet']}.")
                st.session_state.pop("plan_result", None)
                st.rerun()
            except Exception as exc:
                flash_error(exc, "Import impossible")
    imp = active_import()
    if imp:
        st.markdown("#### Base active")
        st.json({k: imp[k] for k in ["filename", "sha256", "source_sheet", "imported_at", "imported_by", "row_count", "valid_rows", "invalid_rows"]})
    df = load_active_orders_df()
    if df.empty:
        st.info("Aucune commande active.")
        return
    show = df[[
        "id", "num_commande", "date_creation", "nom_client", "article", "article_int", "couleur",
        "reste_a_livrer", "num_of", "prod_statut", "reserver_br", "stock_physique",
        "lancement", "relaquage", "poids_t", "poudre", "barre_bal", "nbre_bal", "temps_h",
        "date_livraison", "destination", "status",
    ]].copy()
    search = st.text_input("Recherche commande / client / article")
    if search:
        mask = show.astype(str).apply(lambda c: c.str.contains(search, case=False, na=False)).any(axis=1)
        show = show[mask]
    st.dataframe(show, hide_index=True, use_container_width=True, height=560)


def planning_config_from_ui() -> PlannerConfig:
    y, w, start, end = next_planning_period()
    st.caption(f"Periode automatique proposee: S{w}/{y} · {start.strftime('%d/%m/%Y')} -> {end.strftime('%d/%m/%Y')}")
    c1, c2, c3, c4 = st.columns(4)
    year = int(c1.number_input("Annee", 2025, 2035, y, 1))
    week = int(c2.number_input("Semaine", 1, 53, w, 1))
    capacity = float(c3.number_input("Capacite/jour (h)", 1.0, 24.0, DEFAULT_CAPACITY_H, 0.5))
    cleaning = int(c4.number_input("Changement couleur (min)", 0, 120, DEFAULT_CLEANING_MIN, 5))
    c5, c6 = st.columns(2)
    force = tuple(x.strip().upper() for x in re.split(r"[,;\n]+", c5.text_area("Forcer commandes", placeholder="VTE2601234, VTE2605678")) if x.strip())
    exclude = tuple(x.strip().upper() for x in re.split(r"[,;\n]+", c6.text_area("Exclure commandes", placeholder="VTE2609999")) if x.strip())
    return PlannerConfig(year=year, week=week, capacity_h=capacity, cleaning_min=cleaning, force_commands=force, exclude_commands=exclude)


def show_plan_result(result: Dict[str, Any]) -> None:
    m = result["metrics"]
    cols = st.columns(5)
    values = [
        ("Confiance regles", f"{result['confidence']}%"), ("Charge", f"{m['total_load_h']:.1f} h"),
        ("Capacite", f"{m['capacity_h']:.1f} h"), ("Utilisation", f"{m['utilization_pct']:.1f}%"),
        ("Backlog", m["backlog"]),
    ]
    for c, (k, v) in zip(cols, values): c.metric(k, v)
    st.caption(f"Moteur: {result['engine']} · calcul {result['elapsed_s']:.2f}s")
    if result["hard_errors"]:
        for e in result["hard_errors"]: st.error(e)
    metrics_df = pd.DataFrame(m["days"])
    st.dataframe(metrics_df, hide_index=True, use_container_width=True)
    tabs = st.tabs([x.title() for x in DAYS])
    for d, tab in enumerate(tabs):
        with tab:
            df = result["days"][d]
            st.dataframe(result_day_export_df(df), hide_index=True, use_container_width=True, height=440)
    if not result["unscheduled"].empty:
        with st.expander(f"Backlog - {len(result['unscheduled'])} ligne(s)"):
            st.dataframe(result_day_export_df(result["unscheduled"]), hide_index=True, use_container_width=True, height=420)
    excel = export_plan_excel(result)
    c1, c2 = st.columns(2)
    c1.download_button("Telecharger Excel", excel, file_name=f"Planning_IA_S{result['config'].week}_{result['config'].year}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)
    if REPORTLAB_AVAILABLE:
        pdf = export_plan_pdf(result)
        c2.download_button("Telecharger PDF", pdf, file_name=f"Planning_IA_S{result['config'].week}_{result['config'].year}.pdf", mime="application/pdf", use_container_width=True)


def page_planning() -> None:
    hero("Planning Laquage IA", "Generation deterministe + optimisation mathematique + versionnement du planning officiel")
    if load_active_orders_df().empty:
        st.warning("Importez d'abord une base AX.")
        return
    username = st.session_state.get("username", "")
    tab_generate, tab_official = st.tabs(["Proposition IA", "Planning officiel / historique"])
    with tab_generate:
        cfg = planning_config_from_ui()
        sig = json_dumps(cfg.__dict__) + "|" + norm_text((active_import() or {}).get("sha256"))
        c1, c2 = st.columns([1, 3])
        generate = c1.button("Generer / regenerer", type="primary", use_container_width=True)
        if generate or st.session_state.get("plan_signature") != sig:
            try:
                with st.spinner("Optimisation du planning..."):
                    st.session_state["plan_result"] = generate_plan(cfg)
                    st.session_state["plan_signature"] = sig
            except Exception as exc:
                flash_error(exc, "Generation impossible")
                return
        result = st.session_state.get("plan_result")
        if result:
            show_plan_result(result)
            st.markdown("#### Publication")
            notes = st.text_input("Note de publication", placeholder="Planning valide par responsable production")
            if st.button("Publier comme planning officiel", type="primary", disabled=bool(result.get("hard_errors"))):
                try:
                    vid = publish_plan(result, username, notes)
                    st.success(f"Planning publie - version #{vid}. Les preparations Magasin J-2 ont ete creees.")
                    st.rerun()
                except Exception as exc:
                    flash_error(exc, "Publication impossible")

    with tab_official:
        v = current_published_version()
        if not v:
            st.info("Aucun planning officiel publie.")
        else:
            st.markdown(f"**S{v['week']}/{v['year']} · V{v['version_no']} · publie {v['published_at']}**")
            entries = published_entries(v["id"])
            st.dataframe(entries[["id", "planned_date", "sequence_no", "num_commande", "nom_client", "article_int", "couleur", "num_of", "planned_qty", "duration_h", "status"]], hide_index=True, use_container_width=True, height=480)
            st.download_button("Exporter le planning officiel", export_published_excel(v["id"]), file_name=f"Planning_Officiel_S{v['week']}_{v['year']}_V{v['version_no']}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            if not entries.empty:
                st.markdown("#### Modifier une ligne avec creation automatique d'une nouvelle version")
                choices = {f"#{int(r['id'])} · {r['planned_date']} · {r['num_commande']} · {r['article_int']} · {r['couleur']}": int(r["id"]) for _, r in entries.iterrows()}
                selected = st.selectbox("Ligne", list(choices.keys()))
                old = entries[entries["id"] == choices[selected]].iloc[0]
                new_date = st.date_input("Nouvelle date", value=date.fromisoformat(str(old["planned_date"])))
                reason = st.text_input("Motif obligatoire")
                if st.button("Creer et publier la nouvelle version"):
                    try:
                        new_vid = revise_planning_entry(v["id"], choices[selected], new_date, username, reason)
                        st.success(f"Nouvelle version publiee #{new_vid}.")
                        st.rerun()
                    except Exception as exc:
                        st.error(str(exc))


def build_balancelle_register(entries: pd.DataFrame) -> pd.DataFrame:
    """Décompose le planning officiel en balancelles opérationnelles.

    Le registre privilégie la sécurité: la capacité d'une BAL provient du référentiel
    article/famille. Si le nombre de BAL stocké dans le planning diffère du nombre
    recalculé, les BAL restent générées sans surcharge et sont marquées « À contrôler ».
    """
    columns = [
        "N° BAL", "Balancelle", "Date", "Jour", "Séquence ligne", "BAL dans ligne",
        "Commande", "Client", "Article", "Article/int", "Couleur", "Num OF",
        "Qté dans BAL", "Capacité BAL", "Remplissage %", "Statut", "Contenu", "Entry ID",
    ]
    if entries is None or entries.empty:
        return pd.DataFrame(columns=columns)
    rows: List[Dict[str, Any]] = []
    per_day_counter: Dict[str, int] = defaultdict(int)
    work = entries.copy()
    if "planned_date" in work.columns:
        work = work.sort_values(["planned_date", "sequence_no", "id"], na_position="last")
    for _, r in work.iterrows():
        planned_date = norm_text(r.get("planned_date"))
        try:
            d = date.fromisoformat(planned_date)
            day_label = DAYS[d.weekday()].title() if 0 <= d.weekday() < len(DAYS) else d.strftime("%A")
        except Exception:
            day_label = "—"
        article_int = norm_text(r.get("article_int"))
        qty_total = max(0, int(math.ceil(to_float(r.get("planned_qty")) - 1e-9)))
        cap = max(1, int(infer_bars_per_bal(article_int)))
        expected_nbal = int(math.ceil(qty_total / cap)) if qty_total > 0 else 0
        stored_nbal = max(0, int(round(to_float(r.get("nbre_bal")))))
        mismatch = stored_nbal > 0 and stored_nbal != expected_nbal
        remaining = qty_total
        for bal_in_line in range(1, expected_nbal + 1):
            per_day_counter[planned_date] += 1
            qty = min(cap, remaining)
            remaining -= qty
            fill = round((qty / cap) * 100, 1) if cap else 0.0
            status = "À contrôler" if mismatch else ("Complète" if qty == cap else "Partielle")
            bal_number = per_day_counter[planned_date]
            rows.append({
                "N° BAL": bal_number,
                "Balancelle": f"BAL {bal_number:03d}",
                "Date": planned_date,
                "Jour": day_label,
                "Séquence ligne": int(to_float(r.get("sequence_no"), 0)),
                "BAL dans ligne": bal_in_line,
                "Commande": norm_text(r.get("num_commande")),
                "Client": norm_text(r.get("nom_client")),
                "Article": norm_text(r.get("article")),
                "Article/int": article_int,
                "Couleur": norm_text(r.get("couleur")),
                "Num OF": norm_text(r.get("num_of")),
                "Qté dans BAL": qty,
                "Capacité BAL": cap,
                "Remplissage %": fill,
                "Statut": status,
                "Contenu": f"{norm_text(r.get('num_commande')) or 'Sans commande'} · {article_int or norm_text(r.get('article'))} · {norm_text(r.get('couleur'))}",
                "Entry ID": int(to_float(r.get("id"), 0)),
            })
    return pd.DataFrame(rows, columns=columns)


def export_balancelles_excel(df: pd.DataFrame) -> bytes:
    out = io.BytesIO()
    wb = Workbook(); ws = wb.active; ws.title = "Balancelles"
    headers = list(df.columns)
    navy, white, line = "14365A", "FFFFFF", Side(style="thin", color="D0D5DD")
    for ci, h in enumerate(headers, 1):
        c = ws.cell(1, ci, h); c.fill = PatternFill("solid", fgColor=navy); c.font = Font(color=white, bold=True); c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for ri, (_, row) in enumerate(df.iterrows(), 2):
        for ci, h in enumerate(headers, 1):
            c = ws.cell(ri, ci, row.get(h)); c.border = Border(bottom=line)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(max(1, len(headers)))}{max(1, len(df)+1)}"
    widths = {"Balancelle": 14, "Date": 13, "Jour": 12, "Commande": 18, "Client": 25, "Article": 25, "Article/int": 18, "Couleur": 13, "Num OF": 17, "Contenu": 48}
    for ci, h in enumerate(headers, 1): ws.column_dimensions[get_column_letter(ci)].width = widths.get(h, 14)
    wb.save(out)
    return out.getvalue()


def page_balancelles() -> None:
    hero("Balancelles", "Registre atelier BAL 001…N dérivé du planning officiel, avec capacité, remplissage et contrôle de cohérence")
    v = current_published_version()
    if not v:
        st.info("Publiez d'abord un planning officiel pour générer le registre des balancelles.")
        return
    entries = published_entries(v["id"])
    register = build_balancelle_register(entries)
    if register.empty:
        st.info("Aucune balancelle dans le planning publié.")
        return

    dates = sorted([x for x in register["Date"].dropna().astype(str).unique().tolist() if x])
    default_date = app_today().isoformat() if app_today().isoformat() in dates else dates[0]
    c1, c2, c3 = st.columns([1.15, 1.35, 2.2])
    selected_date = c1.selectbox("Jour", dates, index=dates.index(default_date) if default_date in dates else 0)
    all_colors = sorted(register["Couleur"].dropna().astype(str).unique().tolist())
    colors = c2.multiselect("Couleurs", all_colors, default=all_colors)
    query = c3.text_input("Recherche", placeholder="Commande, client, article, OF…")

    filtered = register[register["Date"] == selected_date].copy()
    if colors:
        filtered = filtered[filtered["Couleur"].isin(colors)]
    if query.strip():
        q = query.strip().lower()
        search_cols = ["Commande", "Client", "Article", "Article/int", "Couleur", "Num OF", "Contenu"]
        mask = pd.Series(False, index=filtered.index)
        for col in search_cols:
            mask = mask | filtered[col].astype(str).str.lower().str.contains(re.escape(q), regex=True, na=False)
        filtered = filtered[mask]

    full_count = int((filtered["Statut"] == "Complète").sum()) if not filtered.empty else 0
    partial_count = int((filtered["Statut"] == "Partielle").sum()) if not filtered.empty else 0
    check_count = int((filtered["Statut"] == "À contrôler").sum()) if not filtered.empty else 0
    avg_fill = float(pd.to_numeric(filtered["Remplissage %"], errors="coerce").fillna(0).mean()) if not filtered.empty else 0.0
    m1, m2, m3, m4, m5 = st.columns(5)
    m1.metric("BAL", len(filtered)); m2.metric("Complètes", full_count); m3.metric("Partielles", partial_count); m4.metric("À contrôler", check_count); m5.metric("Remplissage moyen", f"{avg_fill:.1f}%")

    if not filtered.empty:
        colors_txt = " · ".join(dict.fromkeys(filtered["Couleur"].astype(str).tolist()))
        st.markdown(f"<div class='bal-summary'><span class='bal-chip'><strong>{esc(selected_date)}</strong></span><span class='bal-chip'>Couleurs: <strong>{esc(colors_txt or '—')}</strong></span><span class='bal-chip'>Planning: <strong>S{int(v['week'])}/{int(v['year'])} · V{int(v['version_no'])}</strong></span></div>", unsafe_allow_html=True)
    display_cols = ["Balancelle", "Contenu", "Commande", "Client", "Article/int", "Couleur", "Num OF", "Qté dans BAL", "Capacité BAL", "Remplissage %", "Statut", "BAL dans ligne"]
    st.dataframe(filtered[display_cols], hide_index=True, use_container_width=True, height=610)

    if check_count:
        st.warning(f"{check_count} BAL marquée(s) À contrôler: le Nbre BAL publié ne correspond pas au calcul Barre/bal du référentiel.")
    c1, c2 = st.columns(2)
    c1.download_button("⬇ Exporter CSV", filtered[display_cols].to_csv(index=False).encode("utf-8-sig"), file_name=f"Balancelles_{selected_date}.csv", mime="text/csv", use_container_width=True)
    c2.download_button("⬇ Exporter Excel", export_balancelles_excel(filtered[display_cols]), file_name=f"Balancelles_{selected_date}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)
    st.caption("Lecture: la dernière BAL d'une ligne peut être partielle. 'À contrôler' signale une incohérence entre la quantité, Nbre BAL et la capacité Barre/bal calculée.")


def page_magasin() -> None:
    hero("Magasin J-2", "Preparation automatique deux jours ouvrables avant le laquage")
    df = magasin_rows()
    if df.empty:
        st.info("Aucune preparation active.")
        return
    today = app_today().isoformat()
    due = df[df["prep_date"] <= today]
    cols = st.columns(4)
    cols[0].metric("Total", len(df)); cols[1].metric("A traiter", len(due)); cols[2].metric("Bloquees", int((df["status"] == "BLOQUE").sum())); cols[3].metric("Pretes", int((df["status"] == "PRET").sum()))
    display = df[["id", "prep_date", "planned_date", "num_commande", "nom_client", "article_int", "couleur", "num_of", "required_qty", "reserved_qty", "missing_qty", "location", "status"]]
    st.dataframe(display, hide_index=True, use_container_width=True, height=500)
    st.markdown("#### Mise a jour")
    options = {f"#{int(r['id'])} · {r['prep_date']} · {r['num_commande']} · {r['article_int']}": int(r["id"]) for _, r in df.iterrows()}
    sel = st.selectbox("Preparation", list(options.keys()))
    row = df[df["id"] == options[sel]].iloc[0]
    c1, c2, c3 = st.columns(3)
    status = c1.selectbox("Statut", MAGASIN_STATUSES, index=MAGASIN_STATUSES.index(row["status"]) if row["status"] in MAGASIN_STATUSES else 0)
    reserved = c2.number_input("Quantite reservee", min_value=0.0, max_value=float(row["required_qty"]), value=float(row["reserved_qty"]), step=1.0)
    location = c3.text_input("Emplacement", value=norm_text(row["location"]))
    if st.button("Enregistrer preparation", type="primary"):
        try:
            update_magasin(options[sel], status, reserved, location, st.session_state["username"])
            st.success("Preparation mise a jour.")
            st.rerun()
        except Exception as exc:
            st.error(str(exc))


def page_laquage() -> None:
    hero("Execution Laquage", "Vue atelier simple: matiere prete -> en cours -> controle qualite")
    df = laquage_rows()
    if df.empty:
        st.info("Aucun planning officiel.")
        return
    day_filter = st.date_input("Jour", value=app_today())
    show = df[df["planned_date"] == day_filter.isoformat()].copy()
    if show.empty:
        st.info("Aucune ligne planifiee pour cette date.")
        return
    st.dataframe(show[["id", "sequence_no", "num_commande", "article_int", "couleur", "num_of", "planned_qty", "nbre_bal", "duration_h", "magasin_status", "status"]], hide_index=True, use_container_width=True, height=480)
    options = {f"#{int(r['id'])} · {r['num_commande']} · {r['article_int']} · {r['couleur']}": int(r["id"]) for _, r in show.iterrows()}
    sel = st.selectbox("Operation", list(options.keys()))
    action = st.selectbox("Action", ["EN_COURS_LAQUAGE", "CONTROLE_QUALITE", "BLOQUE"])
    if st.button("Valider l'action", type="primary"):
        try:
            update_laquage_status(options[sel], action, st.session_state["username"])
            st.success("Statut mis a jour.")
            st.rerun()
        except Exception as exc:
            st.error(str(exc))


def page_quality() -> None:
    hero("Controle Qualite", "Conforme -> pret expedition · Non conforme -> re-laquage automatique")
    df = quality_pending_rows()
    if df.empty:
        st.info("Aucun controle en attente.")
        return
    st.dataframe(df[["id", "planned_date", "num_commande", "nom_client", "article_int", "couleur", "planned_qty", "status"]], hide_index=True, use_container_width=True)
    options = {f"#{int(r['id'])} · {r['num_commande']} · {r['article_int']} · {r['couleur']}": int(r["id"]) for _, r in df.iterrows()}
    sel = st.selectbox("Ligne a controler", list(options.keys()))
    r = df[df["id"] == options[sel]].iloc[0]
    result = st.radio("Resultat", ["CONFORME", "NON_CONFORME"], horizontal=True)
    defect = st.selectbox("Defaut", ["", "Couleur", "Nuance", "Rayure", "Aspect", "Adherence", "Pollution", "Poudre", "Preparation", "Autre"], disabled=result == "CONFORME")
    qty = st.number_input("Quantite affectee", 0.0, float(r["planned_qty"]), float(r["planned_qty"]) if result == "NON_CONFORME" else 0.0, 1.0, disabled=result == "CONFORME")
    comment = st.text_area("Commentaire")
    if st.button("Enregistrer controle", type="primary"):
        try:
            record_quality(options[sel], result, defect, qty, comment, st.session_state["username"])
            st.success("Controle enregistre.")
            st.rerun()
        except Exception as exc:
            st.error(str(exc))


def page_relaquage() -> None:
    hero("Re-laquage", "Les non-conformites ouvertes sont reinjectees dans le prochain planning avec priorite forte")
    df = relaquage_rows()
    if df.empty:
        st.info("Aucun re-laquage.")
        return
    st.dataframe(df, hide_index=True, use_container_width=True, height=560)
    st.info("Les lignes au statut A_PLANIFIER sont automatiquement ajoutees au prochain calcul Planning IA.")


def page_logistics() -> None:
    hero("Logistique", "Camions, chauffeurs, chargement et tournees avec controle strict de capacite")
    username = st.session_state["username"]
    tab_plan, tab_trips, tab_master = st.tabs(["Creer une tournee", "Tournees", "Camions / chauffeurs"])
    with tab_plan:
        ready = ready_for_shipping()
        if ready.empty:
            st.info("Aucune marchandise conforme prete a expedier.")
        else:
            ready = ready.copy()
            ready["label"] = ready.apply(lambda r: f"#{int(r['id'])} · {r['num_commande']} · {r['nom_client']} · {to_float(r['poids_t']):.0f} kg · {norm_text(r.get('destination')) or 'destination ?'}", axis=1)
            st.dataframe(ready[["id", "num_commande", "nom_client", "article_int", "couleur", "poids_t", "destination", "status"]], hide_index=True, use_container_width=True)
            trucks = pd.DataFrame(q_all("SELECT * FROM trucks WHERE active=1 ORDER BY code"))
            drivers = pd.DataFrame(q_all("SELECT * FROM drivers WHERE active=1 ORDER BY name"))
            if trucks.empty:
                st.warning("Ajoutez d'abord un camion dans Camions / chauffeurs.")
            else:
                selected_labels = st.multiselect("Commandes/lignes a charger", ready["label"].tolist())
                selected_ids = ready[ready["label"].isin(selected_labels)]["id"].astype(int).tolist()
                selected_weight = float(ready[ready["id"].isin(selected_ids)]["poids_t"].sum()) if selected_ids else 0.0
                truck_labels = {f"{r['code']} · {to_float(r['capacity_kg']):.0f} kg": int(r["id"]) for _, r in trucks.iterrows()}
                truck_sel = st.selectbox("Camion", list(truck_labels.keys()))
                truck_row = trucks[trucks["id"] == truck_labels[truck_sel]].iloc[0]
                capacity = to_float(truck_row["capacity_kg"])
                st.metric("Chargement selectionne", f"{selected_weight:.0f} / {capacity:.0f} kg", f"{(selected_weight/capacity*100 if capacity else 0):.1f}%")
                driver_labels = {"- Aucun -": None}
                if not drivers.empty:
                    driver_labels.update({f"{r['name']} · {norm_text(r['phone'])}": int(r["id"]) for _, r in drivers.iterrows()})
                c1, c2, c3 = st.columns(3)
                driver_sel = c1.selectbox("Chauffeur", list(driver_labels.keys()))
                trip_date = c2.date_input("Date depart", value=app_today() + timedelta(days=1))
                default_dest = norm_text(ready[ready["id"].isin(selected_ids)]["destination"].mode().iloc[0]) if selected_ids and not ready[ready["id"].isin(selected_ids)]["destination"].dropna().empty else ""
                destination = c3.text_input("Destination", value=default_dest)
                if st.button("Creer la tournee", type="primary"):
                    try:
                        tid = create_trip(trip_date, truck_labels[truck_sel], driver_labels[driver_sel], destination, selected_ids, username)
                        st.success(f"Tournee #{tid} creee.")
                        st.rerun()
                    except Exception as exc:
                        st.error(str(exc))
    with tab_trips:
        df = trips_df()
        if df.empty:
            st.info("Aucune tournee.")
        else:
            st.dataframe(df, hide_index=True, use_container_width=True)
            options = {f"#{int(r['id'])} · {r['trip_date']} · {r['destination']} · {r['truck_code']} · {to_float(r['fill_pct']):.1f}%": int(r["id"]) for _, r in df.iterrows()}
            sel = st.selectbox("Tournee a mettre a jour", list(options.keys()))
            status = st.selectbox("Nouveau statut", TRIP_STATUSES)
            if st.button("Mettre a jour la tournee"):
                try:
                    update_trip_status(options[sel], status, username)
                    st.success("Tournee mise a jour.")
                    st.rerun()
                except Exception as exc:
                    st.error(str(exc))
    with tab_master:
        c1, c2 = st.columns(2)
        with c1:
            st.markdown("#### Ajouter camion")
            with st.form("add_truck"):
                code = st.text_input("Code camion")
                plate = st.text_input("Immatriculation")
                capacity = st.number_input("Capacite kg", 100.0, 50000.0, 4000.0, 100.0)
                if st.form_submit_button("Ajouter camion"):
                    try:
                        add_truck(code, plate, capacity, username); st.success("Camion ajoute."); st.rerun()
                    except Exception as exc: st.error(str(exc))
            st.dataframe(pd.DataFrame(q_all("SELECT * FROM trucks ORDER BY code")), hide_index=True, use_container_width=True)
        with c2:
            st.markdown("#### Ajouter chauffeur")
            with st.form("add_driver"):
                name = st.text_input("Nom chauffeur")
                phone = st.text_input("Telephone")
                if st.form_submit_button("Ajouter chauffeur"):
                    try:
                        add_driver(name, phone, username); st.success("Chauffeur ajoute."); st.rerun()
                    except Exception as exc: st.error(str(exc))
            st.dataframe(pd.DataFrame(q_all("SELECT * FROM drivers ORDER BY name")), hide_index=True, use_container_width=True)


def page_analysis() -> None:
    hero("Analyse IA", "Alertes et recommandations explicables basees sur les donnees et regles du systeme")
    recs = pd.DataFrame(ai_recommendations())
    st.dataframe(recs, hide_index=True, use_container_width=True)
    st.markdown("#### Qualite / Pareto des defauts")
    qdf = pd.DataFrame(q_all("SELECT defect_reason,COUNT(*) AS incidents,SUM(qty_affected) AS qte FROM quality_incidents WHERE result='NON_CONFORME' GROUP BY defect_reason ORDER BY incidents DESC"))
    if qdf.empty: st.caption("Pas encore de non-conformite enregistree.")
    else: st.dataframe(qdf, hide_index=True, use_container_width=True)
    st.markdown("#### Performance logistique")
    tdf = trips_df()
    if tdf.empty: st.caption("Pas encore de tournee.")
    else:
        st.dataframe(tdf[["trip_date", "destination", "truck_code", "capacity_kg", "load_kg", "fill_pct", "status"]], hide_index=True, use_container_width=True)


def page_notifications() -> None:
    hero("Notifications", "Evenements inter-services")
    role = st.session_state["role"]
    username = st.session_state["username"]
    rows = q_all("SELECT * FROM notifications WHERE user_role IN (?, 'ADMIN') ORDER BY id DESC LIMIT 300", (role,)) if role != "ADMIN" else q_all("SELECT * FROM notifications ORDER BY id DESC LIMIT 300")
    df = pd.DataFrame(rows)
    if df.empty:
        st.info("Aucune notification.")
        return
    st.dataframe(df[["id", "created_at", "user_role", "event_type", "title", "message", "read_at"]], hide_index=True, use_container_width=True, height=560)
    unread = df[df["read_at"].isna()]["id"].astype(int).tolist() if "read_at" in df.columns else []
    if unread and st.button("Marquer les notifications visibles comme lues"):
        with db_conn() as conn:
            now = app_now().isoformat(timespec="seconds")
            for nid in unread: conn.execute("UPDATE notifications SET read_at=? WHERE id=?", (now, nid))
        audit(username, "READ_NOTIFICATIONS", "notification", "bulk", None, {"count": len(unread)})
        st.rerun()


def page_admin() -> None:
    hero("Administration", "Utilisateurs, audit et parametres techniques")
    username = st.session_state["username"]
    tab_users, tab_audit, tab_system = st.tabs(["Utilisateurs", "Audit", "Systeme"])
    with tab_users:
        with st.form("new_user"):
            c1, c2, c3 = st.columns(3)
            new_user = c1.text_input("Utilisateur")
            new_role = c2.selectbox("Role", ROLES)
            new_pwd = c3.text_input("Mot de passe initial", type="password")
            if st.form_submit_button("Creer utilisateur"):
                try:
                    create_user(new_user, new_pwd, new_role, username)
                    st.success("Utilisateur cree."); st.rerun()
                except Exception as exc: st.error(str(exc))
        users = pd.DataFrame(q_all("SELECT id,username,role,active,created_at FROM users ORDER BY username"))
        st.dataframe(users, hide_index=True, use_container_width=True)
        if not users.empty:
            selected = st.selectbox("Utilisateur a activer/desactiver", users["username"].tolist())
            row = users[users["username"] == selected].iloc[0]
            if selected != username and st.button("Basculer actif/inactif"):
                with db_conn() as conn:
                    conn.execute("UPDATE users SET active=? WHERE id=?", (0 if int(row["active"]) else 1, int(row["id"])))
                audit(username, "TOGGLE_USER", "user", int(row["id"]), {"active": int(row["active"])}, {"active": 0 if int(row["active"]) else 1})
                st.rerun()
    with tab_audit:
        audit_df = pd.DataFrame(q_all("SELECT * FROM audit_log ORDER BY id DESC LIMIT 1000"))
        st.dataframe(audit_df, hide_index=True, use_container_width=True, height=600)
    with tab_system:
        st.json({
            "version": VERSION, "database": str(DB_PATH), "timezone": APP_TIMEZONE,
            "OR-Tools": ORTOOLS_AVAILABLE, "ReportLab": REPORTLAB_AVAILABLE,
            "client_portal": bool(CLIENT_ACCESS_CODE),
        })
        imp = active_import()
        if imp: st.json(imp)
        st.warning("Sauvegardez regulierement alluco.db. En production multi-utilisateurs, migrez vers PostgreSQL.")

# =============================================================================
# 16. APP ROUTER
# =============================================================================
def render_internal_app() -> None:
    if not users_exist():
        bootstrap_admin_ui(); return
    if not st.session_state.get("username"):
        login_ui(); return
    nav = sidebar_navigation()
    pages = {
        "dashboard": page_dashboard, "orders": page_orders, "planning": page_planning,
        "balancelles": page_balancelles, "magasin": page_magasin, "laquage": page_laquage, "quality": page_quality,
        "relaquage": page_relaquage, "logistics": page_logistics, "analysis": page_analysis,
        "notifications": page_notifications, "admin": page_admin,
    }
    fn = pages.get(nav, page_dashboard)
    fn()


def render_app() -> None:
    if st is None:
        raise RuntimeError("Streamlit n'est pas installe. Lancez: pip install -r requirements.txt")
    init_db()
    st.set_page_config(page_title=APP_NAME, page_icon="A", layout="wide", initial_sidebar_state="expanded")
    dark_mode = bool(st.session_state.get("ui_dark_mode", False))
    st.markdown(app_css(dark_mode), unsafe_allow_html=True)

    # Avant authentification, conserver une identité visuelle cohérente et le choix du thème.
    if not st.session_state.get("username"):
        with st.sidebar:
            st.markdown(brand_html(dark_mode, compact=True), unsafe_allow_html=True)
            st.toggle("Mode sombre", key="ui_dark_mode", help="Basculer entre le thème clair et le thème sombre")
            st.divider()
            st.caption(f"Version {VERSION}")
        top = st.radio("Espace", ["Interne", "Portail client"], horizontal=True, label_visibility="collapsed")
        if top == "Portail client":
            client_portal_ui(); return
    render_internal_app()


# =============================================================================
# 17. SELF TESTS / CLI
# =============================================================================
def self_test() -> None:
    print(f"ALLUCO ERP {VERSION} self-test")
    assert infer_bars_per_bal("EC40100") == 14
    y, w, start, end = next_planning_period(date(2026, 10, 7))
    assert (y, w, start, end) == (2026, 42, date(2026, 10, 12), date(2026, 10, 16))
    assert subtract_business_days(date(2026, 10, 15), 2) == date(2026, 10, 13)
    assert white_black_conflict(["BLC", "NOIR"])
    pwd = "TestPassword123!"
    h = make_password_hash(pwd)
    assert verify_password_hash(pwd, h) and not verify_password_hash("bad", h)

    # Test du moteur sans toucher a la base reelle: monkey patch de la source active.
    synthetic = pd.DataFrame([
        {"id": 1, "num_commande": "C1", "date_creation": "2026-09-01", "nom_client": "A", "article": "EC40100-R7016", "article_int": "EC40100", "couleur": "R7016", "nuance": 27, "qte_commandee": 140, "reste_a_livrer": 140, "preleve": 0, "reservation_brut": "oui", "num_of": "OF1", "prod_statut": "cree", "qte_commencee": 0, "qte_restante": 140, "qte_recue": 140, "reserver_br": 140, "stock_physique": 140, "reserver": 140, "lancement": 140, "relaquage": 0, "poids_un": 1.0, "poids_t": 140, "poudre": 7.28, "barre_bal": 14, "nbre_bal": 10, "temps_h": 10*4/60, "stock_brut": 140, "date_livraison": "2026-10-14", "destination": "Sfax", "status": "A_PLANIFIER", "active": 1},
        {"id": 2, "num_commande": "C2", "date_creation": "2026-09-02", "nom_client": "B", "article": "FR100-R7016", "article_int": "FR100", "couleur": "R7016", "nuance": 27, "qte_commandee": 130, "reste_a_livrer": 130, "preleve": 0, "reservation_brut": "oui", "num_of": "OF2", "prod_statut": "cree", "qte_commencee": 0, "qte_restante": 130, "qte_recue": 130, "reserver_br": 130, "stock_physique": 130, "reserver": 130, "lancement": 130, "relaquage": 0, "poids_un": 1.2, "poids_t": 156, "poudre": 8.1, "barre_bal": 13, "nbre_bal": 10, "temps_h": 10*4/60, "stock_brut": 130, "date_livraison": "2026-10-15", "destination": "Sfax", "status": "A_PLANIFIER", "active": 1},
    ])
    global load_active_orders_df, q_all
    old_loader = load_active_orders_df
    old_q_all = q_all
    try:
        load_active_orders_df = lambda: synthetic.copy()
        def fake_q_all(sql: str, params: Sequence[Any] = ()):
            if "relaquage_orders" in sql:
                return []
            return old_q_all(sql, params)
        q_all = fake_q_all
        cfg = PlannerConfig(2026, 42, capacity_h=8.0, solver_seconds=3)
        result = generate_plan(cfg)
        assert result["confidence"] == 100, result["hard_errors"]
        assert sum(len(x) for x in result["days"].values()) == 2
        assert all(x["Charge totale h"] <= 8.0 + 1e-6 for x in result["metrics"]["days"])
        bal_test = build_balancelle_register(pd.DataFrame([{
            "id": 1, "planned_date": "2026-10-12", "sequence_no": 1, "num_commande": "C1", "nom_client": "A",
            "article": "EC40100-R7016", "article_int": "EC40100", "couleur": "R7016", "num_of": "OF1",
            "planned_qty": 140, "nbre_bal": 10,
        }]))
        assert len(bal_test) == 10 and int(bal_test["Qté dans BAL"].sum()) == 140
        assert set(bal_test["Statut"]) == {"Complète"}
    finally:
        load_active_orders_df = old_loader
        q_all = old_q_all
    print("[OK] helpers, securite et moteur planning")


if __name__ == "__main__":
    if "--self-test" in sys.argv:
        self_test()
    elif "--init-db" in sys.argv:
        init_db(); print(f"Base initialisee: {DB_PATH}")
    else:
        if st is None:
            print("Installez les dependances puis lancez: streamlit run alluco_erp.py")
        else:
            render_app()
