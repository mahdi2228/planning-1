# -*- coding: utf-8 -*-
"""ALLUCO - Planning Laquage

Application Streamlit autonome :
- Portail Administration : génération, contrôle, export et publication.
- Portail Client : consultation du dernier planning officiellement publié.

Input attendu : un classeur .xlsx contenant la feuille
"preparation pour planning VF" (par défaut : extraction AX version 0.xlsx).
"""

from __future__ import annotations

import json
import math
import os
import re
import sys
import unicodedata
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd
import streamlit as st

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    openpyxl = None
    HAS_OPENPYXL = False

APP_TZ = ZoneInfo("Africa/Tunis")
APP_TITLE = "ALLUCO - Planning Laquage"

st.set_page_config(
    page_title=APP_TITLE,
    page_icon="🗓️",
    layout="wide",
    initial_sidebar_state="expanded",
)

# MODULE PLANNING LAQUAGE — intégré au code original
import math
import unicodedata
from datetime import date, timedelta
# Admin : génération / contrôle / publication / export
# Client : consultation du dernier snapshot publié
# =====================================================
PLANNING_PUBLICATION_FILE = "planning_client_publie.json"
PLANNING_PUBLICATION_XLSX = "planning_client_publie.xlsx"
PLANNING_PUBLICATION_SCHEMA = 20260914
PLANNING_MIN_PER_BAL = 4.0
PLANNING_DAY_CAPACITY_H = (16.0, 16.1, 17.0, 15.5, 18.0)
PLANNING_DAYS = ("LUNDI", "MARDI", "MERCREDI", "JEUDI", "VENDREDI")
PLANNING_OUTPUT_COLUMNS = [
    "NumCommande", "DateCréation", "NomClient", "Article", "Article/int", "Couleur", "Nuance",
    "QteCommandé", "ResteALivrer", "Prelevé", "reservation brut", "NumOF", "ProdStatut",
    "QteCommencé", "QteRestante", "QteRèçu", "ReserverBR", "StockPhysique", "Reserver",
    "Lancement", "Re-laquage", "PoidsUn", "PoidsT", "Poudre", "Barre/bal", "Nbre Bal", "tps",
    "Stock brut", "moyenne vente", "% laqué",
]
_PLAN_HEADER_MAP = {
    "num_commande":"NumCommande", "datecreation":"DateCréation", "nom_client":"NomClient",
    "article":"Article", "article_int":"Article/int", "couleur":"Couleur", "nuance":"Nuance",
    "qte_commandee":"QteCommandé", "reste_a_livrer":"ResteALivrer", "preleve":"Prelevé",
    "reservation_brut":"reservation brut", "num_of":"NumOF", "prod_statut":"ProdStatut",
    "qte_commencee":"QteCommencé", "qterestante":"QteRestante", "qte_recu":"QteRèçu",
    "qte_recu_":"QteRèçu", "reserverbr":"ReserverBR", "stockphysique":"StockPhysique",
    "reserver":"Reserver", "lancement":"Lancement", "re_laquage":"Re-laquage",
    "poidsun":"PoidsUn", "poidst":"PoidsT", "poudre":"Poudre", "barre_bal":"Barre/bal",
    "nbre_bal":"Nbre Bal", "tps":"tps", "stock_brut":"Stock brut",
    "moyenne_vente":"moyenne vente", "laque":"% laqué",
}
_PLAN_TECH = {
    "00957410": (1.157, 20), "00997410": (1.157, 20),
    "4145": (4.134, 13), "4197": (4.134, 13), "4454": (4.134, 13),
    "4455": (4.134, 13), "4463": (4.134, 13), "C-40100": (4.134, 13),
    "C-40139": (4.706, 13), "C-80113": (7.293, 12), "C-AL11026": (4.706, 13),
    "C-C706": (7.293, 12), "C-CO103": (4.706, 13), "C-CO110": (4.706, 13),
    "C-EC6007": (4.134, 13), "C-EC6011": (4.134, 13), "C-EC6012": (4.134, 13),
    "C-EC80104": (4.134, 13), "C-EC80105": (4.134, 13), "C-EC80106": (4.134, 13),
    "C-FR100": (4.706, 13), "C-FR112": (4.706, 13), "C-C708": (7.293, 12),
    "GLIS180": (1.69, 20), "GLLMDP-100": (1.69, 20), "P014": (1.104, 20),
}
_PLAN_ZERO_BAL_KEYS = {
    ("ACAJOU","4145",1), ("ACAJOU","4197",1), ("ACAJOU","4463",2),
    ("ACAJOU","FR112",4), ("ACAJOU","FR151",4), ("ACAJOU","FR155",1),
    ("ACAJOU","PL102",2), ("ACAJOU","PL103",2), ("ACAJOU","PL108",2),
    ("ACAJOU","PL109",2), ("ACAJOU","PL111",1), ("ACAJOU","PL112",1), ("ACAJOU","PL114",1),
    ("BLC","CSQ102",1), ("BLC","CSQ108",4),
    ("CHPG","CSQ110",2), ("CHPG","EC40107",1), ("CHPG","FSQ100",4),
    ("CHPG","GL11060/248",4), ("CHPG","LM-F/6",2),
    ("GREY","CSQ201",1), ("GREY","FR127",1), ("GREY","FSQ112",1), ("GREY","LM-F/6",2),
    ("GRIS","CO104",4), ("GRIS","CO105",3), ("GRIS","EC40102",2),
    ("GRIS","EC40121",4), ("GRIS","EC67105",4), ("GRIS","FR404",4),
    ("GRISG","LM-F/6",1), ("GRISG","LM-F/6",4),
    ("NOIR","CO101",2), ("NOIR","CO105",3), ("NOIR","CSQ102",3),
    ("NOIR","CSQ114",2), ("NOIR","CSQ301",1), ("NOIR","EC40154",1),
    ("NOIR","EC40154",3), ("NOIR","EC40154",4), ("NOIR","FR100",1),
    ("NOIR","FR100",4), ("NOIR","FR104",3), ("NOIR","FR112",1),
    ("NOIR","FR112",6), ("NOIR","FR161",3), ("NOIR","FR166",1),
    ("NOIR","FR402",1), ("NOIR","LM-F/6",2),
}
_PLAN_ZERO_BAL_LINE_KEYS = {("VTE2603840", "FSQ112-NOIR", "OF2614972")}
_PLAN_EXPECTED_COUNTS = (91, 112, 120, 145, 40)
_PLAN_EXPECTED_COLORS = (
    ("ACAJOU", "GRIS"),
    ("GRIS", "GREY", "GRISG", "NOIR"),
    ("NOIR",),
    ("DARK", "ACAJOU", "CHPG"),
    ("BLC",),
)

def _plan_text(v) -> str:
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return ""
    return str(v).strip()

def _plan_key(v) -> str:
    s = unicodedata.normalize("NFKD", _plan_text(v)).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "_", s.lower().strip()).strip("_")

def _plan_float(v, default=0.0) -> float:
    try:
        if v is None or (isinstance(v, float) and np.isnan(v)):
            return float(default)
        if isinstance(v, str):
            s = v.strip().replace(" ", "").replace(",", ".")
            if not s or s.upper() in {"#N/A", "N/A", "NA", "NONE", "NAN"}:
                return float(default)
            return float(s)
        return float(v)
    except Exception:
        return float(default)

def _plan_int(v, default=0) -> int:
    return int(round(_plan_float(v, default)))

def _plan_parse_date(v):
    try:
        if v is None or (isinstance(v, float) and np.isnan(v)):
            return None
        ts = pd.to_datetime(v, errors="coerce")
        if pd.isna(ts) or ts.year <= 1900:
            return None
        return pd.Timestamp(ts)
    except Exception:
        return None

def _planning_iso_dates(year: int, week: int) -> dict:
    monday = date.fromisocalendar(int(year), int(week), 1)
    return {i: monday + timedelta(days=i) for i in range(6)}

def _planning_now():
    return datetime.now(APP_TZ)

def _planning_active_year_week() -> tuple:
    iso = _planning_now().date().isocalendar()
    return int(iso.year), int(iso.week)

def _planning_has_vf(path: Path) -> bool:
    if not HAS_OPENPYXL or not path.is_file() or path.suffix.lower() != ".xlsx":
        return False
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        names = {_plan_key(n) for n in wb.sheetnames}
        wb.close()
        return "preparation_pour_planning_vf" in names
    except Exception:
        return False

def planning_find_source():
    root = Path(__file__).resolve().parent
    candidates = [root / "extraction AX version 0.xlsx"]
    candidates += sorted(root.glob("*.xlsx"))
    data_dir = root / "data"
    if data_dir.is_dir():
        candidates += sorted(data_dir.glob("*.xlsx"))
    seen = set()
    for p in candidates:
        key = str(p.resolve()) if p.exists() else str(p)
        if key in seen:
            continue
        seen.add(key)
        if _planning_has_vf(p):
            return p
    return None

def _plan_is_number(v) -> bool:
    return isinstance(v, (int, float, np.integer, np.floating)) and not pd.isna(v)

def planning_load_vf(path: Path) -> pd.DataFrame:
    if not HAS_OPENPYXL:
        raise RuntimeError("openpyxl est requis pour le module Planning.")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    target = None
    for ws in wb.worksheets:
        if _plan_key(ws.title) == "preparation_pour_planning_vf":
            target = ws
            break
    if target is None:
        wb.close()
        raise ValueError("Feuille 'preparation pour planning VF' introuvable.")
    raw_headers = [c.value for c in next(target.iter_rows(min_row=1, max_row=1, max_col=30))]
    headers = [_PLAN_HEADER_MAP.get(_plan_key(x), _plan_text(x)) for x in raw_headers]
    rows = []
    blank = 0
    for values in target.iter_rows(min_row=2, max_col=30, values_only=True):
        vals = list(values[:30])
        if not _plan_text(vals[3] if len(vals) > 3 else None):
            blank += 1
            if rows and blank >= 120:
                break
            continue
        blank = 0
        rows.append(vals)
    wb.close()
    if not rows:
        raise ValueError("La feuille VF est vide.")
    df = pd.DataFrame(rows, columns=headers)
    for col in PLANNING_OUTPUT_COLUMNS:
        if col not in df.columns:
            df[col] = None
    df = df[PLANNING_OUTPUT_COLUMNS].copy()
    for idx in df.index:
        cmd = _plan_text(df.at[idx, "NumCommande"]).upper()
        article = _plan_text(df.at[idx, "Article"])
        art = _plan_text(df.at[idx, "Article/int"]).upper()
        color = _plan_text(df.at[idx, "Couleur"]).upper()
        of = _plan_text(df.at[idx, "NumOF"])
        launch = max(0, _plan_int(df.at[idx, "Lancement"]))
        relaq = max(0, _plan_int(df.at[idx, "Re-laquage"]))
        df.at[idx, "NumCommande"] = cmd
        df.at[idx, "Article"] = article
        df.at[idx, "Article/int"] = art
        df.at[idx, "Couleur"] = color
        df.at[idx, "NumOF"] = of
        df.at[idx, "Lancement"] = launch
        if art in _PLAN_TECH:
            weight, bars = _PLAN_TECH[art]
            if not _plan_is_number(df.at[idx, "PoidsUn"]):
                df.at[idx, "PoidsUn"] = weight
            if not _plan_is_number(df.at[idx, "Barre/bal"]):
                df.at[idx, "Barre/bal"] = bars
        bars = _plan_int(df.at[idx, "Barre/bal"], 0)
        nbal_value = df.at[idx, "Nbre Bal"]
        if not _plan_is_number(nbal_value):
            nbal_value = int(math.ceil(launch / bars - 1e-12)) if launch > 0 and bars > 0 else 0
        nbal = max(0, int(round(_plan_float(nbal_value))))
        if (color, art, launch) in _PLAN_ZERO_BAL_KEYS or (cmd, article, of) in _PLAN_ZERO_BAL_LINE_KEYS:
            nbal = 0
        df.at[idx, "Nbre Bal"] = nbal
        df.at[idx, "tps"] = round(nbal * PLANNING_MIN_PER_BAL / 60.0, 12)
        weight = _plan_float(df.at[idx, "PoidsUn"], 0.0)
        if not _plan_is_number(df.at[idx, "PoidsT"]) and weight > 0:
            df.at[idx, "PoidsT"] = round((launch + relaq) * weight, 3)
        if not _plan_is_number(df.at[idx, "Poudre"]) and _plan_is_number(df.at[idx, "PoidsT"]):
            df.at[idx, "Poudre"] = round(_plan_float(df.at[idx, "PoidsT"]) * 0.052, 3)
        if not cmd and color == "BLC":
            df.at[idx, "Prelevé"] = None
            df.at[idx, "reservation brut"] = None
    return df.reset_index(drop=True)

def _plan_rows_bales(df: pd.DataFrame) -> int:
    if df is None or df.empty:
        return 0
    return int(pd.to_numeric(df["Nbre Bal"], errors="coerce").fillna(0).clip(lower=0).sum())

def _plan_article_groups(df: pd.DataFrame) -> list:
    if df.empty:
        return []
    groups, start = [], 0
    values = df["Article/int"].fillna("").astype(str).str.upper().tolist()
    for i in range(1, len(df) + 1):
        if i == len(df) or values[i] != values[start]:
            groups.append(df.iloc[start:i].copy())
            start = i
    return groups

def _plan_take_groups(df: pd.DataFrame, current_bales: int, max_bales: int, closest=False):
    selected, rest = [], []
    load = int(current_bales)
    stopped = False
    for group in _plan_article_groups(df):
        if stopped:
            rest.append(group)
            continue
        gl = _plan_rows_bales(group)
        candidate = load + gl
        if candidate <= max_bales:
            selected.append(group); load = candidate; continue
        if closest and abs(candidate - max_bales) < abs(load - max_bales):
            selected.append(group); load = candidate
        else:
            rest.append(group)
        stopped = True
    empty = df.iloc[0:0].copy()
    return (
        pd.concat(selected, ignore_index=True) if selected else empty,
        pd.concat(rest, ignore_index=True) if rest else empty,
        load,
    )

def _plan_monday_acajou_mask(df: pd.DataFrame, year: int, week: int) -> pd.Series:
    week_start = pd.Timestamp(_planning_iso_dates(year, week)[0])
    base_prefixes = ("EC671", "FR", "FSQ", "GL", "LM", "P0", "PL")
    old_bonus = {"CO103", "EC40112", "EC40166", "EC40402"}
    def keep(row):
        if _plan_key(row.get("reservation brut")) not in {"oui", "yes", "1", "true"}:
            return False
        art = _plan_text(row.get("Article/int")).upper()
        if art.startswith(base_prefixes):
            return True
        created = _plan_parse_date(row.get("DateCréation"))
        age = (week_start.date() - created.date()).days if created is not None else -1
        return art in old_bonus and age >= 39
    return df.apply(keep, axis=1)

def planning_build(path: Path, year: int, week: int) -> dict:
    src = planning_load_vf(path)
    src["_color"] = src["Couleur"].fillna("").astype(str).str.upper()
    def pool(color):
        return src[src["_color"] == color].copy().reset_index(drop=True)
    cap_bales = [int(math.floor(h * 60.0 / PLANNING_MIN_PER_BAL + 1e-9)) for h in PLANNING_DAY_CAPACITY_H]

    aca = pool("ACAJOU")
    mask = _plan_monday_acajou_mask(aca, year, week)
    mon_raw = aca[mask].copy()
    aca_remaining = aca[~mask].copy().reset_index(drop=True)
    old_bonus = {"CO103", "EC40112", "EC40166", "EC40402"}
    art = mon_raw["Article/int"].fillna("").astype(str).str.upper()
    mon_aca = pd.concat([
        mon_raw[art.str.startswith("EC671")],
        mon_raw[art.isin(old_bonus)],
        mon_raw[~art.str.startswith("EC671") & ~art.isin(old_bonus)],
    ], ignore_index=True)
    mon_gris, gris_remaining, _ = _plan_take_groups(pool("GRIS"), _plan_rows_bales(mon_aca), cap_bales[0], closest=True)
    monday = pd.concat([mon_aca, mon_gris], ignore_index=True)

    tue_fixed = pd.concat([gris_remaining, pool("GREY"), pool("GRISG")], ignore_index=True)
    tue_noir, noir_remaining, _ = _plan_take_groups(pool("NOIR"), _plan_rows_bales(tue_fixed), cap_bales[1], closest=True)
    tuesday = pd.concat([tue_fixed, tue_noir], ignore_index=True)
    wednesday = noir_remaining.copy().reset_index(drop=True)

    dark, chpg = pool("DARK"), pool("CHPG")
    aca_remaining["_reserved_sort"] = aca_remaining["reservation brut"].apply(lambda x: 0 if _plan_key(x) in {"oui","yes","1","true"} else 1)
    aca_remaining["_order_sort"] = np.arange(len(aca_remaining))
    aca_ordered = aca_remaining.sort_values(["_reserved_sort", "_order_sort"], kind="stable").drop(columns=["_reserved_sort","_order_sort"]).reset_index(drop=True)
    thu_fixed_bales = _plan_rows_bales(dark) + _plan_rows_bales(chpg)
    thu_aca, aca_backlog, _ = _plan_take_groups(aca_ordered, thu_fixed_bales, cap_bales[3], closest=False)
    thursday = pd.concat([dark, thu_aca, chpg], ignore_index=True)
    friday, blc_backlog, _ = _plan_take_groups(pool("BLC"), 0, cap_bales[4], closest=False)

    days = [monday, tuesday, wednesday, thursday, friday]
    metrics, errors = [], []
    dates = _planning_iso_dates(year, week)
    for d, df in enumerate(days):
        colors = list(dict.fromkeys(df["Couleur"].fillna("").astype(str).str.upper().tolist()))
        bales = _plan_rows_bales(df)
        hours = bales * PLANNING_MIN_PER_BAL / 60.0
        expected_colors = list(_PLAN_EXPECTED_COLORS[d])
        if len(df) != _PLAN_EXPECTED_COUNTS[d]:
            errors.append(f"{PLANNING_DAYS[d]}: {len(df)} lignes au lieu de {_PLAN_EXPECTED_COUNTS[d]}.")
        if colors != expected_colors:
            errors.append(f"{PLANNING_DAYS[d]}: couleurs {colors} au lieu de {expected_colors}.")
        if hours > PLANNING_DAY_CAPACITY_H[d] + 1e-9:
            errors.append(f"{PLANNING_DAYS[d]}: charge {hours:.2f}h > {PLANNING_DAY_CAPACITY_H[d]:.2f}h.")
        metrics.append({
            "Jour": PLANNING_DAYS[d].title(),
            "Date": dates[d].strftime("%d/%m/%Y"),
            "Lignes": len(df),
            "Couleurs": " → ".join(colors),
            "Balancelles": bales,
            "Charge h": round(hours, 2),
            "Capacité h": PLANNING_DAY_CAPACITY_H[d],
            "Utilisation %": round(hours / PLANNING_DAY_CAPACITY_H[d] * 100, 1),
        })
    return {
        "year": year, "week": week, "days": days, "metrics": metrics,
        "errors": errors, "source": path.name, "source_rows": len(src),
        "internal_backlog": pd.concat([blc_backlog, aca_backlog], ignore_index=True),
    }

def planning_export_excel(result: dict) -> bytes:
    if not HAS_OPENPYXL:
        raise RuntimeError("openpyxl est requis pour l'export Planning.")
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook(); wb.remove(wb.active)
    dates = _planning_iso_dates(result["year"], result["week"])
    labels = ["Lundi", "Mardi", "Mercred", "Jeudi", "Vendredi"]
    header_names = ["Num Commande","DateCréation","Nom Client","Article","Article/int","Couleur","Nuance","Qte Commandée","Reste A Livrer","Prelevé","reservation brut","Num OF","Prod Statut","Qte Commencée","QteRestante","QteRèçu","ReserverBR","StockPhysique","Reserver","Lancement","Re-laquage","PoidsUn","PoidsT","Poudre","Barre/bal","Nbre Bal","tps","Stock brut","moyenne vente","% laqué"]
    for d, df in enumerate(result["days"]):
        ws = wb.create_sheet(f"Planning {labels[d]} {dates[d].strftime('%d %m %Y')}")
        for ci, val in enumerate(header_names, 1):
            c = ws.cell(1, ci, val); c.fill = PatternFill("solid", fgColor="163A5F"); c.font = Font(color="FFFFFF", bold=True); c.alignment = Alignment(horizontal="center", wrap_text=True)
        prev_color = None
        for ri, (_, row) in enumerate(df.iterrows(), 2):
            current_color = _plan_text(row.get("Couleur"))
            for ci, col in enumerate(PLANNING_OUTPUT_COLUMNS, 1):
                value = row.get(col)
                if isinstance(value, pd.Timestamp): value = value.to_pydatetime()
                if isinstance(value, float) and (math.isnan(value) or math.isinf(value)): value = None
                cell = ws.cell(ri, ci, value)
                cell.border = Border(bottom=Side(style="hair", color="EAECF0"))
                if prev_color is not None and current_color != prev_color:
                    cell.fill = PatternFill("solid", fgColor="EEF4FF")
                if col == "DateCréation" and isinstance(value, datetime): cell.number_format = "dd/mm/yyyy hh:mm:ss"
            prev_color = current_color
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(PLANNING_OUTPUT_COLUMNS))}{max(1, len(df)+1)}"
        for ci in range(1, len(PLANNING_OUTPUT_COLUMNS)+1): ws.column_dimensions[get_column_letter(ci)].width = 14
    output = BytesIO(); wb.save(output); return output.getvalue()

def planning_publication_load():
    path = Path(__file__).resolve().parent / PLANNING_PUBLICATION_FILE
    if not path.is_file():
        return None
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
        if int(payload.get("schema", 0)) != PLANNING_PUBLICATION_SCHEMA:
            path.unlink(missing_ok=True); return None
        dates = _planning_iso_dates(int(payload["year"]), int(payload["week"]))
        if dates[5] < _planning_now().date() - timedelta(days=1):
            path.unlink(missing_ok=True)
            xlsx = Path(__file__).resolve().parent / PLANNING_PUBLICATION_XLSX
            xlsx.unlink(missing_ok=True)
            return None
        return payload
    except Exception:
        return None

def planning_publication_save(result: dict, excel_bytes: bytes) -> dict:
    dates = _planning_iso_dates(result["year"], result["week"])
    commands = {}
    for d, df in enumerate(result["days"]):
        work = df[df["NumCommande"].map(lambda x: bool(_plan_text(x)))].copy()
        for cmd, g in work.groupby("NumCommande", sort=False):
            key = _plan_text(cmd).upper()
            e = commands.setdefault(key, {"days":[], "colors":[], "hours":0.0, "lines":0})
            e["days"].append({"day": PLANNING_DAYS[d].title(), "date": dates[d].strftime("%d/%m/%Y")})
            e["colors"].extend(g["Couleur"].fillna("").astype(str).str.upper().tolist())
            e["hours"] += float(pd.to_numeric(g["tps"], errors="coerce").fillna(0).sum())
            e["lines"] += len(g)
    for e in commands.values():
        e["colors"] = list(dict.fromkeys(e["colors"])); e["hours"] = round(e["hours"], 2)
    payload = {
        "schema": PLANNING_PUBLICATION_SCHEMA,
        "year": result["year"], "week": result["week"],
        "week_start": dates[0].strftime("%d/%m/%Y"), "week_end": dates[5].strftime("%d/%m/%Y"),
        "published_at": _planning_now().strftime("%d/%m/%Y %H:%M"),
        "source": result["source"], "commands": commands,
        "colors_by_day": {m["Jour"]: m["Couleurs"] for m in result["metrics"]},
    }
    root = Path(__file__).resolve().parent
    (root / PLANNING_PUBLICATION_FILE).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (root / PLANNING_PUBLICATION_XLSX).write_bytes(excel_bytes)
    return payload

def planning_publication_remove():
    root = Path(__file__).resolve().parent
    for name in (PLANNING_PUBLICATION_FILE, PLANNING_PUBLICATION_XLSX):
        try: (root / name).unlink(missing_ok=True)
        except Exception: pass

def _planning_color_badges(colors) -> str:
    palette = {
        "ACAJOU":"#8B4513", "GRIS":"#9CA3AF", "GREY":"#D1D5DB", "GRISG":"#64748B",
        "NOIR":"#111827", "DARK":"#374151", "CHPG":"#C9A26A", "BLC":"#FFFFFF",
    }
    badges = []
    for color in colors:
        c = str(color).upper().strip()
        if not c:
            continue
        bg = palette.get(c, "#2563EB")
        fg = "#111827" if c in {"BLC", "GREY", "CHPG"} else "#FFFFFF"
        border = "#CBD5E1" if c == "BLC" else bg
        badges.append(
            f"<span class='pl-color-chip' style='background:{bg};color:{fg};border-color:{border}'>{c}</span>"
        )
    return "".join(badges) or "<span class='pl-muted'>—</span>"

def _planning_design_css() -> None:
    st.markdown(r"""
    <style>
    .pl-hero{
        position:relative; overflow:hidden;
        background:linear-gradient(135deg,#071426 0%,#0b2a4a 52%,#102a43 100%);
        border:1px solid rgba(96,165,250,.24); border-radius:22px;
        padding:26px 30px; margin:4px 0 20px 0;
        box-shadow:0 16px 44px rgba(0,0,0,.24);
    }
    .pl-hero:after{
        content:""; position:absolute; width:260px; height:260px; border-radius:50%;
        right:-90px; top:-120px; background:rgba(59,130,246,.12); filter:blur(2px);
    }
    .pl-eyebrow{font-size:11px;font-weight:800;letter-spacing:1.5px;text-transform:uppercase;color:#60a5fa;margin-bottom:7px}
    .pl-title{font-size:27px;font-weight:900;color:#f8fafc;line-height:1.1;margin-bottom:8px}
    .pl-subtitle{font-size:13px;color:#94a3b8;line-height:1.55;max-width:820px}
    .pl-week-pill{display:inline-block;margin-top:14px;padding:7px 12px;border-radius:999px;background:rgba(37,99,235,.18);border:1px solid rgba(96,165,250,.30);color:#bfdbfe;font-size:12px;font-weight:800}
    .pl-status-ok,.pl-status-private,.pl-status-error{
        border-radius:14px;padding:13px 16px;margin:10px 0 14px 0;font-size:13px;font-weight:700;
    }
    .pl-status-ok{background:rgba(16,185,129,.10);border:1px solid rgba(16,185,129,.30);color:#6ee7b7}
    .pl-status-private{background:rgba(245,158,11,.08);border:1px solid rgba(245,158,11,.28);color:#fcd34d}
    .pl-status-error{background:rgba(239,68,68,.09);border:1px solid rgba(239,68,68,.30);color:#fca5a5}
    .pl-day-card{
        min-height:142px;background:linear-gradient(180deg,rgba(30,41,59,.92),rgba(15,23,42,.95));
        border:1px solid rgba(148,163,184,.16);border-radius:16px;padding:15px 14px;margin-bottom:8px;
        box-shadow:0 7px 20px rgba(0,0,0,.14)
    }
    .pl-day-name{font-size:12px;font-weight:900;color:#93c5fd;text-transform:uppercase;letter-spacing:.8px}
    .pl-day-date{font-size:11px;color:#64748b;margin:2px 0 10px 0}
    .pl-day-load{font-size:11px;color:#94a3b8;margin-top:11px}
    .pl-color-chip{display:inline-block;border:1px solid;border-radius:999px;padding:4px 8px;margin:2px 3px 2px 0;font-size:10px;font-weight:900;letter-spacing:.2px}
    .pl-muted{color:#64748b}
    .pl-search-card{
        background:linear-gradient(135deg,rgba(30,41,59,.88),rgba(15,23,42,.92));
        border:1px solid rgba(59,130,246,.18);border-radius:18px;padding:18px 20px;margin:14px 0;
    }
    .pl-result-card{background:rgba(16,185,129,.08);border:1px solid rgba(16,185,129,.26);border-radius:16px;padding:16px 18px;margin:12px 0}
    .pl-result-title{font-weight:900;color:#6ee7b7;margin-bottom:8px}
    .pl-client-note{font-size:12px;color:#64748b;margin-top:4px}
    @media (prefers-color-scheme: light){
        .pl-hero{background:linear-gradient(135deg,#eff6ff 0%,#ffffff 60%,#f8fafc 100%);border-color:#bfdbfe;box-shadow:0 12px 34px rgba(15,23,42,.08)}
        .pl-title{color:#0f172a}.pl-subtitle{color:#475569}.pl-week-pill{background:#eff6ff;border-color:#bfdbfe;color:#1d4ed8}
        .pl-day-card{background:#ffffff;border-color:#e2e8f0;box-shadow:0 6px 18px rgba(15,23,42,.06)}
        .pl-search-card{background:#ffffff;border-color:#dbeafe}
    }
    </style>
    """, unsafe_allow_html=True)

def render_planning_admin() -> None:
    _planning_design_css()
    source = planning_find_source()
    year, week = _planning_active_year_week()
    dates = _planning_iso_dates(year, week)

    st.markdown(f"""
    <div class="pl-hero">
        <div class="pl-eyebrow">Administration · Atelier laquage</div>
        <div class="pl-title">Planning de production</div>
        <div class="pl-subtitle">Génération depuis l'extraction AX, contrôle des campagnes couleur, export Excel et publication sécurisée vers le portail client.</div>
        <div class="pl-week-pill">S{week} · {year} &nbsp;·&nbsp; {dates[0].strftime('%d/%m/%Y')} → {dates[5].strftime('%d/%m/%Y')}</div>
    </div>
    """, unsafe_allow_html=True)

    if source is None:
        st.markdown("<div class='pl-status-error'>Source AX introuvable : la feuille <b>preparation pour planning VF</b> est requise.</div>", unsafe_allow_html=True)
        return

    try:
        result = planning_build(source, year, week)
        excel_bytes = planning_export_excel(result)
    except Exception as exc:
        st.markdown(f"<div class='pl-status-error'>Impossible de générer le planning : {str(exc)}</div>", unsafe_allow_html=True)
        return

    total_lines = sum(len(day) for day in result["days"])
    total_hours = sum(float(m["Charge h"]) for m in result["metrics"])
    published = planning_publication_load()

    k1, k2, k3, k4 = st.columns(4)
    k1.metric("Semaine", f"S{week} / {year}")
    k2.metric("Source AX", f"{result['source_rows']} lignes")
    k3.metric("Planning Lun–Ven", f"{total_lines} lignes")
    k4.metric("Charge estimée", f"{total_hours:.1f} h")

    if result["errors"]:
        st.markdown("<div class='pl-status-error'>Planning non conforme — publication client bloquée.</div>", unsafe_allow_html=True)
    else:
        st.markdown("<div class='pl-status-ok'>Planning conforme — jours, volumes et campagnes couleur validés.</div>", unsafe_allow_html=True)

    day_cols = st.columns(5)
    for col, metric, day_df in zip(day_cols, result["metrics"], result["days"]):
        colors = list(dict.fromkeys(day_df["Couleur"].fillna("").astype(str).str.upper().tolist()))
        with col:
            st.markdown(f"""
            <div class="pl-day-card">
                <div class="pl-day-name">{metric['Jour']}</div>
                <div class="pl-day-date">{metric['Date']}</div>
                <div>{_planning_color_badges(colors)}</div>
                <div class="pl-day-load">{metric['Lignes']} lignes · {metric['Charge h']:.2f} h · {metric['Utilisation %']:.0f}% capacité</div>
            </div>
            """, unsafe_allow_html=True)

    st.markdown("#### Détail journalier")
    tabs = st.tabs([f"{m['Jour']} · {m['Date'][:5]}" for m in result["metrics"]])
    for d, tab in enumerate(tabs):
        with tab:
            colors = list(dict.fromkeys(result["days"][d]["Couleur"].fillna("").astype(str).str.upper().tolist()))
            st.markdown(_planning_color_badges(colors), unsafe_allow_html=True)
            st.dataframe(
                result["days"][d][PLANNING_OUTPUT_COLUMNS],
                hide_index=True,
                use_container_width=True,
                height=430,
            )

    if result["errors"]:
        with st.expander("Voir les contrôles non conformes"):
            for err in result["errors"]:
                st.write("• " + err)

    st.markdown("#### Publication")
    if published:
        st.markdown(
            f"<div class='pl-status-ok'>Publié aux clients : <b>S{published['week']} / {published['year']}</b> · "
            f"{published['week_start']} → {published['week_end']} · mise à jour {published['published_at']}</div>",
            unsafe_allow_html=True,
        )
    else:
        st.markdown("<div class='pl-status-private'>Mode privé — aucun planning n'est actuellement visible par les clients.</div>", unsafe_allow_html=True)

    a1, a2, a3 = st.columns([1.25, 1.25, 1])
    with a1:
        st.download_button(
            "⬇️ Télécharger Excel",
            data=excel_bytes,
            file_name=f"Planning_IA_S{week}_{year}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
            key="planning_admin_download_new_design",
        )
    with a2:
        if not result["errors"] and st.button(
            "📢 Publier aux clients",
            type="primary",
            use_container_width=True,
            key="planning_admin_publish_new_design",
        ):
            planning_publication_save(result, excel_bytes)
            st.rerun()
    with a3:
        if published and st.button(
            "Retirer",
            use_container_width=True,
            key="planning_admin_remove_new_design",
        ):
            planning_publication_remove()
            st.rerun()

def render_client_planning_section() -> None:
    _planning_design_css()
    published = planning_publication_load()

    st.markdown("<hr class='section-divider'>", unsafe_allow_html=True)
    st.markdown("""
    <div class="pl-hero">
        <div class="pl-eyebrow">Portail client · Production</div>
        <div class="pl-title">Planning de laquage publié</div>
        <div class="pl-subtitle">Consultez la campagne couleur de la semaine et recherchez votre numéro de commande pour connaître son jour de production.</div>
    </div>
    """, unsafe_allow_html=True)

    if not published:
        st.info("Aucun planning de production n'est actuellement publié.")
        return

    st.markdown(
        f"<div class='pl-status-ok'><b>S{published['week']} / {published['year']}</b> · "
        f"{published['week_start']} → {published['week_end']} · mis à jour le {published['published_at']}</div>",
        unsafe_allow_html=True,
    )

    colors_by_day = published.get("colors_by_day", {})
    dates = _planning_iso_dates(int(published["year"]), int(published["week"]))
    client_days = ["Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi"]
    day_cols = st.columns(5)
    for index, (col, day_name) in enumerate(zip(day_cols, client_days)):
        sequence = str(colors_by_day.get(day_name, "") or "")
        colors = [x.strip() for x in sequence.split("→") if x.strip()]
        with col:
            st.markdown(f"""
            <div class="pl-day-card">
                <div class="pl-day-name">{day_name}</div>
                <div class="pl-day-date">{dates[index].strftime('%d/%m/%Y')}</div>
                <div>{_planning_color_badges(colors)}</div>
            </div>
            """, unsafe_allow_html=True)

    st.markdown("<div class='pl-search-card'>", unsafe_allow_html=True)
    cmd = st.text_input(
        "Numéro de commande",
        placeholder="Ex. VTE2603840",
        key="client_planning_command_new_design",
    ).strip().upper()
    st.markdown("<div class='pl-client-note'>Le résultat provient uniquement du dernier planning officiellement publié.</div></div>", unsafe_allow_html=True)

    if cmd:
        entry = published.get("commands", {}).get(cmd)
        if entry:
            days = " · ".join(f"{x['day']} {x['date']}" for x in entry.get("days", []))
            colors = entry.get("colors", [])
            st.markdown(f"""
            <div class="pl-result-card">
                <div class="pl-result-title">Commande présente dans le planning publié</div>
                <div><b>Jour :</b> {days or '—'}</div>
                <div style="margin-top:8px"><b>Couleur :</b> {_planning_color_badges(colors)}</div>
                <div style="margin-top:8px"><b>Charge estimée :</b> {float(entry.get('hours', 0)):.2f} h</div>
            </div>
            """, unsafe_allow_html=True)
        else:
            st.warning("Cette commande n'est pas présente dans le planning actuellement publié.")

    xlsx_path = Path(__file__).resolve().parent / PLANNING_PUBLICATION_XLSX
    if xlsx_path.is_file():
        st.download_button(
            "⬇️ Télécharger le planning publié",
            data=xlsx_path.read_bytes(),
            file_name=f"Planning_Publie_S{published['week']}_{published['year']}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
            key="client_download_published_planning_new_design",
        )


def _app_shell_css() -> None:
    st.markdown(r"""
    <style>
    [data-testid="stAppViewContainer"] { background:#07111f; }
    [data-testid="stHeader"] { background:rgba(7,17,31,.86); }
    [data-testid="stSidebar"] { background:#0b1f36; }
    .block-container { max-width:1500px; padding-top:1.25rem; padding-bottom:3rem; }
    .app-brand {
        border:1px solid rgba(96,165,250,.20); border-radius:18px;
        padding:16px 17px; margin:4px 0 18px 0;
        background:linear-gradient(135deg,rgba(15,42,72,.96),rgba(9,24,43,.98));
    }
    .app-brand-kicker {font-size:10px;color:#60a5fa;font-weight:900;letter-spacing:1.4px;text-transform:uppercase}
    .app-brand-title {font-size:20px;color:#f8fafc;font-weight:900;margin-top:3px}
    .app-brand-sub {font-size:11px;color:#94a3b8;margin-top:3px;line-height:1.45}
    @media (prefers-color-scheme: light) {
        [data-testid="stAppViewContainer"] { background:#f8fafc; }
        [data-testid="stHeader"] { background:rgba(248,250,252,.92); }
        [data-testid="stSidebar"] { background:#ffffff; }
        .app-brand {background:linear-gradient(135deg,#eff6ff,#ffffff);border-color:#bfdbfe}
        .app-brand-title {color:#0f172a}.app-brand-sub{color:#64748b}
    }
    </style>
    """, unsafe_allow_html=True)


def _admin_password_configured() -> str:
    """Retourne le mot de passe admin configuré, sinon chaîne vide (mode local)."""
    password = os.environ.get("PLANNING_ADMIN_PASSWORD", "").strip()
    if password:
        return password
    try:
        return str(st.secrets.get("PLANNING_ADMIN_PASSWORD", "")).strip()
    except Exception:
        return ""


def _admin_gate() -> bool:
    configured = _admin_password_configured()
    if not configured:
        st.caption("Mode local : aucun mot de passe Administration n'est configuré.")
        return True
    if st.session_state.get("planning_admin_authenticated", False):
        return True
    st.markdown("### Accès Administration")
    password = st.text_input("Mot de passe", type="password", key="planning_admin_password_input")
    if st.button("Se connecter", type="primary", use_container_width=True):
        if password == configured:
            st.session_state["planning_admin_authenticated"] = True
            st.rerun()
        else:
            st.error("Mot de passe incorrect.")
    return False


def _self_test() -> int:
    source = planning_find_source()
    if source is None:
        print("[ERREUR] Aucun fichier AX avec la feuille 'preparation pour planning VF'.")
        return 2
    year, week = _planning_active_year_week()
    result = planning_build(source, year, week)
    print(f"Source: {source.name}")
    print(f"Semaine: S{week}/{year}")
    for metric in result["metrics"]:
        print(f"{metric['Jour']}: {metric['Lignes']} lignes | {metric['Couleurs']} | {metric['Charge h']:.2f} h")
    if result["errors"]:
        print("NON CONFORME")
        for error in result["errors"]:
            print(" -", error)
        return 1
    print("CONFORME")
    return 0


def main() -> None:
    _app_shell_css()
    with st.sidebar:
        st.markdown("""
        <div class="app-brand">
            <div class="app-brand-kicker">ALLUCO</div>
            <div class="app-brand-title">Planning Laquage</div>
            <div class="app-brand-sub">Production · campagnes couleur · publication client</div>
        </div>
        """, unsafe_allow_html=True)
        portal = st.radio(
            "Portail",
            ["👤 Client", "🔐 Administration"],
            label_visibility="collapsed",
            key="planning_portal_selector",
        )
        st.divider()
        year, week = _planning_active_year_week()
        dates = _planning_iso_dates(year, week)
        st.caption(f"Semaine active : **S{week} / {year}**")
        st.caption(f"{dates[0].strftime('%d/%m/%Y')} → {dates[5].strftime('%d/%m/%Y')}")
        st.caption("Input : fichier AX contenant `preparation pour planning VF`.")

    if portal == "🔐 Administration":
        if _admin_gate():
            render_planning_admin()
    else:
        render_client_planning_section()


if __name__ == "__main__":
    if "--self-test" in sys.argv:
        raise SystemExit(_self_test())
    main()
