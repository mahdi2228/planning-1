# -*- coding: utf-8 -*-
"""
ALLUCO - Planning Laquage Lundi -> Vendredi
=============================================

Objectif
--------
- Entrée: un classeur Excel contenant une feuille "extraction ax" (ou une feuille
  avec les colonnes NumCommande + Article).
- Sortie: planning automatique Lundi -> Vendredi + backlog + préparation triée.
- Utilisable en CLI, et en interface Streamlit si Streamlit est installé.

Exemples
--------
CLI:
    python alluco_planning_lun_ven.py extraction.xlsx -o planning.xlsx --year 2026 --week 38

Interface:
    streamlit run alluco_planning_lun_ven.py

Notes métier reproduites depuis le classeur fourni:
- priorité de nuance provenant de la feuille/base historique;
- 4 minutes par balancelle par défaut;
- poudre = PoidsT * 0.052;
- Nbre Bal = arrondi supérieur(Lancement / Barre-balance);
- référentiel article (barres/balancelle, poids, stock brut, moyenne ventes, % laqué)
  embarqué à partir du cache du classeur historique.

Les règles "OF prêt", capacité, maximum 2 couleurs/jour et BLANC <-> NOIR/DARK
sont paramétrables. Le re-laquage à partir de StockPhysique reste désactivé par
 défaut car sa règle exacte n'est pas déductible de façon fiable du seul fichier.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import os
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from statistics import median
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:
    from ortools.sat.python import cp_model
    ORTOOLS_AVAILABLE = True
except Exception:
    cp_model = None
    ORTOOLS_AVAILABLE = False

try:
    import streamlit as st
except Exception:
    st = None

APP_NAME = "ALLUCO - Planning Laquage Lun-Ven"
DAYS = ["LUNDI", "MARDI", "MERCREDI", "JEUDI", "VENDREDI"]
N_DAYS = len(DAYS)
DEFAULT_CAPACITY_H = 16.0
DEFAULT_MIN_PER_BAL = 4.0
DEFAULT_POWDER_COEFF = 0.052
DEFAULT_CLEANING_MIN = 15
DEFAULT_TARGET_UTIL = 0.94
MAX_COLORS_PER_DAY = 2
WHITE_ALIASES = {"BLC", "BLANC", "WHITE", "R9016"}
BLACK_ALIASES = {"NOIR", "DARK", "BLACK", "R9005"}

OUTPUT_COLUMNS = [
    "Num Commande", "DateCréation", "Nom Client", "Article", "Article/int", "Couleur", "Nuance",
    "Qte Commandée", "Reste A Livrer", "Prelevé", "reservation brut", "Num OF", "Prod Statut",
    "Qte Commencée", "QteRestante", "QteRèçu", "ReserverBR", "StockPhysique", "Reserver",
    "Lancement", "Re-laquage", "PoidsUn", "PoidsT", "Poudre", "Barre/bal", "Nbre Bal", "tps",
    "Stock brut", "moyenne vente", "% laqué",
]

REFERENCE = json.loads('{"colors":{"ABRONZE":11.4,"ACAJOU":15.1,"BALSA":5.0,"BEAN":33.0,"BEECH":15.3,"BEL9010":1.2,"BLC":1.0,"BRZ":11.5,"CEDRE":9.0,"CHENE":10.3,"CHPG":11.0,"COOL":30.0,"CORINA":19.0,"DARK":38.0,"FRENE":0.5,"GALET":23.0,"GAYA":39.0,"GREEN":26.0,"GREY":24.0,"GRIS":22.0,"GRISG":25.0,"IROKO":21.0,"JAUNE":10.4,"NOCE":36.0,"NOIR":37.0,"NOYER":15.0,"OTARIE":29.0,"PALMIER":20.0,"PIN":10.1,"PURITY":4.0,"QUARTZ":28.0,"R1013":6.0,"R1015":8.0,"R1019":14.0,"R1021":15.55,"R3020":36.7,"R5012":36.4,"R5015":36.5,"R6002":26.4,"R6005":26.5,"R7016":31.0,"R7021":31.5,"R7023G":11.5,"R7044":23.5,"R8019":32.0,"R9005BR":37.5,"R9006":22.2,"R9007":22.1,"R9016":3.0,"RAL 9006":22.2,"RAL9001":2.5,"SAND":13.0,"SAPIN":10.0,"SEA":27.0,"SIPO":18.0,"SKEY":12.0,"SUN":7.0,"SWEET":35.0,"TECK":10.2,"TRESOR":0.6,"WAPA":17.0,"WENGE":16.0},"tech":{"00120410U":{"bars":10,"weight":4.0},"1076410":{"bars":10,"weight":4.0},"2563005":{"bars":20,"weight":1.0},"2563410":{"bars":10,"weight":4.0},"2563500":{"bars":10,"weight":4.0},"381/60":{"bars":20,"weight":0.115},"45AL1111":{"bars":13,"weight":4.0},"45AL1113":{"bars":12,"weight":4.0},"45AL1114":{"bars":12,"weight":4.0},"45AL1115":{"bars":12,"weight":4.0},"45AL1210":{"bars":13,"weight":4.0},"45AL1211":{"bars":12,"weight":4.0},"45AL1212":{"bars":12,"weight":4.0},"4600-101":{"bars":20,"weight":1.0},"4600-103":{"bars":20,"weight":1.0},"4600-201":{"bars":20,"weight":1.0},"4600-901":{"bars":20,"weight":1.0},"4600-905":{"bars":20,"weight":1.0},"4600-906":{"bars":20,"weight":1.0},"4600-907":{"bars":20,"weight":1.0},"4600-908":{"bars":20,"weight":1.0},"6004410":{"bars":10,"weight":4.0},"60AL2110":{"bars":12,"weight":4.0},"60AL2111":{"bars":12,"weight":4.0},"60AL2112":{"bars":14,"weight":4.0},"60AL2210":{"bars":12,"weight":4.0},"60AL2211":{"bars":13,"weight":4.0},"60AL2212":{"bars":12,"weight":4.0},"60AL2213":{"bars":12,"weight":4.0},"60AL2214":{"bars":12,"weight":4.0},"60AL2215":{"bars":12,"weight":4.0},"80114":{"bars":20,"weight":2.0},"80114/1":{"bars":20,"weight":4.0},"957410":{"bars":10,"weight":4.0},"997410":{"bars":10,"weight":4.0},"997500":{"bars":20,"weight":0.29},"AL11026":{"bars":10,"weight":5.772},"AL11144/1,1":{"bars":20,"weight":2.0},"AL11144/1,2":{"bars":20,"weight":2.0},"AL11144/1,3":{"bars":20,"weight":2.0},"AL11144/1,5":{"bars":20,"weight":2.0},"AL15401":{"bars":6,"weight":19.761},"AL15403":{"bars":10,"weight":12.163},"AL15404":{"bars":14,"weight":7.621},"AL15410":{"bars":5,"weight":10.162},"AL15411":{"bars":14,"weight":3.961},"AL15412":{"bars":14,"weight":3.666},"AL15413":{"bars":14,"weight":3.318},"AL15415":{"bars":14,"weight":3.213},"AL15416":{"bars":20,"weight":1.58},"AL15418":{"bars":20,"weight":0.986},"AL15420":{"bars":5,"weight":19.761},"AL15422":{"bars":5,"weight":19.761},"AL244/1,2":{"bars":13,"weight":4.0},"AL349/1,2":{"bars":20,"weight":2.0},"AL40122":{"bars":9,"weight":7.156499999999999},"AL40123":{"bars":9,"weight":8.775},"AL40133":{"bars":20,"weight":0.9359999999999999},"AL40165":{"bars":13,"weight":4.628},"AL40200":{"bars":13,"weight":5.07},"AL40202":{"bars":13,"weight":4.8165},"AL80104":{"bars":13,"weight":4.173},"AL80105":{"bars":13,"weight":3.523},"AL80106":{"bars":15,"weight":3.9194999999999998},"AL80114":{"bars":20,"weight":2.0},"AL80114/1":{"bars":20,"weight":1.3455},"AL80120":{"bars":12,"weight":4.914},"AL80121":{"bars":12,"weight":5.8045},"ALL11026":{"bars":10,"weight":4.0},"ALL11142/1.1":{"bars":20,"weight":1.034},"ALL11144/1.2":{"bars":20,"weight":0.826},"ALL144/3":{"bars":10,"weight":4.0},"ALL145/2.8":{"bars":17,"weight":1.75},"ALL213/1.5":{"bars":20,"weight":6.162},"ALL244/1.2":{"bars":13,"weight":3.263},"ALL349/1.5":{"bars":10,"weight":4.0},"BL771":{"bars":6,"weight":14.7},"BL772":{"bars":6,"weight":17.38},"BL774":{"bars":20,"weight":1.005},"BL783":{"bars":20,"weight":0.42},"BL784":{"bars":20,"weight":0.285},"BL785":{"bars":20,"weight":0.285},"C-04340":{"bars":10,"weight":4.0},"C-11132":{"bars":20,"weight":1.0},"C-15286":{"bars":13,"weight":3.0},"C-40104":{"bars":13,"weight":3.0},"C-40108":{"bars":20,"weight":4.0},"C-40112":{"bars":14,"weight":3.0},"C-40121":{"bars":9,"weight":2.0},"C-40123":{"bars":20,"weight":4.0},"C-40125":{"bars":20,"weight":4.0},"C-40149":{"bars":10,"weight":4.0},"C-40150":{"bars":13,"weight":2.0},"C-40154":{"bars":9,"weight":2.0},"C-40161":{"bars":14,"weight":2.184},"C-40162":{"bars":14,"weight":3.0},"C-40163":{"bars":14,"weight":3.0},"C-40402":{"bars":13,"weight":2.0},"C-40406":{"bars":13,"weight":2.0},"C-45AL1110":{"bars":10,"weight":4.0},"C-45AL1210":{"bars":10,"weight":4.0},"C-45AL1512":{"bars":10,"weight":4.0},"C-52003":{"bars":5,"weight":6.0},"C-52006":{"bars":5,"weight":6.0},"C-52021":{"bars":10,"weight":4.0},"C-60AL2110":{"bars":10,"weight":4.0},"C-67103":{"bars":12,"weight":2.0},"C-67104":{"bars":13,"weight":3.0},"C-67105":{"bars":13,"weight":3.0},"C-67106":{"bars":15,"weight":3.0},"C-67107":{"bars":20,"weight":2.0},"C-67108":{"bars":20,"weight":2.0},"C-67114":{"bars":20,"weight":2.0},"C-6735":{"bars":14,"weight":3.0},"C-67350":{"bars":14,"weight":3.0},"C-6736":{"bars":14,"weight":2.0},"C-AL40122":{"bars":9,"weight":2.0},"C-C701":{"bars":10,"weight":4.0},"C-C702":{"bars":10,"weight":4.0},"C-C704":{"bars":10,"weight":4.0},"C-C707":{"bars":10,"weight":4.0},"C-CACHE GS":{"bars":20,"weight":1.0},"C-CO107":{"bars":10,"weight":4.0},"C-CO108":{"bars":10,"weight":4.0},"C-COR20*20":{"bars":10,"weight":4.0},"C-COR25X25":{"bars":10,"weight":4.0},"C-COR30X30":{"bars":10,"weight":4.0},"C-COR40X40":{"bars":10,"weight":4.0},"C-EX451111":{"bars":10,"weight":4.0},"C-EX451113":{"bars":10,"weight":4.0},"C-EX451120":{"bars":14,"weight":4.0},"C-EX451121":{"bars":14,"weight":4.0},"C-EX451124":{"bars":10,"weight":4.0},"C-EX451313":{"bars":10,"weight":4.0},"C-EX451315":{"bars":10,"weight":4.0},"C-EX451320":{"bars":10,"weight":4.0},"C-EX451411":{"bars":10,"weight":4.0},"C-EX602112":{"bars":10,"weight":4.0},"C-EX602210":{"bars":15,"weight":4.0},"C-EX602211":{"bars":10,"weight":4.0},"C-EX602212":{"bars":13,"weight":3.0},"C-FERMETURE":{"bars":10,"weight":4.0},"C-FR104":{"bars":10,"weight":4.0},"C-FR108":{"bars":20,"weight":4.0},"C-FR150":{"bars":13,"weight":6.4025},"C-FR402":{"bars":13,"weight":5.7395},"C-FSQ124":{"bars":13,"weight":4.439500000000001},"C-FSQ403":{"bars":13,"weight":7.111000000000001},"C-GC GRAND MODEL":{"bars":20,"weight":4.0},"C-GC PETIT MODEL":{"bars":20,"weight":1.69},"C-GL058":{"bars":20,"weight":4.0},"C-LM-757":{"bars":17,"weight":4.0},"C-LM-79":{"bars":14,"weight":3.0},"C-PA7714":{"bars":14,"weight":3.0},"C-PAUM":{"bars":800,"weight":0.5},"C-REMPLISSAGE":{"bars":10,"weight":4.0},"C-SAF704":{"bars":13,"weight":3.0},"C-SAF706":{"bars":5,"weight":9.0},"C-TPR101":{"bars":12,"weight":2.0},"C-TPR754":{"bars":10,"weight":4.0},"C-TUBE 16":{"bars":13,"weight":2.0},"C-TUBE C 40X50":{"bars":13,"weight":2.0},"CL3002":{"bars":20,"weight":0.0},"CL3003":{"bars":20,"weight":0.0},"CL6419":{"bars":10,"weight":4.0},"CO101":{"bars":12,"weight":6.077500000000001},"CO102":{"bars":12,"weight":6.7665},"CO103":{"bars":12,"weight":7.293000000000001},"CO104":{"bars":13,"weight":4.7059999999999995},"CO105":{"bars":13,"weight":3.5685000000000002},"CO106":{"bars":15,"weight":4.303},"CO107":{"bars":13,"weight":5.2844999999999995},"CO108":{"bars":13,"weight":6.2139999999999995},"CO110":{"bars":6,"weight":10.6795},"CO114":{"bars":20,"weight":0.9035000000000001},"CO124":{"bars":20,"weight":1.742},"CSQ101":{"bars":12,"weight":5.863},"CSQ102":{"bars":12,"weight":6.506499999999999},"CSQ103":{"bars":12,"weight":6.0255},"CSQ104":{"bars":13,"weight":4.7775},"CSQ105":{"bars":13,"weight":3.6659999999999995},"CSQ106":{"bars":15,"weight":4.147},"CSQ107":{"bars":13,"weight":7.228000000000001},"CSQ108":{"bars":13,"weight":6.6365},"CSQ110":{"bars":6,"weight":9.4315},"CSQ114":{"bars":20,"weight":0.9100000000000001},"CSQ124":{"bars":20,"weight":1.3259999999999998},"CSQ125":{"bars":20,"weight":0.52},"CSQ126":{"bars":20,"weight":1.5859999999999999},"CSQ153":{"bars":12,"weight":1.848},"CSQ201":{"bars":12,"weight":4.9595},"CSQ202":{"bars":12,"weight":5.6095},"CSQ203":{"bars":12,"weight":5.1285},"CSQ210":{"bars":12,"weight":8.411},"CSQ300":{"bars":20,"weight":1.1895},"CSQ301":{"bars":20,"weight":1.3130000000000002},"CSQ302":{"bars":20,"weight":1.157},"CSQ303":{"bars":20,"weight":1.079},"CSQ304":{"bars":20,"weight":1.157},"CSQ305":{"bars":20,"weight":1.56},"CSQ306":{"bars":20,"weight":1.404},"CSQ537":{"bars":10,"weight":6.195},"CSQ538":{"bars":12,"weight":5.278},"DS001":{"bars":20,"weight":1.281},"DS004":{"bars":20,"weight":0.672},"DS004-120":{"bars":20,"weight":0.672},"E4600-522":{"bars":20,"weight":0.185},"EC22104":{"bars":10,"weight":4.667},"EC22106":{"bars":13,"weight":4.407},"EC22112":{"bars":20,"weight":1.599},"EC40100":{"bars":13,"weight":4.134},"EC40102":{"bars":13,"weight":3.2825},"EC40104":{"bars":13,"weight":5.7005},"EC40107":{"bars":20,"weight":0.6759999999999999},"EC40112":{"bars":14,"weight":4.16},"EC40121":{"bars":9,"weight":8.7555},"EC40148":{"bars":13,"weight":4.199},"EC40154":{"bars":9,"weight":8.71},"EC40165":{"bars":13,"weight":4.0625},"EC40166":{"bars":20,"weight":1.482},"EC40402":{"bars":13,"weight":4.2250000000000005},"EC40404":{"bars":13,"weight":4.888},"EC40404/1":{"bars":13,"weight":4.888},"EC404042":{"bars":14,"weight":3.965},"EC6007":{"bars":14,"weight":1.0},"EC67101":{"bars":12,"weight":4.9855},"EC67102":{"bars":12,"weight":5.6355},"EC67103":{"bars":12,"weight":6.045},"EC67104":{"bars":13,"weight":3.8739999999999997},"EC67105":{"bars":13,"weight":2.9445},"EC67106":{"bars":15,"weight":3.5425000000000004},"EC80104":{"bars":13,"weight":4.0105},"EC80105":{"bars":13,"weight":3.2825},"EC80106":{"bars":15,"weight":3.6075000000000004},"EC80114":{"bars":20,"weight":1.04},"EC80120":{"bars":12,"weight":4.888},"EC80121":{"bars":12,"weight":5.538},"ECGL248":{"bars":20,"weight":2.0},"F50-200":{"bars":10,"weight":5.02},"FR100":{"bars":13,"weight":5.420999999999999},"FR102":{"bars":13,"weight":4.3875},"FR104":{"bars":13,"weight":7.1695},"FR107":{"bars":20,"weight":0.9684999999999999},"FR108":{"bars":20,"weight":0.8905000000000001},"FR110":{"bars":20,"weight":1.703},"FR111":{"bars":20,"weight":1.703},"FR112":{"bars":14,"weight":4.5175},"FR121":{"bars":9,"weight":11.1735},"FR123":{"bars":20,"weight":1.0},"FR127":{"bars":20,"weight":1.9955},"FR130":{"bars":14,"weight":4.3875},"FR133":{"bars":20,"weight":1.43},"FR135":{"bars":20,"weight":1.5795},"FR139":{"bars":20,"weight":1.794},"FR147":{"bars":13,"weight":4.953},"FR148":{"bars":13,"weight":5.3755},"FR149":{"bars":13,"weight":5.447},"FR150":{"bars":13,"weight":6.4025},"FR151":{"bars":13,"weight":6.3635},"FR153":{"bars":20,"weight":2.0475},"FR154":{"bars":9,"weight":11.206},"FR155":{"bars":9,"weight":12.915500000000002},"FR156":{"bars":9,"weight":10.224499999999999},"FR161":{"bars":14,"weight":2.184},"FR162":{"bars":14,"weight":3.5999999999999996},"FR164":{"bars":14,"weight":3.5945000000000005},"FR166":{"bars":20,"weight":1.6705},"FR168":{"bars":20,"weight":1.5859999999999999},"FR401":{"bars":13,"weight":5.876},"FR402":{"bars":13,"weight":5.7395},"FR403":{"bars":13,"weight":7.2475},"FR404":{"bars":13,"weight":5.9670000000000005},"FR406":{"bars":13,"weight":5.746},"FR410":{"bars":20,"weight":1.4755},"FR411":{"bars":20,"weight":1.2025},"FSQ100":{"bars":13,"weight":4.992},"FSQ102":{"bars":13,"weight":4.303},"FSQ104":{"bars":13,"weight":7.039499999999999},"FSQ107":{"bars":20,"weight":0.969},"FSQ108":{"bars":20,"weight":1.846},"FSQ110":{"bars":20,"weight":1.8784999999999998},"FSQ111":{"bars":20,"weight":1.8459999999999999},"FSQ112":{"bars":12,"weight":5.187},"FSQ121":{"bars":9,"weight":9.0545},"FSQ122":{"bars":20,"weight":1.9565},"FSQ124":{"bars":13,"weight":4.439500000000001},"FSQ130":{"bars":14,"weight":4.3875},"FSQ131":{"bars":9,"weight":5.2844999999999995},"FSQ132":{"bars":17,"weight":2.262},"FSQ139":{"bars":20,"weight":2.002},"FSQ148":{"bars":14,"weight":5.012},"FSQ149":{"bars":13,"weight":5.265000000000001},"FSQ150":{"bars":13,"weight":5.9735000000000005},"FSQ153":{"bars":20,"weight":1.9955},"FSQ156":{"bars":9,"weight":10.179},"FSQ163":{"bars":13,"weight":4.1925},"FSQ164":{"bars":20,"weight":1.9369999999999998},"FSQ165":{"bars":13,"weight":5.213},"FSQ401":{"bars":13,"weight":5.6355},"FSQ402":{"bars":13,"weight":5.5445},"FSQ403":{"bars":13,"weight":7.111000000000001},"FSQ404":{"bars":13,"weight":5.707},"FSQ405":{"bars":13,"weight":7.111000000000001},"FSQ406":{"bars":13,"weight":5.46},"FSQ408":{"bars":13,"weight":5.395},"FSQ460":{"bars":20,"weight":2.418},"FSQ501":{"bars":14,"weight":3.0},"FSQ502":{"bars":14,"weight":5.844},"FSQ510":{"bars":14,"weight":6.487},"FSQ520":{"bars":14,"weight":3.0},"FSQ521":{"bars":12,"weight":5.629},"FSQ530":{"bars":14,"weight":3.0},"FSQ532":{"bars":14,"weight":3.0},"FSQ534":{"bars":20,"weight":2.022},"FSQ535":{"bars":14,"weight":4.5695},"FSQ536":{"bars":12,"weight":5.5315},"FSQ540":{"bars":12,"weight":7.456},"FSQ541":{"bars":20,"weight":2.022},"GL-MB":{"bars":12,"weight":3.6919999999999997},"GL058":{"bars":20,"weight":1.898},"GL11060/248":{"bars":20,"weight":2.4635},"GL180":{"bars":20,"weight":1.69},"GLLM100.6.5":{"bars":10,"weight":10.05},"INNER PROFILE":{"bars":10,"weight":4.0},"KL7401":{"bars":14,"weight":1.179},"KL7402":{"bars":9,"weight":11.453},"KL7403":{"bars":9,"weight":10.16},"KL7404":{"bars":9,"weight":11.265},"KL7405":{"bars":9,"weight":13.689},"KL7406":{"bars":10,"weight":8.736},"KL7407":{"bars":9,"weight":17.934},"KL7412":{"bars":13,"weight":5.233},"KL7417":{"bars":9,"weight":9.822},"KL7420":{"bars":9,"weight":1.885},"KLHS7421":{"bars":20,"weight":0.1},"LM-100":{"bars":10,"weight":4.0},"LM-100.5":{"bars":12,"weight":4.43},"LM-100.6":{"bars":12,"weight":4.43},"LM-100.A":{"bars":10,"weight":4.0},"LM-100.MA":{"bars":10,"weight":4.0},"LM-40":{"bars":20,"weight":1.653},"LM-48":{"bars":17,"weight":1.825},"LM-55":{"bars":17,"weight":3.569},"LM-55A":{"bars":17,"weight":2.0},"LM-60":{"bars":17,"weight":3.07},"LM-60A":{"bars":17,"weight":2.0},"LM-62":{"bars":17,"weight":2.0},"LM-757":{"bars":17,"weight":2.756},"LM-79":{"bars":14,"weight":4.548},"LM-79A":{"bars":14,"weight":4.548},"LM-F":{"bars":17,"weight":2.8405},"LM-F/6":{"bars":17,"weight":2.8405},"LM-F100":{"bars":12,"weight":6.32},"LM-F41":{"bars":17,"weight":1.859},"LM-F45":{"bars":20,"weight":2.19},"LM-F50":{"bars":17,"weight":2.8},"LM-F79/100":{"bars":10,"weight":6.318},"LM-FR":{"bars":17,"weight":2.0},"LM-FVS":{"bars":17,"weight":2.7039999999999997},"LM-S40":{"bars":20,"weight":0.842},"LM-S48":{"bars":20,"weight":1.034},"LM-S758":{"bars":20,"weight":1.3259999999999998},"LM-SR":{"bars":17,"weight":2.0},"LMDP-100":{"bars":14,"weight":7.02},"LMDPF-100":{"bars":20,"weight":1.846},"LMMO-757":{"bars":17,"weight":2.756},"LMMO-S758":{"bars":20,"weight":1.3259999999999998},"MR002":{"bars":9,"weight":11.934000000000001},"MR003":{"bars":5,"weight":6.0},"MR006":{"bars":10,"weight":11.8},"MR009":{"bars":12,"weight":14.7},"MR009_P6500":{"bars":8,"weight":6.0},"MR011":{"bars":12,"weight":8.6},"MR013":{"bars":10,"weight":8.561},"MR014":{"bars":8,"weight":10.667},"MR014_P6500":{"bars":8,"weight":10.6015},"MR021":{"bars":14,"weight":4.925},"MR030":{"bars":12,"weight":5.148000000000001},"MR031":{"bars":10,"weight":5.005},"MR032":{"bars":10,"weight":6.981},"MR033":{"bars":10,"weight":6.195},"MR036":{"bars":10,"weight":6.062},"MR038":{"bars":10,"weight":8.593},"MR045":{"bars":20,"weight":1.424},"OK4716":{"bars":20,"weight":1.846},"OK4717":{"bars":20,"weight":1.859},"OK4718":{"bars":20,"weight":2.789},"OK4719":{"bars":20,"weight":0.91},"OK4720":{"bars":14,"weight":3.562},"OK4721":{"bars":20,"weight":1.775},"OK4722":{"bars":20,"weight":1.872},"OK4725":{"bars":20,"weight":1.63},"P006":{"bars":14,"weight":1.0},"P007":{"bars":20,"weight":1.58},"P010":{"bars":12,"weight":5.7265},"P013":{"bars":20,"weight":1.104},"P031":{"bars":14,"weight":3.1654999999999998},"P039":{"bars":14,"weight":3.2045},"P057":{"bars":14,"weight":3.0},"P057 PIECE":{"bars":800,"weight":0.021},"P057/25":{"bars":14,"weight":5.9345},"P063":{"bars":14,"weight":3.0},"P063-PIECE":{"bars":10,"weight":4.0},"P063/25":{"bars":14,"weight":9.0285},"P7015410":{"bars":14,"weight":1.01},"PA540-777":{"bars":10,"weight":8.328},"PAILLASSES":{"bars":15,"weight":4.0},"PCN40-45":{"bars":800,"weight":0.5},"PI451":{"bars":20,"weight":1.158},"PI452":{"bars":14,"weight":4.788},"PI453":{"bars":12,"weight":5.142},"PI455":{"bars":14,"weight":4.878},"PL101":{"bars":20,"weight":0.906},"PL102":{"bars":20,"weight":1.04},"PL103":{"bars":20,"weight":0.666},"PL105":{"bars":20,"weight":2.298},"PL106":{"bars":20,"weight":2.772},"PL108":{"bars":10,"weight":10.824},"PL109":{"bars":20,"weight":2.286},"PL110":{"bars":14,"weight":6.3},"PL111":{"bars":14,"weight":5.802},"PL112":{"bars":13,"weight":7.695},"PL113":{"bars":10,"weight":5.63},"PL114":{"bars":20,"weight":1.5},"PR6301":{"bars":9,"weight":11.109},"PR6302":{"bars":10,"weight":11.967},"PR6303":{"bars":6,"weight":15.0},"PR6304":{"bars":10,"weight":8.658},"PR6305":{"bars":12,"weight":7.982},"PR6307":{"bars":13,"weight":5.213},"PR6308":{"bars":6,"weight":15.009},"PR6310":{"bars":20,"weight":2.262},"PR6311":{"bars":20,"weight":1.736},"PRN40-45":{"bars":800,"weight":0.5},"PS6007":{"bars":14,"weight":1.0},"SP6001":{"bars":14,"weight":8.151},"SP6002":{"bars":20,"weight":1.666},"SP6003":{"bars":20,"weight":1.484},"SP6004":{"bars":20,"weight":1.654},"SP6005":{"bars":20,"weight":2.01},"SP6006":{"bars":20,"weight":1.388},"T054-12":{"bars":800,"weight":0.015},"TC40X40X1":{"bars":14,"weight":4.056},"TRINGLE G6005791(T1120020)":{"bars":20,"weight":4.0},"TRINGLE G6005792(T1120025":{"bars":20,"weight":4.0}},"stock":{"AL11026":{"stock_brut":0.0,"stock_couleurs":30.0,"pct_laque":1.0,"moyenne_vente":222.375},"AL15401":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"AL15401P3":{"stock_brut":145.0,"stock_couleurs":-145.0,"pct_laque":null,"moyenne_vente":0.0},"AL15403":{"stock_brut":0.0,"stock_couleurs":243.0,"pct_laque":1.0,"moyenne_vente":0.0},"AL15403P1":{"stock_brut":117.0,"stock_couleurs":-117.0,"pct_laque":null,"moyenne_vente":0.0},"AL15403P2":{"stock_brut":126.0,"stock_couleurs":-126.0,"pct_laque":null,"moyenne_vente":0.0},"AL15404":{"stock_brut":0.0,"stock_couleurs":200.0,"pct_laque":1.0,"moyenne_vente":0.0},"AL15404P1":{"stock_brut":103.0,"stock_couleurs":-103.0,"pct_laque":null,"moyenne_vente":0.0},"AL15404P2":{"stock_brut":97.0,"stock_couleurs":-97.0,"pct_laque":null,"moyenne_vente":0.0},"AL15405P1":{"stock_brut":96.0,"stock_couleurs":-96.0,"pct_laque":null,"moyenne_vente":0.0},"AL15405P2":{"stock_brut":124.0,"stock_couleurs":-124.0,"pct_laque":null,"moyenne_vente":0.0},"AL15410":{"stock_brut":79.0,"stock_couleurs":2.0,"pct_laque":0.024691358024691357,"moyenne_vente":0.0},"AL15413":{"stock_brut":116.0,"stock_couleurs":2.0,"pct_laque":0.01694915254237288,"moyenne_vente":0.0},"AL15415":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"AL15416":{"stock_brut":389.0,"stock_couleurs":10.0,"pct_laque":0.02506265664160401,"moyenne_vente":0.0},"AL15420":{"stock_brut":3.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"AL15421":{"stock_brut":29.0,"stock_couleurs":-29.0,"pct_laque":null,"moyenne_vente":0.0},"AL15422":{"stock_brut":91.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"AL15423":{"stock_brut":149.0,"stock_couleurs":-149.0,"pct_laque":null,"moyenne_vente":0.0},"AL40122":{"stock_brut":149.0,"stock_couleurs":7.0,"pct_laque":0.04487179487179487,"moyenne_vente":27.5},"AL40123":{"stock_brut":0.0,"stock_couleurs":3.0,"pct_laque":1.0,"moyenne_vente":0.0},"AL40133":{"stock_brut":319.0,"stock_couleurs":43.0,"pct_laque":0.11878453038674033,"moyenne_vente":27.0},"AL40165":{"stock_brut":639.0,"stock_couleurs":158.0,"pct_laque":0.19824341279799249,"moyenne_vente":73.44444444444444},"AL40200":{"stock_brut":360.0,"stock_couleurs":47.0,"pct_laque":0.11547911547911548,"moyenne_vente":92.125},"AL40202":{"stock_brut":118.0,"stock_couleurs":34.0,"pct_laque":0.2236842105263158,"moyenne_vente":1.0},"AL80104":{"stock_brut":148.0,"stock_couleurs":38.0,"pct_laque":0.20430107526881722,"moyenne_vente":17.5},"AL80105":{"stock_brut":164.0,"stock_couleurs":16.0,"pct_laque":0.08888888888888889,"moyenne_vente":19.571428571428573},"AL80106":{"stock_brut":307.0,"stock_couleurs":45.0,"pct_laque":0.1278409090909091,"moyenne_vente":28.333333333333332},"AL80114":{"stock_brut":0.0,"stock_couleurs":155.0,"pct_laque":1.0,"moyenne_vente":0.0},"AL80114/1":{"stock_brut":1227.0,"stock_couleurs":62.0,"pct_laque":0.04809930178432894,"moyenne_vente":54.0},"AL80120":{"stock_brut":161.0,"stock_couleurs":381.0,"pct_laque":0.7029520295202952,"moyenne_vente":7.0},"AL80121":{"stock_brut":606.0,"stock_couleurs":43.0,"pct_laque":0.0662557781201849,"moyenne_vente":41.666666666666664},"ALL10093/1":{"stock_brut":11.0,"stock_couleurs":17.0,"pct_laque":0.6071428571428571,"moyenne_vente":20.0},"ALL11142/1.1":{"stock_brut":115.0,"stock_couleurs":172.0,"pct_laque":0.5993031358885017,"moyenne_vente":21.4},"ALL11142/1.5":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"ALL11144/1.2":{"stock_brut":0.0,"stock_couleurs":41.0,"pct_laque":1.0,"moyenne_vente":141.5},"ALL11144/2":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"ALL124/1.5":{"stock_brut":173.0,"stock_couleurs":105.0,"pct_laque":0.3776978417266187,"moyenne_vente":0.0},"ALL144/3":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"ALL145/1.3":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"ALL145/2.8":{"stock_brut":257.0,"stock_couleurs":2.0,"pct_laque":0.007722007722007722,"moyenne_vente":15.6},"ALL145/4":{"stock_brut":23.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":300.5},"ALL213/1.5":{"stock_brut":0.0,"stock_couleurs":4.0,"pct_laque":1.0,"moyenne_vente":47.2},"ALL244/1.2":{"stock_brut":0.0,"stock_couleurs":8.0,"pct_laque":1.0,"moyenne_vente":2144.0},"ALL244/1.5":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"ALL316/6":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":2.0},"ALL349/1.5":{"stock_brut":3.0,"stock_couleurs":24.0,"pct_laque":0.8888888888888888,"moyenne_vente":1781.0},"ALL714/4":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"ALL723/2":{"stock_brut":0.0,"stock_couleurs":25.0,"pct_laque":1.0,"moyenne_vente":0.0},"ALL724/2":{"stock_brut":60.0,"stock_couleurs":65.0,"pct_laque":0.52,"moyenne_vente":150.0},"ALL727/1.2":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"ALL727/2":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"BL770":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"BL771":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"BL772":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"BL773":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"BL774":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"BL775":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"BL776":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"CL3002":{"stock_brut":230.0,"stock_couleurs":-230.0,"pct_laque":null,"moyenne_vente":0.0},"CL3003":{"stock_brut":202.0,"stock_couleurs":-202.0,"pct_laque":null,"moyenne_vente":0.0},"CL3004":{"stock_brut":125.0,"stock_couleurs":-125.0,"pct_laque":null,"moyenne_vente":0.0},"CO101":{"stock_brut":125.0,"stock_couleurs":129.0,"pct_laque":0.5078740157480315,"moyenne_vente":108.55555555555556},"CO102":{"stock_brut":0.0,"stock_couleurs":17.0,"pct_laque":1.0,"moyenne_vente":180.66666666666666},"CO103":{"stock_brut":949.0,"stock_couleurs":105.0,"pct_laque":0.09962049335863378,"moyenne_vente":593.6363636363636},"CO104":{"stock_brut":406.0,"stock_couleurs":2.0,"pct_laque":0.004901960784313725,"moyenne_vente":344.5},"CO105":{"stock_brut":7.0,"stock_couleurs":130.0,"pct_laque":0.948905109489051,"moyenne_vente":375.1},"CO106":{"stock_brut":125.0,"stock_couleurs":13.0,"pct_laque":0.09420289855072464,"moyenne_vente":453.90909090909093},"CO107":{"stock_brut":0.0,"stock_couleurs":43.0,"pct_laque":1.0,"moyenne_vente":121.75},"CO108":{"stock_brut":0.0,"stock_couleurs":8.0,"pct_laque":1.0,"moyenne_vente":98.83333333333333},"CO110":{"stock_brut":1.0,"stock_couleurs":7.0,"pct_laque":0.875,"moyenne_vente":134.0},"CO114":{"stock_brut":2067.0,"stock_couleurs":336.0,"pct_laque":0.13982521847690388,"moyenne_vente":2162.818181818182},"CO116":{"stock_brut":770.0,"stock_couleurs":11880.00072,"pct_laque":0.9391304382471213,"moyenne_vente":851.2857142857143},"CO124":{"stock_brut":207.0,"stock_couleurs":24.0,"pct_laque":0.1038961038961039,"moyenne_vente":37.666666666666664},"CSQ101":{"stock_brut":0.0,"stock_couleurs":432.0,"pct_laque":1.0,"moyenne_vente":44.77777777777778},"CSQ102":{"stock_brut":24.0,"stock_couleurs":1.0,"pct_laque":0.04,"moyenne_vente":266.6666666666667},"CSQ103":{"stock_brut":1502.0,"stock_couleurs":4.0,"pct_laque":0.0026560424966799467,"moyenne_vente":153.0},"CSQ104":{"stock_brut":218.0,"stock_couleurs":3.0,"pct_laque":0.013574660633484163,"moyenne_vente":230.63636363636363},"CSQ105":{"stock_brut":397.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":233.0},"CSQ106":{"stock_brut":17.0,"stock_couleurs":13.0,"pct_laque":0.43333333333333335,"moyenne_vente":254.1},"CSQ107":{"stock_brut":0.0,"stock_couleurs":1.0,"pct_laque":1.0,"moyenne_vente":88.11111111111111},"CSQ108":{"stock_brut":173.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":46.111111111111114},"CSQ110":{"stock_brut":109.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":74.16666666666667},"CSQ114":{"stock_brut":639.0,"stock_couleurs":28.0,"pct_laque":0.041979010494752625,"moyenne_vente":351.3333333333333},"CSQ116":{"stock_brut":24.0,"stock_couleurs":3760.0,"pct_laque":0.9936575052854123,"moyenne_vente":175.55555555555554},"CSQ124":{"stock_brut":308.0,"stock_couleurs":6.0,"pct_laque":0.01910828025477707,"moyenne_vente":6.833333333333333},"CSQ125":{"stock_brut":538.0,"stock_couleurs":7.0,"pct_laque":0.012844036697247707,"moyenne_vente":5.8},"CSQ126":{"stock_brut":51.0,"stock_couleurs":4.0,"pct_laque":0.07272727272727272,"moyenne_vente":29.666666666666668},"CSQ153":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"CSQ201":{"stock_brut":315.0,"stock_couleurs":10.0,"pct_laque":0.03076923076923077,"moyenne_vente":4.0},"CSQ202":{"stock_brut":1021.0,"stock_couleurs":3.0,"pct_laque":0.0029296875,"moyenne_vente":75.33333333333333},"CSQ203":{"stock_brut":899.0,"stock_couleurs":2.0,"pct_laque":0.0022197558268590455,"moyenne_vente":53.2},"CSQ210":{"stock_brut":442.0,"stock_couleurs":1.0,"pct_laque":0.002257336343115124,"moyenne_vente":19.0},"CSQ300":{"stock_brut":603.0,"stock_couleurs":10.0,"pct_laque":0.01631321370309951,"moyenne_vente":4.0},"CSQ301":{"stock_brut":2555.0,"stock_couleurs":21.0,"pct_laque":0.008152173913043478,"moyenne_vente":82.77777777777777},"CSQ302":{"stock_brut":358.0,"stock_couleurs":1.0,"pct_laque":0.002785515320334262,"moyenne_vente":69.0},"CSQ303":{"stock_brut":254.0,"stock_couleurs":14.0,"pct_laque":0.05223880597014925,"moyenne_vente":3.5},"CSQ304":{"stock_brut":487.0,"stock_couleurs":36.0,"pct_laque":0.06883365200764818,"moyenne_vente":5.666666666666667},"CSQ305":{"stock_brut":707.0,"stock_couleurs":7.0,"pct_laque":0.00980392156862745,"moyenne_vente":174.77777777777777},"CSQ306":{"stock_brut":0.0,"stock_couleurs":1.0,"pct_laque":1.0,"moyenne_vente":65.8},"CSQ450":{"stock_brut":139.0,"stock_couleurs":-139.0,"pct_laque":null,"moyenne_vente":0.0},"CSQ451":{"stock_brut":90.0,"stock_couleurs":-90.0,"pct_laque":null,"moyenne_vente":0.0},"CSQ537":{"stock_brut":800.0,"stock_couleurs":16.0,"pct_laque":0.0196078431372549,"moyenne_vente":0.0},"CSQ538":{"stock_brut":256.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"CSQ539":{"stock_brut":210.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"DIFF-SL300-1":{"stock_brut":null,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"DIFF31-1-1":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"DIFF31-1-2":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"DIFF31-2":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"DS001":{"stock_brut":0.0,"stock_couleurs":53.0,"pct_laque":1.0,"moyenne_vente":0.0},"DS002":{"stock_brut":901.0,"stock_couleurs":-901.0,"pct_laque":null,"moyenne_vente":0.0},"DS003":{"stock_brut":1069.0,"stock_couleurs":-1069.0,"pct_laque":null,"moyenne_vente":0.0},"DS004":{"stock_brut":68.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"EC22104":{"stock_brut":276.0,"stock_couleurs":88.0,"pct_laque":0.24175824175824176,"moyenne_vente":109.6},"EC22106":{"stock_brut":712.0,"stock_couleurs":425.0,"pct_laque":0.3737906772207564,"moyenne_vente":180.0},"EC22112":{"stock_brut":443.0,"stock_couleurs":136.0,"pct_laque":0.23488773747841105,"moyenne_vente":56.111111111111114},"EC40100":{"stock_brut":442.0,"stock_couleurs":49.0,"pct_laque":0.09979633401221996,"moyenne_vente":2247.5454545454545},"EC40102":{"stock_brut":6.0,"stock_couleurs":5.0,"pct_laque":0.45454545454545453,"moyenne_vente":824.2222222222222},"EC40104":{"stock_brut":57.0,"stock_couleurs":2.0,"pct_laque":0.03389830508474576,"moyenne_vente":150.71428571428572},"EC40107":{"stock_brut":5246.0,"stock_couleurs":510.0,"pct_laque":0.08860319666435024,"moyenne_vente":423.0},"EC40112":{"stock_brut":19.0,"stock_couleurs":9.0,"pct_laque":0.32142857142857145,"moyenne_vente":268.4},"EC40121":{"stock_brut":91.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":272.25},"EC40148":{"stock_brut":180.0,"stock_couleurs":22.0,"pct_laque":0.10891089108910891,"moyenne_vente":317.5},"EC40154":{"stock_brut":287.0,"stock_couleurs":4.0,"pct_laque":0.013745704467353952,"moyenne_vente":102.83333333333333},"EC40165":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"EC40166":{"stock_brut":245.0,"stock_couleurs":7.0,"pct_laque":0.027777777777777776,"moyenne_vente":2305.5454545454545},"EC40402":{"stock_brut":2283.0,"stock_couleurs":211.0,"pct_laque":0.08460304731355253,"moyenne_vente":1635.090909090909},"EC40404":{"stock_brut":2357.0,"stock_couleurs":127.0,"pct_laque":0.05112721417069243,"moyenne_vente":1288.1},"EC404041":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"EC404042":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"EC67101":{"stock_brut":919.0,"stock_couleurs":196.0,"pct_laque":0.1757847533632287,"moyenne_vente":200.44444444444446},"EC67102":{"stock_brut":9.0,"stock_couleurs":1221.0,"pct_laque":0.9926829268292683,"moyenne_vente":159.33333333333334},"EC67103":{"stock_brut":683.0,"stock_couleurs":62.0,"pct_laque":0.08322147651006712,"moyenne_vente":720.75},"EC67104":{"stock_brut":1227.0,"stock_couleurs":29.0,"pct_laque":0.02308917197452229,"moyenne_vente":409.875},"EC67105":{"stock_brut":78.0,"stock_couleurs":169.0,"pct_laque":0.6842105263157895,"moyenne_vente":520.6363636363636},"EC67106":{"stock_brut":0.0,"stock_couleurs":52.0,"pct_laque":1.0,"moyenne_vente":517.5},"EC80104":{"stock_brut":0.0,"stock_couleurs":358.0,"pct_laque":1.0,"moyenne_vente":5.0},"EC80105":{"stock_brut":169.0,"stock_couleurs":152.0,"pct_laque":0.4735202492211838,"moyenne_vente":0.0},"EC80106":{"stock_brut":114.0,"stock_couleurs":168.0,"pct_laque":0.5957446808510638,"moyenne_vente":0.0},"EC80114":{"stock_brut":238.0,"stock_couleurs":165.0,"pct_laque":0.4094292803970223,"moyenne_vente":0.0},"EC80120":{"stock_brut":0.0,"stock_couleurs":100.0,"pct_laque":1.0,"moyenne_vente":3.0},"EC80121":{"stock_brut":0.0,"stock_couleurs":2.0,"pct_laque":1.0,"moyenne_vente":4.0},"ECGL248":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"FR100":{"stock_brut":263.0,"stock_couleurs":88.0,"pct_laque":0.25071225071225073,"moyenne_vente":812.5},"FR102":{"stock_brut":0.0,"stock_couleurs":135.0,"pct_laque":1.0,"moyenne_vente":201.6},"FR104":{"stock_brut":3.0,"stock_couleurs":13.0,"pct_laque":0.8125,"moyenne_vente":105.625},"FR107":{"stock_brut":67.0,"stock_couleurs":272.0,"pct_laque":0.8023598820058997,"moyenne_vente":15.4},"FR108":{"stock_brut":1058.0,"stock_couleurs":29.0,"pct_laque":0.02667893284268629,"moyenne_vente":67.0},"FR110":{"stock_brut":142.0,"stock_couleurs":21.0,"pct_laque":0.12883435582822086,"moyenne_vente":52.125},"FR111":{"stock_brut":5.0,"stock_couleurs":64.0,"pct_laque":0.927536231884058,"moyenne_vente":341.625},"FR112":{"stock_brut":7.0,"stock_couleurs":17.0,"pct_laque":0.7083333333333334,"moyenne_vente":110.9},"FR121":{"stock_brut":0.0,"stock_couleurs":56.0,"pct_laque":1.0,"moyenne_vente":108.8},"FR127":{"stock_brut":17.0,"stock_couleurs":30.0,"pct_laque":0.6382978723404256,"moyenne_vente":25.571428571428573},"FR133":{"stock_brut":219.0,"stock_couleurs":83.0,"pct_laque":0.27483443708609273,"moyenne_vente":40.25},"FR135":{"stock_brut":12.0,"stock_couleurs":39.0,"pct_laque":0.7647058823529411,"moyenne_vente":99.0},"FR139":{"stock_brut":1978.0,"stock_couleurs":186.0,"pct_laque":0.08595194085027727,"moyenne_vente":1480.0},"FR147":{"stock_brut":4.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"FR148":{"stock_brut":2.0,"stock_couleurs":1.0,"pct_laque":0.3333333333333333,"moyenne_vente":92.28571428571429},"FR149":{"stock_brut":0.0,"stock_couleurs":24.0,"pct_laque":1.0,"moyenne_vente":25.0},"FR150":{"stock_brut":0.0,"stock_couleurs":14.0,"pct_laque":1.0,"moyenne_vente":659.2222222222222},"FR151":{"stock_brut":11.0,"stock_couleurs":29.0,"pct_laque":0.725,"moyenne_vente":32.0},"FR154":{"stock_brut":0.0,"stock_couleurs":3.0,"pct_laque":1.0,"moyenne_vente":72.5},"FR155":{"stock_brut":51.0,"stock_couleurs":65.0,"pct_laque":0.5603448275862069,"moyenne_vente":8.666666666666666},"FR156":{"stock_brut":4.0,"stock_couleurs":50.0,"pct_laque":0.9259259259259259,"moyenne_vente":20.666666666666668},"FR161":{"stock_brut":37.0,"stock_couleurs":48.0,"pct_laque":0.5647058823529412,"moyenne_vente":24.285714285714285},"FR164":{"stock_brut":42.0,"stock_couleurs":61.0,"pct_laque":0.5922330097087378,"moyenne_vente":26.625},"FR166":{"stock_brut":169.0,"stock_couleurs":77.0,"pct_laque":0.3130081300813008,"moyenne_vente":435.6666666666667},"FR168":{"stock_brut":355.0,"stock_couleurs":229.0,"pct_laque":0.3921232876712329,"moyenne_vente":0.0},"FR401":{"stock_brut":553.0,"stock_couleurs":148.0,"pct_laque":0.21112696148359486,"moyenne_vente":38.75},"FR402":{"stock_brut":135.0,"stock_couleurs":21.0,"pct_laque":0.1346153846153846,"moyenne_vente":898.4545454545455},"FR403":{"stock_brut":314.0,"stock_couleurs":102.0,"pct_laque":0.24519230769230768,"moyenne_vente":34.875},"FR404":{"stock_brut":1934.0,"stock_couleurs":171.0,"pct_laque":0.08123515439429929,"moyenne_vente":578.9090909090909},"FR406":{"stock_brut":0.0,"stock_couleurs":70.0,"pct_laque":1.0,"moyenne_vente":57.857142857142854},"FR410":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"FR411":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"FSQ100":{"stock_brut":590.0,"stock_couleurs":20.0,"pct_laque":0.03278688524590164,"moyenne_vente":85.33333333333333},"FSQ102":{"stock_brut":659.0,"stock_couleurs":84.0,"pct_laque":0.11305518169582772,"moyenne_vente":13.857142857142858},"FSQ104":{"stock_brut":435.0,"stock_couleurs":2.0,"pct_laque":0.004576659038901602,"moyenne_vente":13.0},"FSQ108":{"stock_brut":500.0,"stock_couleurs":6.0,"pct_laque":0.011857707509881422,"moyenne_vente":0.0},"FSQ110":{"stock_brut":1063.0,"stock_couleurs":5.0,"pct_laque":0.0046816479400749065,"moyenne_vente":41.0},"FSQ111":{"stock_brut":983.0,"stock_couleurs":3.0,"pct_laque":0.0030425963488843813,"moyenne_vente":103.33333333333333},"FSQ112":{"stock_brut":171.0,"stock_couleurs":3.0,"pct_laque":0.017241379310344827,"moyenne_vente":27.714285714285715},"FSQ121":{"stock_brut":110.0,"stock_couleurs":2.0,"pct_laque":0.017857142857142856,"moyenne_vente":22.88888888888889},"FSQ122":{"stock_brut":30.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":23.375},"FSQ124":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":95.75},"FSQ130":{"stock_brut":56.0,"stock_couleurs":8.0,"pct_laque":0.125,"moyenne_vente":1.6666666666666667},"FSQ131":{"stock_brut":143.0,"stock_couleurs":18.0,"pct_laque":0.11180124223602485,"moyenne_vente":64.6},"FSQ132":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":10.166666666666666},"FSQ139":{"stock_brut":2359.0,"stock_couleurs":9.0,"pct_laque":0.003800675675675676,"moyenne_vente":157.22222222222223},"FSQ148":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"FSQ149":{"stock_brut":40.0,"stock_couleurs":3.0,"pct_laque":0.06976744186046512,"moyenne_vente":7.4},"FSQ150":{"stock_brut":779.0,"stock_couleurs":124.0,"pct_laque":0.13732004429678848,"moyenne_vente":16.666666666666668},"FSQ151":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"FSQ153":{"stock_brut":48.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":4.666666666666667},"FSQ156":{"stock_brut":258.0,"stock_couleurs":1.0,"pct_laque":0.003861003861003861,"moyenne_vente":3.8333333333333335},"FSQ163":{"stock_brut":92.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":4.2},"FSQ164":{"stock_brut":610.0,"stock_couleurs":2.0,"pct_laque":0.0032679738562091504,"moyenne_vente":1.6666666666666667},"FSQ165":{"stock_brut":161.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"FSQ401":{"stock_brut":717.0,"stock_couleurs":2.0,"pct_laque":0.0027816411682892906,"moyenne_vente":104.18181818181819},"FSQ402":{"stock_brut":14.0,"stock_couleurs":5.0,"pct_laque":0.2631578947368421,"moyenne_vente":15.0},"FSQ403":{"stock_brut":281.0,"stock_couleurs":17.0,"pct_laque":0.05704697986577181,"moyenne_vente":95.66666666666667},"FSQ404":{"stock_brut":32.0,"stock_couleurs":1.0,"pct_laque":0.030303030303030304,"moyenne_vente":18.4},"FSQ405":{"stock_brut":292.0,"stock_couleurs":5.0,"pct_laque":0.016835016835016835,"moyenne_vente":10.2},"FSQ406":{"stock_brut":25.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":4.0},"FSQ408":{"stock_brut":257.0,"stock_couleurs":11.0,"pct_laque":0.041044776119402986,"moyenne_vente":8.857142857142858},"FSQ460":{"stock_brut":150.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"FSQ500":{"stock_brut":32.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"FSQ501":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"FSQ502":{"stock_brut":155.0,"stock_couleurs":4.0,"pct_laque":0.025157232704402517,"moyenne_vente":0.0},"FSQ510":{"stock_brut":340.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"FSQ520":{"stock_brut":240.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"FSQ521":{"stock_brut":308.0,"stock_couleurs":4.0,"pct_laque":0.01282051282051282,"moyenne_vente":0.0},"FSQ530":{"stock_brut":244.0,"stock_couleurs":12.0,"pct_laque":0.046875,"moyenne_vente":0.0},"FSQ531":{"stock_brut":134.0,"stock_couleurs":2.0,"pct_laque":0.014705882352941176,"moyenne_vente":0.0},"FSQ532":{"stock_brut":200.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"FSQ534":{"stock_brut":473.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"FSQ535":{"stock_brut":124.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"FSQ536":{"stock_brut":196.0,"stock_couleurs":-176.0,"pct_laque":-8.8,"moyenne_vente":0.0},"FSQ540":{"stock_brut":218.0,"stock_couleurs":1.0,"pct_laque":0.0045662100456621,"moyenne_vente":0.0},"FSQ541":{"stock_brut":38.0,"stock_couleurs":2.0,"pct_laque":0.05,"moyenne_vente":0.0},"GL-MB":{"stock_brut":20.0,"stock_couleurs":1.0,"pct_laque":0.047619047619047616,"moyenne_vente":33.6},"GL058":{"stock_brut":0.0,"stock_couleurs":48.0,"pct_laque":1.0,"moyenne_vente":0.0},"GL11060":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"GL11060/248":{"stock_brut":10587.0,"stock_couleurs":152.0,"pct_laque":0.014154018064996741,"moyenne_vente":2112.7272727272725},"GL180":{"stock_brut":727.0,"stock_couleurs":4.0,"pct_laque":0.005471956224350205,"moyenne_vente":257.7142857142857},"GL248":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"GLLM100.6":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"GLLM79":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"GLMV2100":{"stock_brut":null,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"IN-A 4501":{"stock_brut":71.0,"stock_couleurs":-71.0,"pct_laque":null,"moyenne_vente":0.0},"IN-A 4502":{"stock_brut":63.0,"stock_couleurs":-63.0,"pct_laque":null,"moyenne_vente":0.0},"IN-A 4503":{"stock_brut":52.0,"stock_couleurs":-52.0,"pct_laque":null,"moyenne_vente":0.0},"IN-A 4504":{"stock_brut":1.0,"stock_couleurs":-1.0,"pct_laque":null,"moyenne_vente":0.0},"IN-A 4505":{"stock_brut":72.0,"stock_couleurs":-72.0,"pct_laque":null,"moyenne_vente":0.0},"IN-A 4506":{"stock_brut":86.0,"stock_couleurs":-86.0,"pct_laque":null,"moyenne_vente":0.0},"IN-A 4507":{"stock_brut":79.0,"stock_couleurs":-79.0,"pct_laque":null,"moyenne_vente":0.0},"IN-A 4508":{"stock_brut":91.0,"stock_couleurs":-91.0,"pct_laque":null,"moyenne_vente":0.0},"IN-A 4509":{"stock_brut":535.0,"stock_couleurs":-535.0,"pct_laque":null,"moyenne_vente":0.0},"IN-A 4510":{"stock_brut":97.0,"stock_couleurs":-97.0,"pct_laque":null,"moyenne_vente":0.0},"IN-A 4511":{"stock_brut":83.0,"stock_couleurs":-83.0,"pct_laque":null,"moyenne_vente":0.0},"IN-C 4501":{"stock_brut":71.0,"stock_couleurs":-71.0,"pct_laque":null,"moyenne_vente":0.0},"IN-C 4502":{"stock_brut":63.0,"stock_couleurs":-63.0,"pct_laque":null,"moyenne_vente":0.0},"IN-C 4503":{"stock_brut":52.0,"stock_couleurs":-52.0,"pct_laque":null,"moyenne_vente":0.0},"IN-C 4504":{"stock_brut":1.0,"stock_couleurs":-1.0,"pct_laque":null,"moyenne_vente":0.0},"IN-C 4505":{"stock_brut":72.0,"stock_couleurs":-72.0,"pct_laque":null,"moyenne_vente":0.0},"IN-C 4506":{"stock_brut":86.0,"stock_couleurs":-86.0,"pct_laque":null,"moyenne_vente":0.0},"IN-C 4507":{"stock_brut":79.0,"stock_couleurs":-79.0,"pct_laque":null,"moyenne_vente":0.0},"IN-C 4508":{"stock_brut":91.0,"stock_couleurs":-91.0,"pct_laque":null,"moyenne_vente":0.0},"IN-C 4509":{"stock_brut":565.0,"stock_couleurs":-565.0,"pct_laque":null,"moyenne_vente":0.0},"KL7401":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"KL7402":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"KL7403":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"KL7404":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"KL7405":{"stock_brut":0.0,"stock_couleurs":1.0,"pct_laque":1.0,"moyenne_vente":0.0},"KL7406":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"KL7407":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"KL7409":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"KL7410":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"KL7412":{"stock_brut":0.0,"stock_couleurs":1.0,"pct_laque":1.0,"moyenne_vente":0.0},"KL7413":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"KL7414":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"KL7416":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"KL7417":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"KL7420":{"stock_brut":809.0,"stock_couleurs":5.0,"pct_laque":0.006142506142506142,"moyenne_vente":0.0},"KL7421HS-1.2":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"KL7421HS-2.5":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"KL7421HS-3":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"KLHS7415":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"KLHS7418":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"KLHS7419":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"KLHS7421":{"stock_brut":0.0,"stock_couleurs":24.0,"pct_laque":1.0,"moyenne_vente":0.0},"LM-55":{"stock_brut":1210.0,"stock_couleurs":-1210.0,"pct_laque":null,"moyenne_vente":0.0},"LM-55A":{"stock_brut":389.0,"stock_couleurs":1315.0,"pct_laque":0.7717136150234741,"moyenne_vente":0.0},"LM-60":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"LM-60A":{"stock_brut":0.0,"stock_couleurs":2990.0,"pct_laque":1.0,"moyenne_vente":0.0},"LM-62":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"LM-757":{"stock_brut":16498.0,"stock_couleurs":777.0,"pct_laque":0.044978292329956586,"moyenne_vente":13391.545454545454},"LM-79":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"LM-80":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"LM-F":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"LM-F/6":{"stock_brut":2860.0,"stock_couleurs":-2405.0,"pct_laque":-5.285714285714286,"moyenne_vente":1927.6},"LM-F45":{"stock_brut":13985.0,"stock_couleurs":118.0,"pct_laque":0.00836701411047295,"moyenne_vente":292.46153846153845},"LM-F79/100":{"stock_brut":176.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.23076923076923078},"LM-FR":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"LM-FVS":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"LM-S758":{"stock_brut":415.0,"stock_couleurs":75.0,"pct_laque":0.15306122448979592,"moyenne_vente":1750.6},"LM-SR":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"LMDP-100":{"stock_brut":87.0,"stock_couleurs":1.0,"pct_laque":0.011363636363636364,"moyenne_vente":0.0},"LMDPF-100":{"stock_brut":232.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"LMMO-757":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"LMMO-S758":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"LMOR-19":{"stock_brut":890.0,"stock_couleurs":-818.0,"pct_laque":-11.36111111111111,"moyenne_vente":0.0},"LMOR-24":{"stock_brut":107.0,"stock_couleurs":-105.0,"pct_laque":-52.5,"moyenne_vente":0.0},"LMOR-36":{"stock_brut":447.0,"stock_couleurs":-439.0,"pct_laque":-54.875,"moyenne_vente":0.0},"LMOR-58":{"stock_brut":1152.0,"stock_couleurs":-1083.0,"pct_laque":-15.695652173913043,"moyenne_vente":0.0},"LMOR-59":{"stock_brut":89.0,"stock_couleurs":-87.0,"pct_laque":-43.5,"moyenne_vente":0.0},"LMOR-65":{"stock_brut":147.0,"stock_couleurs":-147.0,"pct_laque":null,"moyenne_vente":0.0},"MR002":{"stock_brut":1456.0,"stock_couleurs":3392.0,"pct_laque":0.6996699669966997,"moyenne_vente":947.8},"MR003":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"MR006":{"stock_brut":123.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"MR014":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"MR014_P6500":{"stock_brut":424.0,"stock_couleurs":71.0,"pct_laque":0.14343434343434344,"moyenne_vente":6.4},"MR020":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"MR021":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"MR030":{"stock_brut":0.0,"stock_couleurs":2.0,"pct_laque":1.0,"moyenne_vente":92.75},"OK4701P1":{"stock_brut":253.0,"stock_couleurs":-253.0,"pct_laque":null,"moyenne_vente":0.0},"OK4701P2":{"stock_brut":549.0,"stock_couleurs":-549.0,"pct_laque":null,"moyenne_vente":0.0},"OK4702P1":{"stock_brut":225.0,"stock_couleurs":-225.0,"pct_laque":null,"moyenne_vente":0.0},"OK4702P2":{"stock_brut":786.0,"stock_couleurs":-786.0,"pct_laque":null,"moyenne_vente":0.0},"OK4703P1":{"stock_brut":67.0,"stock_couleurs":-67.0,"pct_laque":null,"moyenne_vente":0.0},"OK4703P2":{"stock_brut":463.0,"stock_couleurs":-463.0,"pct_laque":null,"moyenne_vente":0.0},"OK4704P1":{"stock_brut":49.0,"stock_couleurs":-49.0,"pct_laque":null,"moyenne_vente":0.0},"OK4704P2":{"stock_brut":471.0,"stock_couleurs":-471.0,"pct_laque":null,"moyenne_vente":0.0},"OK4705P1":{"stock_brut":103.0,"stock_couleurs":-103.0,"pct_laque":null,"moyenne_vente":0.0},"OK4705P2":{"stock_brut":437.0,"stock_couleurs":-437.0,"pct_laque":null,"moyenne_vente":0.0},"OK4706P1":{"stock_brut":24.0,"stock_couleurs":-24.0,"pct_laque":null,"moyenne_vente":0.0},"OK4706P2":{"stock_brut":427.0,"stock_couleurs":-427.0,"pct_laque":null,"moyenne_vente":0.0},"OK4707P1":{"stock_brut":42.0,"stock_couleurs":-42.0,"pct_laque":null,"moyenne_vente":0.0},"OK4707P2":{"stock_brut":217.0,"stock_couleurs":-217.0,"pct_laque":null,"moyenne_vente":0.0},"OK4709P2":{"stock_brut":143.0,"stock_couleurs":-143.0,"pct_laque":null,"moyenne_vente":0.0},"OK4715":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"OK4716":{"stock_brut":1194.0,"stock_couleurs":-1193.0,"pct_laque":-1193.0,"moyenne_vente":0.0},"OK4717":{"stock_brut":322.0,"stock_couleurs":-322.0,"pct_laque":null,"moyenne_vente":0.0},"OK4718":{"stock_brut":159.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"OK4719":{"stock_brut":382.0,"stock_couleurs":-382.0,"pct_laque":null,"moyenne_vente":0.0},"OK4721":{"stock_brut":642.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"OK4722":{"stock_brut":517.0,"stock_couleurs":-498.0,"pct_laque":-26.210526315789473,"moyenne_vente":0.0},"OK4723":{"stock_brut":1677.0,"stock_couleurs":4.0,"pct_laque":0.002379535990481856,"moyenne_vente":0.0},"OK4723_3":{"stock_brut":4.0,"stock_couleurs":-4.0,"pct_laque":null,"moyenne_vente":0.0},"OK4724":{"stock_brut":1025.0,"stock_couleurs":-1025.0,"pct_laque":null,"moyenne_vente":0.0},"OK4724_3":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"OK4725":{"stock_brut":161.0,"stock_couleurs":-161.0,"pct_laque":null,"moyenne_vente":0.0},"P007":{"stock_brut":121.0,"stock_couleurs":133.0,"pct_laque":0.5236220472440944,"moyenne_vente":0.0},"P010":{"stock_brut":92.0,"stock_couleurs":1.0,"pct_laque":0.010752688172043012,"moyenne_vente":56.0},"P013":{"stock_brut":630.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"P031":{"stock_brut":231.0,"stock_couleurs":265.0,"pct_laque":0.5342741935483871,"moyenne_vente":252.875},"P039":{"stock_brut":0.0,"stock_couleurs":12.0,"pct_laque":1.0,"moyenne_vente":28.5},"P057-BARRES":{"stock_brut":100.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":23.0},"P063-BARRES":{"stock_brut":49.0,"stock_couleurs":3.0,"pct_laque":0.057692307692307696,"moyenne_vente":0.0},"P650":{"stock_brut":83.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"P660":{"stock_brut":136.0,"stock_couleurs":1.0,"pct_laque":0.0072992700729927005,"moyenne_vente":0.0},"P662":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"P663":{"stock_brut":100.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"PI451":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"PI452":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"PI453":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"PI454":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"PI455":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"PI456":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"PIV001":{"stock_brut":33.0,"stock_couleurs":-33.0,"pct_laque":null,"moyenne_vente":0.0},"PIV002":{"stock_brut":70.0,"stock_couleurs":-70.0,"pct_laque":null,"moyenne_vente":0.0},"PIV003":{"stock_brut":42.0,"stock_couleurs":-42.0,"pct_laque":null,"moyenne_vente":0.0},"PIV004":{"stock_brut":60.0,"stock_couleurs":-60.0,"pct_laque":null,"moyenne_vente":0.0},"PIV005":{"stock_brut":95.0,"stock_couleurs":-95.0,"pct_laque":null,"moyenne_vente":0.0},"PIV006":{"stock_brut":174.0,"stock_couleurs":-174.0,"pct_laque":null,"moyenne_vente":0.0},"PIV007":{"stock_brut":1210.0,"stock_couleurs":-1210.0,"pct_laque":null,"moyenne_vente":0.0},"PL101":{"stock_brut":263.0,"stock_couleurs":1.0,"pct_laque":0.003787878787878788,"moyenne_vente":0.0},"PL102":{"stock_brut":185.0,"stock_couleurs":2.0,"pct_laque":0.0106951871657754,"moyenne_vente":0.0},"PL103":{"stock_brut":1386.0,"stock_couleurs":12.0,"pct_laque":0.008583690987124463,"moyenne_vente":0.0},"PL104":{"stock_brut":75.0,"stock_couleurs":-74.0,"pct_laque":-74.0,"moyenne_vente":0.0},"PL105":{"stock_brut":665.0,"stock_couleurs":-665.0,"pct_laque":null,"moyenne_vente":0.0},"PL106":{"stock_brut":658.0,"stock_couleurs":-656.0,"pct_laque":-328.0,"moyenne_vente":0.0},"PL107":{"stock_brut":281.0,"stock_couleurs":-263.0,"pct_laque":-14.61111111111111,"moyenne_vente":0.0},"PL108":{"stock_brut":622.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"PL109":{"stock_brut":526.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"PL110":{"stock_brut":486.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"PL111":{"stock_brut":217.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"PL112":{"stock_brut":234.0,"stock_couleurs":1.0,"pct_laque":0.00425531914893617,"moyenne_vente":0.0},"PL113":{"stock_brut":314.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"PL114":{"stock_brut":274.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"PL116":{"stock_brut":64.0,"stock_couleurs":-64.0,"pct_laque":null,"moyenne_vente":0.0},"PR6301":{"stock_brut":0.0,"stock_couleurs":1803.0,"pct_laque":1.0,"moyenne_vente":0.0},"PR6301P1":{"stock_brut":110.0,"stock_couleurs":-110.0,"pct_laque":null,"moyenne_vente":0.0},"PR6301P2":{"stock_brut":55.0,"stock_couleurs":-55.0,"pct_laque":null,"moyenne_vente":0.0},"PR6302P1":{"stock_brut":133.0,"stock_couleurs":-133.0,"pct_laque":null,"moyenne_vente":0.0},"PR6303":{"stock_brut":0.0,"stock_couleurs":96.0,"pct_laque":1.0,"moyenne_vente":0.0},"PR6303P2":{"stock_brut":96.0,"stock_couleurs":-96.0,"pct_laque":null,"moyenne_vente":0.0},"PR6304":{"stock_brut":0.0,"stock_couleurs":98.0,"pct_laque":1.0,"moyenne_vente":0.0},"PR6304P1":{"stock_brut":96.0,"stock_couleurs":-96.0,"pct_laque":null,"moyenne_vente":0.0},"PR6305":{"stock_brut":0.0,"stock_couleurs":58.0,"pct_laque":1.0,"moyenne_vente":0.0},"PR6305P1":{"stock_brut":53.0,"stock_couleurs":-53.0,"pct_laque":null,"moyenne_vente":0.0},"PR6307":{"stock_brut":0.0,"stock_couleurs":148.0,"pct_laque":1.0,"moyenne_vente":0.0},"PR6307P1":{"stock_brut":140.0,"stock_couleurs":-140.0,"pct_laque":null,"moyenne_vente":0.0},"PR6308":{"stock_brut":0.0,"stock_couleurs":162.0,"pct_laque":1.0,"moyenne_vente":0.0},"PR6308P2":{"stock_brut":162.0,"stock_couleurs":-162.0,"pct_laque":null,"moyenne_vente":0.0},"PR6309":{"stock_brut":76.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"PR6310":{"stock_brut":66.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"PR6311":{"stock_brut":387.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"PR6314":{"stock_brut":125.0,"stock_couleurs":0.0,"pct_laque":0.0,"moyenne_vente":0.0},"PR6315":{"stock_brut":182.0,"stock_couleurs":-182.0,"pct_laque":null,"moyenne_vente":0.0},"SP6001":{"stock_brut":0.0,"stock_couleurs":null,"pct_laque":null,"moyenne_vente":0.0},"SP6001P1":{"stock_brut":278.0,"stock_couleurs":-278.0,"pct_laque":null,"moyenne_vente":0.0},"SP6002":{"stock_brut":0.0,"stock_couleurs":2.0,"pct_laque":1.0,"moyenne_vente":0.0},"SP6002P1":{"stock_brut":612.0,"stock_couleurs":-612.0,"pct_laque":null,"moyenne_vente":0.0},"SP6003":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"SP6003P1":{"stock_brut":534.0,"stock_couleurs":-534.0,"pct_laque":null,"moyenne_vente":0.0},"SP6004P1":{"stock_brut":392.0,"stock_couleurs":-392.0,"pct_laque":null,"moyenne_vente":0.0},"SP6005P1":{"stock_brut":123.0,"stock_couleurs":-123.0,"pct_laque":null,"moyenne_vente":0.0},"SP6006P1":{"stock_brut":343.0,"stock_couleurs":-343.0,"pct_laque":null,"moyenne_vente":0.0},"SP6007":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"SP6007P1":{"stock_brut":155.0,"stock_couleurs":-155.0,"pct_laque":null,"moyenne_vente":0.0},"SP6009P1":{"stock_brut":98.0,"stock_couleurs":-98.0,"pct_laque":null,"moyenne_vente":0.0},"SP6009P2":{"stock_brut":103.0,"stock_couleurs":-103.0,"pct_laque":null,"moyenne_vente":0.0},"SP6010P2":{"stock_brut":56.0,"stock_couleurs":-56.0,"pct_laque":null,"moyenne_vente":0.0},"SP6012P1":{"stock_brut":113.0,"stock_couleurs":-113.0,"pct_laque":null,"moyenne_vente":0.0},"SP6015P1":{"stock_brut":329.0,"stock_couleurs":-329.0,"pct_laque":null,"moyenne_vente":0.0},"SP6015P1-1.2":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"SP6015P1-2.5":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"SP6015P1-3":{"stock_brut":0.0,"stock_couleurs":0.0,"pct_laque":null,"moyenne_vente":0.0},"SP6016P2":{"stock_brut":120.0,"stock_couleurs":-120.0,"pct_laque":null,"moyenne_vente":0.0},"SP6017P2":{"stock_brut":1209.0,"stock_couleurs":-1209.0,"pct_laque":null,"moyenne_vente":0.0},"SP6018P1":{"stock_brut":688.0,"stock_couleurs":-688.0,"pct_laque":null,"moyenne_vente":0.0},"SP6018P2":{"stock_brut":671.0,"stock_couleurs":-671.0,"pct_laque":null,"moyenne_vente":0.0},"SP6019P1":{"stock_brut":195.0,"stock_couleurs":-195.0,"pct_laque":null,"moyenne_vente":0.0}}}')
COLOR_PRIORITY: Dict[str, float] = {str(k).upper(): float(v) for k, v in REFERENCE["colors"].items()}
TECH_REFERENCE: Dict[str, Dict[str, Any]] = {str(k).upper(): v for k, v in REFERENCE["tech"].items()}
STOCK_REFERENCE: Dict[str, Dict[str, Any]] = {str(k).upper(): v for k, v in REFERENCE["stock"].items()}


def norm_text(v: Any) -> str:
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except Exception:
        pass
    return str(v).strip()


def norm_key(v: Any) -> str:
    s = unicodedata.normalize("NFKD", norm_text(v)).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")


def to_float(v: Any, default: float = 0.0) -> float:
    try:
        if v is None or (isinstance(v, float) and math.isnan(v)):
            return default
        if isinstance(v, str):
            s = v.strip().replace(" ", "").replace(",", ".")
            if not s or s.upper() in {"#N/A", "N/A", "NA", "NONE", "NAN", "-"}:
                return default
            return float(s)
        return float(v)
    except Exception:
        return default


def to_int(v: Any, default: int = 0) -> int:
    return int(round(to_float(v, float(default))))


def parse_date(v: Any) -> Optional[pd.Timestamp]:
    if v is None:
        return None
    try:
        if isinstance(v, (int, float)):
            n = float(v)
            if 20000 <= n <= 80000:
                return pd.Timestamp("1899-12-30") + pd.to_timedelta(n, unit="D")
        ts = pd.to_datetime(v, errors="coerce")
        if pd.isna(ts) or ts.year <= 1900:
            return None
        return pd.Timestamp(ts)
    except Exception:
        return None


def iso_week_dates(year: int, week: int) -> Dict[int, date]:
    monday = date.fromisocalendar(int(year), int(week), 1)
    return {i: monday + timedelta(days=i) for i in range(N_DAYS)}


def automatic_planning_week(reference_date: Optional[date] = None) -> Tuple[int, int]:
    """Lun-Mer: semaine courante. Jeu-Dim: semaine suivante."""
    ref = reference_date or date.today()
    target = ref if ref.weekday() <= 2 else ref + timedelta(days=7)
    iso = target.isocalendar()
    return int(iso.year), int(iso.week)


def _ref_key(v: Any) -> str:
    return norm_text(v).upper().replace(" ", " ").strip()


def _ref_candidates(v: Any) -> List[str]:
    k = _ref_key(v)
    out = [k]
    if re.fullmatch(r"0+\d+", k):
        out.append(k.lstrip("0") or "0")
    if re.fullmatch(r"\d+", k):
        out.append(k.lstrip("0") or "0")
    return list(dict.fromkeys(out))


def article_family(article_internal: str) -> str:
    m = re.match(r"([A-Z]+)", _ref_key(article_internal))
    return m.group(1) if m else ""


def _build_family_fallbacks() -> Tuple[Dict[str, int], Dict[str, float]]:
    bars: Dict[str, List[float]] = defaultdict(list)
    weights: Dict[str, List[float]] = defaultdict(list)
    for art, vals in TECH_REFERENCE.items():
        fam = article_family(art)
        if not fam:
            continue
        b = to_float(vals.get("bars"), 0)
        w = to_float(vals.get("weight"), 0)
        if b > 0:
            bars[fam].append(b)
        if w > 0:
            weights[fam].append(w)
    return (
        {fam: max(1, int(round(median(vs)))) for fam, vs in bars.items() if vs},
        {fam: float(median(vs)) for fam, vs in weights.items() if vs},
    )


FAMILY_BARS, FAMILY_WEIGHTS = _build_family_fallbacks()


def lookup_technical(article_internal: str) -> Tuple[int, float, str]:
    for k in _ref_candidates(article_internal):
        if k in TECH_REFERENCE:
            v = TECH_REFERENCE[k]
            bars = max(1, to_int(v.get("bars"), 13))
            weight = max(0.0, to_float(v.get("weight"), 0.0))
            return bars, weight, "article"
    fam = article_family(article_internal)
    if fam in FAMILY_BARS:
        return FAMILY_BARS[fam], max(0.0, FAMILY_WEIGHTS.get(fam, 0.0)), "famille"
    return 13, 0.0, "fallback"


def lookup_stock_master(article_internal: str) -> Dict[str, Optional[float]]:
    for k in _ref_candidates(article_internal):
        if k in STOCK_REFERENCE:
            return STOCK_REFERENCE[k]
    return {}


def stable_unknown_priority(color: str) -> float:
    h = hashlib.sha256(norm_text(color).upper().encode("utf-8")).hexdigest()
    return 45.0 + (int(h[:6], 16) % 1000) / 100.0


def color_priority(color: str) -> Tuple[float, bool]:
    c = norm_text(color).upper()
    if c in COLOR_PRIORITY:
        return float(COLOR_PRIORITY[c]), True
    return stable_unknown_priority(c), False


def _looks_like_color_token(token: str) -> bool:
    t = norm_text(token).upper()
    if not t or t == "BRUT":
        return False
    if re.search(r"\d+(?:[.,]\d+)?X\d+", t) or "/" in t or "." in t:
        return False
    if t in COLOR_PRIORITY:
        return True
    if re.fullmatch(r"RAL?\d{4}", t) or re.fullmatch(r"N\d{2}", t):
        return True
    return bool(re.fullmatch(r"[A-Z]{3,14}", t))


def split_article(article: Any) -> Tuple[str, str]:
    s = norm_text(article)
    if "-" not in s:
        return s, ""
    left, right = s.rsplit("-", 1)
    color = right.strip().upper()
    if not _looks_like_color_token(color):
        return s, ""
    return left.strip(), color


def _color_class(color: Any) -> str:
    c = norm_text(color).upper()
    if c in WHITE_ALIASES:
        return "WHITE"
    if c in BLACK_ALIASES:
        return "BLACK"
    return "OTHER"


def _white_black_conflict(colors_a: Iterable[Any], colors_b: Optional[Iterable[Any]] = None) -> bool:
    a = {_color_class(c) for c in colors_a if norm_text(c)}
    if colors_b is None:
        return "WHITE" in a and "BLACK" in a
    b = {_color_class(c) for c in colors_b if norm_text(c)}
    return ("WHITE" in a and "BLACK" in b) or ("BLACK" in a and "WHITE" in b)


COLUMN_ALIASES = {
    "societe": "Société",
    "numcommande": "NumCommande",
    "etatcommande": "EtatCommande",
    "datecreation": "DateCréation",
    "nomclient": "NomClient",
    "article": "Article",
    "qtecommande": "QteCommandé",
    "restealivrer": "ResteALivrer",
    "preleve": "Prelevé",
    "2": "reservation brut",
    "reservation_brut": "reservation brut",
    "numof": "NumOF",
    "datedebut": "DateDébut",
    "prodstatut": "ProdStatut",
    "qtecommence": "QteCommencé",
    "qterestante": "QteRestante",
    "qterecu": "QteRèçu",
    "reserverbr": "ReserverBR",
    "stockphysique": "StockPhysique",
    "reserver": "Reserver",
    "poidarticle": "PoidArticle",
    "poidsarticle": "PoidArticle",
    "datelivraisonconfirme": "DateLivraisonConfirmé",
    "dateexpeditionconfirme": "DateExpeditionConfirmé",
    "dateexpeditiondemande": "DateExpeditionDemandé",
}

REQUIRED_COLUMNS = ["NumCommande", "DateCréation", "NomClient", "Article", "QteCommandé", "ResteALivrer"]


def _canonicalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    rename = {}
    for c in df.columns:
        k = norm_key(c)
        if k in COLUMN_ALIASES:
            rename[c] = COLUMN_ALIASES[k]
    return df.rename(columns=rename)


def load_extraction_workbook(data: bytes) -> Tuple[pd.DataFrame, str]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    candidate = None
    for ws in wb.worksheets:
        if norm_key(ws.title) == "extraction_ax":
            candidate = ws
            break
    if candidate is None:
        for ws in wb.worksheets:
            first = [norm_key(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1, max_col=min(ws.max_column, 100)))]
            if "numcommande" in first and "article" in first:
                candidate = ws
                break
    if candidate is None:
        wb.close()
        raise ValueError("Feuille 'extraction ax' introuvable et aucune feuille NumCommande/Article détectée.")

    header = [c.value for c in next(candidate.iter_rows(min_row=1, max_row=1, max_col=min(candidate.max_column, 100)))]
    while header and header[-1] is None:
        header.pop()
    rows = []
    blanks = 0
    for row in candidate.iter_rows(min_row=2, max_col=len(header), values_only=True):
        if all(v is None for v in row):
            blanks += 1
            if blanks >= 100:
                break
            continue
        blanks = 0
        rows.append(tuple(row))
    sheet_name = candidate.title
    wb.close()
    df = _canonicalize_columns(pd.DataFrame(rows, columns=header))
    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        raise ValueError("Colonnes obligatoires manquantes: " + ", ".join(missing))
    return df, sheet_name


@dataclass(frozen=True)
class PlannerConfig:
    year: int
    week: int
    capacity_h: float = DEFAULT_CAPACITY_H
    cleaning_min: int = DEFAULT_CLEANING_MIN
    minutes_per_bal: float = DEFAULT_MIN_PER_BAL
    powder_coeff: float = DEFAULT_POWDER_COEFF
    target_utilization: float = DEFAULT_TARGET_UTIL
    max_colors_per_day: int = MAX_COLORS_PER_DAY
    include_without_of: bool = False
    use_stock_for_relaquage: bool = False
    forbid_white_black_adjacent: bool = True
    solver_seconds: float = 12.0


def due_date_from_row(row: pd.Series) -> Optional[pd.Timestamp]:
    for c in ["DateLivraisonConfirmé", "DateExpeditionConfirmé", "DateExpeditionDemandé"]:
        if c in row.index:
            d = parse_date(row.get(c))
            if d is not None:
                return d
    return None


def _priority_score(row: pd.Series, week_start: pd.Timestamp) -> Tuple[float, str]:
    score = 100.0
    reasons: List[str] = []
    prod = norm_key(row.get("ProdStatut"))
    if "commenc" in prod:
        score += 650
        reasons.append("OF commencé")
    elif "cree" in prod:
        score += 500
        reasons.append("OF créé")
    else:
        score -= 250
        reasons.append("OF non créé")

    remaining = max(0.0, to_float(row.get("ResteALivrer")))
    res_flag = norm_key(row.get("reservation brut"))
    reserver_br = max(0.0, to_float(row.get("ReserverBR")))
    if res_flag in {"oui", "yes", "1", "true"}:
        score += 450
        reasons.append("brut réservé")
    elif remaining > 0 and reserver_br <= 0:
        score -= 100

    if remaining > 0 and reserver_br > 0:
        ratio = min(1.5, reserver_br / remaining)
        score += 220 * min(1.0, ratio)
        if ratio >= 1:
            reasons.append("matière suffisante")

    created = parse_date(row.get("DateCréation"))
    if created is not None:
        age = max(0, (week_start.date() - created.date()).days)
        score += min(700, age * 7)
        if age >= 30:
            reasons.append("commande ancienne")

    due = due_date_from_row(row)
    if due is not None:
        delta = (due.date() - week_start.date()).days
        if delta < 0:
            score += 2500 + min(2000, abs(delta) * 120)
            reasons.append("retard")
        elif delta <= 5:
            score += 1200 - delta * 100
            reasons.append("échéance semaine")

    if norm_text(row.get("NumOF")):
        score += 80
    return float(score), " · ".join(reasons[:5])


def build_candidates(source: pd.DataFrame, cfg: PlannerConfig) -> Tuple[pd.DataFrame, Dict[str, Any]]:
    week_start = pd.Timestamp(iso_week_dates(cfg.year, cfg.week)[0])
    excluded = Counter()
    quality = Counter()
    rows: List[Dict[str, Any]] = []

    for src_idx, src in source.iterrows():
        cmd = norm_text(src.get("NumCommande")).upper()
        if not cmd:
            excluded["commande vide"] += 1
            continue
        etat = norm_key(src.get("EtatCommande"))
        if etat and "encours" not in etat:
            excluded["commande non encours"] += 1
            continue
        remaining = max(0.0, to_float(src.get("ResteALivrer")))
        if remaining <= 0:
            excluded["reste nul"] += 1
            continue

        article = norm_text(src.get("Article"))
        article_internal, color = split_article(article)
        if not article_internal or not color:
            excluded["article/couleur non planifiable"] += 1
            continue

        prod_status = norm_text(src.get("ProdStatut"))
        ps = norm_key(prod_status)
        if "termine" in ps:
            excluded["OF terminé"] += 1
            continue
        ready = ("cree" in ps) or ("commenc" in ps)
        if not ready and not cfg.include_without_of:
            excluded["OF non créé"] += 1
            continue

        bars, ref_weight, ref_source = lookup_technical(article_internal)
        direct_weight = max(0.0, to_float(src.get("PoidArticle"))) if "PoidArticle" in source.columns else 0.0
        unit_weight = direct_weight if direct_weight > 0 else ref_weight
        weight_source = "extraction" if direct_weight > 0 else ref_source
        if unit_weight <= 0:
            quality["poids_inconnu"] += 1
        if ref_source == "fallback":
            quality["barres_fallback"] += 1
        elif ref_source == "famille":
            quality["barres_famille"] += 1

        nuance, nuance_known = color_priority(color)
        if not nuance_known:
            quality["nuance_inconnue"] += 1

        stock_master = lookup_stock_master(article_internal)
        score, reason = _priority_score(src, week_start)
        due = due_date_from_row(src)

        rows.append({
            "Num Commande": cmd,
            "DateCréation": parse_date(src.get("DateCréation")),
            "Nom Client": norm_text(src.get("NomClient")),
            "Article": article,
            "Article/int": article_internal,
            "Couleur": color,
            "Nuance": round(float(nuance), 3),
            "Qte Commandée": to_int(src.get("QteCommandé")),
            "Reste A Livrer": int(math.ceil(remaining - 1e-9)),
            "Prelevé": to_int(src.get("Prelevé")),
            "reservation brut": norm_text(src.get("reservation brut")) or norm_text(src.get(" 2")) or "non",
            "Num OF": norm_text(src.get("NumOF")),
            "Prod Statut": prod_status,
            "Qte Commencée": to_int(src.get("QteCommencé")),
            "QteRestante": to_int(src.get("QteRestante")),
            "QteRèçu": to_int(src.get("QteRèçu")),
            "ReserverBR": to_int(src.get("ReserverBR")),
            "StockPhysique": to_int(src.get("StockPhysique")),
            "Reserver": to_int(src.get("Reserver")),
            "Lancement": 0,
            "Re-laquage": 0,
            "PoidsUn": round(float(unit_weight), 4),
            "PoidsT": 0.0,
            "Poudre": 0.0,
            "Barre/bal": int(bars),
            "Nbre Bal": 0,
            "tps": 0.0,
            "Stock brut": to_float(stock_master.get("stock_brut"), 0.0),
            "moyenne vente": to_float(stock_master.get("moyenne_vente"), 0.0),
            "% laqué": to_float(stock_master.get("pct_laque"), 0.0),
            "_source_index": int(src_idx),
            "_score": round(score, 3),
            "_reason": reason,
            "_due": due,
            "_ref_source": ref_source,
            "_weight_source": weight_source,
        })

    df = pd.DataFrame(rows)
    if df.empty:
        return df, {"source_rows": len(source), "excluded": dict(excluded), "quality": dict(quality)}

    df = df.sort_values(["_score", "DateCréation", "Num Commande"], ascending=[False, True, True], na_position="last").reset_index(drop=True)

    # IMPORTANT: StockPhysique est répété sur plusieurs lignes d'un même article dans AX.
    # Si l'option est activée, on alloue donc le stock UNE SEULE FOIS par article complet.
    stock_remaining: Dict[str, float] = {}
    if cfg.use_stock_for_relaquage:
        for article, g in df.groupby(df["Article"].astype(str).str.upper(), sort=False):
            stock_remaining[article] = max(0.0, float(pd.to_numeric(g["StockPhysique"], errors="coerce").fillna(0).max()))

    launch_vals: List[int] = []
    relaq_vals: List[int] = []
    poids_t_vals: List[float] = []
    poudre_vals: List[float] = []
    nbal_vals: List[int] = []
    tps_vals: List[float] = []
    for _, r in df.iterrows():
        remaining = max(0, to_int(r["Reste A Livrer"]))
        relaq = 0
        if cfg.use_stock_for_relaquage:
            k = norm_text(r["Article"]).upper()
            available = max(0.0, stock_remaining.get(k, 0.0))
            relaq = int(min(float(remaining), available))
            stock_remaining[k] = max(0.0, available - relaq)
        launch = max(0, remaining - relaq)
        bars = max(1, to_int(r["Barre/bal"], 13))
        nbal = int(math.ceil(launch / bars)) if launch > 0 else 0
        weight_total = (launch + relaq) * max(0.0, to_float(r["PoidsUn"]))
        powder = weight_total * cfg.powder_coeff
        tps = nbal * cfg.minutes_per_bal / 60.0
        launch_vals.append(launch)
        relaq_vals.append(relaq)
        poids_t_vals.append(round(weight_total, 3))
        poudre_vals.append(round(powder, 3))
        nbal_vals.append(nbal)
        tps_vals.append(tps)
    df["Lancement"] = launch_vals
    df["Re-laquage"] = relaq_vals
    df["PoidsT"] = poids_t_vals
    df["Poudre"] = poudre_vals
    df["Nbre Bal"] = nbal_vals
    df["tps"] = tps_vals
    df["_line_id"] = [hashlib.sha1(f"{r['_source_index']}|{r['Num Commande']}|{r['Article']}|{r['Num OF']}".encode("utf-8")).hexdigest()[:16] for _, r in df.iterrows()]

    return df, {
        "source_rows": int(len(source)),
        "eligible_lines": int(len(df)),
        "excluded": dict(excluded),
        "quality": dict(quality),
    }


@dataclass
class Job:
    job_id: str
    line_indices: List[int]
    command: str
    color: str
    duration_min: int
    score: float
    created: Optional[datetime]


def build_jobs(lines: pd.DataFrame, cfg: PlannerConfig) -> List[Job]:
    jobs: List[Job] = []
    max_day = max(1, int(round(cfg.capacity_h * 60)))
    for (cmd, color), g in lines.groupby(["Num Commande", "Couleur"], sort=False):
        g = g.sort_values(["_score", "DateCréation"], ascending=[False, True], na_position="last")
        chunk: List[int] = []
        chunk_min = 0
        chunk_no = 1
        for idx, row in g.iterrows():
            mins = max(1, int(round(to_float(row["tps"]) * 60)))
            if chunk and chunk_min + mins > max_day:
                sub = lines.loc[chunk]
                created_vals = [x.to_pydatetime() if isinstance(x, pd.Timestamp) else x for x in sub["DateCréation"].tolist() if x is not None and not pd.isna(x)]
                jobs.append(Job(f"{cmd}|{color}|{chunk_no}", list(chunk), str(cmd), str(color), chunk_min, float(sub["_score"].max()), min(created_vals) if created_vals else None))
                chunk_no += 1
                chunk = []
                chunk_min = 0
            chunk.append(int(idx))
            chunk_min += mins
        if chunk:
            sub = lines.loc[chunk]
            created_vals = [x.to_pydatetime() if isinstance(x, pd.Timestamp) else x for x in sub["DateCréation"].tolist() if x is not None and not pd.isna(x)]
            jobs.append(Job(f"{cmd}|{color}|{chunk_no}", list(chunk), str(cmd), str(color), chunk_min, float(sub["_score"].max()), min(created_vals) if created_vals else None))
    return jobs


def _color_target_days(jobs: Sequence[Job]) -> Dict[str, int]:
    colors = sorted({j.color for j in jobs}, key=lambda c: (color_priority(c)[0], c))
    if len(colors) <= 1:
        return {c: 0 for c in colors}
    return {c: int(round(i * (N_DAYS - 1) / (len(colors) - 1))) for i, c in enumerate(colors)}


def _job_unscheduled_penalty(job: Job) -> int:
    return max(10000, int(round(25000 + job.score * 180 + min(job.duration_min, 900) * 10)))


def assign_jobs_ortools(jobs: List[Job], cfg: PlannerConfig) -> Tuple[Dict[str, int], List[str], str]:
    if not ORTOOLS_AVAILABLE or not jobs:
        return {}, [], ""
    colors = sorted({j.color for j in jobs})
    commands = sorted({j.command for j in jobs})
    by_color: Dict[str, List[int]] = defaultdict(list)
    by_command: Dict[str, List[int]] = defaultdict(list)
    for ji, job in enumerate(jobs):
        by_color[job.color].append(ji)
        by_command[job.command].append(ji)
    target_day = _color_target_days(jobs)

    model = cp_model.CpModel()
    x: Dict[Tuple[int, int], Any] = {}
    u: Dict[int, Any] = {}
    y: Dict[Tuple[int, str], Any] = {}
    objective: List[Any] = []
    cap = int(round(cfg.capacity_h * 60))

    for ji, job in enumerate(jobs):
        u[ji] = model.NewBoolVar(f"u_{ji}")
        for d in range(N_DAYS):
            x[(ji, d)] = model.NewBoolVar(f"x_{ji}_{d}")
            objective.append(x[(ji, d)] * abs(d - target_day.get(job.color, d)) * 3500)
        model.Add(sum(x[(ji, d)] for d in range(N_DAYS)) + u[ji] == 1)
        objective.append(u[ji] * _job_unscheduled_penalty(job))

    for d in range(N_DAYS):
        active = []
        for color in colors:
            y[(d, color)] = model.NewBoolVar(f"y_{d}_{norm_key(color)}")
            idxs = by_color[color]
            for ji in idxs:
                model.Add(x[(ji, d)] <= y[(d, color)])
            model.Add(y[(d, color)] <= sum(x[(ji, d)] for ji in idxs))
            active.append(y[(d, color)])
            objective.append(y[(d, color)] * 4000)
        n_colors = sum(active)
        model.Add(n_colors <= cfg.max_colors_per_day)
        second = model.NewBoolVar(f"second_{d}")
        model.Add(n_colors <= 1 + second)
        objective.append(second * 28000)
        prod = sum(jobs[ji].duration_min * x[(ji, d)] for ji in range(len(jobs)))
        model.Add(prod + cfg.cleaning_min * second <= cap)
        target = int(round(cap * cfg.target_utilization))
        load = model.NewIntVar(0, cap, f"load_{d}")
        model.Add(load == prod + cfg.cleaning_min * second)
        dev = model.NewIntVar(0, cap, f"dev_{d}")
        model.Add(dev >= target - load)
        model.Add(dev >= load - target)
        objective.append(dev * 2)

    if cfg.forbid_white_black_adjacent:
        white = [c for c in colors if _color_class(c) == "WHITE"]
        black = [c for c in colors if _color_class(c) == "BLACK"]
        for d in range(N_DAYS):
            for wc in white:
                for bc in black:
                    model.Add(y[(d, wc)] + y[(d, bc)] <= 1)
        for d in range(N_DAYS - 1):
            for wc in white:
                for bc in black:
                    model.Add(y[(d, wc)] + y[(d + 1, bc)] <= 1)
                    model.Add(y[(d, bc)] + y[(d + 1, wc)] <= 1)

    for ci, cmd in enumerate(commands):
        idxs = by_command[cmd]
        if len(idxs) <= 1:
            continue
        presents = []
        for d in range(N_DAYS):
            p = model.NewBoolVar(f"cmd_{ci}_{d}")
            for ji in idxs:
                model.Add(x[(ji, d)] <= p)
            model.Add(p <= sum(x[(ji, d)] for ji in idxs))
            presents.append(p)
        extra = model.NewIntVar(0, N_DAYS - 1, f"split_{ci}")
        model.Add(extra >= sum(presents) - 1)
        objective.append(extra * 14000)

    model.Minimize(sum(objective))
    solver = cp_model.CpSolver()
    solver.parameters.max_time_in_seconds = max(3.0, float(cfg.solver_seconds))
    solver.parameters.num_search_workers = max(1, min(8, os.cpu_count() or 2))
    solver.parameters.random_seed = 37
    status = solver.Solve(model)
    if status not in (cp_model.OPTIMAL, cp_model.FEASIBLE):
        return {}, [], ""

    assignments: Dict[str, int] = {}
    unscheduled: List[str] = []
    for ji, job in enumerate(jobs):
        if solver.Value(u[ji]):
            unscheduled.append(job.job_id)
            continue
        for d in range(N_DAYS):
            if solver.Value(x[(ji, d)]):
                assignments[job.job_id] = d
                break
        if job.job_id not in assignments:
            unscheduled.append(job.job_id)
    return assignments, unscheduled, "OR-Tools CP-SAT"


def _can_place(job: Job, d: int, loads: List[int], colors_by_day: List[set], cfg: PlannerConfig) -> Tuple[bool, int]:
    cap = int(round(cfg.capacity_h * 60))
    new_colors = set(colors_by_day[d]) | {job.color}
    if len(new_colors) > cfg.max_colors_per_day:
        return False, 0
    if cfg.forbid_white_black_adjacent:
        if _white_black_conflict(new_colors):
            return False, 0
        if d > 0 and _white_black_conflict(new_colors, colors_by_day[d - 1]):
            return False, 0
        if d < N_DAYS - 1 and _white_black_conflict(new_colors, colors_by_day[d + 1]):
            return False, 0
    clean_before = cfg.cleaning_min if len(colors_by_day[d]) >= 2 else 0
    clean_after = cfg.cleaning_min if len(new_colors) >= 2 else 0
    prod_before = loads[d] - clean_before
    projected = prod_before + job.duration_min + clean_after
    return projected <= cap, projected


def assign_jobs_fallback(jobs: List[Job], cfg: PlannerConfig) -> Tuple[Dict[str, int], List[str], str]:
    loads = [0] * N_DAYS
    colors_by_day: List[set] = [set() for _ in range(N_DAYS)]
    commands_by_day: List[set] = [set() for _ in range(N_DAYS)]
    target_day = _color_target_days(jobs)
    assignments: Dict[str, int] = {}
    unscheduled: List[str] = []
    ordered = sorted(jobs, key=lambda j: (-j.score, j.created or datetime.max, -j.duration_min, color_priority(j.color)[0]))

    for job in ordered:
        options = []
        for d in range(N_DAYS):
            ok, projected = _can_place(job, d, loads, colors_by_day, cfg)
            if not ok:
                continue
            new_color = job.color not in colors_by_day[d]
            ncolors = len(colors_by_day[d]) + (1 if new_color else 0)
            mono_pen = 0 if job.color in colors_by_day[d] else 18000 if not colors_by_day[d] else 90000
            target_pen = abs(d - target_day.get(job.color, d)) * 5500
            group_bonus = -7000 if job.command in commands_by_day[d] else 0
            balance = abs(cfg.capacity_h * 60 * cfg.target_utilization - projected) * 1.5
            early_pen = d * 20
            cost = mono_pen + target_pen + balance + group_bonus + early_pen - job.score * 8
            options.append((cost, d, projected, ncolors))
        if not options:
            unscheduled.append(job.job_id)
            continue
        _, d, projected, _ = min(options, key=lambda x: (x[0], x[1]))
        colors_by_day[d].add(job.color)
        loads[d] = projected
        commands_by_day[d].add(job.command)
        assignments[job.job_id] = d
    return assignments, unscheduled, "Heuristique déterministe"


def assign_jobs(jobs: List[Job], cfg: PlannerConfig) -> Tuple[Dict[str, int], List[str], str]:
    if ORTOOLS_AVAILABLE:
        a, u, e = assign_jobs_ortools(jobs, cfg)
        if e:
            return a, u, e
    return assign_jobs_fallback(jobs, cfg)


def build_result(source: pd.DataFrame, cfg: PlannerConfig) -> Dict[str, Any]:
    lines, quality = build_candidates(source, cfg)
    if lines.empty:
        raise ValueError("Aucune ligne éligible. Vérifier ProdStatut, ResteALivrer et les suffixes couleur.")
    jobs = build_jobs(lines, cfg)
    assignments, unscheduled_ids, engine = assign_jobs(jobs, cfg)
    job_map = {j.job_id: j for j in jobs}
    day_indices: Dict[int, List[int]] = {d: [] for d in range(N_DAYS)}
    for jid, d in assignments.items():
        if jid in job_map:
            day_indices[d].extend(job_map[jid].line_indices)
    uns_idx: List[int] = []
    for jid in unscheduled_ids:
        if jid in job_map:
            uns_idx.extend(job_map[jid].line_indices)
    planned_idx = {i for ids in day_indices.values() for i in ids}
    uns_idx.extend(i for i in lines.index if i not in planned_idx)
    uns_idx = list(dict.fromkeys(uns_idx))

    days: Dict[int, pd.DataFrame] = {}
    for d in range(N_DAYS):
        if not day_indices[d]:
            days[d] = lines.iloc[0:0].copy()
            continue
        df = lines.loc[day_indices[d]].copy()
        df = df.sort_values(["Nuance", "Couleur", "_score", "DateCréation", "Num Commande"], ascending=[True, True, False, True, True], na_position="last").reset_index(drop=True)
        df["_planned_day"] = DAYS[d]
        days[d] = df
    unscheduled = lines.loc[uns_idx].copy().sort_values(["_score", "Nuance"], ascending=[False, True]).reset_index(drop=True) if uns_idx else lines.iloc[0:0].copy()

    day_metrics = []
    total_load = 0.0
    hard_errors: List[str] = []
    for d in range(N_DAYS):
        df = days[d]
        prod_h = float(pd.to_numeric(df.get("tps", pd.Series(dtype=float)), errors="coerce").fillna(0).sum())
        colors = list(dict.fromkeys(df.get("Couleur", pd.Series(dtype=str)).dropna().astype(str).str.upper().tolist()))
        cleaning_h = cfg.cleaning_min / 60.0 if len(colors) >= 2 else 0.0
        total_h = prod_h + cleaning_h
        total_load += total_h
        if total_h > cfg.capacity_h + 0.01:
            hard_errors.append(f"{DAYS[d]}: surcharge {total_h:.2f}h > {cfg.capacity_h:.2f}h")
        if len(colors) > cfg.max_colors_per_day:
            hard_errors.append(f"{DAYS[d]}: {len(colors)} couleurs > {cfg.max_colors_per_day}")
        day_metrics.append({
            "Jour": DAYS[d], "Date": iso_week_dates(cfg.year, cfg.week)[d].strftime("%d/%m/%Y"),
            "Charge production h": round(prod_h, 2), "Changement couleur h": round(cleaning_h, 2),
            "Charge totale h": round(total_h, 2), "Capacité h": round(cfg.capacity_h, 2),
            "Charge %": round(total_h / cfg.capacity_h * 100, 1) if cfg.capacity_h else 0,
            "Couleurs": " -> ".join(colors), "Nb couleurs": len(colors), "Lignes": int(len(df)),
        })

    if cfg.forbid_white_black_adjacent:
        day_colors = [set(days[d].get("Couleur", pd.Series(dtype=str)).dropna().astype(str).str.upper().tolist()) for d in range(N_DAYS)]
        for d in range(N_DAYS):
            if _white_black_conflict(day_colors[d]):
                hard_errors.append(f"{DAYS[d]}: BLANC avec NOIR/DARK")
        for d in range(N_DAYS - 1):
            if _white_black_conflict(day_colors[d], day_colors[d + 1]):
                hard_errors.append(f"Transition interdite {DAYS[d]} -> {DAYS[d+1]}: BLANC <-> NOIR/DARK")

    prep = lines.sort_values(["Nuance", "Couleur", "_score", "DateCréation"], ascending=[True, True, False, True], na_position="last").reset_index(drop=True)
    return {
        "config": cfg,
        "days": days,
        "unscheduled": unscheduled,
        "preparation": prep,
        "quality": quality,
        "engine": engine,
        "hard_errors": hard_errors,
        "confidence": 100 if not hard_errors else max(0, 100 - 25 * len(hard_errors)),
        "metrics": {
            "days": day_metrics,
            "total_load_h": round(total_load, 2),
            "capacity_h": round(cfg.capacity_h * N_DAYS, 2),
            "utilization_pct": round(total_load / (cfg.capacity_h * N_DAYS) * 100, 1) if cfg.capacity_h else 0,
            "backlog": int(len(unscheduled)),
        },
    }


def _business_df(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for c in OUTPUT_COLUMNS:
        if c not in out.columns:
            out[c] = None
    return out[OUTPUT_COLUMNS]


def _style_header(ws, row: int = 1, fill: str = "17365D") -> None:
    white = "FFFFFF"
    line = Side(style="thin", color="D0D5DD")
    for c in ws[row]:
        if c.value is None:
            continue
        c.fill = PatternFill("solid", fgColor=fill)
        c.font = Font(color=white, bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=line)


def _write_dataframe(ws, df: pd.DataFrame, formula_params: bool = True) -> None:
    for ci, col in enumerate(df.columns, 1):
        ws.cell(1, ci, col)
    _style_header(ws)
    col_idx = {c: i + 1 for i, c in enumerate(df.columns)}
    for ri, (_, row) in enumerate(df.iterrows(), start=2):
        for ci, col in enumerate(df.columns, 1):
            val = row.get(col)
            if isinstance(val, pd.Timestamp):
                val = val.to_pydatetime()
            ws.cell(ri, ci, val)
        if formula_params and all(c in col_idx for c in ["Lancement", "Re-laquage", "PoidsUn", "PoidsT", "Poudre", "Barre/bal", "Nbre Bal", "tps"]):
            T = get_column_letter(col_idx["Lancement"])
            U = get_column_letter(col_idx["Re-laquage"])
            V = get_column_letter(col_idx["PoidsUn"])
            W = get_column_letter(col_idx["PoidsT"])
            X = get_column_letter(col_idx["Poudre"])
            Y = get_column_letter(col_idx["Barre/bal"])
            Z = get_column_letter(col_idx["Nbre Bal"])
            AA = get_column_letter(col_idx["tps"])
            ws[f"{W}{ri}"] = f"=({T}{ri}+{U}{ri})*{V}{ri}"
            ws[f"{X}{ri}"] = f"={W}{ri}*'Paramètres'!$B$6"
            ws[f"{Z}{ri}"] = f"=IFERROR(ROUNDUP({T}{ri}/{Y}{ri},0),0)"
            ws[f"{AA}{ri}"] = f"={Z}{ri}*'Paramètres'!$B$5/60"
    ws.freeze_panes = "A2"
    if df.shape[1] > 0:
        ws.auto_filter.ref = f"A1:{get_column_letter(df.shape[1])}{max(1, len(df)+1)}"
    for ci, col in enumerate(df.columns, 1):
        width = 14
        if col in {"Nom Client", "Article"}:
            width = 24
        elif col in {"Num Commande", "DateCréation", "Prod Statut", "reservation brut"}:
            width = 17
        elif col in {"Article/int"}:
            width = 18
        ws.column_dimensions[get_column_letter(ci)].width = width
    for row in ws.iter_rows(min_row=2, max_row=max(2, len(df)+1)):
        for cell in row:
            if cell.column == col_idx.get("DateCréation") and isinstance(cell.value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm"
    for c in ["Nuance", "PoidsUn", "PoidsT", "Poudre", "tps", "moyenne vente", "% laqué"]:
        if c in col_idx:
            for r in range(2, len(df)+2):
                ws.cell(r, col_idx[c]).number_format = "0.000"


def export_excel(result: Dict[str, Any], source_sheet_name: str = "extraction ax") -> bytes:
    cfg: PlannerConfig = result["config"]
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    navy = "17365D"; blue = "155EEF"; green = "E2F0D9"; amber = "FFF2CC"; red = "F4CCCC"; white = "FFFFFF"

    ws = wb.create_sheet("Résumé")
    ws["A1"] = "ALLUCO - Planning Laquage Lundi -> Vendredi"
    ws["A1"].fill = PatternFill("solid", fgColor=navy); ws["A1"].font = Font(color=white, bold=True, size=15)
    ws.merge_cells("A1:D1")
    metrics = result["metrics"]
    summary = [
        ("Semaine", f"S{cfg.week} / {cfg.year}"),
        ("Source", source_sheet_name),
        ("Moteur", result["engine"]),
        ("Confiance règles", f"{result['confidence']}%"),
        ("Charge semaine", f"{metrics['total_load_h']:.2f} h"),
        ("Capacité", f"{metrics['capacity_h']:.2f} h"),
        ("Utilisation", f"{metrics['utilization_pct']:.1f}%"),
        ("Backlog", metrics["backlog"]),
        ("Politique", f"Lun-Ven · max {cfg.max_colors_per_day} couleurs/jour · {cfg.minutes_per_bal:g} min/bal"),
        ("Re-laquage stock", "ACTIF" if cfg.use_stock_for_relaquage else "DÉSACTIVÉ (à valider métier)"),
    ]
    for i, (k, v) in enumerate(summary, start=3):
        ws.cell(i, 1, k).font = Font(bold=True, color=navy)
        ws.cell(i, 2, v)
    ws.column_dimensions["A"].width = 28; ws.column_dimensions["B"].width = 44
    r0 = 15
    headers = ["Jour", "Date", "Charge prod h", "Nettoyage h", "Charge totale h", "Capacité h", "%", "Couleurs", "Lignes"]
    for c, h in enumerate(headers, 1): ws.cell(r0, c, h)
    _style_header(ws, r0, blue)
    for rr, dm in enumerate(metrics["days"], r0 + 1):
        vals = [dm["Jour"].title(), dm["Date"], dm["Charge production h"], dm["Changement couleur h"], dm["Charge totale h"], dm["Capacité h"], dm["Charge %"], dm["Couleurs"], dm["Lignes"]]
        for cc, v in enumerate(vals, 1): ws.cell(rr, cc, v)
    ws.column_dimensions["H"].width = 35
    qrow = r0 + N_DAYS + 3
    ws.cell(qrow, 1, "Contrôles").font = Font(bold=True, color=navy, size=12)
    qrow += 1
    if result["hard_errors"]:
        for msg in result["hard_errors"]:
            ws.cell(qrow, 1, "ERREUR").fill = PatternFill("solid", fgColor=red); ws.cell(qrow, 2, msg); qrow += 1
    else:
        ws.cell(qrow, 1, "OK").fill = PatternFill("solid", fgColor=green); ws.cell(qrow, 2, "Capacité et contraintes journalières validées."); qrow += 1
    quality = result["quality"]
    for k, v in quality.get("quality", {}).items():
        if v:
            ws.cell(qrow, 1, "INFO").fill = PatternFill("solid", fgColor=amber); ws.cell(qrow, 2, f"{k}: {v}"); qrow += 1

    # Paramètres avant les tableaux, car les formules des feuilles y font référence.
    pws = wb.create_sheet("Paramètres")
    pws.append(["Paramètre", "Valeur", "Commentaire"])
    params = [
        ("Année", cfg.year, "semaine ISO"),
        ("Semaine", cfg.week, "semaine ISO"),
        ("Capacité/jour h", cfg.capacity_h, "Lundi à vendredi"),
        ("Minutes/balancelle", cfg.minutes_per_bal, "historique du classeur: 4 min"),
        ("Coefficient poudre", cfg.powder_coeff, "Poudre = PoidsT × coefficient"),
        ("Nettoyage min", cfg.cleaning_min, "si 2 couleurs dans la journée"),
        ("Max couleurs/jour", cfg.max_colors_per_day, "contrainte dure"),
        ("Inclure OF non créés", cfg.include_without_of, "désactivé par défaut"),
        ("Utiliser StockPhysique en re-laquage", cfg.use_stock_for_relaquage, "option à valider métier"),
    ]
    # Keep exact cell positions expected by formulas: B5 = minutes/bal, B6 = poudre coeff.
    pws.delete_rows(2, pws.max_row)
    pws.append(["Année", cfg.year, "semaine ISO"])
    pws.append(["Semaine", cfg.week, "semaine ISO"])
    pws.append(["Capacité/jour h", cfg.capacity_h, "Lundi à vendredi"])
    pws.append(["Minutes/balancelle", cfg.minutes_per_bal, "historique du classeur: 4 min"])
    pws.append(["Coefficient poudre", cfg.powder_coeff, "Poudre = PoidsT × coefficient"])
    pws.append(["Nettoyage min", cfg.cleaning_min, "si 2 couleurs dans la journée"])
    pws.append(["Max couleurs/jour", cfg.max_colors_per_day, "contrainte dure"])
    pws.append(["Inclure OF non créés", cfg.include_without_of, "désactivé par défaut"])
    pws.append(["Utiliser StockPhysique en re-laquage", cfg.use_stock_for_relaquage, "option à valider métier"])
    _style_header(pws)
    pws.column_dimensions["A"].width = 34; pws.column_dimensions["B"].width = 18; pws.column_dimensions["C"].width = 50

    for d in range(N_DAYS):
        ws_day = wb.create_sheet(f"Planning {DAYS[d].title()}")
        _write_dataframe(ws_day, _business_df(result["days"][d]), formula_params=True)

    bws = wb.create_sheet("Backlog")
    backlog = result["unscheduled"].copy()
    backlog_cols = OUTPUT_COLUMNS + ["_score", "_reason", "_ref_source"]
    for c in backlog_cols:
        if c not in backlog.columns: backlog[c] = None
    _write_dataframe(bws, backlog[backlog_cols], formula_params=True)
    if "_reason" in backlog_cols:
        bws.column_dimensions[get_column_letter(backlog_cols.index("_reason") + 1)].width = 52

    vws = wb.create_sheet("Préparation VF")
    _write_dataframe(vws, _business_df(result["preparation"]), formula_params=True)

    # Référentiel couleur embarqué pour audit/maintenance.
    cws = wb.create_sheet("Base couleur")
    cws.append(["Couleur", "Priorité Nuance"])
    for c, p in sorted(COLOR_PRIORITY.items(), key=lambda kv: (kv[1], kv[0])):
        cws.append([c, p])
    _style_header(cws)
    cws.column_dimensions["A"].width = 20; cws.column_dimensions["B"].width = 18

    out = io.BytesIO(); wb.save(out); return out.getvalue()


def generate_file(input_bytes: bytes, cfg: PlannerConfig) -> Tuple[bytes, Dict[str, Any], str]:
    source, sheet = load_extraction_workbook(input_bytes)
    result = build_result(source, cfg)
    return export_excel(result, sheet), result, sheet


def render_streamlit() -> None:
    if st is None:
        raise RuntimeError("Streamlit non installé")
    st.set_page_config(page_title=APP_NAME, layout="wide")
    st.title(APP_NAME)
    st.caption("Importer extraction AX -> génération automatique Lundi à Vendredi")
    uploaded = st.file_uploader("Extraction AX (.xlsx)", type=["xlsx"])
    auto_y, auto_w = automatic_planning_week(date.today())
    c1, c2, c3, c4 = st.columns(4)
    with c1: year = int(st.number_input("Année", 2024, 2035, auto_y, 1))
    with c2: week = int(st.number_input("Semaine ISO", 1, 53, auto_w, 1))
    with c3: cap = float(st.number_input("Capacité/jour (h)", 1.0, 24.0, DEFAULT_CAPACITY_H, 0.5))
    with c4: min_bal = float(st.number_input("Minutes/balancelle", 0.5, 30.0, DEFAULT_MIN_PER_BAL, 0.5))
    c5, c6, c7 = st.columns(3)
    with c5: include_without = st.checkbox("Inclure OF non créés", value=False)
    with c6: use_stock = st.checkbox("Utiliser StockPhysique en re-laquage", value=False, help="Option à valider avec la règle atelier.")
    with c7: wb_rule = st.checkbox("Interdire BLANC <-> NOIR/DARK consécutifs", value=True)
    if not uploaded:
        st.info("Charge un fichier Excel contenant la feuille 'extraction ax'.")
        return
    cfg = PlannerConfig(year=year, week=week, capacity_h=cap, minutes_per_bal=min_bal, include_without_of=include_without, use_stock_for_relaquage=use_stock, forbid_white_black_adjacent=wb_rule)
    try:
        excel, result, sheet = generate_file(uploaded.getvalue(), cfg)
    except Exception as exc:
        st.error(str(exc)); return
    m = result["metrics"]
    a,b,c,d = st.columns(4)
    a.metric("Lignes éligibles", result["quality"].get("eligible_lines", 0))
    b.metric("Charge", f"{m['total_load_h']:.1f} h / {m['capacity_h']:.1f} h")
    c.metric("Utilisation", f"{m['utilization_pct']:.0f}%")
    d.metric("Backlog", m["backlog"])
    st.caption(f"Source: {sheet} · moteur: {result['engine']} · OR-Tools: {'actif' if ORTOOLS_AVAILABLE else 'fallback'}")
    tabs = st.tabs([DAYS[d].title() for d in range(N_DAYS)])
    for d, tab in enumerate(tabs):
        with tab:
            dm = m["days"][d]
            st.write(f"**{dm['Date']} · {dm['Charge totale h']:.2f} h / {dm['Capacité h']:.2f} h · {dm['Couleurs'] or 'aucune couleur'}**")
            st.dataframe(_business_df(result["days"][d]), hide_index=True, use_container_width=True, height=460)
    st.download_button("Télécharger le planning Excel", data=excel, file_name=f"Planning_Lun_Ven_S{week}_{year}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary")


def _running_in_streamlit() -> bool:
    if st is None:
        return False
    try:
        from streamlit.runtime.scriptrunner import get_script_run_ctx
        return get_script_run_ctx() is not None
    except Exception:
        return False


def cli_main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="Générer le planning ALLUCO Lundi-Vendredi depuis extraction AX.")
    parser.add_argument("input", type=Path, help="Fichier .xlsx contenant extraction ax")
    parser.add_argument("-o", "--output", type=Path, default=Path("Planning_Lun_Ven.xlsx"))
    parser.add_argument("--year", type=int)
    parser.add_argument("--week", type=int)
    parser.add_argument("--capacity", type=float, default=DEFAULT_CAPACITY_H)
    parser.add_argument("--minutes-per-bal", type=float, default=DEFAULT_MIN_PER_BAL)
    parser.add_argument("--cleaning-min", type=int, default=DEFAULT_CLEANING_MIN)
    parser.add_argument("--include-without-of", action="store_true")
    parser.add_argument("--use-stock-for-relaquage", action="store_true")
    parser.add_argument("--allow-white-black-adjacent", action="store_true")
    args = parser.parse_args(argv)
    auto_y, auto_w = automatic_planning_week(date.today())
    cfg = PlannerConfig(
        year=args.year or auto_y,
        week=args.week or auto_w,
        capacity_h=args.capacity,
        minutes_per_bal=args.minutes_per_bal,
        cleaning_min=args.cleaning_min,
        include_without_of=args.include_without_of,
        use_stock_for_relaquage=args.use_stock_for_relaquage,
        forbid_white_black_adjacent=not args.allow_white_black_adjacent,
    )
    excel, result, sheet = generate_file(args.input.read_bytes(), cfg)
    args.output.write_bytes(excel)
    print(f"OK: {args.output}")
    print(f"source={sheet} eligible={result['quality'].get('eligible_lines',0)} planned={sum(len(x) for x in result['days'].values())} backlog={len(result['unscheduled'])}")
    for dm in result["metrics"]["days"]:
        print(f"{dm['Jour']}: {dm['Charge totale h']:.2f}h / {dm['Capacité h']:.2f}h | {dm['Couleurs']} | {dm['Lignes']} lignes")
    return 0


if __name__ == "__main__":
    if _running_in_streamlit():
        render_streamlit()
    else:
        raise SystemExit(cli_main())

